"""에이전트 구성·역할·호출 흐름 다이어그램 (PNG + PPTX). 내용은 agents/roles.py에서 직접 읽는다.

  docs/diagrams/agent_01_roster.png        역할 카드 10장 (입력·출력·도구·상한)
  docs/diagrams/agent_02_dependency.png    에이전트 간 결과 재사용(의존) 그래프 + 사이의 코드 단계
  docs/diagrams/agent_03_orchestration.png 스윔레인: 사람 · 오케스트레이터(코드) · 에이전트 · 저장소 순서
  docs/AGENT_OVERVIEW.pptx

실행: .venv\\Scripts\\python scripts\\render_agent_diagrams.py
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches, Pt

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "scripts"))

from render_flow_diagrams import (  # noqa: E402
    C_AGENT,
    C_CODE,
    C_EXT,
    C_GATE,
    C_HUMAN,
    C_LINE,
    C_MCP,
    C_STORE,
    C_UI,
    LEGEND,
    Canvas,
    font,
    mid_bottom,
    mid_left,
    mid_right,
    mid_top,
)

from proposal_agent.agents.roles import ROLES, AgentRole  # noqa: E402

OUT_DIR = REPO / "docs" / "diagrams"

PURPOSE = {
    "rfp_analyst": "RFP 개요·범위·제약·평가·일정·예산·명시 솔루션·연계 대상 정리",
    "requirement_extractor": "청크 원문에서 요구사항 초안 추출 (ID·출처 검증은 코드)",
    "requirement_verifier": "별도 세션에서 누락·중복·왜곡·출처·모순·범위 검사",
    "tech_researcher": "웹 검색으로 솔루션·프레임워크·호환성 조사, 출처·미확인 기록",
    "architect": "구성 요소·연결·기술 제안 (구성도는 코드가 그림)",
    "planner": "WBS 항목·공수·근거·선후행 제안 (날짜 계산 금지)",
    "proposal_writer": "제출용·발표용 제안서 문안·목차 제안",
    "document_generator": "문서 구성 계획 (파일은 코드가 생성)",
    "qa_reviewer": "산출물 내용 정합성·과장 문구 검토 (판정은 코드·사람)",
    "semantic_differ": "구축 변경 요청을 기존 요구사항과 의미 비교 (규칙 결과에 덧붙임)",
}
PHASE = {
    "semantic_differ": "구축 단계",
    "qa_reviewer": "제안·구축 (선택)",
}
INPUT_SHORT = {
    "rfp_analyst": "추출 블록 청크",
    "requirement_extractor": "청크 + RFP 분석",
    "requirement_verifier": "청크 + 요구사항 목록",
    "tech_researcher": "요구사항·RFP 분석·조사 주제",
    "architect": "요구사항·조사 결과",
    "planner": "요구사항·아키텍처·기간",
    "proposal_writer": "요구사항·조사·아키텍처·WBS",
    "document_generator": "기준 데이터 전체",
    "qa_reviewer": "산출물 텍스트·기준·코드 검증 결과",
    "semantic_differ": "변경 항목·기존 요구사항·규칙 분류",
}
CONSUMER = {  # 결과를 쓰는 코드
    "rfp_analyst": "조사 주제·기간 힌트·목차 결정",
    "requirement_extractor": "RequirementService (ID 발급·출처 검증)",
    "requirement_verifier": "지적 사항 → Gate 1 (코드 판정)",
    "tech_researcher": "출처 재분류·구현 가능성 저장·QA-TECH",
    "architect": "GenerationService (구성 요소·구성도)",
    "planner": "PlanningService (영업일 일정·공수 계산)",
    "proposal_writer": "PPT 문안·목차",
    "document_generator": "문서 계획 참고",
    "qa_reviewer": "QA-AI-nn 사람 검토 항목",
    "semantic_differ": "불일치 시 AMBIGUOUS로 낮춤",
}
ORDER = [
    "rfp_analyst",
    "requirement_extractor",
    "requirement_verifier",
    "tech_researcher",
    "architect",
    "planner",
    "proposal_writer",
    "document_generator",
    "qa_reviewer",
    "semantic_differ",
]


def tools_text(r: AgentRole) -> str:
    parts: list[str] = []
    if r.builtin_tools:
        parts.append("내장 " + "/".join(r.builtin_tools))
    n = len(r.mcp_tool_patterns)
    if n == 0:
        parts.append("MCP 없음")
    elif n <= 2:
        parts.append("pwa " + ", ".join(p.split("__")[-1] for p in r.mcp_tool_patterns))
    else:
        parts.append(f"pwa 읽기 {n}종")
    return " · ".join(parts)


# --- 1. 역할 카드 -------------------------------------------------------------------------------
def diagram_roster() -> Canvas:
    c = Canvas.new(2400, 1560)
    c.title(
        "에이전트 구성 — 역할 10종 (agents/roles.py 기준)",
        "모든 역할은 JSON 구조화 출력만 돌려주고 파일 쓰기 도구가 없다. 스킬·서브에이전트 없음. 실제 모델 응답은 크레딧 부족으로 미검증",
    )
    cw, ch = 452, 600
    for i, key in enumerate(ORDER):
        r = ROLES[key]
        col, row = i % 5, i // 5
        x, y = 40 + col * (cw + 20), 110 + row * (ch + 20)
        fill = C_AGENT if key != "semantic_differ" else "#FDE2C8"
        c.draw.rounded_rectangle(
            (x, y, x + cw, y + ch), radius=16, fill=fill, outline=C_LINE, width=3
        )
        c.draw.rounded_rectangle(
            (x, y, x + cw, y + 74), radius=16, fill="#F4B183", outline=C_LINE, width=3
        )
        c.draw.rectangle((x, y + 40, x + cw, y + 74), fill="#F4B183")
        c.draw.text((x + 16, y + 10), f"{i + 1}. {key}", font=font(24, True), fill="#1F1F1F")
        c.draw.text((x + 16, y + 42), r.title, font=font(19), fill="#333333")
        badge = PHASE.get(key, "제안 단계")
        bw = c.draw.textlength(badge, font=font(16)) + 20
        c.draw.rounded_rectangle(
            (x + cw - bw - 14, y + 12, x + cw - 14, y + 38),
            radius=8,
            fill="#FFFFFF",
            outline=C_LINE,
            width=1,
        )
        c.draw.text((x + cw - bw - 4, y + 15), badge, font=font(16), fill="#333333")
        yy = y + 90
        rows = [
            ("목적", PURPOSE[key]),
            ("입력", INPUT_SHORT[key]),
            ("출력", r.output_schema.__name__),
            ("도구", tools_text(r)),
            ("재사용", ", ".join(r.reuses) or "없음"),
            (
                "상한",
                f"{r.max_turns}턴 · ${r.max_budget_usd:g} · {r.timeout_s}s · 재시도 {r.retries}",
            ),
            ("effort", f"{r.effort} · 모델 {r.model_key}"),
            ("결과 사용", CONSUMER[key]),
            ("저장", f"{r.output_subdir}/{r.output_stem}_vX.Y.json"),
        ]
        for label, val in rows:
            c.draw.text((x + 16, yy), label, font=font(17, True), fill="#7F3F00")
            lines = _wrap_px(c, val, font(16), cw - 130)
            for ln in lines:
                c.draw.text((x + 110, yy), ln, font=font(16), fill="#1F1F1F")
                yy += 23
            yy += 8
    c.draw.text(
        (40, 1400),
        "공통 통제: permission_mode=dontAsk · 허용 목록 밖 도구는 PreToolUse 훅이 거부 · Bash/Write/Edit/NotebookEdit/Agent/Skill 금지 · 예산·턴·시간 초과는 즉시 실패(재시도는 일시 오류만) · 출력은 Pydantic 스키마 검증",
        font=font(18),
        fill="#555555",
    )
    return c


def _wrap_px(c: Canvas, text: str, f, max_px: float) -> list[str]:  # type: ignore[no-untyped-def]
    """픽셀 폭 기준 줄바꿈 (공백·구분자 우선, 없으면 글자 단위)."""
    lines: list[str] = []
    cur = ""
    for token in text.replace("/", "/\u200b").split(" "):
        cand = (cur + " " + token).strip() if cur else token
        if c.draw.textlength(cand.replace("\u200b", ""), font=f) <= max_px:
            cur = cand
            continue
        if cur:
            lines.append(cur.replace("\u200b", ""))
        # 긴 토큰은 글자 단위로 자른다
        piece = ""
        for ch in token.replace("\u200b", ""):
            if c.draw.textlength(piece + ch, font=f) > max_px:
                lines.append(piece)
                piece = ch
            else:
                piece += ch
        cur = piece
    if cur:
        lines.append(cur.replace("\u200b", ""))
    return lines[:4]


def _wrap(text: str, n: int) -> list[str]:
    words = text.split(" ")
    lines: list[str] = []
    cur = ""
    for w in words:
        if len(cur) + len(w) + 1 > n and cur:
            lines.append(cur)
            cur = w
        else:
            cur = (cur + " " + w).strip()
    if cur:
        lines.append(cur)
    return lines[:3]


# --- 2. 의존(재사용) 그래프 -----------------------------------------------------------------------
def diagram_dependency() -> Canvas:
    c = Canvas.new(2400, 1560)
    c.title(
        "에이전트 간 호출 Flow — 결과 재사용(reuses)과 사이의 코드 단계",
        "에이전트는 서로 직접 호출하지 않는다. 오케스트레이터(코드)가 순서대로 호출하고, 이전 역할의 JSON을 <prior role> 블록으로 다음 역할 프롬프트에 넣는다",
    )
    bw, bh = 300, 96
    pos = {
        "rfp_analyst": (120, 230),
        "requirement_extractor": (560, 230),
        "requirement_verifier": (1000, 230),
        "tech_researcher": (120, 540),
        "architect": (560, 540),
        "planner": (1000, 540),
        "proposal_writer": (560, 850),
        "document_generator": (1000, 850),
        "qa_reviewer": (1440, 850),
    }
    boxes = {}
    for key, (x, y) in pos.items():
        r = ROLES[key]
        boxes[key] = c.box(
            x, y, bw, bh, f"{key}\n→ {r.output_schema.__name__}", fill=C_AGENT, size=20
        )
    # 코드 단계 (회색 노트)
    code = {}
    code["extract"] = c.box(
        480, 120, 460, 56, "코드: 문서 추출 → 블록 → 청크 (Gate 0)", fill=C_CODE, size=18
    )
    code["ids"] = c.box(
        1400,
        230,
        420,
        96,
        "코드: ID 발급·출처 검증\n→ DB 저장 (요구사항 목록)",
        fill=C_CODE,
        size=18,
    )
    code["gate1"] = c.diamond(2000, 278, 220, 110, "Gate 1\n사람 승인", size=18)
    code["src"] = c.box(
        60, 680, 420, 56, "코드: 출처 신뢰도 재분류·미확인 강등", fill=C_CODE, size=18
    )
    code["sched"] = c.box(
        1400, 560, 420, 56, "코드: 영업일 일정·공수·인력 계산", fill=C_CODE, size=18
    )
    code["gen"] = c.box(
        560, 990, 740, 56, "코드: 구성도·Excel 4종·PPT 2종 생성 (_vX.Y)", fill=C_CODE, size=18
    )
    code["iqa"] = c.box(
        1440, 990, 560, 56, "코드: 통합 QA 30항목 → Gate 2·3·4·5", fill=C_CODE, size=18
    )
    code["final"] = c.box(2060, 830, 300, 136, "사람 승인\nGate 5\n→ 최종본", fill=C_HUMAN, size=19)
    # 재사용 화살표 (reuses)
    edges = [
        ("rfp_analyst", "requirement_extractor", "RfpAnalysis 재사용"),
        ("rfp_analyst", "tech_researcher", ""),
        ("rfp_analyst", "architect", ""),
        ("rfp_analyst", "planner", ""),
        ("tech_researcher", "architect", "ResearchReport 재사용"),
        ("architect", "planner", "ArchitectureDraft 재사용"),
        ("architect", "proposal_writer", ""),
        ("planner", "proposal_writer", "WBSDraft 재사용"),
        ("tech_researcher", "proposal_writer", ""),
        ("rfp_analyst", "proposal_writer", ""),
        ("proposal_writer", "document_generator", "ProposalContent 재사용"),
        ("planner", "document_generator", ""),
        ("architect", "document_generator", ""),
        ("document_generator", "qa_reviewer", "DocumentPlan 재사용"),
        ("proposal_writer", "qa_reviewer", ""),
    ]
    for a, b, label in edges:
        ba, bb = boxes[a], boxes[b]
        same_row = abs(ba[1] - bb[1]) < 5
        if same_row:
            c.arrow(
                mid_right(ba), mid_left(bb), label=label, label_offset=(0, -26), dashed=not label
            )
        elif ba[0] == bb[0]:
            c.arrow(
                mid_bottom(ba), mid_top(bb), label=label, label_offset=(120, 0), dashed=not label
            )
        else:
            start = (ba[0] + bw * 0.5, ba[3]) if ba[1] < bb[1] else mid_bottom(ba)
            end = (bb[0] + bw * 0.5, bb[1])
            c.arrow(start, end, label=label, label_offset=(0, -14), dashed=not label, width=2)
    # 코드 단계 연결
    c.arrow(mid_bottom(code["extract"]), (270, 230), width=2)
    c.arrow((710, 176), (710, 230), width=2)
    # 추출기 → (코드: ID 발급) → DB → 검증기 (독립 검증은 DB 목록을 받음)
    c.draw.line(
        [mid_bottom(boxes["requirement_extractor"]), (710, 360), (1610, 360)], fill=C_LINE, width=2
    )
    c.arrow(
        (1610, 360),
        mid_bottom(code["ids"]),
        width=2,
        label="초안 → 출처 검증·ID 발급",
        label_offset=(-230, 16),
    )
    c.arrow(
        mid_left(code["ids"]),
        mid_right(boxes["requirement_verifier"]),
        width=2,
        label="DB 요구사항 목록",
        label_offset=(0, -24),
    )
    c.draw.line([mid_right(code["ids"]), (1880, 278)], fill=C_LINE, width=3)
    c.arrow((1880, 278), mid_left(code["gate1"]), label="findings")
    c.arrow(mid_bottom(boxes["tech_researcher"]), mid_top(code["src"]), width=2)
    c.arrow(mid_right(boxes["planner"]), mid_left(code["sched"]), width=2)
    c.arrow(mid_bottom(boxes["document_generator"]), (1150, 990), width=2)
    c.arrow(mid_right(code["gen"]), mid_left(code["iqa"]))
    c.arrow(
        mid_top(code["iqa"]),
        (1590, 946),
        width=2,
        dashed=True,
        label="산출물 텍스트",
        label_offset=(130, 0),
    )
    c.arrow(mid_right(code["iqa"]), (2210, 966), width=2)
    c.arrow(
        mid_right(boxes["qa_reviewer"]),
        mid_left(code["final"]),
        dashed=True,
        label="QA-AI-nn (검토 항목)",
        label_offset=(0, -26),
    )
    # 구축 단계 레인
    c.draw.rounded_rectangle(
        (60, 1100, 2340, 1290), radius=16, fill="#FBF3EC", outline="#9E9E9E", width=2
    )
    c.draw.text((80, 1110), "구축 단계 (별도 흐름)", font=font(22, True), fill="#555555")
    b1 = c.box(
        90,
        1160,
        420,
        90,
        "코드: 규칙 의미 비교 (semantic.py)\nADD/MODIFY/DELETE/NO_CHANGE/AMBIGUOUS",
        fill=C_CODE,
        size=18,
    )
    b2 = c.box(
        600, 1160, 380, 90, "semantic_differ (선택)\n→ SemanticDiffDraft", fill="#FDE2C8", size=20
    )
    b3 = c.box(
        1070,
        1160,
        420,
        90,
        "코드: 불일치 → AMBIGUOUS\n영향 분석(impact.py) · 승인 후 반영",
        fill=C_CODE,
        size=18,
    )
    b4 = c.box(1580, 1160, 300, 90, "사람 결정\n후보 수락·승인/반려", fill=C_HUMAN, size=19)
    b5 = c.box(
        1970, 1160, 340, 90, "코드: 문서 재생성 → 통합 QA\n(qa_reviewer 선택)", fill=C_CODE, size=18
    )
    c.arrow(mid_right(b1), mid_left(b2), dashed=True, label="규칙 결과 입력", label_offset=(0, -26))
    c.arrow(mid_right(b2), mid_left(b3), dashed=True)
    c.arrow(mid_right(b3), mid_left(b4))
    c.arrow(mid_right(b4), mid_left(b5))
    c.draw.text(
        (60, 1320),
        "읽는 법: 실선 = 이름 붙은 주요 재사용, 점선 = 추가 재사용(프롬프트에 함께 포함). requirement_verifier는 추출 결과를 프롬프트로 받지 않고 DB에 저장된 요구사항 목록을 받는다(독립 검증).\n"
        "에이전트 사이에는 항상 코드 단계가 있다: 출처 검증·ID 발급, 신뢰도 재분류, 일정 계산, 파일 생성, 통합 QA. 게이트 통과와 최종 확정은 사람이 한다.",
        font=font(19),
        fill="#555555",
    )
    c.legend(60, 1420, LEGEND[1:4])
    c.legend(600, 1420, LEGEND[6:])
    return c


# --- 3. 스윔레인 오케스트레이션 ------------------------------------------------------------------
def diagram_orchestration() -> Canvas:
    c = Canvas.new(3400, 1200)
    c.title(
        "오케스트레이션 스윔레인 — 사람 · 코드 · 에이전트 · 저장소 (제안 단계, 16단계)",
        "StepRunner가 단계를 순서대로 실행한다. 에이전트는 호출될 때만 동작하고, 게이트·파일·저장은 코드와 사람이 맡는다",
    )
    lanes = [
        ("사람 (UI)", C_HUMAN),
        ("코드: 오케스트레이터·서비스", C_CODE),
        ("에이전트 (Claude)", C_AGENT),
        ("저장소: SQLite · Output", C_STORE),
    ]
    top, lane_h = 110, 205
    for i, (name, color) in enumerate(lanes):
        y = top + i * lane_h
        c.draw.rectangle(
            (40, y, 3360, y + lane_h),
            fill="#FBFBFB" if i % 2 else "#FFFFFF",
            outline="#BDBDBD",
            width=2,
        )
        c.draw.rounded_rectangle(
            (50, y + 12, 240, y + lane_h - 12), radius=12, fill=color, outline=C_LINE, width=2
        )
        c.text_center(145, y + lane_h / 2, name, size=20, bold=True)

    def ly(i: int) -> float:
        return top + i * lane_h + lane_h / 2

    x0, step, bw, bh = 280, 192, 176, 90
    items = [
        (0, "RFP 업로드", "ui"),
        (1, "추출(코드)\nGate 0 판정", "code"),
        (0, "Gate 0 승인\n(검토 필요 시)", "human"),
        (2, "rfp_analyst", "agent"),
        (2, "requirement_\nextractor ×청크", "agent"),
        (1, "ID 발급·출처\n검증·캐시", "code"),
        (2, "requirement_\nverifier ×창", "agent"),
        (1, "Gate 1 판정", "code"),
        (0, "Gate 1 승인", "human"),
        (2, "tech_researcher\n(웹 검색)", "agent"),
        (2, "architect →\nplanner", "agent"),
        (1, "일정·공수·인력\n계산 (코드)", "code"),
        (1, "산출물 생성\nExcel·PPT·PNG", "code"),
        (1, "통합 QA 30항목\nGate 2·3·4·5", "code"),
        (2, "qa_reviewer\n(선택)", "agent"),
        (0, "Gate 5 승인\n→ 최종본", "human"),
    ]
    fills = {"ui": C_UI, "code": C_CODE, "human": C_HUMAN, "agent": C_AGENT}
    boxes = []
    for i, (lane, label, kind) in enumerate(items):
        x = x0 + i * step
        y = ly(lane) - bh / 2
        b = c.box(x, y, bw, bh, label, fill=fills[kind], size=17)
        c.draw.text((x + 4, y - 26), f"{i + 1}", font=font(19, True), fill="#7F3F00")
        boxes.append(b)
    for i in range(len(boxes) - 1):
        a, b = boxes[i], boxes[i + 1]
        if abs(a[1] - b[1]) < 5:
            c.arrow(mid_right(a), mid_left(b), width=2)
        else:
            c.arrow(mid_right(a), mid_left(b), width=2)
    store = {
        1: "ExtractedBlock\nworkspace/chunks",
        3: "rfp_analysis 02_\n(JSON·AgentOutput)",
        4: "requirement_\nextraction 03_",
        5: "Requirement\nSourceRef (DB)",
        6: "verification 03_",
        7: "GateEvaluation",
        9: "research 04_",
        10: "architecture 05_\nwbs_draft 06_",
        11: "WBSItem·Dependency\nStaffRole·Allocation",
        12: "Deliverable\n06_~10_ 파일",
        13: "QaRun · 보고서 11_\nGateEvaluation",
        14: "qa_report 11_",
        15: "Approval\n13_Final_",
    }
    for i, text in store.items():
        b = boxes[i]
        cx = (b[0] + b[2]) / 2
        tb = c.box(cx - bw / 2, ly(3) - 42, bw, 84, text, fill=C_STORE, size=15, bold_first=False)
        c.draw.line([(cx, b[3]), (cx, tb[1])], fill="#9E9E9E", width=2)
        c.arrow((cx, tb[1] - 14), (cx, tb[1]), width=2, color="#9E9E9E")
    c.draw.text(
        (40, 950),
        "게이트: Gate 0(추출 품질)·Gate 1(요구사항)·Gate 2(기술)·Gate 3(계획)·Gate 4(파일)·Gate 5(최종 독립 QA)는 코드가 판정하고 NEEDS_HUMAN_REVIEW는 사람 승인으로만 통과한다. "
        "3·9·16은 사람 결정이며 FAIL은 승인으로 통과시킬 수 없다.\n"
        "에이전트 호출 상한: 역할별 턴·예산·시간, 단계당 호출 수(PWA_MAX_AGENT_CALLS_PER_STEP), 프로젝트 누적 비용(PWA_MAX_BUDGET_USD_PER_PROJECT). 같은 청크는 캐시로 재사용하고, 신호 없는 청크는 호출하지 않는다.\n"
        "구축 단계는 별도 흐름(기준 적재 → 변경 요청 → 규칙 비교(+semantic_differ 선택) → 영향 분석 → 사람 승인 → 반영·재생성 → 통합 QA). agent_02_dependency.png 하단 참조.",
        font=font(18),
        fill="#555555",
    )
    c.legend(
        40,
        1060,
        [(C_UI, "화면 입력"), (C_HUMAN, "사람 결정"), (C_CODE, "코드"), (C_AGENT, "에이전트")],
    )
    c.legend(
        600,
        1060,
        [(C_STORE, "저장 기록"), (C_GATE, "게이트(코드 판정)"), (C_MCP, "MCP"), (C_EXT, "외부")],
    )
    return c


def build_pptx(images: list[tuple[str, Path]]) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.7), Inches(2.5))
    tf = tb.text_frame
    tf.text = "Proposal Writer Agent — 에이전트 구성·역할·호출 Flow"
    tf.paragraphs[0].runs[0].font.size = Pt(34)
    tf.paragraphs[0].runs[0].font.bold = True
    p = tf.add_paragraph()
    p.text = "역할 10종 · 결과 재사용 그래프 · 오케스트레이션 스윔레인 (2026-10-15, roles.py 기준) · 원본: docs/AGENT_OVERVIEW.md"
    p.runs[0].font.size = Pt(16)
    for title, path in images:
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.6))
        tb.text_frame.text = title
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
        tb.text_frame.paragraphs[0].runs[0].font.bold = True
        with Image.open(path) as im:
            w, h = im.size
        scale = min(Inches(12.7) / w, Inches(6.5) / h)
        pw, ph = int(w * scale), int(h * scale)
        s.shapes.add_picture(
            str(path),
            Emu(int((prs.slide_width - pw) / 2)),
            Inches(0.85),
            width=Emu(pw),
            height=Emu(ph),
        )
    # 역할 표 슬라이드
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.6))
    tb.text_frame.text = "역할별 입력·출력·도구·상한 (roles.py)"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True
    rows = len(ORDER) + 1
    table = s.shapes.add_table(rows, 6, Inches(0.4), Inches(0.9), Inches(12.5), Inches(6.0)).table
    heads = ["역할", "입력", "출력", "도구", "재사용", "상한(턴·$·초)"]
    for j, h in enumerate(heads):
        table.cell(0, j).text = h
    for i, key in enumerate(ORDER, start=1):
        r = ROLES[key]
        vals = [
            key,
            INPUT_SHORT[key],
            r.output_schema.__name__,
            tools_text(r),
            ", ".join(r.reuses) or "-",
            f"{r.max_turns} · {r.max_budget_usd:g} · {r.timeout_s}",
        ]
        for j, v in enumerate(vals):
            table.cell(i, j).text = v
    for row in table.rows:
        for cell in row.cells:
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(11)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    jobs = [
        ("agent_01_roster.png", "1. 에이전트 구성 — 역할 카드 10종", diagram_roster),
        (
            "agent_02_dependency.png",
            "2. 에이전트 간 호출 Flow — 결과 재사용 그래프",
            diagram_dependency,
        ),
        (
            "agent_03_orchestration.png",
            "3. 오케스트레이션 스윔레인 (제안 단계)",
            diagram_orchestration,
        ),
    ]
    images = []
    for name, title, fn in jobs:
        path = OUT_DIR / name
        fn().save(path)
        images.append((title, path))
        print("wrote", path)
    out = REPO / "docs" / "AGENT_OVERVIEW.pptx"
    out.write_bytes(build_pptx(images))
    print("wrote", out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
