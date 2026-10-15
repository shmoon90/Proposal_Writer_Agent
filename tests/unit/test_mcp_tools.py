"""인프로세스 MCP 도구(pwa): 핸들러 응답 스키마, 경로 제한, 파일 검증."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

from fixtures import builders as bld
from proposal_agent.agents.tools import TOOL_NAMES, TOOL_SPECS, ToolContext, verify_file
from proposal_agent.agents.tools.project_tools import (
    tool_get_prior_output,
    tool_get_project,
    tool_get_requirement,
    tool_get_source_block,
    tool_list_findings,
    tool_list_gates,
    tool_list_requirements,
    tool_list_research,
    tool_search_source_text,
    tool_verify_deliverable_file,
)
from proposal_agent.models import DocumentKind, Project, Requirement, RequirementType
from proposal_agent.services import Services


@pytest.fixture
def ctx(services: Services, proposal_project: Project) -> ToolContext:
    services.documents.register_upload(
        proposal_project, kind=DocumentKind.RFP, file_name="rfp.docx", data=bld.build_docx()
    )
    services.extraction.extract_all(proposal_project)
    services.extraction.evaluate_gate0(proposal_project)
    services.ctx.repos.requirements.insert(
        Requirement(
            id="FR-001",
            project_id=proposal_project.id,
            name="SSO 인증",
            type=RequirementType.FUNCTIONAL,
        )
    )
    return ToolContext(services.ctx.repos, services.ctx.paths, proposal_project)


def test_tool_names_and_specs() -> None:
    assert len(TOOL_SPECS) == 10 and all(n.startswith("mcp__pwa__") for n in TOOL_NAMES)
    for _name, description, schema, fn in TOOL_SPECS:
        assert description and callable(fn) and all(t in (str, int, bool) for t in schema.values())


def test_project_and_requirement_tools_are_json_serializable(ctx: ToolContext) -> None:
    p = tool_get_project(ctx, {})
    assert p["name"] == "제안 테스트 사업" and json.dumps(p, default=str)
    lst = tool_list_requirements(ctx, {"type": "functional"})
    assert lst["count"] == 1 and lst["requirements"][0]["id"] == "FR-001"
    one = tool_get_requirement(ctx, {"id": "FR-001"})
    assert one["name"] == "SSO 인증" and "history" not in one
    assert "error" in tool_get_requirement(ctx, {"id": "FR-999"})


def test_source_search_and_block(ctx: ToolContext) -> None:
    found = tool_search_source_text(ctx, {"query": "SSO", "limit": 5})
    assert found["count"] >= 1 and "SSO" in found["blocks"][0]["text"]
    block = tool_get_source_block(ctx, {"block_id": found["blocks"][0]["id"]})
    assert block["text"] and block["loc"]
    assert "error" in tool_search_source_text(ctx, {"query": "x"})
    assert "error" in tool_get_source_block(ctx, {"block_id": "blk_nope"})


def test_lists_and_gates(ctx: ToolContext) -> None:
    assert tool_list_research(ctx, {})["count"] == 0
    assert tool_list_findings(ctx, {})["count"] == 0
    gates = tool_list_gates(ctx, {})
    assert gates["gate0"]["status"] == "PASS" and gates["gate1"] is None
    assert tool_get_prior_output(ctx, {"role": "rfp_analyst"})["found"] is False


def test_verify_deliverable_file_structures_and_path_guard(
    ctx: ToolContext, tmp_path: Path
) -> None:
    project = ctx.project
    out = ctx.paths.project_output_dir(project.output_dir)
    (out / "09_Excel_Deliverables" / "a.docx").write_bytes(bld.build_docx())
    (out / "09_Excel_Deliverables" / "a.pdf").write_bytes(bld.build_text_pdf())
    from openpyxl import Workbook
    from pptx import Presentation

    wb = Workbook()
    wb.active.title = "요구사항목록"
    wb.active.append(["ID", "명"])
    buf = io.BytesIO()
    wb.save(buf)
    (out / "09_Excel_Deliverables" / "a.xlsx").write_bytes(buf.getvalue())
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5])
    buf2 = io.BytesIO()
    prs.save(buf2)
    (out / "07_Submission_Proposal_PPT" / "a.pptx").write_bytes(buf2.getvalue())
    (out / "09_Excel_Deliverables" / "bad.xlsx").write_bytes(b"not a workbook")

    assert tool_verify_deliverable_file(ctx, {"relative_path": "09_Excel_Deliverables/a.xlsx"})[
        "sheets"
    ] == {"요구사항목록": 1}
    assert (
        tool_verify_deliverable_file(ctx, {"relative_path": "07_Submission_Proposal_PPT/a.pptx"})[
            "slides"
        ]
        == 1
    )
    assert (
        tool_verify_deliverable_file(ctx, {"relative_path": "09_Excel_Deliverables/a.docx"})[
            "paragraphs"
        ]
        > 5
    )
    assert (
        tool_verify_deliverable_file(ctx, {"relative_path": "09_Excel_Deliverables/a.pdf"})["pages"]
        >= 1
    )
    bad = tool_verify_deliverable_file(ctx, {"relative_path": "09_Excel_Deliverables/bad.xlsx"})
    assert bad["ok"] is False and bad["error"]
    missing = tool_verify_deliverable_file(
        ctx, {"relative_path": "09_Excel_Deliverables/none.xlsx"}
    )
    assert missing["exists"] is False and missing["ok"] is False
    escape = tool_verify_deliverable_file(ctx, {"relative_path": "../../.env"})
    assert "error" in escape and "허용 범위" in escape["error"]
    assert verify_file(tmp_path / "nope.json")["ok"] is False


def test_build_project_server_wraps_handlers(ctx: ToolContext) -> None:
    from proposal_agent.agents.tools import build_project_server

    server = build_project_server(ctx)
    assert server["type"] == "sdk" and server["name"] == "pwa"
