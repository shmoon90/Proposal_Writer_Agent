"""Phase 6 구축 단계 변경관리: 14개 시나리오 (기준 적재, 매핑, 분류, 영향, 승인 흐름, 롤백, 문서 버전)."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fixtures import builders as bld
from fixtures import excel_builders as xb
from proposal_agent.agents import FakeAgentRunner
from proposal_agent.agents.schemas import SemanticDiffCandidate, SemanticDiffDraft
from proposal_agent.changes.impact import analyze_impact
from proposal_agent.changes.semantic import classify_request, similarity
from proposal_agent.config import IngestError, ValidationError
from proposal_agent.ingest.excel_reader import profile_workbook
from proposal_agent.models import (
    ChangeStatus,
    DocumentKind,
    InputKind,
    Project,
    ProjectMode,
    Requirement,
    RequirementStatus,
    RequirementType,
    StaffingScenario,
)
from proposal_agent.models.change import CandidateClass
from proposal_agent.models.planning import TaskStatus
from proposal_agent.planning.calendar import BusinessCalendar
from proposal_agent.services import AppContext, Services


@pytest.fixture
def services(ctx: AppContext) -> Services:
    return Services.from_context(ctx, runner=FakeAgentRunner({}))


def _upload_all(
    services: Services,
    p: Project,
    *,
    staffing: bool = True,
    wbs: bytes | None = None,
    spec: bytes | None = None,
) -> None:
    services.documents.register_upload(
        p,
        kind=DocumentKind.REQUIREMENT_SPEC,
        file_name="요구사항정의서.xlsx",
        data=spec or xb.build_requirement_spec(),
    )
    services.documents.register_upload(
        p, kind=DocumentKind.TRACE_MATRIX, file_name="추적표.xlsx", data=xb.build_trace_matrix()
    )
    services.documents.register_upload(
        p, kind=DocumentKind.WBS, file_name="WBS.xlsx", data=wbs or xb.build_wbs()
    )
    if staffing:
        services.documents.register_upload(
            p,
            kind=DocumentKind.STAFFING_PLAN,
            file_name="투입계획.xlsx",
            data=xb.build_staffing_plan(),
        )


@pytest.fixture
def loaded(services: Services, build_project: Project) -> Project:
    _upload_all(services, build_project)
    services.baseline.load(build_project, actor="tester")
    return services.ctx.repos.projects.get(build_project.id)


def _reqs() -> list[Requirement]:
    return [
        Requirement(
            id=r["id"],
            project_id="p",
            name=r["name"],
            description=r["desc"],
            type=RequirementType.FUNCTIONAL
            if r["type"] == "기능"
            else RequirementType.NON_FUNCTIONAL,
        )
        for r in xb.BASE_REQUIREMENTS
    ]


def _out(services: Services, rel: str) -> Path:
    return services.ctx.paths.resolve_output(rel)


# 1. 기존 산출물 업로드·원본 보존·기준 적재(ID 유지)·Revision v1.0 --------------------------------
def test_s01_baseline_load_preserves_originals_and_ids(
    services: Services, build_project: Project
) -> None:
    _upload_all(services, build_project)
    result = services.baseline.load(build_project, actor="tester")
    repos = services.ctx.repos
    ids = [r.id for r in repos.requirements.find(build_project.id, order_by="id")]
    assert ids == ["SFR-001", "SFR-002", "SFR-003", "SFR-004", "SFR-005", "SNR-001"]
    assert result.revision.number == "v1.0" and result.wbs_items == 16 and result.dependencies == 7
    assert result.screens == 5 and result.unit_tests == 5 and result.integration_tests == 3
    assert result.staffing_plan_present and result.roles == 5 and result.id_scheme == "SFR-000"
    originals = [d for d in repos.documents.find(build_project.id)]
    assert len(originals) == 4 and all(_out(services, d.stored_path).exists() for d in originals)
    assert _out(services, result.snapshot_path or "").exists()
    p = repos.projects.get(build_project.id)
    assert p.baseline_revision_id == result.revision.id
    # 재적재해도 ID는 바뀌지 않고 이력만 남는다
    services.baseline.load(p, actor="tester")
    assert [r.id for r in repos.requirements.find(p.id, order_by="id")] == ids
    assert len(repos.revisions.find(p.id)) == 2


# 2. 비표준 양식: 동의어·헤더 행 자동 매핑 ---------------------------------------------------------
def test_s02_nonstandard_headers_mapped_without_confirmation(
    services: Services, build_project: Project
) -> None:
    prof = profile_workbook(xb.build_nonstandard_spec(), DocumentKind.REQUIREMENT_SPEC)
    sheet = prof.sheets[0]
    assert sheet.header_row == 3 and not prof.needs_confirmation
    mapped = {c.field: c.header for c in sheet.columns if c.column_index is not None}
    assert (
        mapped["requirement_id"] == "번호"
        and mapped["name"] == "명칭"
        and mapped["description"] == "내용"
        and mapped["type"] == "구분"
    )
    _upload_all(services, build_project, spec=xb.build_nonstandard_spec())
    result = services.baseline.load(build_project, actor="tester")
    assert result.requirements == 6
    assert (
        services.ctx.repos.requirements.get("SNR-001", build_project.id).type
        == RequirementType.NON_FUNCTIONAL
    )


# 3. 확신 없는 매핑 → 사용자 확정 필요, 확정 후 적재 -------------------------------------------------
def test_s03_uncertain_mapping_requires_user_confirmation(
    services: Services, build_project: Project
) -> None:
    _upload_all(services, build_project, spec=xb.build_ambiguous_spec())
    profiles = services.baseline.profiles(build_project)
    spec_prof = profiles[DocumentKind.REQUIREMENT_SPEC]
    assert spec_prof.blocking and not spec_prof.confirmed
    with pytest.raises(ValidationError, match="컬럼 매핑"):
        services.baseline.load(build_project, actor="tester")
    services.baseline.confirm_mapping(
        build_project, spec_prof.document, {"Sheet1": {"name": 1, "description": 2}}, actor="tester"
    )
    doc = services.ctx.repos.documents.get(spec_prof.document.id)
    assert (
        doc.column_mapping
        and doc.column_mapping["confirmed"]
        and doc.column_mapping["confirmed_by"] == "tester"
    )
    result = services.baseline.load(build_project, actor="tester")
    assert result.requirements == 6
    assert services.ctx.repos.requirements.get("SFR-002", build_project.id).name == "민원 접수"


# 4. 손상 파일·필수 문서 누락 -------------------------------------------------------------------
def test_s04_corrupted_and_missing_inputs(services: Services, build_project: Project) -> None:
    with pytest.raises(IngestError):
        profile_workbook(xb.build_corrupted_xlsx(), DocumentKind.WBS)
    services.documents.register_upload(
        build_project,
        kind=DocumentKind.REQUIREMENT_SPEC,
        file_name="spec.xlsx",
        data=xb.build_requirement_spec(),
    )
    with pytest.raises(ValidationError, match="기준 문서가 부족"):
        services.baseline.load(build_project, actor="tester")
    with pytest.raises(ValidationError, match="기준 데이터가 적재되지"):
        services.changes.create_request(
            build_project,
            input_kind=InputKind.NATURAL_LANGUAGE,
            request_text="x",
            requester="a",
            reason="b",
        )


# 5. 자연어 분류: ADD / MODIFY / DELETE_REQUEST / NO_CHANGE / AMBIGUOUS ---------------------------
def test_s05_semantic_classification_rules() -> None:
    reqs = _reqs()
    one = lambda t: classify_request(t, reqs)[0]  # noqa: E731
    add = one("모바일 앱에서 푸시 알림을 수신하는 기능이 필요하다.")
    assert (
        add.classification == CandidateClass.ADD
        and add.requirement_id is None
        and add.needs_confirmation
    )
    mod = one("사용자 로그인은 아이디/비밀번호 외에 간편인증(카카오)도 지원하도록 변경 바랍니다.")
    assert (
        mod.classification == CandidateClass.MODIFY
        and mod.requirement_id == "SFR-001"
        and mod.before
        and mod.after
    )
    dele = one("팩스 접수 기능은 더 이상 필요 없으니 삭제해 주세요.")
    assert dele.classification == CandidateClass.DELETE_REQUEST and dele.requirement_id == "SFR-005"
    same = one("사용자는 아이디와 비밀번호로 로그인한다. 5회 실패 시 계정을 잠근다.")
    assert same.classification == CandidateClass.NO_CHANGE and same.requirement_id == "SFR-001"
    num = one("주요 화면은 2초 이내에 응답해야 한다.")
    assert (
        num.classification == CandidateClass.MODIFY and num.requirement_id == "SNR-001"
    )  # 수치가 다르면 '같음'이 아니다
    amb = one("접수 기능을 삭제")
    assert amb.classification == CandidateClass.AMBIGUOUS and amb.needs_confirmation
    multi = classify_request("1. 모바일 푸시 알림 수신 기능 추가\n2. 팩스 접수 삭제", reqs)
    assert [c.classification for c in multi] == [CandidateClass.ADD, CandidateClass.DELETE_REQUEST]
    assert all(c.rationale and 0.0 <= c.confidence <= 1.0 for c in multi)
    assert similarity("민원 접수", "민원 접수") == 1.0 and similarity("", "x") == 0.0


# 6. 버튼 입력(선택/추가/수정/삭제) + 중복 추가 감지, 사용자 출처 구분 ----------------------------------
def test_s06_button_inputs_and_duplicate_add() -> None:
    reqs = _reqs()
    mod = classify_request(
        "로그인 5회 실패 시 10분간 잠금",
        reqs,
        input_kind=InputKind.MODIFY_BUTTON,
        selected_ids=["SFR-001"],
    )
    assert (
        len(mod) == 1 and mod[0].classification == CandidateClass.MODIFY and mod[0].source == "user"
    )
    dele = classify_request(
        "", reqs, input_kind=InputKind.DELETE_BUTTON, selected_ids=["SFR-005", "SFR-004"]
    )
    assert [c.requirement_id for c in dele] == ["SFR-005", "SFR-004"] and all(
        c.classification == CandidateClass.DELETE_REQUEST for c in dele
    )
    dup = classify_request(
        "사용자는 아이디와 비밀번호로 로그인한다. 5회 실패 시 계정을 잠근다.",
        reqs,
        input_kind=InputKind.ADD_BUTTON,
    )[0]
    assert dup.classification == CandidateClass.AMBIGUOUS and dup.requirement_id == "SFR-001"
    new = classify_request("모바일 푸시 알림 수신", reqs, input_kind=InputKind.ADD_BUTTON)[0]
    assert new.classification == CandidateClass.ADD and new.source == "user"


# 7. 문서 업로드 변경 요청 + AI 의미 비교 덧붙이기(불일치 시 모호 처리) -------------------------------
def test_s07_document_request_and_agent_overlay(ctx: AppContext, build_project: Project) -> None:
    def differ(request):  # type: ignore[no-untyped-def]
        import re

        assert "<rule_result>" in request.prompt and "SFR-005" in request.prompt
        push = int(re.search(r"\[(\d+)\] [^\n]*푸시", request.prompt).group(1))
        fax = int(re.search(r"\[(\d+)\] [^\n]*팩스", request.prompt).group(1))
        return SemanticDiffDraft(
            candidates=[
                SemanticDiffCandidate(
                    item_index=push, classification="ADD", rationale="새 채널"
                ),  # 규칙(ADD)과 일치
                SemanticDiffCandidate(
                    item_index=fax,
                    classification="MODIFY",
                    requirement_id="SFR-002",
                    rationale="접수 채널 축소로 봄",
                    confidence=0.6,
                ),  # 규칙(DELETE)과 불일치
            ]
        )

    services = Services.from_context(ctx, runner=FakeAgentRunner({"semantic_differ": differ}))
    _upload_all(services, build_project)
    services.baseline.load(build_project, actor="tester")
    p = ctx.repos.projects.get(build_project.id)
    doc = bld.build_docx(
        sections=[
            bld.Section(
                "변경 요청",
                paragraphs=[
                    "1. 모바일 앱에서 푸시 알림을 수신하는 기능 추가",
                    "2. 팩스 접수 기능 삭제",
                ],
            )
        ],
        title="변경요청서",
    )
    cr = services.changes.create_request(
        p,
        input_kind=InputKind.DOCUMENT,
        request_text="",
        requester="고객",
        reason="변경요청서",
        attachments=[("변경요청서.docx", doc)],
    )
    assert cr.attachment_document_ids and "[첨부" in cr.request_text
    cr = services.changes.analyze(p, cr, use_agent=True)
    assert cr.analysis_sources == ["rule", "ai"]
    assert any(c.source == "rule+ai" for c in cr.candidates)
    c1 = next(c for c in cr.candidates if "푸시" in c.evidence[0])
    assert (
        c1.classification == CandidateClass.ADD
        and c1.source == "rule+ai"
        and "규칙과 일치" in c1.rationale
    )
    c2 = next(c for c in cr.candidates if "팩스" in c.evidence[0])
    assert (
        c2.classification == CandidateClass.AMBIGUOUS
        and "AI 의견" in c2.rationale
        and c2.needs_confirmation
    )
    # AI 실패 시 규칙 결과만 사용하고 미해결 사항에 기록
    services2 = Services.from_context(ctx, runner=FakeAgentRunner({}))
    cr2 = services2.changes.create_request(
        p,
        input_kind=InputKind.NATURAL_LANGUAGE,
        request_text="팩스 접수 기능 삭제",
        requester="고객",
        reason="r",
    )
    cr2 = services2.changes.analyze(p, cr2, use_agent=True)
    assert (
        cr2.analysis_sources == ["rule"]
        and cr2.unresolved
        and cr2.candidates[0].classification == CandidateClass.DELETE_REQUEST
    )


# 8. 영향 분석: 신규 Task·공수 범위·테스트 추가, 여유(slack)가 있으면 후행을 밀지 않음 --------------------
def test_s08_impact_add_task_and_slack_preserved(services: Services, loaded: Project) -> None:
    cr = services.changes.create_request(
        loaded,
        input_kind=InputKind.NATURAL_LANGUAGE,
        request_text="모바일 앱에서 푸시 알림을 수신하는 기능이 필요하다.",
        requester="홍",
        reason="r",
    )
    cr = services.changes.analyze(loaded, cr, use_agent=False)
    imp = cr.impact
    assert imp and not imp.error
    new = [t for t in imp.tasks if t.action == "add"]
    assert len(new) == 1 and new[0].effort_added_md > 0 and new[0].after_start is not None
    assert (
        imp.effort.estimate_status == "범위"
        and imp.effort.range_min_md is not None
        and imp.effort.basis
    )
    assert (
        imp.tests_added
        and imp.effort.added_md == new[0].effort_added_md
        and imp.effort.after_md == imp.effort.before_md + imp.effort.net_md
    )
    # 여유가 있는 Task 수정: B(5일, 여유 5일)에 3 M/D 추가 → 후행 C는 움직이지 않는다
    rows = [
        ("1", "", "", "분석", "PL", "", date(2026, 11, 2), date(2026, 11, 27), 20, 30),
        ("1.1", "1", "", "조사", "PL", "", date(2026, 11, 2), date(2026, 11, 13), 10, 15),
        ("1.1.1", "1.1", "SFR-001", "A", "PL", "", date(2026, 11, 2), date(2026, 11, 13), 10, 10),
        ("1.1.2", "1.1", "SFR-002", "B", "PL", "", date(2026, 11, 2), date(2026, 11, 6), 5, 5),
        ("1.2", "1", "", "설계", "PL", "", date(2026, 11, 16), date(2026, 11, 27), 10, 15),
        (
            "1.2.1",
            "1.2",
            "SFR-003",
            "C",
            "PL",
            "1.1.1, 1.1.2",
            date(2026, 11, 16),
            date(2026, 11, 27),
            10,
            15,
        ),
    ]
    p2 = services.projects.create(name="여유 테스트", mode=ProjectMode.BUILD)
    _upload_all(services, p2, staffing=False, wbs=xb.build_wbs(rows))
    services.baseline.load(p2, actor="t")
    p2 = services.ctx.repos.projects.get(p2.id)
    cr2 = services.changes.create_request(
        p2,
        input_kind=InputKind.MODIFY_BUTTON,
        request_text="민원 접수에 첨부 용량 제한 추가",
        requester="홍",
        reason="r",
        selected_ids=["SFR-002"],
    )
    cr2 = services.changes.analyze(p2, cr2, use_agent=False)
    cr2 = services.changes.decide_candidate(
        p2, cr2, cr2.candidates[0].id, decision="accept", overrides={"effort_md": 3}
    )
    imp2 = cr2.impact
    assert imp2 and not imp2.error
    b = next(t for t in imp2.tasks if t.wbs_code == "1.1.2")
    assert b.after_end == date(2026, 11, 11) and b.effort_added_md == Decimal("3")
    assert imp2.schedule_delta_days == 0 and all(
        d.startswith("1.1.2 ") for d in imp2.changed_task_dates
    )  # 후행 C·상위 항목은 불변
    assert imp2.effort.estimate_status == "확정"


# 9. 삭제 영향: 전용 Task 제거(보류)·공유 Task 부분 감소·의존 우회 ------------------------------------
def test_s09_delete_impact_shared_and_exclusive_tasks(services: Services, loaded: Project) -> None:
    cr = services.changes.create_request(
        loaded,
        input_kind=InputKind.DELETE_BUTTON,
        request_text="",
        requester="홍",
        reason="범위 축소",
        selected_ids=["SFR-005", "SFR-003"],
    )
    cr = services.changes.analyze(loaded, cr, use_agent=False)
    imp = cr.impact
    assert imp and not imp.error
    removed = {t.wbs_code: t for t in imp.tasks if t.action == "remove"}
    assert "2.2.2" in removed and removed["2.2.2"].effort_reduced_md == Decimal("20")
    shared = next(
        t for t in imp.tasks if t.wbs_code == "1.1.1"
    )  # SFR-001/002/003 공유 → 1/3만 감소
    assert shared.action == "modify" and shared.effort_reduced_md == Decimal("6.7")
    assert "SFR-001" in imp.affected_requirement_ids and "SFR-002" in imp.affected_requirement_ids
    assert any("2.2.2" in d for d in imp.dependency_changes)
    assert imp.effort.net_md < 0 and imp.staffing_scenario == StaffingScenario.S5_OFFSET


# 10. 순환 의존 → 계산 중단, 승인 차단 ---------------------------------------------------------------
def test_s10_cycle_stops_impact_and_blocks_approval(
    services: Services, build_project: Project
) -> None:
    _upload_all(services, build_project, wbs=xb.build_wbs(extra_deps=[("3.1.1", "2.1.1", "FS", 0)]))
    services.baseline.load(build_project, actor="t")
    p = services.ctx.repos.projects.get(build_project.id)
    cr = services.changes.create_request(
        p,
        input_kind=InputKind.NATURAL_LANGUAGE,
        request_text="모바일 앱에서 푸시 알림을 수신하는 기능이 필요하다.",
        requester="홍",
        reason="r",
    )
    cr = services.changes.analyze(p, cr, use_agent=False)
    assert cr.impact and cr.impact.error and "순환" in cr.impact.error
    cr = services.changes.decide_candidate(p, cr, cr.candidates[0].id, decision="accept")
    blockers = services.changes.approval_blockers(cr)
    assert any("영향 분석 오류" in b for b in blockers)
    with pytest.raises(ValidationError, match="승인 조건"):
        services.changes.approve(p, cr, approver="PM", comments="x")
    assert services.ctx.repos.requirements.count(p.id) == 6


# 11. 인력 시나리오: 투입계획 없음(6), 새 역할(4), 기간 연장(3), 상쇄(5) ---------------------------------
def test_s11_staffing_scenarios(services: Services, loaded: Project) -> None:
    # 6: 투입계획 없음
    p0 = services.projects.create(name="투입계획 없음", mode=ProjectMode.BUILD)
    _upload_all(services, p0, staffing=False)
    services.baseline.load(p0, actor="t")
    p = services.ctx.repos.projects.get(p0.id)
    cr = services.changes.create_request(
        p,
        input_kind=InputKind.NATURAL_LANGUAGE,
        request_text="모바일 앱에서 푸시 알림을 수신하는 기능이 필요하다.",
        requester="홍",
        reason="r",
    )
    cr = services.changes.analyze(p, cr, use_agent=False)
    assert (
        cr.impact
        and cr.impact.staffing_scenario == StaffingScenario.S6_UNKNOWN
        and not cr.impact.staffing_plan_present
    )
    # 4: 새 역할
    cr2 = services.changes.create_request(
        loaded,
        input_kind=InputKind.ADD_BUTTON,
        request_text="AI 챗봇 상담 기능",
        requester="홍",
        reason="r",
    )
    cr2 = services.changes.analyze(loaded, cr2, use_agent=False)
    cr2 = services.changes.decide_candidate(
        loaded,
        cr2,
        cr2.candidates[0].id,
        decision="accept",
        overrides={"role": "AI엔지니어", "effort_md": 10},
    )
    assert cr2.impact and cr2.impact.staffing_scenario == StaffingScenario.S4_NEW_ROLE
    assert any(r.role == "AI엔지니어" and not r.exists_in_plan for r in cr2.impact.roles)
    # 3: 기간 연장 (큰 추가 공수, 기존 역할)
    cr3 = services.changes.create_request(
        loaded,
        input_kind=InputKind.MODIFY_BUTTON,
        request_text="성능 시험 범위를 전 화면으로 확대",
        requester="홍",
        reason="r",
        selected_ids=["SNR-001"],
    )
    cr3 = services.changes.analyze(loaded, cr3, use_agent=False)
    cr3 = services.changes.decide_candidate(
        loaded, cr3, cr3.candidates[0].id, decision="accept", overrides={"effort_md": 20}
    )
    assert (
        cr3.impact
        and cr3.impact.schedule_delta_days > 0
        and cr3.impact.staffing_scenario == StaffingScenario.S3_EXTEND
    )
    assert cr3.impact.cost_delta is None  # QA 단가 없음 → 비용 미계산
    # 비용: 영향 역할 모두 단가가 있으면 비용 산출
    for role in services.ctx.repos.staff_roles.find(loaded.id):
        if role.unit_cost_per_mm is None:
            services.planning.set_unit_cost(loaded, role.id, Decimal("7000000"), actor="t")
    cr4 = services.changes.create_request(
        loaded,
        input_kind=InputKind.MODIFY_BUTTON,
        request_text="접수번호 조회에 QR 지원",
        requester="홍",
        reason="r",
        selected_ids=["SFR-003"],
    )
    cr4 = services.changes.analyze(loaded, cr4, use_agent=False)
    cr4 = services.changes.decide_candidate(
        loaded, cr4, cr4.candidates[0].id, decision="accept", overrides={"effort_md": 6}
    )
    assert (
        cr4.impact is not None and cr4.impact.cost_delta is not None and cr4.impact.cost_delta > 0
    )


# 12. 승인 흐름: 승인 전 기준 불변 → 승인 후 새 ID·이력·추적 재검토·버전 파일·Revision 연결, 반려는 무변경 ---
def test_s12_approval_flow_and_versioned_outputs(services: Services, loaded: Project) -> None:
    repos = services.ctx.repos
    cr = services.changes.create_request(
        loaded,
        input_kind=InputKind.NATURAL_LANGUAGE,
        request_text="1. 모바일 앱에서 푸시 알림을 수신하는 기능이 필요하다.\n2. 주요 화면은 2초 이내에 응답해야 한다.\n3. 팩스 접수 기능은 더 이상 필요 없으니 삭제해 주세요.\n4. 접수 기능을 삭제",
        requester="홍길동",
        reason="고객 요청",
    )
    cr = services.changes.analyze(loaded, cr, use_agent=False)
    assert [c.classification for c in cr.candidates] == [
        CandidateClass.ADD,
        CandidateClass.MODIFY,
        CandidateClass.DELETE_REQUEST,
        CandidateClass.AMBIGUOUS,
    ]
    # 승인 전: 기준 데이터 불변
    assert repos.requirements.count(loaded.id) == 6 and repos.requirements.get(
        "SNR-001", loaded.id
    ).description.startswith("주요 화면은 3초")
    with pytest.raises(ValidationError, match="모호"):
        services.changes.decide_candidate(loaded, cr, "C04", decision="accept")
    cr = services.changes.decide_candidate(loaded, cr, "C04", decision="reject")
    for cid in ("C01", "C02", "C03"):
        cr = services.changes.decide_candidate(loaded, cr, cid, decision="accept")
    assert cr.candidates[0].user_decision == "accept" and cr.candidates[0].source == "rule"
    assert services.changes.approval_blockers(cr) == []
    res = services.changes.approve(loaded, cr, approver="PM", comments="승인")
    cr = res.request
    assert (
        cr.status == ChangeStatus.APPLIED
        and cr.approval_id
        and cr.revision_id == res.revision.id
        and res.revision.number == "v2.0"
    )
    assert res.new_ids == {"C01": "SFR-006"}  # 기존 체계(SFR-)로 승인 후 발급
    new = repos.requirements.get("SFR-006", loaded.id)
    assert (
        new.added_by == "홍길동" and new.added_on == cr.requested_on and new.change_state == "added"
    )
    snr = repos.requirements.get("SNR-001", loaded.id)
    assert (
        snr.change_state == "modified"
        and "2초" in snr.description
        and snr.previous_versions[0]["description"].startswith("주요 화면은 3초")
        and snr.version == 2
    )
    fax = repos.requirements.get("SFR-005", loaded.id)
    assert (
        fax.change_state == "deleted" and fax.status != RequirementStatus.DELETED
    )  # 물리 삭제 없음
    traces = repos.traces.find(loaded.id, requirement_id="SFR-005")
    assert traces and all(t.review_status == "재검토 필요" for t in traces)
    held = next(w for w in repos.wbs.tree(loaded.id) if w.code == "2.2.2")
    assert held.status == TaskStatus.ON_HOLD
    assert any(
        w.code == "2.1.4" and w.related_requirement_ids == ["SFR-006"]
        for w in repos.wbs.tree(loaded.id)
    )
    assert (
        res.revision.approval_id == cr.approval_id
        and res.revision.change_ids == [cr.id]
        and res.revision.parent_revision_id == loaded.baseline_revision_id
    )
    assert cr.before_files and len(cr.after_files) == 4 and not res.qa_errors
    names = [Path(f).name for f in cr.after_files]
    assert (
        names[0].startswith("Requirements_Definition_")
        and names[0].endswith(".xlsx")
        and len(names[0]) == len("Requirements_Definition_YYYYMMDDHHMM.xlsx")
    )
    assert all(_out(services, f).exists() for f in cr.after_files)
    wb = load_workbook(_out(services, cr.after_files[0]))
    assert "변경관리대장" in wb.sheetnames
    ws = wb["기능요구사항"]
    header = [c.value for c in ws[1]]
    assert "추가 요청자" in header and "추가일" in header and "변경 상태" in header
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    assert (
        rows["SFR-006"][header.index("추가 요청자")] == "홍길동"
        and rows["SFR-005"][header.index("변경 상태")] == "삭제"
    )
    tm = load_workbook(_out(services, cr.after_files[1]))["추적표"]
    th = [c.value for c in tm[1]]
    trows = {r[0]: r for r in tm.iter_rows(min_row=2, values_only=True) if r[0]}
    assert (
        trows["SFR-005"][th.index("검토 상태")] == "재검토 필요"
        and trows["SFR-005"][th.index("화면 ID")] == "SCR-005"
    )  # 링크 유지
    # 같은 분(minute)에 두 번째 반영 → _1 접미사
    cr2 = services.changes.create_request(
        loaded,
        input_kind=InputKind.MODIFY_BUTTON,
        request_text="로그인 5회 실패 시 10분간 잠금",
        requester="김",
        reason="보안",
        selected_ids=["SFR-001"],
    )
    cr2 = services.changes.analyze(loaded, cr2, use_agent=False)
    cr2 = services.changes.decide_candidate(
        loaded, cr2, cr2.candidates[0].id, decision="accept", overrides={"effort_md": 3}
    )
    res2 = services.changes.approve(loaded, cr2, approver="PM", comments="ok")
    assert res2.revision.number == "v3.0" and res2.revision.parent_revision_id == res.revision.id
    assert (
        all(Path(f).stem.endswith("_1") for f in res2.request.after_files)
        or Path(res2.request.after_files[0]).stem != Path(cr.after_files[0]).stem
    )
    # 반려: 아무것도 바뀌지 않는다
    cr3 = services.changes.create_request(
        loaded,
        input_kind=InputKind.DELETE_BUTTON,
        request_text="",
        requester="김",
        reason="r",
        selected_ids=["SFR-004"],
    )
    cr3 = services.changes.analyze(loaded, cr3, use_agent=False)
    with pytest.raises(ValidationError, match="반려 사유"):
        services.changes.reject(loaded, cr3, approver="PM", comments="")
    cr3 = services.changes.reject(loaded, cr3, approver="PM", comments="범위 유지")
    assert (
        cr3.status == ChangeStatus.REJECTED
        and repos.requirements.get("SFR-004", loaded.id).change_state == "none"
    )
    assert len(repos.revisions.find(loaded.id)) == 3
    log = services.changes.change_log_rows(loaded)
    assert any(r[0] == "CR-001" and r[1] == "ADD" and r[9] == "applied" for r in log) and any(
        r[0] == "CR-003" and r[9] == "rejected" for r in log
    )


# 13. 반영 중 오류 → 트랜잭션 롤백, 상태 error, 재시도 가능 ------------------------------------------
def test_s13_apply_error_rolls_back(
    services: Services, loaded: Project, monkeypatch: pytest.MonkeyPatch
) -> None:
    repos = services.ctx.repos
    cr = services.changes.create_request(
        loaded,
        input_kind=InputKind.NATURAL_LANGUAGE,
        request_text="모바일 앱에서 푸시 알림을 수신하는 기능이 필요하다.",
        requester="홍",
        reason="r",
    )
    cr = services.changes.analyze(loaded, cr, use_agent=False)
    cr = services.changes.decide_candidate(loaded, cr, "C01", decision="accept")

    def boom(*args, **kwargs):  # type: ignore[no-untyped-def]
        raise RuntimeError("디스크 오류 시뮬레이션")

    monkeypatch.setattr(services.changes, "_record_revision", boom)
    with pytest.raises(RuntimeError):
        services.changes.approve(loaded, cr, approver="PM", comments="x")
    cr = services.changes.get(loaded, cr.id)
    assert cr.status == ChangeStatus.ERROR and cr.error and cr.approval_id
    assert repos.requirements.count(loaded.id) == 6 and not any(
        w.code == "2.1.4" for w in repos.wbs.tree(loaded.id)
    )
    assert len(repos.revisions.find(loaded.id)) == 1
    monkeypatch.undo()
    res = services.changes.apply(loaded, cr, actor="PM")
    assert res.request.status == ChangeStatus.APPLIED and repos.requirements.count(loaded.id) == 7


# 14. 투입계획 유무에 따른 문서 갱신: 없으면 투입계획을 만들지 않고 안내, 있으면 Staffing_Plan 생성 -----
def test_s14_staffing_plan_optional_regeneration(services: Services, loaded: Project) -> None:
    p0 = services.projects.create(name="투입계획 없음", mode=ProjectMode.BUILD)
    _upload_all(services, p0, staffing=False)
    services.baseline.load(p0, actor="t")
    p = services.ctx.repos.projects.get(p0.id)
    cr = services.changes.create_request(
        p,
        input_kind=InputKind.ADD_BUTTON,
        request_text="모바일 푸시 알림 수신",
        requester="홍",
        reason="r",
    )
    cr = services.changes.analyze(p, cr, use_agent=False)
    cr = services.changes.decide_candidate(
        p, cr, cr.candidates[0].id, decision="accept", overrides={"effort_md": 4}
    )
    res = services.changes.approve(p, cr, approver="PM", comments="ok")
    kinds = [d.kind.value for d in res.deliverables]
    assert "staffing_plan_xlsx" not in kinds and len(kinds) == 3
    assert any("투입계획" in u for u in res.request.unresolved)
    # 투입계획이 있는 프로젝트: 갱신 포함 + 월별 투입 재계산
    cr2 = services.changes.create_request(
        loaded,
        input_kind=InputKind.ADD_BUTTON,
        request_text="모바일 푸시 알림 수신",
        requester="홍",
        reason="r",
    )
    cr2 = services.changes.analyze(loaded, cr2, use_agent=False)
    cr2 = services.changes.decide_candidate(
        loaded, cr2, cr2.candidates[0].id, decision="accept", overrides={"effort_md": 4}
    )
    res2 = services.changes.approve(loaded, cr2, approver="PM", comments="ok")
    assert "staffing_plan_xlsx" in [d.kind.value for d in res2.deliverables]
    assert services.ctx.repos.allocations.find(loaded.id)
    wb = load_workbook(
        _out(
            services,
            next(d.file_path for d in res2.deliverables if d.kind.value == "staffing_plan_xlsx"),
        )
    )
    assert "투입계획" in wb.sheetnames and wb["투입계획"].max_row >= 6
    # 선택적으로 갱신 끄기
    cr3 = services.changes.create_request(
        loaded,
        input_kind=InputKind.ADD_BUTTON,
        request_text="관리자 공지사항 게시",
        requester="홍",
        reason="r",
    )
    cr3 = services.changes.analyze(loaded, cr3, use_agent=False)
    cr3 = services.changes.decide_candidate(
        loaded, cr3, cr3.candidates[0].id, decision="accept", overrides={"effort_md": 2}
    )
    res3 = services.changes.approve(
        loaded, cr3, approver="PM", comments="ok", update_staffing=False
    )
    assert "staffing_plan_xlsx" not in [d.kind.value for d in res3.deliverables]


def test_impact_without_wbs_reports_error() -> None:
    imp = analyze_impact(
        [],
        requirements=[],
        wbs=[],
        dependencies=[],
        roles=[],
        staffing_plan_present=False,
        project_start=None,
        project_end=None,
        calendar=BusinessCalendar(),
    )
    assert imp.error
