"""통합 개발 지시문 단계 3 — 제안 전략·상세 기술 검토.

전략 항목 보강·STR 저장(§7.1), 요구사항 분석 14항목(§4.1), 대안 비교·비기능 설계(§4.2·§8), 출처 확인 날짜,
덱 반영(전략 표·대안 비교·비기능 설계), QA-TECH-08 설계 재검증(§4.4)."""

from __future__ import annotations

import io
from datetime import date
from types import SimpleNamespace

from pptx import Presentation

from proposal_agent.agents.schemas import (
    ANALYSIS_FIELDS,
    AlternativeComparison,
    AlternativeOptionDetail,
    ApproachItem,
    DesignReport,
    RequirementAnalysis,
    StrategyBrief,
    StrategyItem,
)
from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.pptx.decks import build_submission_deck
from proposal_agent.generators.pptx.outline import decide_outline
from proposal_agent.generators.pptx.req_groups import (
    design_comparison_blocks,
    design_nonfunctional_rows,
    tech_choice_rows,
)
from proposal_agent.models import (
    AcceptanceStatus,
    ArchitectureComponent,
    Difficulty,
    Project,
    ProjectMode,
    ProposalStrategy,
    Requirement,
    RequirementCategory,
    RequirementType,
)
from proposal_agent.services import Services
from proposal_agent.validation.qa_checks.tech import _design_reverification_issues


def _project() -> Project:
    return Project(
        name="전략",
        mode=ProjectMode.PROPOSAL,
        output_dir="Project_s3",
        workspace_dir="ws_s3",
        start_date=date(2026, 11, 2),
        end_date=date(2027, 4, 30),
        created_by="t",
    )


def _req(p: Project, rid: str, name: str, *, analysis: dict | None = None) -> Requirement:
    return Requirement(
        id=rid,
        project_id=p.id,
        name=name,
        description=name,
        type=RequirementType.FUNCTIONAL,
        category=RequirementCategory.BUSINESS,
        rfp_requirement_id="SFR-001",
        created_by="t",
        acceptance_status=AcceptanceStatus.ACCEPTED,
        implementation_approach="- 접수 폼과 상태 전이를 구현",
        difficulty=Difficulty("medium"),
        verification_method="시연",
        acceptance_criteria="상태 전이가 규칙대로 동작",
        analysis=analysis or {},
    )


def _design() -> dict:
    return {
        "title": "민원 접수·배정 워크플로",
        "requirement_ids": ["SFR-001"],
        "key_message": "상태 전이 한 흐름으로 처리한다",
        "overview": ["접수·배정을 상태 기계로 묶는다"],
        "structure": [
            {"name": "접수 화면", "layer": "client", "technology": "HTML5", "role": "입력"},
            {"name": "워크플로 서비스", "layer": "app", "technology": "Spring 5.3", "role": "전이"},
        ],
        "flows": [
            {"title": "접수", "actor": "민원인", "action": "입력"},
            {
                "title": "배정",
                "actor": "서비스",
                "action": "규칙 조회",
                "decision": "규칙 없음 → 수동",
            },
            {"title": "종결", "actor": "담당자", "action": "처리"},
        ],
        "tech_choices": [
            {
                "area": "배치",
                "choice": "ShedLock DB 락",
                "reason": "중복 실행 방지",
                "source": "https://example.org/shedlock",
                "verified": True,
                "checked_on": "2026-10-16",
            },
            {
                "area": "업무 처리",
                "choice": "상태 전이 규칙 테이블",
                "reason": "표만 수정",
                "source": "설계 제안",
                "verified": False,
            },
        ],
        "data_interfaces": ["CASE(상태 코드 5종)"],
        "comparisons": [
            {
                "topic": "배치 중복 실행 방지",
                "options": [
                    AlternativeOptionDetail(
                        name="ShedLock",
                        pros=["코드 변경 최소"],
                        cons=["이력 없음"],
                        cost_complexity="낮음",
                        fit_conditions="주기 작업 몇 개",
                    ).model_dump(),
                    AlternativeOptionDetail(
                        name="Quartz 클러스터",
                        pros=["정확히 1회"],
                        cons=["설정 큼"],
                        cost_complexity="중",
                    ).model_dump(),
                ],
                "recommended": "ShedLock",
                "reason": "작업이 적다",
                "poc_needed": True,
                "poc_scope": "WAS 2대에서 중복 발송 0건",
                "source": "설계 제안",
            }
        ],
        "performance_scalability": ["목록 조회 인덱스"],
        "failure_recovery": ["재시도 2회 후 수동 큐"],
        "operations_monitoring": ["실패율 로그"],
        "auth_security": ["소속 기반 조회 제한"],
        "test_approach": ["전이 규칙 전수 시험"],
    }


def test_schemas_strategy_analysis_comparison() -> None:
    s = StrategyItem(
        title="t",
        description="d",
        differentiator="x",
        client_problem="문제",
        approach="방법",
        technologies=["Spring"],
        expected_effects=["효과(RFP 근거)"],
        risks=["위험"],
        unverified_claims=["미확인"],
    )
    assert s.client_problem == "문제" and len(ANALYSIS_FIELDS) == 14
    a = RequirementAnalysis(purpose=["p"], impact=["난이도 중"])
    assert a.filled_count() == 2
    item = ApproachItem(requirement_id="FR-001", approach="- a", analysis=a)
    assert item.analysis is not None and item.analysis.purpose == ["p"]
    rep = DesignReport.model_validate({"designs": [_design()]})
    d = rep.designs[0]
    assert d.comparisons[0].recommended == "ShedLock" and d.comparisons[0].poc_needed
    assert d.tech_choices[0].checked_on == "2026-10-16" and d.test_approach
    AlternativeComparison(topic="x", options=[], recommended="")


def test_strategy_persisted_as_str_entities(services: Services, proposal_project: Project) -> None:
    brief = StrategyBrief(
        rfp_summary="요약",
        strategies=[
            StrategyItem(
                title="상태 기반 워크플로로 누락을 없앤다",
                description="d",
                differentiator="x",
                client_problem="수기 처리 누락",
                approach="상태 기계",
                technologies=["Spring"],
                expected_effects=["누락 0건(RFP 목표)"],
                evidence=["RFP: 추진배경"],
                related_requirement_ids=["FR-001"],
                unverified_claims=["처리 기간 단축 폭"],
            ),
            StrategyItem(title="공통 표 컴포넌트", description="d2", differentiator="y"),
        ],
    )
    kept = services.strategy.persist_strategies(proposal_project, brief, actor="t")
    assert [s.code for s in kept] == ["STR-001", "STR-002"] and kept[
        0
    ].client_problem == "수기 처리 누락"
    assert kept[0].rfp_basis == ["RFP: 추진배경"] and kept[0].unverified_claims == [
        "처리 기간 단축 폭"
    ]
    # 같은 제목은 갱신(버전 ↑), 빠진 전략은 dropped
    brief2 = StrategyBrief(
        rfp_summary="요약",
        strategies=[
            StrategyItem(
                title="상태 기반 워크플로로 누락을 없앤다",
                description="d'",
                differentiator="x'",
                approach="상태 기계 v2",
            )
        ],
    )
    kept2 = services.strategy.persist_strategies(proposal_project, brief2, actor="t")
    assert (
        kept2[0].code == "STR-001" and kept2[0].version == 2 and kept2[0].approach == "상태 기계 v2"
    )
    all_rows = {s.code: s for s in services.ctx.repos.strategies.find(proposal_project.id)}
    assert all_rows["STR-002"].status == "dropped" and [
        s.code for s in services.strategy.strategies(proposal_project)
    ] == ["STR-001"]


def _deck_text(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    parts = []
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts.extend(c.text for r in sh.table.rows for c in r.cells)
    return "\n".join(parts)


def test_deck_renders_strategy_table_comparison_and_nonfunctional() -> None:
    p = _project()
    strat = ProposalStrategy(
        project_id=p.id,
        code="STR-001",
        title="상태 기반 워크플로로 누락을 없앤다",
        client_problem="수기 처리 누락",
        approach="상태 기계",
        technologies=["Spring 5.3"],
        differentiator="전이 규칙 표",
        expected_effects=["누락 0건"],
        evidence=["RFP: 추진배경"],
        unverified_claims=["처리 기간 단축 폭"],
    )
    ctx = GenerationContext(
        project=p,
        requirements=[_req(p, "FR-001", "민원 접수 화면")],
        designs=[_design()],
        strategies=[strat],
        proposal_content={
            "profile": "submission",
            "sections": [
                {"slide_group": "understanding", "title": "사업 이해", "bullets": ["배경"]}
            ],
        },
        period_start=date(2026, 11, 2),
        period_end=date(2027, 4, 30),
    )
    text = _deck_text(build_submission_deck(ctx, decide_outline(ctx, [])))
    for needle in (
        "STR-001",
        "수기 처리 누락",
        "발주처 문제·목표",
        "미확인 주장",
        "대안 비교: 배치 중복 실행 방지",
        "Quartz 클러스터",
        "추천: ShedLock",
        "PoC 필요",
        "비기능 설계",
        "재시도 2회 후 수동 큐",
        "(확인 2026-10-16)",
    ):
        assert needle in text, needle
    rows = tech_choice_rows(_design())
    assert rows[0][3] == "example.org (확인 2026-10-16)"
    assert design_nonfunctional_rows(_design())[0][0] == "성능·확장"
    blocks = design_comparison_blocks(_design())
    assert blocks[0]["kind"] == "table" and blocks[0]["headers"] == [
        "항목",
        "ShedLock",
        "Quartz 클러스터",
    ]
    assert blocks[1]["kind"] == "bullets" and blocks[1]["lines"][1].startswith("PoC 필요")


def test_qa_tech_08_design_reverification() -> None:
    p = _project()
    comps = [
        ArchitectureComponent(
            project_id=p.id,
            id="APP01",
            layer="app",
            name="워크플로",
            technology="Spring 5.3 / ShedLock",
        )
    ]
    research = [SimpleNamespace(confirmed_facts=[], recommended_stack=[], sources=[], summary="")]
    reqs = [_req(p, "FR-001", "접수", analysis={"purpose": ["p"]}), _req(p, "FR-002", "배정")]
    fails, revs, ids = _design_reverification_issues([_design()], reqs, comps, research)
    assert fails == []  # 추천안이 대안 중 하나
    assert (
        any("PoC 필요" in x for x in revs)
        and any("14항목 기술 검토(analysis) 없음" in x and "FR-002" in x for x in revs)
        and ids == ["FR-002"]
    )
    assert not any(
        "ShedLock DB 락" in x and "없음" in x for x in revs
    )  # 구성 요소 기술에 ShedLock 이 있음
    bad = _design()
    bad["comparisons"][0]["recommended"] = "없는 대안"
    bad["tech_choices"].append(
        {
            "area": "DB",
            "choice": "MongoDB 7",
            "reason": "r",
            "source": "https://x",
            "verified": False,
        }
    )
    fails2, revs2, _ = _design_reverification_issues([bad], reqs, comps, research)
    assert fails2 and "없는 대안" in fails2[0]
    assert any("MongoDB 7" in x for x in revs2)
