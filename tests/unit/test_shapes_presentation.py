"""Phase 13: 편집 가능 도형 도식, 발표용 덱 구조, 에이전트 visual 반영, 편집 가능성·구현 근거 QA, 조사 선행."""

from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.util import Cm

from fixtures import builders as bld
from proposal_agent.agents import AgentRequest, FakeAgentRunner
from proposal_agent.config import ProjectPaths, Settings, ValidationError
from proposal_agent.generators.pptx import shapes
from proposal_agent.generators.pptx.template import Box
from proposal_agent.models import (
    AcceptanceStatus,
    DeliverableKind,
    DocumentKind,
    Project,
    ProjectMode,
    QaResult,
)
from proposal_agent.orchestrator import StepRunner
from proposal_agent.services import AppContext, Services
from proposal_agent.validation import QaOutcome, check_pptx, pptx_texts
from unit.test_proposal_template import _researcher
from unit.test_requirements_phase2 import extractor_from_text, verifier_all_ok

STYLE = shapes.ShapeStyle()
BOX = Box(left=1.0, top=2.0, width=28.0, height=14.0)


def _blank() -> tuple[Any, Any]:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _kinds(slide: Any) -> dict[str, int]:
    out: dict[str, int] = {}
    for sh in slide.shapes:
        k = str(sh.shape_type).split(".")[-1].split(" ")[0] if sh.shape_type is not None else "?"
        if getattr(sh, "has_table", False) and sh.has_table:
            k = "TABLE"
        out[k] = out.get(k, 0) + 1
    return out


def _texts(slide: Any) -> str:
    return "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


# --- 도식 -----------------------------------------------------------------------------------
def test_each_visual_kind_draws_editable_shapes_without_pictures() -> None:
    cases: list[tuple[str, dict[str, Any]]] = [
        (
            "cards",
            {"kind": "cards", "items": [{"title": f"목표 {i}", "text": "설명"} for i in range(3)]},
        ),
        (
            "flow",
            {"kind": "flow", "items": [{"title": f"단계 {i}", "text": "작업"} for i in range(4)]},
        ),
        (
            "two_column",
            {
                "kind": "two_column",
                "left_title": "현황 및 이슈",
                "right_title": "방안",
                "items": [
                    {"title": "수작업", "group": "left"},
                    {"title": "자동화", "group": "right"},
                ],
            },
        ),
        (
            "dataflow",
            {
                "kind": "dataflow",
                "items": [
                    {"id": "ui", "title": "화면", "group": "클라이언트"},
                    {"id": "api", "title": "API", "group": "서버"},
                    {"id": "db", "title": "DB", "group": "데이터"},
                ],
                "edges": [
                    {"source": "ui", "target": "api", "label": "HTTPS"},
                    {"source": "api", "target": "db", "label": "JDBC"},
                ],
            },
        ),
        (
            "screen",
            {
                "kind": "screen",
                "title": "사건 접수 화면",
                "items": [
                    {"title": "메뉴", "group": "nav"},
                    {"title": "접수 목록", "text": "그리드"},
                    {"title": "상세", "text": "폼"},
                ],
            },
        ),
        (
            "kpis",
            {
                "kind": "kpis",
                "items": [{"title": "기간", "text": "8개월"}, {"title": "인력", "text": "5명"}],
            },
        ),
        (
            "compare",
            {
                "kind": "compare",
                "headers": ["구분", "AS-IS", "TO-BE"],
                "rows": [["접수", "수기", "시스템"]],
            },
        ),
    ]
    for kind, spec in cases:
        _prs, slide = _blank()
        drawn = shapes.draw_visual(slide, BOX, spec, style=STYLE)
        assert drawn == kind
        kinds = _kinds(slide)
        assert "PICTURE" not in kinds, kind
        assert sum(kinds.values()) >= (1 if kind == "compare" else 2), (kind, kinds)
        text = _texts(slide)
        if kind == "dataflow":
            assert kinds.get("LINE", 0) + kinds.get("CONNECTOR", 0) >= 2 and "HTTPS" in text
        if kind == "compare":
            tbl = next(
                sh for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table
            )
            cells = " ".join(c.text for row in tbl.table.rows for c in row.cells)
            assert kinds.get("TABLE") == 1 and "TO-BE" in cells
        if kind == "screen":
            assert "사건 접수 화면" in text and "접수 목록" in text
        if kind == "kpis":
            assert "8개월" in text
    # 모든 도형은 영역 안
    _prs, slide = _blank()
    shapes.draw_cards(slide, BOX, [{"title": "a", "text": "b"}] * 8, style=STYLE)
    for sh in slide.shapes:
        assert sh.left >= Cm(BOX.left) - 10 and sh.top >= Cm(BOX.top) - 10
        assert sh.left + sh.width <= Cm(BOX.left + BOX.width) + 10


def test_timeline_bars_and_architecture_connectors() -> None:
    _prs, slide = _blank()
    shapes.draw_timeline(
        slide,
        BOX,
        [
            ("1 분석", date(2026, 11, 2), date(2026, 12, 15)),
            ("2 개발", date(2026, 12, 16), date(2027, 4, 30)),
        ],
        date(2026, 11, 2),
        date(2027, 6, 30),
        style=STYLE,
    )
    kinds = _kinds(slide)
    assert kinds.get("AUTO_SHAPE", 0) >= 2 and "26.11" in _texts(slide) and "PICTURE" not in kinds

    class Conn:
        def __init__(self, to: str, proto: str) -> None:
            self.to_component_id, self.protocol = to, proto

    class Comp:
        def __init__(self, cid: str, layer: str, name: str, conns: list[Conn]) -> None:
            self.id, self.layer, self.name, self.technology, self.connections = (
                cid,
                layer,
                name,
                "Spring Boot",
                conns,
            )

    comps = [
        Comp("WEB01", "web", "웹 서버", [Conn("APP01", "HTTP")]),
        Comp("APP01", "app", "업무 서버", [Conn("DB01", "JDBC")]),
        Comp("DB01", "data", "DBMS", []),
    ]
    _prs, slide = _blank()
    shapes.draw_architecture(slide, BOX, comps, style=STYLE)
    kinds = _kinds(slide)
    assert kinds.get("LINE", 0) + kinds.get("CONNECTOR", 0) == 2
    assert "애플리케이션" in _texts(slide) and "JDBC" in _texts(slide)


# --- 덱 · QA ----------------------------------------------------------------------------------
WRITER_VISUAL_SECTIONS: list[dict[str, Any]] = [
    {
        "slide_group": "understanding",
        "title": "제안 배경",
        "chapter": "Ⅰ. 사업의 이해",
        "number": "1.1",
        "lead_message": "리드: 민원 처리 자동화로 인적 오류 제거",
        "bullets": ["수작업 접수의 누락", "회원사 일원화"],
        "visual": {
            "kind": "two_column",
            "left_title": "현황 및 이슈",
            "right_title": "시스템 구성 방안",
            "items": [
                {"title": "민원 접수 수작업", "group": "left"},
                {"title": "접수 일원화 시스템", "group": "right"},
            ],
        },
        "feasibility_basis": ["사건 접수·배정 워크플로 구현"],
    },
    {
        "slide_group": "architecture",
        "title": "호출 절차",
        "chapter": "Ⅱ. 사업 수행 방안",
        "number": "2.1",
        "visual": {
            "kind": "process",
            "items": [
                {"id": "a", "title": "접수"},
                {"id": "b", "title": "배정"},
                {"id": "c", "title": "종결"},
            ],
        },
    },
]


def _writer_visual(request: AgentRequest) -> dict[str, Any]:
    return {
        "profile": "presentation",
        "sections": WRITER_VISUAL_SECTIONS,
        "key_messages": ["편집 가능 도형"],
        "unverified_claims": [],
    }


@pytest.fixture
def proj(tmp_path: Path) -> Iterator[tuple[Services, Project, FakeAgentRunner]]:
    ctx = AppContext.create(Settings(), paths=ProjectPaths(root=tmp_path), configure_logging=False)
    runner = FakeAgentRunner(
        {
            "requirement_extractor": extractor_from_text,
            "requirement_verifier": verifier_all_ok,
            "tech_researcher": _researcher,
            "proposal_writer": _writer_visual,
        }
    )
    s = Services.from_context(ctx, runner=runner)
    p = s.projects.create(
        name="발표 도식",
        mode=ProjectMode.PROPOSAL,
        start_date=date(2026, 11, 2),
        end_date=date(2027, 6, 30),
        client_name="OO기관",
    )
    s.documents.register_upload(
        p, kind=DocumentKind.RFP, file_name="rfp.docx", data=bld.build_docx()
    )
    r = StepRunner(s)
    r.run_extraction(p)
    r.run_requirement_extraction(p)
    r.run_verification(p)
    yield s, p, runner
    ctx.close()


def test_presentation_deck_structure_is_editable_and_uses_agent_visuals(
    proj: tuple[Services, Project, FakeAgentRunner],
) -> None:
    s, p, _ = proj
    r = StepRunner(s)
    s.strategy.evaluate(p)  # Phase 18: 목차·전략(Gate P) 승인 후에만 문안 작성
    s.gates.approve_human_review(p, "gateP", approver="pm", comments="목차 확인")
    with pytest.raises(ValidationError, match="기술 조사 기록이 없어"):
        r.run_proposal_writing(p)
    r.run_tech_research(p)
    r.run_proposal_writing(p)
    r.run_generation(p)
    latest = s.generation.latest_deliverables(p)
    pres_path = s.ctx.paths.resolve_output(latest[DeliverableKind.PRESENTATION_PPTX].file_path)
    prs = Presentation(str(pres_path))
    text = "\n".join(pptx_texts(pres_path))
    assert "CONTENTS" in text and "핵심 메시지" in text and "사업의 필요성" in text
    assert "Ⅰ. 사업의 이해" in text and "Ⅱ. 사업 수행 방안" in text  # 에이전트 지정 장
    assert (
        "리드: 민원 처리 자동화로 인적 오류 제거" in text
        and "민원 접수 수작업" in text
        and "접수 일원화 시스템" in text
    )
    assert "1. 접수" in text and "3. 종결" in text  # process 도식
    for sl in prs.slides:  # 그림 없음 → 모든 요소 편집 가능
        assert not any("PICTURE" in str(sh.shape_type) for sh in sl.shapes)
    sub_path = s.ctx.paths.resolve_output(latest[DeliverableKind.SUBMISSION_PPTX].file_path)
    sub_prs = Presentation(str(sub_path))
    assert not any("PICTURE" in str(sh.shape_type) for sl in sub_prs.slides for sh in sl.shapes)
    sub_text = "\n".join(pptx_texts(sub_path))
    assert "시스템 구성도" in sub_text and "애플리케이션" in sub_text  # 구성도가 도형으로
    run = s.qa.run(p, max_fix_rounds=0).qa_run
    assert run is not None
    by_id = {c.check_id: c for c in run.checks}
    assert by_id["QA-FILE-12"].executed and by_id["QA-FILE-12"].result is QaResult.PASS, by_id[
        "QA-FILE-12"
    ].issues
    assert by_id["QA-FILE-04"].result is QaResult.PASS, by_id["QA-FILE-04"].issues
    assert by_id["QA-TECH-07"].executed


def test_qa_file_12_fails_on_embedded_picture(tmp_path: Path) -> None:
    from PIL import Image

    img = tmp_path / "x.png"
    Image.new("RGB", (40, 20), "white").save(img)
    prs = Presentation()
    for i in range(3):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        tb = sl.shapes.add_textbox(Cm(1), Cm(1), Cm(5), Cm(1))
        tb.text_frame.text = f"슬라이드 {i}"
        if i == 1:
            sl.shapes.add_picture(str(img), Cm(2), Cm(3))
    path = tmp_path / "deck.pptx"
    buf = io.BytesIO()
    prs.save(buf)
    path.write_bytes(buf.getvalue())
    out = QaOutcome()
    check_pptx(path, out, "제출용")
    pic = [f for f in out.findings if f.rule == "pptx.picture"]
    assert len(pic) == 1 and "2" in pic[0].message


def test_tech_07_flags_accepted_functional_requirements_without_research(
    proj: tuple[Services, Project, FakeAgentRunner],
) -> None:
    s, p, _ = proj
    for req in s.requirements.list_requirements(p):
        if req.id.startswith("FR-"):
            req.acceptance_status = AcceptanceStatus.ACCEPTED
            s.ctx.repos.requirements.update(req)
    StepRunner(s).run_generation(p)
    run = s.qa.run(p, max_fix_rounds=0).qa_run
    assert run is not None
    c = next(c for c in run.checks if c.check_id == "QA-TECH-07")
    assert c.result is QaResult.NEEDS_HUMAN_REVIEW
    assert any("근거 조사 없음" in i for i in c.issues) and any(
        "구현 방안 미" in i
        for i in c.issues  # 미작성 또는 미확정(규칙 초안 [자동 초안])
    )
    assert "FR-002" in c.requirement_ids
