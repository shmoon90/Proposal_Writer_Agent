"""통합 개발 지시문 단계 4 — 제안서·제출 산출물.

비용 산정(단가 × 공수, 예산 → 마진, 없으면 산정 불가), 인력투입계획 '비용산정' 시트, 추적표의 전략·API·DB·설계·코드·
테스트케이스·기술검토 열, 정의서 '기술검토(14항목)' 열, GenerationContext.design_links, 분량 모드(요약)·묶음당 장 수 상한,
배포·네트워크 구조 소절, 의사결정 표 도식."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Cm

from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.excel.requirement_spec import (
    COLUMNS as SPEC_COLUMNS,
)
from proposal_agent.generators.excel.requirement_spec import (
    SHEET_FUNCTIONAL,
    build_requirement_spec,
)
from proposal_agent.generators.excel.staffing_plan import SHEET_COST, build_staffing_plan
from proposal_agent.generators.excel.style import HEADER_ROW
from proposal_agent.generators.excel.trace_matrix import HEADERS, SHEET_MAIN, build_trace_matrix
from proposal_agent.generators.pptx.decks import build_submission_deck
from proposal_agent.generators.pptx.outline import decide_outline
from proposal_agent.generators.pptx.shapes import (
    AGENT_VISUAL_KINDS,
    deployment_items_from_components,
    draw_decision_table,
    draw_deployment,
)
from proposal_agent.generators.pptx.shapes.base import ShapeStyle
from proposal_agent.generators.pptx.template import Box
from proposal_agent.models import (
    AcceptanceStatus,
    ApiSpec,
    ArchitectureComponent,
    CodeModule,
    DbTable,
    DesignItem,
    Difficulty,
    Project,
    ProjectMode,
    ProjectSettings,
    ProposalStrategy,
    Requirement,
    RequirementCategory,
    RequirementType,
    StaffingAllocation,
    StaffRole,
)
from proposal_agent.models import TestCase as TcModel
from proposal_agent.models.design import ComponentConnection
from proposal_agent.planning.cost import estimate_cost, parse_krw
from proposal_agent.services import Services
from unit.test_stage3_strategy_tech import _deck_text, _design


def _project() -> Project:
    return Project(
        name="단계4",
        mode=ProjectMode.PROPOSAL,
        output_dir="Project_s4",
        workspace_dir="ws_s4",
        start_date=date(2026, 11, 2),
        end_date=date(2027, 4, 30),
        created_by="t",
    )


def _req(p: Project, rid: str, name: str, **kw: object) -> Requirement:
    return Requirement(
        id=rid,
        project_id=p.id,
        name=name,
        description=name,
        type=RequirementType.FUNCTIONAL,
        category=RequirementCategory.BUSINESS,
        rfp_requirement_id="SFR-001",
        created_by="t",
        acceptance_status=AcceptanceStatus.ACCEPTED,
        implementation_approach="- 구현",
        difficulty=Difficulty("medium"),
        verification_method="시연",
        acceptance_criteria="동작",
        **kw,  # type: ignore[arg-type]
    )


def _roles_allocs(p: Project) -> tuple[list[StaffRole], list[StaffingAllocation]]:
    pm = StaffRole(project_id=p.id, role_name="PM", unit_cost_per_mm=Decimal("12000000"))
    dev = StaffRole(project_id=p.id, role_name="개발자")
    allocs = [
        StaffingAllocation(
            project_id=p.id, wbs_item_id="w1", staff_role_id=pm.id, effort_md=Decimal("20")
        ),
        StaffingAllocation(
            project_id=p.id, wbs_item_id="w1", staff_role_id=dev.id, effort_md=Decimal("40")
        ),
    ]
    return [pm, dev], allocs


# --- ③ 비용 산정 ------------------------------------------------------------------------------------------


def test_parse_krw_variants() -> None:
    assert parse_krw("총 사업비 450,000,000원(부가세 포함)") == Decimal("450000000")
    assert parse_krw("약 4.5억") == Decimal("450000000")
    assert parse_krw("3천만원") == Decimal("30000000")
    assert parse_krw("12개월, 3명") is None  # 단위 없는 작은 수는 금액 아님


def test_estimate_cost_requires_rates_and_budget() -> None:
    p = _project()
    roles, allocs = _roles_allocs(p)
    # 개발자 단가 없음 → 산정 불가, 이유 기록
    est = estimate_cost(ProjectSettings(), roles, allocs)
    assert est.status == "unavailable" and est.missing_rates == ["개발자"]
    assert est.total_cost is None and any("단가 없는 역할" in a for a in est.assumptions)
    # 설정 단가(역할명) + RFP 예산 문구 → 원가·마진
    st = ProjectSettings(cost_rates={"개발자": Decimal("8000000")}, overhead_ratio=Decimal("0.1"))
    est2 = estimate_cost(st, roles, allocs, budget_hints=["사업 예산은 1억원 이내"])
    assert est2.status == "estimated" and est2.revenue == Decimal("100000000")
    assert est2.revenue_source.startswith("RFP 예산 문구")
    # 20MD→1MM×1200만 + 40MD→2MM×800만 = 2800만, 간접비 10% → 3080만
    assert est2.direct_cost == Decimal("28000000") and est2.total_cost == Decimal("30800000")
    assert est2.margin == Decimal("69200000") and est2.margin_ratio == Decimal("0.692")
    d = est2.to_dict()
    assert d["status"] == "estimated" and d["cost"] == "30800000" and len(d["lines"]) == 2
    # 설정 예산이 RFP 문구보다 우선, 예산 없으면 마진 없음
    est3 = estimate_cost(
        st.model_copy(update={"budget_amount": Decimal("50000000")}), roles, allocs
    )
    assert est3.revenue_source == "설정 예산" and est3.margin == Decimal("19200000")
    est4 = estimate_cost(st, roles, allocs)
    assert est4.status == "unavailable" and est4.total_cost == Decimal("30800000")
    assert any("예산 정보 없음" in a for a in est4.assumptions)


def test_staffing_plan_cost_sheet() -> None:
    p = _project()
    roles, allocs = _roles_allocs(p)
    st = ProjectSettings(
        cost_rates={"개발자": Decimal("8000000")}, budget_amount=Decimal("50000000")
    )
    ctx = GenerationContext(
        project=p,
        requirements=[],
        roles=roles,
        allocations=allocs,
        settings=st,
        period_start=date(2026, 11, 2),
        period_end=date(2027, 4, 30),
    )
    wb = load_workbook(io.BytesIO(build_staffing_plan(ctx)))
    ws = wb[SHEET_COST]
    rows = {str(r[0]): r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    assert rows["PM"][1] == 1.0 and rows["PM"][2] == 12000000.0 and rows["PM"][3] == 12000000.0
    assert rows["원가 합계"][3] == 28000000.0 and rows["매출(예산)"][3] == 50000000.0
    assert rows["마진"][3] == 22000000.0 and "마진율 44.0%" in str(rows["마진"][4])
    # 단가 없으면 '산정 불가' 행만
    ctx2 = GenerationContext(
        project=p,
        requirements=[],
        roles=roles,
        allocations=allocs,
        settings=ProjectSettings(),
        period_start=date(2026, 11, 2),
        period_end=date(2027, 4, 30),
    )
    ws2 = load_workbook(io.BytesIO(build_staffing_plan(ctx2)))[SHEET_COST]
    labels = [str(r[0]) for r in ws2.iter_rows(min_row=2, values_only=True) if r[0]]
    assert "산정 불가" in labels and "원가 합계" not in labels


# --- ④ 추적표·정의서 열 -------------------------------------------------------------------------------------


def test_trace_matrix_design_link_columns() -> None:
    p = _project()
    r1 = _req(p, "FR-001", "접수", analysis={"purpose": ["p"], "business_rules": ["r"]})
    r2 = _req(p, "FR-002", "배정")
    ctx = GenerationContext(
        project=p,
        requirements=[r1, r2],
        design_links={
            "FR-001": {
                "strategy": ["STR-001"],
                "api": ["API-001", "API-002"],
                "db_table": ["TBL-001"],
                "design_item": ["DSG-001"],
                "code_module": ["MOD-001"],
                "test_case": ["TC-001"],
            }
        },
    )
    for h in (
        "전략 ID",
        "API ID",
        "DB 테이블 ID",
        "설계 항목 ID",
        "기술검토(14항목)",
        "코드 모듈 ID",
        "테스트케이스 ID",
    ):
        assert h in HEADERS, h
    wb = load_workbook(io.BytesIO(build_trace_matrix(ctx, [], [], [], [])))
    ws = wb[SHEET_MAIN]
    head = [c.value for c in ws[HEADER_ROW]]
    idc = head.index("요구사항 ID")
    body = {r[idc]: r for r in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True) if r[idc]}
    row1, row2 = body["FR-001"], body["FR-002"]
    assert row1[head.index("API ID")] == "API-001, API-002"
    assert (
        row1[head.index("전략 ID")] == "STR-001" and row1[head.index("테스트케이스 ID")] == "TC-001"
    )
    assert row1[head.index("기술검토(14항목)")] == "2/14"
    assert row2[head.index("API ID")] == "미정" and row2[head.index("기술검토(14항목)")] in (
        "",
        None,
    )


def test_requirement_spec_analysis_column() -> None:
    p = _project()
    r = _req(
        p,
        "FR-001",
        "접수",
        analysis={
            "purpose": ["민원 접수 자동화"],
            "api_integration": ["연계 2건"],
            "impact": ["중"],
        },
    )
    assert "기술검토(14항목)" in [c.header for c in SPEC_COLUMNS]
    wb = load_workbook(
        io.BytesIO(build_requirement_spec(GenerationContext(project=p, requirements=[r])))
    )
    ws = wb[SHEET_FUNCTIONAL]
    hdr = [c.value for c in ws[HEADER_ROW]]
    col, idc = hdr.index("기술검토(14항목)"), hdr.index("요구사항 ID")
    cell = next(
        row[col]
        for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True)
        if row[idc] == "FR-001"
    )
    assert str(cell).startswith("3/14 항목") and "목적: 민원 접수 자동화" in str(cell)


def test_build_context_design_links_from_entities(
    services: Services, proposal_project: Project
) -> None:
    repos = services.ctx.repos
    pid = proposal_project.id
    repos.requirements.insert(_req(proposal_project, "FR-001", "접수"))
    repos.strategies.insert(
        ProposalStrategy(
            project_id=pid, code="STR-001", title="전략", related_requirement_ids=["FR-001"]
        )
    )
    repos.strategies.insert(
        ProposalStrategy(
            project_id=pid,
            code="STR-002",
            title="버림",
            status="dropped",
            related_requirement_ids=["FR-001"],
        )
    )
    repos.api_specs.insert(
        ApiSpec(project_id=pid, code="API-001", name="접수 API", related_requirement_ids=["FR-001"])
    )
    repos.db_tables.insert(
        DbTable(project_id=pid, code="TBL-001", name="CASE", related_requirement_ids=["FR-001"])
    )
    repos.design_items.insert(
        DesignItem(
            project_id=pid, code="DSG-001", name="접수 화면", related_requirement_ids=["FR-001"]
        )
    )
    repos.code_modules.insert(
        CodeModule(
            project_id=pid, code="MOD-001", name="case-api", related_requirement_ids=["FR-001"]
        )
    )
    repos.test_cases.insert(
        TcModel(project_id=pid, code="TC-001", name="접수 정상", related_requirement_ids=["FR-001"])
    )
    ctx = services.generation.build_context(proposal_project)
    links = ctx.design_links["FR-001"]
    assert links["strategy"] == ["STR-001"] and links["api"] == ["API-001"]
    assert links["db_table"] == ["TBL-001"] and links["design_item"] == ["DSG-001"]
    assert links["code_module"] == ["MOD-001"] and links["test_case"] == ["TC-001"]
    assert ctx.settings is not None and ctx.settings.proposal_detail_level == "detailed"


# --- ① 분량 모드·장 수 상한, ② 배포 구조·의사결정 표 ---------------------------------------------------------


def _deck_ctx(
    p: Project, settings: ProjectSettings, *, components: list[ArchitectureComponent] | None = None
) -> GenerationContext:
    return GenerationContext(
        project=p,
        requirements=[_req(p, "FR-001", "민원 접수 화면")],
        designs=[_design()],
        components=components or [],
        settings=settings,
        proposal_content={
            "profile": "submission",
            "sections": [
                {"slide_group": "understanding", "title": "사업 이해", "bullets": ["배경"]}
            ],
        },
        period_start=date(2026, 11, 2),
        period_end=date(2027, 4, 30),
    )


def _slide_titles(data: bytes) -> list[str]:
    prs = Presentation(io.BytesIO(data))
    out = []
    for sl in prs.slides:
        texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text]
        out.append(texts[0] if texts else "")
    return out


def test_summary_mode_and_group_slide_cap() -> None:
    p = _project()
    detailed = _deck_text(
        build_submission_deck(
            _deck_ctx(p, ProjectSettings()), decide_outline(_deck_ctx(p, ProjectSettings()), [])
        )
    )
    assert (
        "처리 절차 상세" in detailed
        and "대안 비교: 배치 중복 실행 방지" in detailed
        and "비기능 설계" in detailed
    )
    st = ProjectSettings(proposal_detail_level="summary")
    brief = _deck_text(
        build_submission_deck(_deck_ctx(p, st), decide_outline(_deck_ctx(p, st), []))
    )
    assert "처리 절차 상세" not in brief and "대안 비교: 배치 중복 실행 방지" not in brief
    assert "비기능 설계" not in brief and "구현 구조" in brief and "적용 기술" in brief
    # 장 수 상한: 상세 모드 묶음 소절이 2장 이상 → 1장으로 자르고 생략 표기
    capped_st = ProjectSettings(max_slides_per_group=1)
    capped = _deck_text(
        build_submission_deck(_deck_ctx(p, capped_st), decide_outline(_deck_ctx(p, capped_st), []))
    )
    assert "분량 상한(1장)으로" in capped and "장 생략" in capped
    assert len(
        _slide_titles(
            build_submission_deck(
                _deck_ctx(p, capped_st), decide_outline(_deck_ctx(p, capped_st), [])
            )
        )
    ) < len(
        _slide_titles(
            build_submission_deck(
                _deck_ctx(p, ProjectSettings()), decide_outline(_deck_ctx(p, ProjectSettings()), [])
            )
        )
    )


def test_deployment_diagram_section_and_items() -> None:
    p = _project()
    comps = [
        ArchitectureComponent(
            project_id=p.id, id="C1", layer="client", name="민원인 브라우저", technology="HTML5"
        ),
        ArchitectureComponent(
            project_id=p.id,
            id="W1",
            layer="web",
            name="웹 서버",
            technology="Nginx",
            connections=[ComponentConnection(to_component_id="A1", protocol="HTTP")],
        ),
        ArchitectureComponent(
            project_id=p.id, id="A1", layer="app", name="WAS", technology="Tomcat"
        ),
        ArchitectureComponent(
            project_id=p.id,
            id="D1",
            layer="data",
            name="DB",
            technology="PostgreSQL",
            extra={"zone": "데이터·인프라"},
        ),
    ]
    items, edges = deployment_items_from_components(comps)
    assert [i["group"] for i in items] == ["외부·사용자", "DMZ", "내부망", "데이터·인프라"]
    assert edges == [{"source": "W1", "target": "A1", "label": "HTTP"}]
    assert "deployment" in AGENT_VISUAL_KINDS and "decision" in AGENT_VISUAL_KINDS
    ctx = _deck_ctx(p, ProjectSettings(), components=comps)
    outline = decide_outline(ctx, [])
    assert any(s.key == "deployment_diagram" for s in outline)
    text = _deck_text(build_submission_deck(ctx, outline))
    assert "배포·네트워크 구조" in text and "영역 4개 · 구성 요소 4개 · 연결 1건" in text
    # 구성 요소 없으면 소절을 건너뛴다
    ctx0 = _deck_ctx(p, ProjectSettings())
    assert "배포·네트워크 구조" not in _deck_text(
        build_submission_deck(ctx0, decide_outline(ctx0, []))
    )


def test_draw_deployment_and_decision_table_shapes() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(29.7), Cm(21)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    style = ShapeStyle()
    draw_deployment(
        slide,
        Box(left=1, top=1, width=27, height=12),
        [
            {"id": "c", "title": "브라우저", "group": "외부·사용자"},
            {"id": "w", "title": "웹", "group": "DMZ"},
            {"id": "a", "title": "WAS", "group": "내부망"},
        ],
        [
            {"source": "c", "target": "w", "label": "HTTPS"},
            {"source": "w", "target": "a", "label": "AJP"},
        ],
        style=style,
    )
    n_before = len(slide.shapes)
    assert n_before >= 3 * 2 + 3 + 2  # 영역 머리·바탕 + 노드 + 연결선(+라벨)
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert "DMZ" in texts and "WAS" in texts and "HTTPS" in texts
    draw_decision_table(
        slide,
        Box(left=1, top=14, width=27, height=6),
        ["기준", "가중치", "A안", "B안"],
        [["비용", "3", "5", "3"], ["성능", "2", "3", "5"], ["설명", "-", "텍스트", "텍스트"]],
        style=style,
    )
    tbl = next(sh for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table)
    cells = [c.text for r in tbl.table.rows for c in r.cells]
    assert "가중 합계" in cells and "21.0" in cells and "19.0" in cells
