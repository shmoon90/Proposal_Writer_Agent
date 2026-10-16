"""Phase 15: WBS(상세·약식·내장 간트·마일스톤·개발 조직/개발자/진척율), 투입계획(구성원별 월별 공수·합계),
planner 초안의 team·milestones 반영, 고객 양식(WBS ID 없음·월별 투입계획) 재적재."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from fixtures import excel_builders as xb
from proposal_agent.agents.schemas import WBSDraft
from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.excel.staffing_plan import SHEET_PLAN, SHEET_ROLES, split_members
from proposal_agent.generators.excel.style import HEADER_ROW
from proposal_agent.generators.excel.wbs_gantt import (
    SHEET_REPORT,
    SHEET_WBS,
    build_wbs_gantt,
    gantt_first_col,
    read_week_header,
    week_starts,
)
from proposal_agent.models import DocumentKind, Project, ProjectMode, Requirement, RequirementType
from proposal_agent.planning.calendar import BusinessCalendar
from proposal_agent.services import Services

DRAFT = {
    "items": [
        {
            "code": "1",
            "level": "phase",
            "name": "분석·설계",
            "estimate_basis": "하위 합계",
            "category": "공통",
        },
        {
            "code": "1.1",
            "parent_code": "1",
            "level": "activity",
            "name": "요구사항 분석",
            "estimate_basis": "하위 합계",
            "category": "공통",
        },
        {
            "code": "1.1.1",
            "parent_code": "1.1",
            "level": "task",
            "name": "업무 분석",
            "effort_md_estimate": 10,
            "estimate_basis": "요구사항 5건 × 2MD",
            "roles": [{"role": "PL", "ratio": "1"}],
            "related_requirement_ids": ["FR-001"],
            "category": "공통",
            "organization": "수행사",
            "assignee": "홍길동",
            "quantity": "요구사항 5건",
        },
        {
            "code": "2",
            "level": "phase",
            "name": "구현",
            "estimate_basis": "하위 합계",
            "category": "민원",
        },
        {
            "code": "2.1",
            "parent_code": "2",
            "level": "activity",
            "name": "민원 기능 개발",
            "estimate_basis": "하위 합계",
            "category": "민원",
        },
        {
            "code": "2.1.1",
            "parent_code": "2.1",
            "level": "task",
            "name": "접수 화면 개발",
            "effort_md_estimate": 40,
            "estimate_basis": "화면 8개 × 5MD",
            "roles": [{"role": "백엔드", "ratio": "1"}],
            "related_requirement_ids": ["FR-001"],
            "category": "민원",
            "organization": "협력사",
            "quantity": "화면 8개",
        },
    ],
    "dependencies": [
        {"predecessor_code": "1.1.1", "successor_code": "2.1.1", "type": "FS", "lag_days": 0}
    ],
    "team": [
        {"organization": "수행사", "role": "PL", "grade": "특급", "name": "홍길동", "headcount": 1},
        {
            "organization": "협력사",
            "role": "백엔드",
            "grade": "고급",
            "headcount": 2,
            "duties": "화면·API 개발",
        },
        {"organization": "수행사", "role": "QA", "grade": "중급", "headcount": 1},
    ],
    "milestones": [
        {"name": "착수", "timing": "start"},
        {"name": "중간보고", "timing": "after:1.1.1"},
        {"name": "종료", "timing": "end"},
    ],
    "assumptions": [],
    "risks": [],
}


def _project(services: Services, name: str = "WBS 양식") -> Project:
    p = services.projects.create(
        name=name,
        mode=ProjectMode.PROPOSAL,
        start_date=date(2026, 11, 2),
        end_date=date(2027, 2, 26),
    )
    services.ctx.repos.requirements.insert(
        Requirement(id="FR-001", project_id=p.id, name="민원 접수", type=RequirementType.FUNCTIONAL)
    )
    return p


def _ctx(services: Services, p: Project) -> GenerationContext:
    wbs, deps, roles, allocs = services.planning.current_plan(p)
    return GenerationContext(
        project=p,
        requirements=list(services.ctx.repos.requirements.find(p.id)),
        wbs=wbs,
        dependencies=deps,
        roles=roles,
        allocations=allocs,
        period_start=p.start_date,
        period_end=p.end_date,
    )


def _fill(cell) -> str | None:  # type: ignore[no-untyped-def]
    return cell.fill.fgColor.rgb[-6:] if cell.fill and cell.fill.fill_type == "solid" else None


# 1. 계획 수립: team·milestones·구분/조직/개발자/수량 반영 -------------------------------------------
def test_build_plan_applies_team_milestones_and_task_fields(services: Services) -> None:
    p = _project(services)
    result = services.planning.build_plan(p, WBSDraft.model_validate(DRAFT), None)
    wbs, _, roles, _ = services.planning.current_plan(p)
    by_code = {w.code: w for w in wbs}
    t = by_code["1.1.1"]
    assert (t.category, t.organization, t.assignee, t.quantity) == (
        "공통",
        "수행사",
        "홍길동",
        "요구사항 5건",
    )
    ms = [w for w in wbs if w.milestone and w.parent_id]
    assert [m.name for m in ms] == ["착수", "중간보고", "종료"] and by_code["0"].milestone
    cal = BusinessCalendar()
    assert ms[0].planned_start == date(2026, 11, 2)
    assert ms[1].planned_start == cal.next_business_day(
        by_code["1.1.1"].planned_end + __import__("datetime").timedelta(days=1)
    )  # type: ignore[operator]
    assert ms[2].planned_start == date(2027, 2, 26) and all(m.effort_md == 0 for m in ms)
    assert not any("마일스톤" in n and "기본값" in n for n in result.review_notes)
    by_role = {r.role_name: r for r in roles}
    assert by_role["PL"].grade == "특급" and by_role["PL"].members == ["홍길동"]
    assert by_role["백엔드"].headcount_planned == 2 and by_role["백엔드"].organization == "협력사"
    assert by_role["백엔드"].duties == "화면·API 개발" and by_role["백엔드"].member_names() == [
        "백엔드1",
        "백엔드2",
    ]
    assert (
        "QA" in by_role and "조직 구성에만" in by_role["QA"].notes
    )  # 배분 없는 역할도 행으로 남긴다
    assert result.total_effort_md == Decimal("50")


# 2. WBS 파일: 상세(묶음 머리글+간트) · 약식(병합·마일스톤 라벨) ---------------------------------------
def test_wbs_workbook_detail_gantt_and_report_sheet(services: Services) -> None:
    p = _project(services, "WBS 파일")
    services.planning.build_plan(p, WBSDraft.model_validate(DRAFT), None)
    ctx = _ctx(services, p)
    wb = load_workbook(io.BytesIO(build_wbs_gantt(ctx)))
    ws = wb[SHEET_WBS]
    headers = [c.value for c in ws[HEADER_ROW]]
    groups = [c.value for c in ws[1] if c.value]
    assert groups[:7] == ["식별", "개발 항목", "담당", "일정", "공수", "관계", "비고"]
    first = gantt_first_col(ws)
    assert first == headers.index("비고") + 2
    weeks = read_week_header(ws, first)
    assert weeks == week_starts(
        min(w.planned_start for w in ctx.wbs if w.planned_start),
        max(w.planned_end for w in ctx.wbs if w.planned_end),
    )
    assert weeks[0].weekday() == 0 and str(ws.cell(row=1, column=first).value).startswith("2026-1")
    rows = {
        ws.cell(row=r, column=1).value: r
        for r in range(3, ws.max_row + 1)
        if ws.cell(row=r, column=1).value is not None
    }
    task_row = rows["1.1.1"]
    rec = dict(zip(headers, [c.value for c in ws[task_row]][: len(headers)], strict=False))
    assert rec["구분"] == "공통" and rec["개발 조직"] == "수행사" and rec["개발자"] == "홍길동"
    assert rec["수량"] == "요구사항 5건" and rec["진척율(%)"] == 0 and rec["단계"] == "작업"
    assert (
        dict(zip(headers, [c.value for c in ws[rows["2.1.1"]]], strict=False))["개발자"] == "(미정)"
    )
    filled_task = [c for c in ws[task_row][first - 1 : first - 1 + len(weeks)] if _fill(c)]
    assert filled_task and all(_fill(c) == "BDD7EE" for c in filled_task)
    assert _fill(ws.cell(row=rows["1"], column=first)) == "2E75B6"  # 단계는 진한 파랑
    ms_row = rows["0.1"]
    ms_cells = [c for c in ws[ms_row][first - 1 : first - 1 + len(weeks)] if c.value]
    assert (
        len(ms_cells) == 1
        and str(ms_cells[0].value).startswith("◆")
        and _fill(ms_cells[0]) == "F4B183"
    )
    # 약식WBS
    rp = wb[SHEET_REPORT]
    assert [c.value for c in rp[HEADER_ROW]][:5] == ["구분", "단계", "항목", "시작일", "종료일"]
    assert rp.cell(row=3, column=1).value == "사업관리" and "착수" in str(
        rp.cell(row=3, column=3).value
    )
    first_r = gantt_first_col(rp)
    assert first_r == 6
    labels = [c.value for c in rp[3][first_r - 1 : first_r - 1 + len(weeks)] if c.value]
    assert labels[0] == "착수" and labels[-1] == "종료"
    merged = {str(m) for m in rp.merged_cells.ranges}
    assert any(m.startswith("A") for m in merged)  # 구분 세로 병합
    names = [
        rp.cell(row=r, column=3).value
        for r in range(4, rp.max_row + 1)
        if rp.cell(row=r, column=3).value
    ]
    assert "업무 분석" in names and "접수 화면 개발" in names
    assert "간트" not in wb.sheetnames and {"의존관계", "산정근거"} <= set(wb.sheetnames)


# 3. 투입계획: 구성원별 월별 공수·계·합계 ----------------------------------------------------------
def test_staffing_plan_member_rows_and_totals(services: Services) -> None:
    p = _project(services, "투입계획")
    services.planning.build_plan(p, WBSDraft.model_validate(DRAFT), None)
    ctx = _ctx(services, p)
    from proposal_agent.generators.excel.staffing_plan import build_staffing_plan

    wb = load_workbook(io.BytesIO(build_staffing_plan(ctx)))
    ws = wb[SHEET_PLAN]
    headers = [c.value for c in ws[HEADER_ROW]]
    groups = [c.value for c in ws[1] if c.value]
    assert groups[0] == "수행 조직·구성원" and "월별 공수(MM)" in groups
    month_cols = [h for h in headers if isinstance(h, str) and len(h) == 7 and h[4] == "-"]
    assert month_cols and month_cols[0] == "2026-11"
    members = [
        ws.cell(row=r, column=3).value
        for r in range(3, ws.max_row + 1)
        if ws.cell(row=r, column=2).value and ws.cell(row=r, column=1).value != "합계"
    ]
    assert members == ["홍길동", "백엔드1", "백엔드2", "QA"]
    total_col = headers.index("계(MM)") + 1
    assert str(ws.cell(row=3, column=total_col).value).startswith("=SUM(")
    total_row = next(
        r for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=1).value == "합계"
    )
    assert total_row == 3 + len(members)
    m1 = headers.index(month_cols[0]) + 1
    assert (
        str(ws.cell(row=total_row, column=m1).value)
        == f"=SUM({ws.cell(row=3, column=m1).column_letter}3:{ws.cell(row=3, column=m1).column_letter}{total_row - 1})"
    )
    # 백엔드 2명 분배 합 = 역할 월별 MM
    be_rows = [r for r in range(3, total_row) if ws.cell(row=r, column=2).value == "백엔드"]
    role_sheet = wb[SHEET_ROLES]
    rh = [c.value for c in role_sheet[1]]
    be_role_row = next(
        r
        for r in range(2, role_sheet.max_row + 1)
        if role_sheet.cell(row=r, column=2).value == "백엔드"
    )
    for m in month_cols:
        col = headers.index(m) + 1
        split_sum = sum(Decimal(str(ws.cell(row=r, column=col).value or 0)) for r in be_rows)
        role_val = Decimal(str(role_sheet.cell(row=be_role_row, column=rh.index(m) + 1).value or 0))
        assert split_sum == role_val
    notes = [ws.cell(row=r, column=headers.index("비고") + 1).value for r in be_rows]
    assert all("구성원 이름 미정" in str(n) for n in notes)
    assert split_members({"2026-11": Decimal("1.00")}, 3) == [
        {"2026-11": Decimal("0.33")},
        {"2026-11": Decimal("0.33")},
        {"2026-11": Decimal("0.34")},
    ]
    assert {"공수요약", "월별투입"} <= set(wb.sheetnames)


# 4. 고객 양식 재적재: WBS ID 없는 상세 WBS(구분·항목·수량·담당 + 주 단위 열), 월별 투입계획 -------------
def _customer_wbs() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(
        [None, "구분", "항목", None, "수량", "담당", "7", None, None, None, "8", None, None, None]
    )
    ws.append([None, None, None, None, None, None, 1, 2, 3, 4, 1, 2, 3, 4])
    rows = [
        ("공시", "큰 글씨", "Mobile", 15, "김개발"),
        (None, None, "PC", 20, None),
        (None, "요약", "PC", 34, "이개발"),
        ("통계", "조회", "PC", 12, "박개발"),
    ]
    for cat, item, sub, qty, who in rows:
        ws.append([None, cat, item, sub, qty, who])
    wb.save(buf := io.BytesIO())
    return buf.getvalue()


def _customer_staffing() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["* OO 사업 월별 투입계획"])
    ws.append(["구분", "등급", "6", "7", "8", "계", "비고"])
    ws.append(["PM", "특급", 0.5, 1, 1, 2.5, ""])  # 실제 파일은 SUM 수식의 캐시 값이 읽힌다
    ws.append(["개발", "고급", 1, 1, 1, 3, "* 3MM -> 3.5MM"])
    ws.append(["개발", "중급", 0.5, 1, 1, 2.5, ""])
    ws.append(["계", "", 2, 3, 3, 8, ""])
    wb.save(buf := io.BytesIO())
    return buf.getvalue()


def test_baseline_loads_customer_wbs_without_ids_and_monthly_staffing(services: Services) -> None:
    p = services.projects.create(name="고객 WBS 적재", mode=ProjectMode.BUILD)
    services.documents.register_upload(
        p,
        kind=DocumentKind.REQUIREMENT_SPEC,
        file_name="spec.xlsx",
        data=xb.build_requirement_spec(),
    )
    services.documents.register_upload(
        p, kind=DocumentKind.TRACE_MATRIX, file_name="trace.xlsx", data=xb.build_trace_matrix()
    )
    services.documents.register_upload(
        p, kind=DocumentKind.WBS, file_name="wbs.xlsx", data=_customer_wbs()
    )
    services.documents.register_upload(
        p, kind=DocumentKind.STAFFING_PLAN, file_name="staff.xlsx", data=_customer_staffing()
    )
    profiles = services.baseline.profiles(p)
    wbs_prof = profiles[DocumentKind.WBS].profile
    assert wbs_prof.sheets and not wbs_prof.needs_confirmation
    mapped = {c.field: c.header for c in wbs_prof.sheets[0].columns if c.header}
    assert (
        mapped["category"] == "구분"
        and mapped["assignee"] == "담당"
        and mapped["quantity"] == "수량"
    )
    result = services.baseline.load(p, actor="t")
    assert any("코드(1, 1.1, 1.1.1)를 발급" in w for w in result.warnings)
    items = {w.code: w for w in services.ctx.repos.wbs.tree(p.id)}
    assert {"1", "1.1", "1.1.1", "1.1.2", "1.2", "1.2.1", "2", "2.1", "2.1.1"} <= set(items)
    assert (
        items["1"].name == "공시"
        and items["1.1"].name == "큰 글씨"
        and items["1.1.1"].name == "Mobile"
    )
    assert items["1.1.1"].assignee == "김개발" and items["1.1.1"].quantity == "15"
    assert items["1.1.2"].name == "PC" and items["1.1.2"].parent_id == items["1.1"].id
    assert items["2.1.1"].assignee == "박개발" and items["2"].name == "통계"
    roles = {r.role_name: r for r in services.ctx.repos.staff_roles.find(p.id)}
    assert set(roles) == {"PM", "개발"} and roles["PM"].grade == "특급"
    assert (
        roles["개발"].headcount_planned == 1
    )  # 구성원 이름 열이 없으면 인원은 행 수로 세지 않는다
    assert "기준 공수" in roles["개발"].notes
