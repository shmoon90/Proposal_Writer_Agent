"""Phase 19: 슬라이드별 명세(결론형 제목·목적·독자·근거·연결·레이아웃), 기술 상세·대안 비교, 용어집,
도식 13종(raci·data_model·procedure·security_layers·layers 추가), QA-CON-10(추상 표현·근거 없는 수치)·QA-CON-11(약어 미설명)."""

from __future__ import annotations

import io
from datetime import date
from types import SimpleNamespace
from typing import Any

from pptx import Presentation

from proposal_agent.agents.schemas import LAYOUT_TYPES, ProposalContent
from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.pptx import shapes
from proposal_agent.generators.pptx.decks import build_presentation_deck, build_submission_deck
from proposal_agent.generators.pptx.outline import decide_outline
from proposal_agent.generators.pptx.slide_spec import build_slide_spec
from proposal_agent.models import (
    AcceptanceStatus,
    Project,
    ProjectMode,
    Requirement,
    RequirementType,
)
from proposal_agent.validation.independent_qa import (
    _abstract_and_number_issues,
    _unexplained_acronyms,
)

BOX = shapes.Box(left=1.0, top=1.0, width=22.0, height=12.0)
STYLE = shapes.ShapeStyle()


def _blank() -> Any:
    prs = Presentation()
    return prs.slides.add_slide(prs.slide_layouts[6])


def _flat(shapes: Any) -> list[Any]:
    out: list[Any] = []
    for sh in shapes:
        if "GROUP" in str(sh.shape_type):
            out.extend(_flat(sh.shapes))
        else:
            out.append(sh)
    return out


def _kinds(slide: Any) -> dict[str, int]:
    out: dict[str, int] = {}
    for sh in _flat(slide.shapes):
        k = str(sh.shape_type).split(".")[-1].split(" ")[0]
        out[k] = out.get(k, 0) + 1
    return out


def test_new_visual_kinds_are_editable_shapes() -> None:
    cases = [
        {
            "kind": "raci",
            "headers": ["활동", "PM", "개발", "고객"],
            "rows": [["요구 확정", "A", "C", "R"], ["설계", "R", "R", "I"]],
        },
        {
            "kind": "data_model",
            "items": [
                {"id": "u", "title": "사용자", "text": "ID\n이름"},
                {"id": "c", "title": "사건", "text": "번호\n상태"},
            ],
            "edges": [{"source": "u", "target": "c", "label": "1:N"}],
        },
        {
            "kind": "procedure",
            "items": [
                {"title": "장애 감지", "group": "운영"},
                {"title": "보고", "group": "PM"},
                {"title": "복구", "group": "운영"},
            ],
        },
        {
            "kind": "security_layers",
            "items": [
                {"title": "방화벽", "group": "네트워크"},
                {"title": "계정·권한", "group": "응용", "text": "역할 기반"},
                {"title": "암호화", "group": "데이터"},
            ],
        },
        {
            "kind": "layers",
            "items": [{"title": "브라우저", "group": "사용자"}, {"title": "WAS", "group": "응용"}],
        },
    ]
    assert len(shapes.AGENT_VISUAL_KINDS) == 16  # Phase 20: chart, 단계 4: deployment·decision 추가
    for spec in cases:
        slide = _blank()
        drawn = shapes.draw_visual(slide, BOX, spec, style=STYLE)
        assert drawn == spec["kind"]
        kinds = _kinds(slide)
        assert "PICTURE" not in kinds, spec["kind"]
        assert sum(kinds.values()) >= 2, (spec["kind"], kinds)
        if spec["kind"] == "raci":
            tbl = next(
                sh for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table
            )
            cells = [c.text for row in tbl.table.rows for c in row.cells]
            assert "A" in cells and "요구 확정" in cells
        if spec["kind"] in ("data_model", "procedure"):
            assert kinds.get("LINE", 0) + kinds.get("CONNECTOR", 0) >= 1


def _project() -> Project:
    return Project(
        name="사건관리",
        mode=ProjectMode.PROPOSAL,
        output_dir="Project_x",
        workspace_dir="ws_x",
        start_date=date(2026, 11, 2),
        end_date=date(2027, 4, 30),
        created_by="t",
    )


def _req(p: Project, rid: str, name: str, desc: str) -> Requirement:
    return Requirement(
        id=rid,
        project_id=p.id,
        name=name,
        description=desc,
        type=RequirementType.FUNCTIONAL,
        created_by="t",
        acceptance_status=AcceptanceStatus.ACCEPTED,
        implementation_approach="웹 화면과 API로 구현",
    )


CONTENT: dict[str, Any] = {
    "profile": "submission",
    "sections": [
        {
            "slide_group": "architecture",
            "title": "시스템 구성",
            "chapter": "Ⅱ. 사업 수행 방안",
            "section": "1. 시스템 구성",
            "number": "1.1",
            "key_message": "접수·배정·종결을 한 화면에서 처리해 처리 기간을 줄인다",
            "lead_message": "3계층 웹 구조",
            "purpose": "목표 시스템 구조와 기술 선택 근거를 설명",
            "audience": "technical",
            "layout_type": "tech_detail",
            "bullets": ["브라우저 — WAS — DB 3계층", "SSO 연계로 로그인 1회"],
            "evidence": [
                {"kind": "requirement", "ref": "FR-001", "text": "접수 화면 요구"},
                {"kind": "research", "ref": "WAS 비교", "text": "공개 WAS 운영 사례"},
            ],
            "links_prev": ["Ⅰ.2"],
            "links_next": ["1.2"],
            "terms": ["SSO", "WAS"],
            "tech_detail": {
                "components": ["웹 서버", "WAS(공개 SW)", "관계형 DB"],
                "interfaces": ["SSO 연계(SAML)"],
                "data": ["사건·접수·처리 이력"],
                "operations": ["일 단위 백업", "장애 시 대기 서버 전환"],
                "security": ["역할 기반 권한", "전송 구간 암호화"],
                "alternatives": [
                    {
                        "name": "상용 WAS",
                        "kind": "commercial",
                        "pros": ["벤더 지원"],
                        "cons": ["라이선스 비용"],
                        "cost_note": "연 라이선스",
                        "basis": "WAS 비교",
                    },
                    {
                        "name": "공개 WAS",
                        "kind": "open_source",
                        "pros": ["비용 없음", "운영 사례 다수"],
                        "cons": ["자체 운영 역량 필요"],
                        "selected": True,
                        "basis": "WAS 비교",
                    },
                ],
                "selection_reason": "운영 사례와 비용을 고려해 공개 WAS를 선정",
            },
            "visual": {
                "kind": "layers",
                "items": [
                    {"title": "브라우저", "group": "사용자"},
                    {"title": "WAS", "group": "응용"},
                    {"title": "DB", "group": "데이터"},
                ],
            },
            "data_refs": ["FR-001"],
            "requirement_ids": ["FR-001"],
        },
        {
            "slide_group": "risk",
            "title": "위험 관리",
            "number": "3.1",
            "bullets": ["연계 지연 위험 → 사전 인터페이스 합의"],
        },
    ],
    "key_messages": ["한 화면 처리"],
    "glossary": [
        {
            "term": "SSO",
            "full_form": "Single Sign-On",
            "meaning": "한 번 로그인으로 여러 시스템을 쓰는 방식",
            "first_section": "1.1",
        },
        {
            "term": "WAS",
            "full_form": "Web Application Server",
            "meaning": "화면 요청을 처리하는 서버 프로그램",
        },
    ],
}


def test_schema_accepts_phase19_fields() -> None:
    pc = ProposalContent.model_validate(CONTENT)
    sec = pc.sections[0]
    assert sec.key_message.startswith("접수") and sec.audience == "technical"
    assert sec.tech_detail is not None and sec.tech_detail.alternatives[1].selected
    assert len(pc.glossary) == 2 and sec.layout_type in LAYOUT_TYPES


def _deck_texts(data: bytes) -> list[str]:
    prs = Presentation(io.BytesIO(data))
    out = []
    for slide in prs.slides:
        parts = []
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts.extend(c.text for r in sh.table.rows for c in r.cells)
        out.append("\n".join(parts))
    return out


def test_decks_render_key_message_tech_detail_glossary_and_slide_spec() -> None:
    p = _project()
    ctx = GenerationContext(
        project=p,
        requirements=[_req(p, "FR-001", "사건 접수", "사건을 접수한다")],
        proposal_content=CONTENT,
        period_start=date(2026, 11, 2),
        period_end=date(2027, 4, 30),
    )
    outline = decide_outline(ctx, [])
    data = build_submission_deck(ctx, outline)
    texts = _deck_texts(data)
    joined = "\n".join(texts)
    assert "접수·배정·종결을 한 화면에서 처리해 처리 기간을 줄인다" in joined  # 결론형 제목 → 리드
    assert "대안 비교" in joined and "공개 WAS" in joined and "● 선정" in joined
    assert "기술 상세" in joined and "용어 설명" in joined and "Single Sign-On" in joined
    prs = Presentation(io.BytesIO(data))
    sec_slide = next(s for s, t in zip(prs.slides, texts, strict=True) if "처리 기간을 줄인다" in t)
    notes = sec_slide.notes_slide.notes_text_frame.text
    assert "[목적]" in notes and "[독자] technical" in notes and "[연결]" in notes
    assert "[리드] 3계층 웹 구조" in notes
    spec = build_slide_spec(data, ctx, "submission")
    entries = ctx.slide_specs["submission"]
    arch = next(e for e in entries if e["key"] == "architecture")
    assert (
        arch["slide_to"] >= arch["slide_from"] + 1
    )  # 복합 장표(도식+기술 상세) + 대안 비교 (Phase 22)
    assert arch["visual_kind"] == "layers" and arch["has_tech_detail"]
    assert any(e["key"] == "glossary" for e in entries)
    first = next(s for s in spec["slides"] if s["slide_no"] == arch["slide_from"])
    assert first["key_message"].startswith("접수") and first["rendered_layout"] in (
        "visual",
        "visual_table",
    )
    assert "FR-001" in first["requirement_ids"] and first["has_notes"]
    assert spec["summary"]["sections_without_key_message"] == ["3.1"]
    assert spec["summary"]["glossary_terms"] == 2 and spec["summary"]["pictures"] == 0
    # 발표용: 대안 비교 표 + 용어 설명(요약)만, 기술 상세 카드는 없음
    pres = build_presentation_deck(ctx, outline)
    pj = "\n".join(_deck_texts(pres))
    assert "대안 비교" in pj and "용어 설명" in pj and "기술 상세" not in pj
    assert ctx.slide_specs["presentation"]


def _inp(ctx: GenerationContext) -> Any:
    return SimpleNamespace(
        all_requirements=list(ctx.requirements),
        requirements=list(ctx.requirements),
        research=[],
        blocks=[],
        ctx=ctx,
    )


def test_qa_con10_flags_abstract_words_and_unsupported_numbers() -> None:
    p = _project()
    ctx = GenerationContext(
        project=p,
        requirements=[_req(p, "NFR-001", "응답 시간", "주요 화면은 3초 이내 응답한다.")],
    )
    sections = [
        {
            "number": "2.1",
            "title": "성능",
            "key_message": "최적의 구조로 처리 시간을 30% 줄인다",
            "bullets": ["화면 응답 3초 이내", "동시 사용자 500명 처리"],
            "evidence": [{"kind": "requirement", "ref": "NFR-001", "text": "3초 이내"}],
        },
        {"number": "2.2", "title": "운영", "bullets": ["백업은 매일 4회 수행"], "evidence": []},
    ]
    issues = _abstract_and_number_issues(_inp(ctx), sections, "")
    assert any("[2.1] 추상 표현: 최적의" in x for x in issues)
    unsupported = next(x for x in issues if x.startswith("[2.1] 근거 없는 수치"))
    assert "30%" in unsupported and "500명" in unsupported and "3초" not in unsupported
    assert any(x.startswith("[2.2] 근거 없는 수치: 4회") and "(evidence 없음)" in x for x in issues)
    # 제안 문안이 없으면 PPT 본문의 추상 표현만
    only = _abstract_and_number_issues(_inp(ctx), [], "완벽한 품질 보증 체계")
    assert len(only) == 1 and "완벽한" in only[0]
    assert _abstract_and_number_issues(_inp(ctx), [], "일정 12개월") == []


def test_qa_con11_flags_unexplained_acronyms() -> None:
    content = {"glossary": [{"term": "SSO", "full_form": "Single Sign-On", "meaning": "x"}]}
    sections = [
        {
            "number": "1.1",
            "title": "구성",
            "bullets": ["SSO 연계", "JWT 토큰 인증", "RBAC(역할 기반 접근 제어)"],
        },
        {
            "number": "1.2",
            "title": "운영",
            "bullets": ["FR-001 대응", "Ⅱ장 참조", "WBS 기준 PM 보고"],
        },
        {"number": "1.3", "title": "연계", "bullets": ["전자문서 교환(EDI) 연계", "OCR 적용"]},
    ]
    issues = _unexplained_acronyms(content, sections, "")
    assert issues and "JWT(1.1)" in issues[0] and "OCR(1.3)" in issues[0]
    for ok in ("SSO", "RBAC", "EDI", "FR", "WBS", "PM"):
        assert f"{ok}(" not in issues[0] and f" {ok}," not in issues[0]
    assert _unexplained_acronyms({}, [], "요구사항 FR-001 대응, WBS 기준") == []
    ppt_only = _unexplained_acronyms({}, [], "SAML 연계 방식")
    assert ppt_only and "제출용 PPT 본문" in ppt_only[0] and "SAML" in ppt_only[0]
