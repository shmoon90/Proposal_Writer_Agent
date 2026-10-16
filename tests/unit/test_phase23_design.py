"""Phase 23: 요구사항 묶음별 구현 설계(design_writer) — 스키마, 설계 기준 묶음, 덱의 구조·절차 도식·적용 기술 표,
페이지 배분, QA-CON-13, 브리프·파이프라인 단계."""

from __future__ import annotations

import io
from datetime import date
from types import SimpleNamespace
from typing import Any

from pptx import Presentation

from proposal_agent.agents.schemas import DesignReport
from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.pptx.builder import DeckBuilder
from proposal_agent.generators.pptx.decks import (
    _paginate_blocks,
    build_presentation_deck,
    build_submission_deck,
)
from proposal_agent.generators.pptx.outline import decide_outline
from proposal_agent.generators.pptx.req_groups import (
    design_flow_spec,
    design_structure_spec,
    design_texts,
    flow_table_rows,
    group_requirements,
    tech_choice_rows,
)
from proposal_agent.models import (
    AcceptanceStatus,
    Difficulty,
    Project,
    ProjectMode,
    Requirement,
    RequirementCategory,
    RequirementType,
)
from proposal_agent.requirements.design import (
    design_min_issues,
    design_targets,
    resolve_design_ids,
)
from proposal_agent.services.pipeline_service import PROPOSAL_STEPS
from proposal_agent.validation.independent_qa import _design_depth_issues


def _project() -> Project:
    return Project(
        name="설계",
        mode=ProjectMode.PROPOSAL,
        output_dir="Project_d",
        workspace_dir="ws_d",
        start_date=date(2026, 11, 2),
        end_date=date(2027, 4, 30),
        created_by="t",
    )


def _req(
    p: Project, rid: str, name: str, approach: str, *, rfp: str | None, cat: str = "business"
) -> Requirement:
    return Requirement(
        id=rid,
        project_id=p.id,
        name=name,
        description=name,
        type=RequirementType.FUNCTIONAL,
        category=RequirementCategory(cat),
        rfp_requirement_id=rfp,
        created_by="t",
        acceptance_status=AcceptanceStatus.ACCEPTED,
        implementation_approach=approach,
        difficulty=Difficulty("medium"),
        verification_method="시연",
        acceptance_criteria="상태 전이가 규칙대로 동작",
    )


def _reqs(p: Project) -> list[Requirement]:
    return [
        _req(p, "FR-001", "민원 접수 화면", "- 접수 폼과 임시저장 상태 전이를 구현", rfp="SFR-001"),
        _req(p, "FR-002", "민원 배정", "- 담당자 배정 규칙 테이블로 자동 배정", rfp="SFR-001"),
        _req(
            p,
            "FR-003",
            "통계 조회",
            "- 집계 쿼리와 네이티브 차트로 월별 통계 표시",
            rfp="SFR-002",
            cat="report",
        ),
        _req(
            p,
            "FR-004",
            "알림 발송",
            "- 이벤트별 알림 템플릿 발송\n- 전제: 템플릿 승인은 발주처와 협의",
            rfp=None,
            cat="interface",
        ),
    ]


def _design() -> dict[str, Any]:
    return {
        "title": "민원 접수·배정·알림 워크플로",
        "requirement_ids": ["SFR-001", "FR-004"],
        "key_message": "접수부터 알림까지 상태 전이 한 흐름으로 처리한다",
        "overview": [
            "접수·배정·알림을 상태 기계로 묶어 한 트랜잭션에서 처리",
            "알림은 발송 큐로 분리해 실패 재시도",
        ],
        "structure": [
            {
                "name": "접수·배정 화면",
                "layer": "client",
                "technology": "HTML5·표준 CSS",
                "role": "접수 폼·배정 목록",
                "requirement_ids": ["FR-001"],
            },
            {
                "name": "워크플로 서비스",
                "layer": "app",
                "technology": "Spring 5.3 서비스 계층",
                "role": "상태 전이·배정 규칙 적용",
                "requirement_ids": ["FR-001", "FR-002"],
            },
            {
                "name": "알림 발송 모듈",
                "layer": "integration",
                "technology": "메시징 대행사 REST API",
                "role": "템플릿 발송·결과 콜백",
                "requirement_ids": ["FR-004"],
            },
            {
                "name": "민원·알림 이력 테이블",
                "layer": "data",
                "technology": "Oracle 11.2",
                "role": "상태 이력·발송 이력",
            },
        ],
        "flows": [
            {"title": "접수 입력", "actor": "민원인", "action": "폼 입력 → 임시저장 → 접수 요청"},
            {
                "title": "자동 배정",
                "actor": "워크플로 서비스",
                "action": "배정 규칙 테이블 조회 → 담당자 결정",
                "decision": "규칙 없음 → 수동 배정 대기",
            },
            {
                "title": "알림 발송",
                "actor": "알림 발송 모듈",
                "action": "템플릿 치환 → 발송 요청 → 결과 저장",
                "decision": "실패 → 3회 재시도 후 문자 대체",
            },
            {"title": "처리 확인", "actor": "담당자", "action": "배정 목록에서 처리 → 종결"},
        ],
        "tech_choices": [
            {
                "area": "배치",
                "choice": "ShedLock DB 락",
                "reason": "WAS 2대에서 알림 재시도 배치가 중복 실행되지 않도록",
                "alternatives": ["Quartz 클러스터 — 설정·DB 의존이 커서 배제"],
                "source": "https://github.com/lukas-krecan/ShedLock",
                "verified": True,
            },
            {
                "area": "연계",
                "choice": "메시징 REST API",
                "reason": "템플릿 사전 승인·대체 발송 지원",
                "source": "https://example.org/messaging-api",
                "verified": False,
            },
            {
                "area": "업무 처리",
                "choice": "상태 전이 규칙 테이블",
                "reason": "상태 추가 시 표만 수정",
                "source": "설계 제안",
                "verified": False,
            },
        ],
        "data_interfaces": [
            "민원(상태 코드 7종)·배정·알림 이력 테이블",
            "POST /api/notify {템플릿, 수신자, 변수}",
        ],
        "premises": ["알림 템플릿 승인 기간은 발주처와 협의"],
        "risks": ["배정 규칙 미정 → 착수 2주 내 확정"],
    }


def _deck_text(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    parts = []
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts.extend(c.text for r in sh.table.rows for c in r.cells)
    return "\n".join(parts)


def test_schema_and_min_requirements() -> None:
    rep = DesignReport.model_validate({"designs": [_design()], "research_needed": [], "notes": ""})
    d = rep.designs[0].model_dump(mode="json")
    assert design_min_issues(d) == []
    weak = dict(
        d, structure=d["structure"][:1], flows=d["flows"][:2], tech_choices=[], data_interfaces=[]
    )
    assert len(design_min_issues(weak)) == 4
    p = _project()
    reqs = _reqs(p)
    assert resolve_design_ids(d, reqs) == [
        "FR-001",
        "FR-002",
        "FR-004",
    ]  # RFP 번호 → 세부 요구 전부
    assert [r.id for r in design_targets(reqs)] == ["FR-001", "FR-002", "FR-003", "FR-004"]


def test_groups_follow_designs_then_code() -> None:
    p = _project()
    ctx = GenerationContext(project=p, requirements=_reqs(p), designs=[_design()])
    groups = group_requirements(ctx)
    assert groups[0].source == "design" and set(groups[0].ids) == {"FR-001", "FR-002", "FR-004"}
    assert groups[0].key_message.startswith("접수부터") and groups[0].notes == [
        "알림 템플릿 승인 기간은 발주처와 협의"
    ]
    assert any(g.source == "code" and g.ids == ["FR-003"] for g in groups[1:])
    # 설계를 끄면 코드 묶음
    assert all(g.source == "code" for g in group_requirements(ctx, use_designs=False))
    spec = design_structure_spec(_design())
    assert spec["kind"] == "layers" and [i["group"] for i in spec["items"]][:2] == [
        "사용자·화면",
        "응용(WAS)",
    ]
    flow = design_flow_spec(_design())
    assert flow["kind"] == "procedure" and flow["items"][1]["text"] == "◆ 분기"
    frows = flow_table_rows(_design())
    assert frows[1][1] == "워크플로 서비스" and "규칙 없음" in frows[1][3] and frows[0][3] == "-"
    rows = tech_choice_rows(_design())
    assert (
        rows[0][1] == "ShedLock DB 락"
        and "대안: Quartz" in rows[0][2]
        and rows[0][3] == "github.com (확인)"
    )
    assert (
        rows[1][3] == "example.org (미확인)" and rows[2][3] == "자체 설계"
    )  # 자체 설계는 확인 대상 아님


def test_decks_render_structure_flow_tech_and_qa_con13() -> None:
    p = _project()
    ctx = GenerationContext(
        project=p,
        requirements=_reqs(p),
        designs=[_design()],
        proposal_content={
            "profile": "submission",
            "sections": [
                {"slide_group": "understanding", "title": "사업 이해", "bullets": ["배경"]}
            ],
        },
        period_start=date(2026, 11, 2),
        period_end=date(2027, 4, 30),
    )
    outline = decide_outline(ctx, [])
    sub = build_submission_deck(ctx, outline)
    text = _deck_text(sub)
    for needle in (
        "구현 구조",
        "처리 절차",
        "적용 기술·선정 이유",
        "워크플로 서비스",
        "자동 배정",
        "ShedLock DB 락",
        "대안: Quartz",
        "POST /api/notify",
        "접수 폼과 임시저장 상태 전이를 구현",
    ):
        assert needle in text, needle
    # 모든 도형이 슬라이드 안
    prs = Presentation(io.BytesIO(sub))
    for sl in prs.slides:
        for sh in sl.shapes:
            assert (
                sh.top + sh.height <= prs.slide_height + 1000
                and sh.left + sh.width <= prs.slide_width + 1000
            )
    inp = SimpleNamespace(requirements=_reqs(p), ctx=ctx)
    fails, revs, ids = _design_depth_issues(inp, text)
    assert fails == [] and any("출처 미확인" in x for x in revs)
    assert any("설계에 없는 요구사항" in x for x in revs) and ids == [
        "FR-003"
    ]  # FR-003 은 설계 밖 → 검토
    # 설계 문구가 덱에 없으면 실패, 설계가 없으면 검토
    fails2, _r2, ids2 = _design_depth_issues(inp, "아무 내용 없음")
    assert fails2 and "FR-001" in ids2
    inp0 = SimpleNamespace(
        requirements=_reqs(p), ctx=GenerationContext(project=p, requirements=_reqs(p))
    )
    f0, r0, _ = _design_depth_issues(inp0, text)
    assert f0 == [] and r0 and "구현 설계 없음" in r0[0]
    # 발표용: 묶음마다 구조 도식 + 절차/기술 장
    ptext = _deck_text(build_presentation_deck(ctx, outline))
    assert "구현 구조" in ptext and "처리 절차" in ptext and "ShedLock DB 락" in ptext
    assert design_texts(_design())[0] == "접수·배정 화면"


def test_paginate_blocks_splits_tables_by_rows() -> None:
    b = DeckBuilder()
    long_rows = [
        [f"R-{i:03d}\n요구 {i}", "• " + "구현 방안 설명 " * 12, "중", "시연"] for i in range(14)
    ]
    blocks: list[dict[str, Any]] = [
        {"kind": "bullets", "title": "", "lines": ["요점 1", "요점 2"]},
        {"kind": "visual", "title": "구현 구조", "spec": design_structure_spec(_design())},
        {
            "kind": "table",
            "title": "요구사항별 구현 방안",
            "headers": ["요구사항", "구현 방안", "난이도", "검증"],
            "rows": long_rows,
            "col_widths": [2.2, 5.5, 0.8, 2.5],
        },
    ]
    pages = _paginate_blocks(b, blocks, has_lead=True)
    assert len(pages) >= 3
    assert (
        pages[0][0]["kind"] == "bullets" and pages[0][1]["kind"] == "visual"
    )  # 요점과 구조 도식은 같은 장
    drawn = sum(len(blk["rows"]) for pg in pages for blk in pg if blk.get("kind") == "table")
    assert drawn == 14  # 행 유실 없음
    assert any(
        blk.get("title", "").endswith("(계속)")
        for pg in pages[1:]
        for blk in pg
        if blk.get("kind") == "table"
    )


def test_pipeline_has_design_step_after_architecture() -> None:
    keys = [s.key for s in PROPOSAL_STEPS]
    assert keys.index("impl_design") == keys.index("architecture") + 1
    assert keys.index("impl_design") < keys.index("proposal_ppt")
    step = next(s for s in PROPOSAL_STEPS if s.key == "impl_design")
    assert step.agent_role == "design_writer" and not step.requires_gate
