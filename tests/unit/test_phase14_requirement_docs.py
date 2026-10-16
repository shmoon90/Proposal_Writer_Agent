"""Phase 14: 요구사항 정의서·추적표 양식(묶음 머리글·색·변경 행), 요구사항별 구현 방안 자동 작성(에이전트+규칙 초안),
문맥 비교·난이도 가중 공수, 고객 양식 정의서·추적표 재적재."""

from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from fixtures import builders as bld
from fixtures import excel_builders as xb
from proposal_agent.agents import AgentRequest, FakeAgentRunner
from proposal_agent.changes.impact import _effort_estimate
from proposal_agent.changes.semantic import (
    change_difficulty,
    classify_request,
    context_delta,
)
from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.excel.requirement_spec import (
    HEADER_ROW,
    REQUIRED_COLUMNS,
    SHEET_FUNCTIONAL,
    build_requirement_spec,
)
from proposal_agent.generators.excel.style import GROUP_FILLS, ROW_FILLS
from proposal_agent.generators.excel.trace_matrix import SHEET_MAIN, SHEET_NFR, build_trace_matrix
from proposal_agent.models import (
    AcceptanceStatus,
    DeliverableKind,
    Difficulty,
    DocumentKind,
    InputKind,
    Project,
    ProjectMode,
    Requirement,
    RequirementCategory,
    RequirementType,
)
from proposal_agent.models.change import CandidateClass, ChangeCandidate
from proposal_agent.orchestrator import StepRunner
from proposal_agent.requirements.approach import RULE_MARK, draft_rule_based
from proposal_agent.requirements.difficulty import estimate_difficulty
from proposal_agent.services import AppContext, Services
from unit.test_requirements_phase2 import extractor_from_text, verifier_all_ok


def _req(p: Project, rid: str, name: str, desc: str, **kw: Any) -> Requirement:
    base: dict[str, Any] = dict(
        id=rid,
        project_id=p.id,
        name=name,
        description=desc,
        type=RequirementType.FUNCTIONAL,
        created_by="t",
        acceptance_status=AcceptanceStatus.ACCEPTED,
    )
    base.update(kw)
    return Requirement(**base)


def _fill(cell: Any) -> str | None:
    return cell.fill.fgColor.rgb[-6:] if cell.fill and cell.fill.fill_type == "solid" else None


# 1. 정의서 양식: 묶음 머리글·색·변경 상태 행·목록 시트 -------------------------------------------------
def test_spec_layout_groups_colors_and_change_rows(services: Services) -> None:
    p = services.projects.create(name="정의서 양식", mode=ProjectMode.PROPOSAL)
    reqs = [
        _req(
            p,
            "FR-001",
            "로그인",
            "사용자는 아이디·비밀번호로 로그인한다.",
            rfp_requirement_id="SFR-001",
            implementation_approach="- 표준 로그인 모듈 적용\n- 5회 실패 잠금",
            approach_source="agent",
            approach_basis=["TR-0001"],
            difficulty=Difficulty.LOW,
            acceptance_criteria="- 로그인 성공/실패 시나리오 통과",
        ),
        _req(
            p,
            "FR-002",
            "비밀번호 찾기",
            "이메일 인증으로 비밀번호를 재설정한다.",
            rfp_requirement_id="SFR-001",
            change_state="added",
            change_summary="CR-0001 추가",
            added_by="홍",
            added_on=date(2026, 10, 15),
        ),
        _req(
            p,
            "FR-003",
            "민원 접수",
            "민원인은 민원을 접수한다.",
            rfp_requirement_id="SFR-002",
            change_state="modified",
            change_summary="CR-0002 변경 — 첨부 용량 제한",
        ),
        _req(
            p,
            "FR-004",
            "팩스 접수",
            "팩스 민원 수기 등록",
            rfp_requirement_id="SFR-003",
            change_state="deleted",
            change_summary="CR-0003 삭제",
        ),
        _req(
            p,
            "NFR-001",
            "응답 시간",
            "주요 화면은 3초 이내 응답한다.",
            type=RequirementType.NON_FUNCTIONAL,
            category=RequirementCategory.PERFORMANCE,
        ),
    ]
    reqs[2].add_history("홍", "change.modify", changes={"cr": "CR-0002"}, note="첨부 용량 제한")
    ctx = GenerationContext(project=p, requirements=reqs)
    wb = load_workbook(io.BytesIO(build_requirement_spec(ctx)))
    assert wb.sheetnames[:3] == ["표지", "개정이력", "요구사항목록"]
    ws = wb[SHEET_FUNCTIONAL]
    groups = [c.value for c in ws[1] if c.value]
    assert groups == [
        "식별",
        "요구사항 내용",
        "구현 방안",
        "수용·우선순위",
        "변경 정보 (구축 단계 기준)",
        "협의·비고",
    ]
    headers = [c.value for c in ws[HEADER_ROW]]
    assert all(h in headers for h in REQUIRED_COLUMNS)
    for h in (
        "RFP 고유번호",
        "대분류",
        "중분류",
        "구현 난이도",
        "난이도 근거",
        "구현 근거(조사·출처)",
        "검사 기준",
        "요구사항 변경여부",
        "변경 일자",
        "변경 근거",
        "검토 및 협의 내용",
    ):
        assert h in headers, h
    # 묶음 색: 구현 묶음 머리글은 황갈, 변경 정보는 연녹
    assert (
        _fill(ws.cell(row=HEADER_ROW, column=headers.index("구현방안") + 1)) == GROUP_FILLS["impl"]
    )
    assert (
        _fill(ws.cell(row=HEADER_ROW, column=headers.index("변경 내용") + 1))
        == GROUP_FILLS["change"]
    )
    rows = {
        ws.cell(row=r, column=2).value: r
        for r in range(3, ws.max_row + 1)
        if ws.cell(row=r, column=2).value
    }
    assert _fill(ws.cell(row=rows["FR-002"], column=1)) == ROW_FILLS["added"]
    assert _fill(ws.cell(row=rows["FR-003"], column=1)) == ROW_FILLS["modified"]
    assert _fill(ws.cell(row=rows["FR-004"], column=1)) == ROW_FILLS["deleted"]
    assert ws.cell(row=rows["FR-004"], column=2).font.strike is True
    state_col = headers.index("요구사항 변경여부") + 1
    assert [
        ws.cell(row=rows[i], column=state_col).value
        for i in ("FR-001", "FR-002", "FR-003", "FR-004")
    ] == [
        "신규",
        "추가",
        "변경",
        "삭제",
    ]
    assert (
        ws.cell(row=rows["FR-003"], column=headers.index("변경 근거") + 1).value == "첨부 용량 제한"
    )
    assert ws.cell(row=rows["FR-003"], column=headers.index("변경 일자") + 1).value
    basis = ws.cell(row=rows["FR-001"], column=headers.index("구현 근거(조사·출처)") + 1).value
    assert "에이전트" in basis and "TR-0001" in basis
    assert ws.cell(row=rows["FR-001"], column=headers.index("구현 난이도") + 1).value == "하"
    # 요구사항목록: 구분별 요약
    lst = wb["요구사항목록"]
    first_col = [lst.cell(row=r, column=1).value for r in range(2, lst.max_row + 1)]
    assert any(str(v).startswith("기능 (4건)") for v in first_col if v)
    assert any(str(v).startswith("비기능 (1건)") for v in first_col if v)
    guide = wb["작성가이드"]
    assert "구현방안" in [guide.cell(row=r, column=1).value for r in range(1, guide.max_row + 1)]
    assert ws.print_title_rows in ("1:2", "$1:$2") and ws.freeze_panes == "C3"


# 2. 추적표 양식: 묶음·비기능 검사기준 시트 ---------------------------------------------------------
def test_trace_matrix_groups_and_nfr_sheet(services: Services) -> None:
    p = services.projects.create(name="추적표 양식", mode=ProjectMode.PROPOSAL)
    reqs = [
        _req(p, "FR-001", "로그인", "로그인한다", acceptance_criteria="- 로그인 시나리오 통과"),
        _req(
            p,
            "NFR-001",
            "응답 시간",
            "3초 이내",
            type=RequirementType.NON_FUNCTIONAL,
            implementation_approach="- 캐시·인덱스 튜닝",
            acceptance_criteria="- 부하 시험 3초 이내",
        ),
        _req(p, "FR-009", "제외 기능", "범위 외", acceptance_status=AcceptanceStatus.REJECTED),
    ]
    ctx = GenerationContext(project=p, requirements=reqs)
    wb = load_workbook(io.BytesIO(build_trace_matrix(ctx, [], [], [], [])))
    ws = wb[SHEET_MAIN]
    groups = [c.value for c in ws[1] if c.value]
    assert groups == ["요구사항", "수용·대응", "설계", "구현", "테스트", "검사 기준", "검토"]
    headers = [c.value for c in ws[HEADER_ROW]]
    assert (
        _fill(ws.cell(row=HEADER_ROW, column=headers.index("통합테스트 ID") + 1))
        == GROUP_FILLS["test"]
    )
    rows = {ws.cell(row=r, column=4).value: r for r in range(3, ws.max_row + 1)}
    assert (
        ws.cell(row=rows["FR-001"], column=headers.index("판정 기준") + 1).value
        == "- 로그인 시나리오 통과"
    )
    assert ws.cell(row=rows["FR-001"], column=headers.index("완료여부") + 1).value == "미착수"
    assert ws.cell(row=rows["FR-009"], column=headers.index("완료여부") + 1).value == "대상제외"
    nfr = wb[SHEET_NFR]
    nfr_headers = [c.value for c in nfr[HEADER_ROW]]
    assert nfr_headers[3:9] == [
        "요구사항 ID",
        "요구사항 명",
        "요구사항 내용",
        "구현 내용",
        "검사 기준",
        "수행 결과",
    ]
    assert (
        nfr.cell(row=3, column=4).value == "NFR-001"
        and nfr.cell(row=3, column=7).value == "- 캐시·인덱스 튜닝"
    )
    assert nfr.max_row == 3  # 비기능 1건만
    untraced = wb["미추적"]
    assert {untraced.cell(row=r, column=1).value for r in range(2, untraced.max_row + 1)} == {
        "FR-001",
        "NFR-001",
    }


# 3. 구현 방안 자동 작성: 에이전트 + 규칙 초안 + 사람 작성 보존 + 조사 기록 ---------------------------------
def approach_responder(request: AgentRequest) -> dict[str, Any]:
    ids = re.findall(r'<req id="([A-Z]+-\d+)"', request.prompt)
    first = ids[0]
    return {
        "items": [
            {
                "requirement_id": first,
                "approach": "표준 로그인 모듈로 인증 처리\n5회 실패 시 계정 잠금, 잠금 해제는 관리자 화면\n전제: 비밀번호 정책은 고객 보안 규정 준수",
                "difficulty": "하",
                "difficulty_basis": "공통 모듈 재사용",
                "verification_method": "시연",
                "acceptance_criteria": "로그인 성공/실패/잠금 시나리오 통과",
                "research_refs": [],
                "sources": [
                    {
                        "url": "https://docs.example.org/auth/guide",
                        "title": "인증 모듈 개발 가이드",
                        "product": "예시 프레임워크",
                        "version": "4.x",
                        "trust": "vendor_guide",
                        "excerpt": "로그인 실패 횟수 제한을 설정으로 지원한다.",
                    }
                ],
                "feasibility": "feasible",
                "notes": "",
            },
            {"requirement_id": "FR-999", "approach": "- 없는 ID", "difficulty": "중"},
        ],
        "research_needed": [f"{ids[-1]}: 외부 연계 방식"],
        "notes": "",
    }


@pytest.fixture
def approach_services(ctx: AppContext) -> Services:
    return Services.from_context(
        ctx,
        runner=FakeAgentRunner(
            {
                "requirement_extractor": extractor_from_text,
                "requirement_verifier": verifier_all_ok,
                "approach_writer": approach_responder,
            }
        ),
    )


def _proposal(services: Services) -> Project:
    p = services.projects.create(
        name="구현 방안",
        mode=ProjectMode.PROPOSAL,
        start_date=date(2026, 11, 2),
        end_date=date(2027, 8, 31),
    )
    services.documents.register_upload(
        p, kind=DocumentKind.RFP, file_name="rfp.docx", data=bld.build_docx()
    )
    r = StepRunner(services)
    r.run_extraction(p)
    r.run_requirement_extraction(p)
    r.run_verification(p)
    return p


def test_approach_agent_rule_fallback_and_user_preserved(approach_services: Services) -> None:
    s = approach_services
    p = _proposal(s)
    reqs = s.requirements.list_requirements(p)
    assert len(reqs) >= 3
    user_req = reqs[1]
    s.requirements.update_requirement(
        p,
        user_req.id,
        {"implementation_approach": "- 사람이 직접 쓴 방안", "approach_source": "user"},
        actor="pm",
    )
    outcome = StepRunner(s).run_approach_writing(p, actor="t", only_missing=False)
    assert outcome.run.status.value == "approved"
    after = {r.id: r for r in s.requirements.list_requirements(p)}
    first = after[reqs[0].id]
    assert first.implementation_approach is not None and first.implementation_approach.startswith(
        "- 표준 로그인 모듈"
    )
    assert first.implementation_approach.count("\n") == 2  # 3줄, 글머리 정규화
    assert first.approach_source == "agent" and first.difficulty == Difficulty.LOW
    assert first.verification_method == "시연" and first.acceptance_criteria.startswith("- 로그인")
    research = s.research.list_research(p)
    assert len(research) == 1 and research[0].related_requirement_ids == [first.id]
    assert research[0].sources[0].url == "https://docs.example.org/auth/guide"
    assert (
        research[0].id in first.approach_basis
        and research[0].sources[0].url in first.approach_basis
    )
    assert after[user_req.id].implementation_approach == "- 사람이 직접 쓴 방안"
    assert after[user_req.id].approach_source == "user"
    others = [
        r
        for rid, r in after.items()
        if rid not in (first.id, user_req.id) and r.type.value not in ("evaluation", "submission")
    ]
    assert others and all(RULE_MARK in (r.implementation_approach or "") for r in others)
    assert all(
        r.approach_source in ("rule", "research") and r.difficulty is not None for r in others
    )
    assert all(r.acceptance_criteria for r in others)
    assert "추가 조사" in outcome.summary and "FR-999" not in {r.id for r in after.values()}
    # 산출물 생성 후 QA-CON-09 실행: 에이전트 작성 항목은 통과 대상, 규칙 초안은 검토 대상
    runner = StepRunner(s)
    runner.run_generation(p)
    runner.run_qa(p)
    run = s.qa.latest_run(p)
    assert run is not None
    con09 = next(c for c in run.checks if c.check_id == "QA-CON-09")
    assert con09.executed and con09.final_result.value != "FAIL"
    assert any("[자동 초안]" in i or "규칙 초안" in i for i in con09.issues)
    spec = s.generation.latest_deliverables(p)[DeliverableKind.REQUIREMENT_SPEC_XLSX]
    wb = load_workbook(io.BytesIO(s.ctx.paths.resolve_output(spec.file_path).read_bytes()))
    ws = wb[SHEET_FUNCTIONAL]
    headers = [c.value for c in ws[HEADER_ROW]]
    col = headers.index("구현방안") + 1
    idc = headers.index("요구사항 ID") + 1
    assert all(
        ws.cell(row=r, column=col).value not in (None, "", "검토 필요")
        for r in range(3, ws.max_row + 1)
        if ws.cell(row=r, column=idc).value
    )


def test_rule_draft_uses_project_data_and_marks_itself(services: Services) -> None:
    p = services.projects.create(name="규칙 초안", mode=ProjectMode.PROPOSAL)
    req = _req(
        p,
        "FR-001",
        "민원 연계",
        "외부 기관 API로 민원을 실시간 전송한다.",
        category=RequirementCategory.INTERFACE,
    )
    text, basis, source, _ = draft_rule_based(req, [], [], None)
    assert text.startswith("- 연계 방식") and RULE_MARK in text and source == "rule" and basis == []
    level, why = estimate_difficulty(req.description, category=req.category)
    assert level == Difficulty.HIGH and "복잡도 신호" in why and "인터페이스" in why
    low, _ = estimate_difficulty("안내 문구를 바꾼다", category=RequirementCategory.UI)
    assert low == Difficulty.LOW


# 4. 문맥 비교·난이도 가중 공수 ---------------------------------------------------------------------
def test_context_delta_and_difficulty_weighted_effort(services: Services) -> None:
    before = "민원인은 민원 유형을 선택하고 내용과 첨부파일을 등록하여 민원을 접수한다."
    after = "민원인은 민원을 접수하고 외부 기관 API로 연계 전송한다. 응답은 3초 이내여야 한다."
    aspects = context_delta(before, after)
    assert {"interface", "number", "quality", "io"} <= set(aspects)
    assert context_delta(before, before) == []
    level, why = change_difficulty(after, aspects, None)
    assert level == "상" and "변경 관점" in why
    p = services.projects.create(name="문맥", mode=ProjectMode.BUILD)
    existing = [_req(p, "SFR-002", "민원 접수", before, difficulty=Difficulty.MEDIUM)]
    cand = classify_request(
        after, existing, input_kind=InputKind.MODIFY_BUTTON, selected_ids=["SFR-002"]
    )[0]
    assert cand.classification == CandidateClass.MODIFY and cand.difficulty == "상"
    assert "문맥 차이" in cand.rationale and "연계" in cand.rationale
    # 같은 유사도라도 난이도가 높으면 공수 추정이 커진다
    base = ChangeCandidate(
        id="C1", classification=CandidateClass.MODIFY, similarity=0.6, change_aspects=["number"]
    )
    low = base.model_copy(update={"difficulty": "하"})
    high = base.model_copy(
        update={"difficulty": "상", "change_aspects": ["interface", "number", "quality"]}
    )
    e_low = _effort_estimate(low, [], Decimal("10"))
    e_high = _effort_estimate(high, [], Decimal("10"))
    assert e_low[0] < e_high[0] and "난이도 가중치" in e_high[4] and "문맥 관점 3개" in e_high[4]
    added = ChangeCandidate(id="C2", classification=CandidateClass.ADD, difficulty="상")
    assert _effort_estimate(added, [], None)[0] == Decimal("14")  # 기본 10 M/D × 1.4
    override = ChangeCandidate(
        id="C3", classification=CandidateClass.ADD, difficulty="상", user_override={"effort_md": 5}
    )
    assert _effort_estimate(override, [], None)[0] == Decimal("5")  # 사용자 입력이 우선


# 5. 고객 양식 정의서·추적표 재적재 (제목 행·묶음 머리글·동의어 컬럼) -------------------------------------
def _customer_style_spec() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "요구사항정의서(1.0)"
    ws.append(["OO 사업 요구사항 정의서"])
    ws.append([])
    ws.append(
        ["식별", None, None, None, None, "내용", None, None, "구현", None, None, None, "관리"]
    )
    ws.append(
        [
            "No",
            "요구사항 ID",
            "대분류",
            "중분류",
            "요구사항명",
            "상세명",
            "요구사항 설명",
            "유형",
            "요구사항 해결 방안",
            "중요도",
            "난이도",
            "담당팀",
            "관련 근거",
            "요구사항변경여부",
            "수용여부",
            "과업대비표 구분",
            "비고",
        ]
    )
    rows = [
        (
            1,
            "SFR-001",
            "채널",
            "인터페이스",
            "로그인",
            "SSO 로그인",
            "사용자는 SSO로 로그인한다.",
            "기능",
            "ㅇ SSO 연계 모듈 적용\n- 세션 관리",
            "상",
            "상",
            "응용1",
            "SFR-001",
            "신규",
            "수용",
            "기능",
            "",
        ),
        (
            2,
            "SFR-002",
            "채널",
            "통계",
            "민원 접수",
            "",
            "민원인은 민원을 접수한다.",
            "기능",
            "- 접수 화면 구현",
            "중",
            "중",
            "응용1",
            "SFR-002",
            "변경",
            "부분수용\n- 제외: 팩스",
            "기능",
            "범위 조정",
        ),
        (
            3,
            "SNR-001",
            "인프라",
            "성능",
            "응답 시간",
            "",
            "주요 화면은 3초 이내 응답한다.",
            "성능",
            "",
            "중",
            "",
            "인프라",
            "PER-001",
            "신규",
            "수용",
            "비기능",
            "",
        ),
    ]
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _customer_style_trace() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "검사기준서(기능)"
    ws.append(["OO 사업"])
    ws.append(["요구사항", None, None, "설계", None, "테스트", None, "검사기준"])
    ws.append(
        [
            "No",
            "분석 요구사항ID",
            "분석 요구사항명",
            "화면/보고서ID",
            "프로그램ID",
            "단위테스트케이스ID",
            "통합테스트케이스ID",
            "판정기준",
            "완료여부",
            "WBS",
        ]
    )
    ws.append(
        [
            1,
            "SFR-001",
            "로그인",
            "UI-001",
            "PRG-001",
            "UT-001",
            "IT-001",
            "SSO 로그인 성공",
            "완료",
            "2.1.1",
        ]
    )
    ws.append(
        [
            2,
            "SFR-002",
            "민원 접수",
            "UI-002",
            "PRG-002",
            "UT-002",
            "IT-001",
            "접수 완료 알림 수신",
            "완료",
            "2.1.2",
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_baseline_loads_customer_style_spec_and_trace(services: Services) -> None:
    p = services.projects.create(name="고객 양식 적재", mode=ProjectMode.BUILD)
    services.documents.register_upload(
        p, kind=DocumentKind.REQUIREMENT_SPEC, file_name="정의서.xlsx", data=_customer_style_spec()
    )
    services.documents.register_upload(
        p, kind=DocumentKind.TRACE_MATRIX, file_name="추적표.xlsx", data=_customer_style_trace()
    )
    services.documents.register_upload(
        p, kind=DocumentKind.WBS, file_name="wbs.xlsx", data=xb.build_wbs()
    )
    profiles = services.baseline.profiles(p)
    spec_prof = profiles[DocumentKind.REQUIREMENT_SPEC]
    assert spec_prof.profile.sheets[0].header_row == 4
    mapped = {c.field: c.header for c in spec_prof.profile.sheets[0].columns if c.header}
    assert mapped["implementation"] == "요구사항 해결 방안" and mapped["difficulty"] == "난이도"
    assert mapped["type"] == "과업대비표 구분" and mapped["category"] == "중분류"
    assert not spec_prof.profile.needs_confirmation
    result = services.baseline.load(p, actor="t")
    assert result.requirements == 3
    reqs = {r.id: r for r in services.ctx.repos.requirements.find(p.id)}
    sso = reqs["SFR-001"]
    assert sso.difficulty == Difficulty.HIGH and sso.category == RequirementCategory.INTERFACE
    assert sso.business_area == "채널" and sso.approach_source == "baseline"
    assert (
        sso.implementation_approach is not None and "SSO 연계 모듈" in sso.implementation_approach
    )
    assert sso.acceptance_criteria == "SSO 로그인 성공"  # 추적표 판정기준 → 검사 기준
    assert sso.change_state == "none" and sso.type == RequirementType.FUNCTIONAL
    civil = reqs["SFR-002"]
    assert civil.acceptance_status == AcceptanceStatus.PARTIALLY_ACCEPTED
    assert civil.change_state == "modified" and civil.category == RequirementCategory.REPORT
    perf = reqs["SNR-001"]
    assert (
        perf.type == RequirementType.NON_FUNCTIONAL
        and perf.category == RequirementCategory.PERFORMANCE
    )
    assert perf.difficulty is None and perf.implementation_approach is None
    traces = services.ctx.repos.traces.find(p.id, requirement_id="SFR-001")
    assert {t.target_type.value for t in traces} >= {
        "screen",
        "unit_test",
        "integration_test",
        "wbs_item",
    }
