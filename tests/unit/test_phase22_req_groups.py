"""Phase 22: 요구사항 묶음별 구현 방안 소절(코드 묶음·에이전트 묶음, 정의서 전문 표, 기능 블록 도식),
복합 장표(밴드 배치), QA-CON-12(정의서 구현방안 ↔ 제안서 반영)."""

from __future__ import annotations

import io
from datetime import date
from types import SimpleNamespace
from typing import Any

from pptx import Presentation

from proposal_agent.agents.schemas import ProposalContent
from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.pptx import shapes
from proposal_agent.generators.pptx.decks import build_presentation_deck, build_submission_deck
from proposal_agent.generators.pptx.outline import ESSENTIAL_KEYS, decide_outline
from proposal_agent.generators.pptx.req_groups import (
    approach_table_rows,
    function_block_spec,
    group_requirements,
)
from proposal_agent.models import (
    AcceptanceStatus,
    Difficulty,
    Project,
    ProjectMode,
    Requirement,
    RequirementCategory,
    RequirementType,
)
from proposal_agent.validation.independent_qa import _approach_coverage_issues


def _project() -> Project:
    return Project(
        name="묶음",
        mode=ProjectMode.PROPOSAL,
        output_dir="Project_g",
        workspace_dir="ws_g",
        start_date=date(2026, 11, 2),
        end_date=date(2027, 4, 30),
        created_by="t",
    )


def _req(
    p: Project,
    rid: str,
    name: str,
    approach: str,
    *,
    rfp: str | None,
    cat: str = "business",
    diff: str = "medium",
) -> Requirement:
    return Requirement(
        id=rid,
        project_id=p.id,
        name=name,
        description=name,
        type=RequirementType.FUNCTIONAL,
        category=RequirementCategory(cat),
        rfp_requirement_id=rfp,
        created_by="t",
        acceptance_status=AcceptanceStatus.ACCEPTED,
        implementation_approach=approach,
        difficulty=Difficulty(diff),
        verification_method="시연",
        acceptance_criteria="상태 전이가 규칙대로 동작",
    )


def _reqs(p: Project) -> list[Requirement]:
    return [
        _req(
            p,
            "FR-001",
            "사건 접수 화면",
            "- 접수 폼과 임시저장 상태 전이를 구현\n- 첨부 파일은 저장소에 보관",
            rfp="SFR-001",
        ),
        _req(
            p,
            "FR-002",
            "사건 배정",
            "- 담당자 배정 규칙 테이블로 자동 배정\n- 수동 재배정 화면 제공",
            rfp="SFR-001",
            diff="high",
        ),
        _req(
            p,
            "FR-003",
            "통계 조회",
            "- 집계 쿼리와 네이티브 차트로 월별 통계 표시",
            rfp="SFR-002",
            cat="report",
        ),
        _req(
            p,
            "FR-004",
            "알림 발송",
            "- 이벤트별 알림톡 템플릿 발송\n- 전제: 템플릿 승인은 발주처와 협의",
            rfp=None,
            cat="interface",
            diff="low",
        ),
        _req(p, "FR-005", "구현방안 없음", "", rfp="SFR-002"),
    ]


def test_code_grouping_by_rfp_id_then_category() -> None:
    p = _project()
    ctx = GenerationContext(project=p, requirements=_reqs(p))
    groups = group_requirements(ctx)
    titles = {g.title: g.ids for g in groups}
    # SFR-001 두 건은 한 묶음, 나머지는 분류별; 구현방안 없는 FR-005 는 제외
    assert any(set(ids) == {"FR-001", "FR-002"} for ids in titles.values())
    assert any(ids == ["FR-003"] for ids in titles.values()) and any(
        ids == ["FR-004"] for ids in titles.values()
    )
    assert "FR-005" not in {i for ids in titles.values() for i in ids}
    g = next(g for g in groups if set(g.ids) == {"FR-001", "FR-002"})
    assert g.key_message and "난이도 상 1건" in g.key_message
    rows = approach_table_rows(g)
    assert (
        rows[0][0].startswith("FR-001 / SFR-001")
        and "임시저장" in rows[0][1]
        and rows[1][2] == "상"
    )
    spec = function_block_spec(g)
    assert [f["title"] for f in spec["functions"]] == ["사건 접수 화면", "사건 배정"] and spec[
        "actors"
    ]


def test_agent_groups_override_and_rfp_id_expansion() -> None:
    p = _project()
    content = {
        "profile": "submission",
        "sections": [],
        "requirement_groups": [
            {
                "title": "사건 처리 워크플로",
                "requirement_ids": ["SFR-001", "FR-004"],
                "key_message": "접수→배정→알림을 한 흐름으로",
                "summary": ["접수·배정·알림 연결"],
                "notes": ["알림 템플릿 협의"],
            },
        ],
    }
    ProposalContent.model_validate(content)
    ctx = GenerationContext(project=p, requirements=_reqs(p), proposal_content=content)
    groups = group_requirements(ctx)
    first = groups[0]
    assert first.source == "agent" and set(first.ids) == {"FR-001", "FR-002", "FR-004"}
    assert first.key_message.startswith("접수") and first.notes == ["알림 템플릿 협의"]
    # 묶이지 않은 FR-003 은 코드 묶음으로 뒤에 붙는다
    assert any(g.source == "code" and g.ids == ["FR-003"] for g in groups[1:])


def _deck_text(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    parts = []
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts.extend(c.text for r in sh.table.rows for c in r.cells)
    return "\n".join(parts)


def test_decks_render_group_pages_with_full_approach_text_and_qa_con12_passes() -> None:
    p = _project()
    ctx = GenerationContext(
        project=p,
        requirements=_reqs(p),
        proposal_content={
            "profile": "submission",
            "sections": [
                {"slide_group": "understanding", "title": "사업 이해", "bullets": ["배경"]}
            ],
        },
        period_start=date(2026, 11, 2),
        period_end=date(2027, 4, 30),
    )
    assert "req_groups" in ESSENTIAL_KEYS
    outline = decide_outline(ctx, [])
    assert any(o.key == "req_groups" for o in outline)
    sub = build_submission_deck(ctx, outline)
    text = _deck_text(sub)
    for line in (
        "접수 폼과 임시저장 상태 전이를 구현",
        "담당자 배정 규칙 테이블로 자동 배정",
        "집계 쿼리와 네이티브 차트로 월별 통계 표시",
    ):
        assert line in text  # 정의서 구현방안 전문이 잘리지 않고 실림
    assert "요구사항별 구현 방안" in text and "기능 구성" in text
    entries = [e for e in ctx.slide_specs["submission"] if e["key"] == "req_groups"]
    assert entries and entries[0]["slide_to"] >= entries[0]["slide_from"]
    # 전제·협의 문구는 ※ 주석으로
    assert "템플릿 승인은 발주처와 협의" in text
    # QA-CON-12: 덱 텍스트 기준 전부 반영 → 실패·검토 없음
    inp = SimpleNamespace(requirements=_reqs(p))
    fails, revs, _ids = _approach_coverage_issues(inp, text)
    assert fails == [] and revs == []
    # 반영이 빠지면 필수 기능 요구는 실패, 일부만 있으면 검토
    fails2, revs2, ids2 = _approach_coverage_issues(inp, "접수 폼과 임시저장 상태 전이를 구현")
    assert (
        any(x.startswith("FR-002") for x in fails2)
        and any(x.startswith("FR-001") for x in revs2)
        and "FR-002" in ids2
    )
    # 발표용: 묶음마다 블록 도식 1장
    pres = build_presentation_deck(ctx, outline)
    ptext = _deck_text(pres)
    assert "기능 구성" in ptext and "사건 접수 화면" in ptext


def test_composite_slide_bands_and_block_sizing() -> None:
    from proposal_agent.generators.pptx.builder import DeckBuilder

    b = DeckBuilder()
    kinds = b.composite_slide(
        "복합",
        [
            {"kind": "bullets", "title": "", "lines": ["첫째", "둘째"]},
            {
                "kind": "visual",
                "title": "구성·흐름",
                "spec": {
                    "kind": "layers",
                    "items": [{"title": "A", "group": "사용자"}, {"title": "B", "group": "응용"}],
                },
            },
            {
                "kind": "table",
                "title": "기술 상세",
                "headers": ["구분", "내용"],
                "rows": [["구성", "• x\n• y"], ["보안", "• z"]],
                "col_widths": [1.6, 8.4],
            },
        ],
        "footer",
        lead="결론",
        footnotes=["전제: 협의"],
    )
    assert kinds == ["bullets", "visual", "table"]
    prs = Presentation(io.BytesIO(b.to_bytes()))
    sl = prs.slides[0]
    texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
    assert any("구성·흐름" in t for t in texts) and any("기술 상세" in t for t in texts)
    tbl = next(sh for sh in sl.shapes if getattr(sh, "has_table", False) and sh.has_table)
    assert tbl.table.cell(1, 1).text == "• x\n• y"
    # 모든 도형이 슬라이드 안
    for sh in sl.shapes:
        assert (
            sh.top + sh.height <= prs.slide_height + 1000
            and sh.left + sh.width <= prs.slide_width + 1000
        )
    # 블록 높이 추정은 양수이며 글머리는 축소 대상이 아님
    est = shapes.estimate_block_height(
        {"kind": "bullets", "lines": ["a" * 100]}, 17.0, shapes.ShapeStyle()
    )
    assert est > 0.5


def test_difficulty_label_in_rows() -> None:
    p = _project()
    g = group_requirements(GenerationContext(project=p, requirements=_reqs(p)))[0]
    labels: list[Any] = [r[2] for r in approach_table_rows(g)]
    assert set(labels) <= {"상", "중", "하", "-"}
