"""Phase 17: RFP 분석 항목 확장(스키마·프롬프트), 페이지 제한·문서 규격 파싱, 근거 구분·대응 전략, 덱 슬라이드 번호 역산과
추적표 '대응 제안서 페이지', 생성 순서(덱 → Excel)."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Cm

from fixtures import builders as bld
from proposal_agent.agents import FakeAgentRunner, load_prompt
from proposal_agent.agents.schemas import RfpAnalysis
from proposal_agent.generators.excel.style import HEADER_ROW
from proposal_agent.generators.pptx.pagemap import format_pages, requirement_page_map
from proposal_agent.ingest.rfp_profile import profile_blocks
from proposal_agent.models import (
    AcceptanceStatus,
    BlockType,
    DeliverableKind,
    DocumentKind,
    ExtractedBlock,
    FileFormat,
    Project,
    ProjectMode,
    Requirement,
    RequirementType,
    ScopeClass,
)
from proposal_agent.models.extraction import BlockLocation
from proposal_agent.orchestrator import StepRunner
from proposal_agent.requirements.evidence import (
    EVIDENCE_LABELS,
    default_evidence_class,
    default_response_strategy,
)
from proposal_agent.services import AppContext, Services
from unit.test_requirements_phase2 import extractor_from_text, verifier_all_ok


def _services(ctx: AppContext) -> Services:
    return Services.from_context(
        ctx,
        runner=FakeAgentRunner(
            {"requirement_extractor": extractor_from_text, "requirement_verifier": verifier_all_ok}
        ),
    )


# 1. RFP 분석 스키마·프롬프트 --------------------------------------------------------------------------
def test_rfp_analysis_schema_has_directive_fields_and_prompt_mentions_them() -> None:
    a = RfpAnalysis(overview="x")
    for f in (
        "current_issues",
        "security_privacy",
        "data_migration_operation",
        "dev_environment",
        "deliverables_schedule",
        "test_acceptance_training",
        "proposal_guidelines",
        "document_format",
        "required_submissions",
    ):
        assert getattr(a, f) == []
    assert a.page_limit == ""
    text = load_prompt("rfp_analyst")
    for key in (
        "current_issues",
        "security_privacy",
        "page_limit",
        "required_submissions",
        "지어내지",
    ):
        assert key in text


# 2. 페이지 제한·문서 규격 파싱 ----------------------------------------------------------------------
def _block(
    bid: str, text: str, btype: BlockType = BlockType.PARAGRAPH, order: int = 0
) -> ExtractedBlock:
    return ExtractedBlock(
        id=bid,
        project_id="p",
        source_document_id="d",
        file_name="rfp.docx",
        file_format=FileFormat.DOCX,
        order=order,
        block_type=btype,
        text=text,
        location=BlockLocation(page=1),
    )


def test_rfp_profile_parses_page_limit_and_format_notes() -> None:
    blocks = [
        _block("b1", "Ⅳ. 제안서 작성요령", BlockType.HEADING, 1),
        _block("b2", "1. 제안서는 A4 용지 100쪽 이내로 작성한다.", order=2),
        _block("b3", "2. 글꼴은 11pt 이상, 제본 10부와 PDF 파일을 제출한다.", order=3),
        _block("b4", "Ⅴ. 평가 기준", BlockType.HEADING, 4),
        _block("b5", "기술 평가 90점, 가격 10점", order=5),
    ]
    prof = profile_blocks(blocks, apply=False)
    assert prof.page_limit == 100 and "100쪽" in prof.page_limit_text
    assert any("11pt" in n for n in prof.format_notes) and any("A4" in n for n in prof.format_notes)
    assert "페이지 제한: 100쪽" in prof.summary_text()
    none = profile_blocks(
        [_block("c1", "사업 개요", BlockType.HEADING, 1), _block("c2", "본 사업은…", order=2)],
        apply=False,
    )
    assert none.page_limit is None and none.format_notes == []


# 3. 근거 구분·대응 전략 기본값 ---------------------------------------------------------------------
def test_evidence_class_and_response_strategy_defaults() -> None:
    base = dict(project_id="p", name="n", type=RequirementType.FUNCTIONAL)
    assert default_evidence_class(Requirement(id="A", rfp_quote="원문", **base)) == "rfp"
    assert (
        default_evidence_class(
            Requirement(id="B", scope_class=ScopeClass.OPTIONAL_PROPOSAL, **base)
        )
        == "design_proposal"
    )
    assert (
        default_evidence_class(Requirement(id="C", approach_source="baseline", **base))
        == "reference_doc"
    )
    assert (
        default_evidence_class(Requirement(id="D", evidence_class="research", **base)) == "research"
    )
    assert set(EVIDENCE_LABELS) == {"rfp", "reference_doc", "research", "design_proposal"}
    r = Requirement(
        id="E",
        acceptance_status=AcceptanceStatus.ACCEPTED,
        implementation_approach="- [자동 초안] x\n- 표준 모듈 재사용\n- 전제",
        **base,
    )
    assert default_response_strategy(r) == "수용 — 표준 모듈 재사용"
    r.response_strategy = "표준 공통컴포넌트 재사용으로 수용"
    assert default_response_strategy(r) == "표준 공통컴포넌트 재사용으로 수용"
    rej = Requirement(id="F", acceptance_status=AcceptanceStatus.REJECTED, notes="범위 외", **base)
    assert default_response_strategy(rej) == "미수용 — 범위 외"


# 4. 슬라이드 번호 역산 ----------------------------------------------------------------------------------
def test_requirement_page_map_scans_text_tables_and_groups() -> None:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    s1.shapes.add_textbox(Cm(1), Cm(1), Cm(10), Cm(1)).text_frame.text = "개요"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    tbl = s2.shapes.add_table(2, 2, Cm(1), Cm(1), Cm(10), Cm(3)).table
    tbl.cell(0, 0).text = "FR-001"
    tbl.cell(1, 0).text = "SFR-002"
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    g = s3.shapes.add_group_shape()
    g.shapes.add_textbox(Cm(1), Cm(1), Cm(6), Cm(1)).text_frame.text = "대응: FR-001, FR-003"
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    s4.shapes.add_textbox(Cm(1), Cm(1), Cm(6), Cm(1)).text_frame.text = "FR-001 상세"
    buf = io.BytesIO()
    prs.save(buf)
    reqs = [
        Requirement(id="FR-001", project_id="p", name="a", type=RequirementType.FUNCTIONAL),
        Requirement(
            id="FR-002",
            project_id="p",
            name="b",
            type=RequirementType.FUNCTIONAL,
            rfp_requirement_id="SFR-002",
        ),
        Requirement(id="FR-003", project_id="p", name="c", type=RequirementType.FUNCTIONAL),
        Requirement(id="FR-009", project_id="p", name="d", type=RequirementType.FUNCTIONAL),
    ]
    pm = requirement_page_map(buf.getvalue(), reqs)
    assert pm == {"FR-001": [2, 3, 4], "FR-002": [2], "FR-003": [3]}
    assert format_pages([2, 3, 4]) == "2-4" and format_pages([3, 5, 6, 9]) == "3, 5-6, 9"
    assert format_pages(None) == "—" and format_pages([7]) == "7"


# 5. 생성 흐름: 덱 뒤에 Excel, 추적표·정의서에 대응 페이지·근거 구분·대응 전략 --------------------------
def test_generation_fills_trace_pages_and_evidence(ctx: AppContext) -> None:
    services = _services(ctx)
    p: Project = services.projects.create(
        name="대응 페이지",
        mode=ProjectMode.PROPOSAL,
        start_date=date(2026, 11, 2),
        end_date=date(2027, 8, 31),
    )
    services.documents.register_upload(
        p, kind=DocumentKind.RFP, file_name="rfp.docx", data=bld.build_docx()
    )
    r = StepRunner(services)
    r.run_extraction(p)
    r.run_requirement_extraction(p)
    r.run_verification(p)
    reqs = services.requirements.list_requirements(p)
    assert all(q.evidence_class == "rfp" for q in reqs)
    for q in reqs:
        q.acceptance_status = AcceptanceStatus.ACCEPTED
        services.ctx.repos.requirements.update(q)
    out = r.run_generation(p)
    kinds = [d.kind for d in services.generation.latest_deliverables(p).values()]
    assert out.run.status.value == "approved"
    deliverables = services.ctx.repos.deliverables.find(p.id, order_by="created_at")
    order = [d.kind for d in deliverables]
    assert order.index(DeliverableKind.SUBMISSION_PPTX) < order.index(
        DeliverableKind.TRACE_MATRIX_XLSX
    )
    assert DeliverableKind.TRACE_MATRIX_XLSX in kinds
    tm = services.generation.latest_deliverables(p)[DeliverableKind.TRACE_MATRIX_XLSX]
    wb = load_workbook(io.BytesIO(services.ctx.paths.resolve_output(tm.file_path).read_bytes()))
    ws = wb["추적표"]
    headers = [c.value for c in ws[HEADER_ROW]]
    for h in ("대응 전략", "근거 구분", "제안서 페이지(제출용)", "제안서 페이지(발표용)"):
        assert h in headers
    idc = headers.index("요구사항 ID") + 1
    rows = {
        ws.cell(row=rr, column=idc).value: rr
        for rr in range(3, ws.max_row + 1)
        if ws.cell(row=rr, column=idc).value
    }
    fr = next(rid for rid in rows if str(rid).startswith("FR-"))
    sub_pages = ws.cell(row=rows[fr], column=headers.index("제안서 페이지(제출용)") + 1).value
    assert sub_pages and sub_pages != "—" and sub_pages[0].isdigit()
    assert ws.cell(row=rows[fr], column=headers.index("근거 구분") + 1).value == "RFP 명시"
    assert str(ws.cell(row=rows[fr], column=headers.index("대응 전략") + 1).value).startswith(
        "수용"
    )
    spec = services.generation.latest_deliverables(p)[DeliverableKind.REQUIREMENT_SPEC_XLSX]
    ws2 = load_workbook(io.BytesIO(services.ctx.paths.resolve_output(spec.file_path).read_bytes()))[
        "기능요구사항"
    ]
    h2 = [c.value for c in ws2[HEADER_ROW]]
    assert "근거 구분" in h2 and "대응 전략" in h2
    # 수동 추가 요구(부가 제안)는 설계 제안
    from proposal_agent.models import RequirementDraft

    added = services.requirements.add_manual(
        p,
        RequirementDraft(
            name="알림 재전송",
            description="실패한 알림을 재전송한다.",
            type=RequirementType.FUNCTIONAL,
            scope_class=ScopeClass.OPTIONAL_PROPOSAL,
            impact_note="공수 2 M/D",
        ),
        actor="pm",
    )
    assert added.evidence_class == "design_proposal"
