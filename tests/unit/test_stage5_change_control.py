"""통합 개발 지시문 단계 5 — 기준선·변경 관리.

회의록 구조화(규칙·에이전트)와 권한 미확인 승인 금지(§10), 기술 협상 → 변경 요청 → 기준선 반영 통제, 수정 문서 비교
(정의서 Excel ID 기준: 문구만/의미/추가/삭제, PPT 판독 불확실; A-5), 변경 상태 8종(승인 대기), 영향 분석 확장(§11.2),
선택 재생성, QA-CON-14."""

from __future__ import annotations

import io
from datetime import date, timedelta
from types import SimpleNamespace

import pytest
from pptx import Presentation

from fixtures import excel_builders as xb
from proposal_agent.agents import FakeAgentRunner
from proposal_agent.changes.compare import (
    KIND_ADD,
    KIND_DELETE,
    KIND_SCREEN_MOVE,
    KIND_SEMANTIC,
    KIND_WORDING,
    change_kind_of,
    classify_document_text,
    compare_spec_workbook,
)
from proposal_agent.config.errors import ValidationError
from proposal_agent.decisions import extract_meeting_items
from proposal_agent.models import (
    ApiSpec,
    ChangeStatus,
    DecisionKind,
    DecisionStatus,
    DeliverableKind,
    DocumentKind,
    NegotiationStatus,
    Project,
    RequirementTrace,
    TraceTargetType,
)
from proposal_agent.models.change import CandidateClass, InputKind
from proposal_agent.models.enums import CHANGE_STATUS_FLOW, CHANGE_STATUS_LABELS
from proposal_agent.services import AppContext, Services
from proposal_agent.validation.qa_checks.consistency import _decision_control_issues
from unit.test_phase6_changes import _reqs, _upload_all

MEETING = """제안 준비 회의록
결정: 제안서 분량은 요약 모드로 한다 (담당: 기술 PL)
요청: 계산 이력 보관 기간을 1년으로 늘려 달라 (SFR-002 관련)
미결: 모바일 지원 범위 확인 필요 — 담당: 영업 PM, 기한 2026-10-20
조치: PoC 결과 공유 (담당: 기술 PL, 기한 2026-10-22)
위험: 서버 사양 미정
그냥 논의한 문장은 항목이 아니다
"""


@pytest.fixture
def loaded(services: Services, build_project: Project) -> Project:
    _upload_all(services, build_project)
    services.baseline.load(build_project, actor="tester")
    return services.ctx.repos.projects.get(build_project.id)


# --- 회의록 구조화·권한 규칙 ----------------------------------------------------------------------


def test_meeting_rule_extraction_and_authority_rule(services: Services, loaded: Project) -> None:
    items = extract_meeting_items(MEETING)
    kinds = [i["kind"] for i in items]
    assert kinds == [
        DecisionKind.OFFICIAL_DECISION,
        DecisionKind.CUSTOMER_REQUEST,
        DecisionKind.OPEN_QUESTION,
        DecisionKind.ACTION_ITEM,
        DecisionKind.RISK_ISSUE,
    ]
    assert items[1]["related_requirement_ids"] == ["SFR-002"]
    assert items[2]["owner"] == "영업" or items[2]["owner"].startswith("영업")
    assert items[2]["due_date"] == date(2026, 10, 20)
    doc = services.documents.register_upload(
        loaded,
        kind=DocumentKind.MEETING_NOTE,
        file_name="회의록.txt",
        data=MEETING.encode("utf-8"),
        actor="t",
    )
    created, warnings = services.decisions.structure_meeting(
        loaded, doc, actor="t", use_agent=False
    )
    assert [d.code for d in created] == ["DEC-001", "DEC-002", "DEC-003", "DEC-004", "DEC-005"]
    assert all(d.status == DecisionStatus.PROPOSED and not d.authority_confirmed for d in created)
    assert created[1].related_requirement_ids == ["SFR-002"]
    official = created[0]
    # 권한 미확인 → 승인 금지
    with pytest.raises(ValidationError):
        services.decisions.confirm_decision(
            loaded,
            official.id,
            status=DecisionStatus.APPROVED,
            approver="PM",
            authority_confirmed=False,
        )
    # 공식 결정이 아닌 항목은 승인 대상이 아님
    with pytest.raises(ValidationError):
        services.decisions.confirm_decision(
            loaded,
            created[2].id,
            status=DecisionStatus.APPROVED,
            approver="PM",
            authority_confirmed=True,
        )
    # 전환은 승인된 공식 결정만
    with pytest.raises(ValidationError):
        services.decisions.decision_to_change_request(loaded, official.id, changes=services.changes)
    d = services.decisions.confirm_decision(
        loaded,
        official.id,
        status=DecisionStatus.APPROVED,
        approver="사업 책임자",
        authority_confirmed=True,
        comments="고객 PM 서면 확인",
    )
    assert d.is_binding and d.approved_by == "사업 책임자"
    cr = services.decisions.decision_to_change_request(
        loaded, official.id, changes=services.changes
    )
    assert cr.decision_ids == [official.id] and cr.status == ChangeStatus.REQUESTED
    assert official.id in {
        x
        for dd in services.decisions.decisions(loaded)
        for x in [dd.id]
        if dd.related_change_request_ids
    }
    # 기한 지난 미결은 overdue
    assert services.decisions.overdue_items(loaded, today=date(2026, 10, 25)) and not (
        services.decisions.overdue_items(loaded, today=date(2026, 10, 1))
    )


def test_meeting_agent_overlay(ctx: AppContext, build_project: Project) -> None:
    runner = FakeAgentRunner(
        {
            "meeting_analyzer": {
                "meeting_title": "회의",
                "items": [
                    {
                        "kind": "official_decision",
                        "title": "제안서 분량은 요약 모드로 한다",
                        "content": "요약 모드",
                        "authority_hint": "사업 책임자",
                        "evidence": "결정: 제안서 분량은 요약 모드로 한다",
                    },
                    {
                        "kind": "action_item",
                        "title": "성능 목표 재협의",
                        "owner": "기술 PL",
                        "due_date": "2026-11-01",
                        "evidence": "서버 사양 확정 후 재협의",
                    },
                ],
                "unresolved": ["모바일 범위는 회의록만으로 판단 불가"],
            }
        },
        cost_per_call=0.0,
    )
    s = Services.from_context(ctx, runner=runner)
    doc = s.documents.register_upload(
        build_project,
        kind=DocumentKind.MEETING_NOTE,
        file_name="회의록.txt",
        data=MEETING.encode("utf-8"),
        actor="t",
    )
    created, warnings = s.decisions.structure_meeting(build_project, doc, actor="t", use_agent=True)
    titles = [d.title for d in created]
    assert "성능 목표 재협의" in titles  # 에이전트 보강 항목
    assert sum(1 for t in titles if t.startswith("제안서 분량")) == 1  # 같은 제목은 합쳐짐
    assert any("모바일 범위" in w for w in warnings)
    assert all(not d.authority_confirmed for d in created)  # 권한 힌트가 있어도 자동 확인하지 않음


# --- 협상 → 변경 요청 → 기준선 반영 통제 -------------------------------------------------------------


def test_negotiation_control_to_baseline(services: Services, loaded: Project) -> None:
    n = services.decisions.create_negotiation(
        loaded,
        title="로그인 실패 잠금 횟수 완화",
        before="5회 실패 시 잠금",
        after="사용자는 아이디와 비밀번호로 로그인한다. 3회 실패 시 계정을 잠근다.",
        reason="고객 보안 정책",
        contract_impact="없음",
        related_requirement_ids=["SFR-001"],
        actor="t",
    )
    assert n.code == "NEG-001" and not n.can_apply_to_baseline
    with pytest.raises(ValidationError):  # 고객 합의 미확인 → 합의 불가
        services.decisions.settle_negotiation(
            loaded, n.id, status=NegotiationStatus.AGREED, customer_agreed=None, approver="PM"
        )
    with pytest.raises(ValidationError):  # 합의 전 전환 불가
        services.decisions.negotiation_to_change_request(loaded, n.id, changes=services.changes)
    with pytest.raises(ValidationError):  # '반영' 은 코드만 설정
        services.decisions.settle_negotiation(
            loaded, n.id, status=NegotiationStatus.APPLIED, customer_agreed=True, approver="PM"
        )
    n = services.decisions.settle_negotiation(
        loaded, n.id, status=NegotiationStatus.AGREED, customer_agreed=True, approver="PM"
    )
    assert n.can_apply_to_baseline
    cr = services.decisions.negotiation_to_change_request(loaded, n.id, changes=services.changes)
    assert cr.negotiation_id == n.id and cr.selected_requirement_ids == ["SFR-001"]
    cr = services.changes.analyze(loaded, cr, use_agent=False)
    assert cr.status == ChangeStatus.ANALYZED
    for c in cr.candidates:
        if c.classification == CandidateClass.NO_CHANGE:
            continue
        services.changes.decide_candidate(
            loaded,
            cr,
            c.id,
            decision="accept",
            classification="MODIFY" if c.classification == CandidateClass.AMBIGUOUS else None,
            overrides={"requirement_id": "SFR-001"} if not c.requirement_id else None,
        )
    cr = services.changes.get(loaded, cr.id)
    assert cr.status == ChangeStatus.PENDING_APPROVAL  # 8종 중 '승인 대기'
    assert services.changes.approval_blockers(cr) == []
    res = services.changes.approve(
        loaded, cr, approver="PM", comments="", regenerate_kinds=["requirement_spec_xlsx"]
    )
    assert res.request.status == ChangeStatus.APPLIED
    # 선택 재생성: 정의서만 다시 만들었다
    assert [d.kind for d in res.deliverables] == [DeliverableKind.REQUIREMENT_SPEC_XLSX]
    neg = services.ctx.repos.negotiations.get(n.id, loaded.id)
    assert neg.status == NegotiationStatus.APPLIED
    req = services.ctx.repos.requirements.get("SFR-001", loaded.id)
    assert "3회" in req.description and req.change_state == "modified"


def test_change_status_eight_and_labels() -> None:
    assert len(ChangeStatus) == 8 and len(CHANGE_STATUS_LABELS) == 8
    assert [s.value for s in CHANGE_STATUS_FLOW] == [
        "requested",
        "analyzed",
        "pending_approval",
        "approved",
        "applied",
    ]
    assert ChangeStatus.PENDING_APPROVAL.label == "승인 대기" and ChangeStatus.APPLIED.order == 4
    assert ChangeStatus.ERROR.order == len(CHANGE_STATUS_FLOW)


# --- 수정 문서 비교 (A-5) -------------------------------------------------------------------------


def _modified_spec() -> bytes:
    rows = [dict(r) for r in xb.BASE_REQUIREMENTS]
    rows[0]["name"] = "사용자 로그인 (계정 인증)"  # 문구만(명칭) 변경
    rows[1]["desc"] = rows[1]["desc"].replace(
        "첨부파일을 등록하여", "첨부파일(최대 3건)을 등록하여"
    )  # 수치 → 의미 변경
    deleted = rows.pop(2)  # 삭제
    rows.append(
        {
            "id": "SFR-099",
            "name": "접수 알림",
            "desc": "민원 접수 시 담당자에게 SMS 로 알린다.",
            "type": "기능",
            "sheet": "기능요구사항",
        }
    )
    assert deleted["id"] == "SFR-003"
    return xb.build_requirement_spec(rows)


def test_compare_spec_workbook_kinds() -> None:
    cands, warnings = compare_spec_workbook(_modified_spec(), _reqs(), file_name="수정정의서.xlsx")
    by_id = {c.requirement_id: c for c in cands if c.requirement_id}
    assert (
        by_id["SFR-001"].change_kind == KIND_WORDING
        and by_id["SFR-001"].classification == CandidateClass.MODIFY
    )
    assert by_id["SFR-001"].confidence == 0.9 and not by_id["SFR-001"].needs_confirmation
    assert (
        by_id["SFR-002"].change_kind == KIND_SEMANTIC
        and "number" in by_id["SFR-002"].change_aspects
    )
    assert (
        by_id["SFR-003"].change_kind == KIND_DELETE
        and by_id["SFR-003"].classification == CandidateClass.DELETE_REQUEST
    )
    added = [c for c in cands if c.classification == CandidateClass.ADD]
    assert len(added) == 1 and added[0].change_kind == KIND_ADD and "SFR-099" in added[0].rationale
    unchanged = [c for c in cands if c.classification == CandidateClass.NO_CHANGE]
    assert unchanged and all(c.user_decision == "reject" for c in unchanged)
    assert all(c.compare_source == "excel" and not c.uncertain for c in cands)
    assert change_kind_of(CandidateClass.MODIFY, ["screen"]) == KIND_SCREEN_MOVE
    assert change_kind_of(CandidateClass.MODIFY, ["screen", "rule"]) == KIND_SEMANTIC


def _pptx_bytes(lines: list[str]) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tb = slide.shapes.add_textbox(0, 0, prs.slide_width, prs.slide_height)
    tb.text_frame.text = "\n".join(lines)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_document_compare_flow_excel_then_pptx_uncertain(
    services: Services, loaded: Project
) -> None:
    cr = services.changes.create_request(
        loaded,
        input_kind=InputKind.DOCUMENT,
        request_text="",
        requester="고객 PM",
        reason="정의서 v1.1 수정본",
        attachments=[("수정정의서.xlsx", _modified_spec())],
        actor="t",
    )
    assert "ID 기준 비교" in cr.request_text
    cr = services.changes.analyze(loaded, cr, use_agent=False)
    assert "document" in cr.analysis_sources
    kinds = {
        c.requirement_id or "NEW": c.change_kind for c in cr.candidates if c.change_kind != "none"
    }
    assert kinds["SFR-001"] == KIND_WORDING and kinds["SFR-002"] == KIND_SEMANTIC
    assert kinds["SFR-003"] == KIND_DELETE and kinds["NEW"] == KIND_ADD
    # 영향 확장: 추적표의 화면·테스트, 연결 엔티티(API), 보안 검토
    services.ctx.repos.api_specs.insert(
        ApiSpec(
            project_id=loaded.id,
            code="API-001",
            name="접수 API",
            related_requirement_ids=["SFR-002"],
        )
    )
    cr = services.changes.analyze(loaded, cr, use_agent=False)
    imp = cr.impact
    assert imp is not None and imp.apis_affected == ["API-001"]
    assert imp.screens_affected  # 기준 추적표에서 SFR-002 화면
    assert imp.security_review_needed and any("API" in r for r in imp.security_reasons)
    # PPT 첨부: 판독 불확실 표시·신뢰도 상한
    cr2 = services.changes.create_request(
        loaded,
        input_kind=InputKind.DOCUMENT,
        request_text="",
        requester="고객 PM",
        reason="제안서 수정본",
        attachments=[
            ("수정제안서.pptx", _pptx_bytes(["민원 접수 시 담당자에게 SMS 알림을 발송한다."]))
        ],
        actor="t",
    )
    cr2 = services.changes.analyze(loaded, cr2, use_agent=False)
    assert cr2.candidates and all(
        c.uncertain and c.compare_source == "pptx" for c in cr2.candidates
    )
    assert all(c.confidence <= 0.6 and c.needs_confirmation for c in cr2.candidates)
    assert any("불확실" in u for u in cr2.unresolved)
    docs = classify_document_text(
        "1. 로그인 실패 3회로 변경", _reqs(), source="docx", file_name="w.docx"
    )
    assert docs[0].uncertain and docs[0].confidence <= 0.7 and docs[0].id.startswith("D")


# --- QA-CON-14 --------------------------------------------------------------------------------------


def test_qa_con_14_decision_control() -> None:
    today = date.today()
    good = SimpleNamespace(
        id="d1",
        code="DEC-001",
        kind=DecisionKind.OFFICIAL_DECISION,
        status=DecisionStatus.APPROVED,
        authority_confirmed=True,
        approved_by="PM",
        due_date=None,
        is_binding=True,
    )
    bad = SimpleNamespace(
        id="d2",
        code="DEC-002",
        kind=DecisionKind.OFFICIAL_DECISION,
        status=DecisionStatus.APPROVED,
        authority_confirmed=False,
        approved_by="",
        due_date=None,
        is_binding=False,
    )
    overdue = SimpleNamespace(
        id="d3",
        code="DEC-003",
        kind=DecisionKind.ACTION_ITEM,
        status=DecisionStatus.PROPOSED,
        authority_confirmed=False,
        approved_by="",
        due_date=today - timedelta(days=3),
        is_binding=False,
    )
    neg_ok = SimpleNamespace(
        id="n1",
        code="NEG-001",
        status=NegotiationStatus.APPLIED,
        customer_agreed=True,
        approved_by="PM",
    )
    neg_bad = SimpleNamespace(
        id="n2",
        code="NEG-002",
        status=NegotiationStatus.APPLIED,
        customer_agreed=None,
        approved_by="",
    )
    crs = [
        SimpleNamespace(id="CR-001", decision_ids=["d1"], negotiation_id="n1"),
        SimpleNamespace(id="CR-002", decision_ids=["d2"], negotiation_id="n2"),
        SimpleNamespace(id="CR-003", decision_ids=["dx"], negotiation_id=None),
    ]
    inp = SimpleNamespace(
        decisions=[good, bad, overdue], negotiations=[neg_ok, neg_bad], change_requests=crs
    )
    fails, revs = _decision_control_issues(inp)
    assert any("DEC-002: 권한 미확인" in f for f in fails) and any("NEG-002" in f for f in fails)
    assert any("CR-002: 원천 결정 DEC-002" in f for f in fails) and any(
        "CR-003: 원천 결정 dx 없음" in f for f in fails
    )
    assert not any("CR-001" in f for f in fails)
    assert revs and "DEC-003" in revs[0]


def test_traces_feed_impact(services: Services, loaded: Project) -> None:
    traces = services.ctx.repos.traces.find(loaded.id, requirement_id="SFR-002")
    assert any(t.target_type == TraceTargetType.SCREEN for t in traces) or isinstance(
        traces[0], RequirementTrace
    )
