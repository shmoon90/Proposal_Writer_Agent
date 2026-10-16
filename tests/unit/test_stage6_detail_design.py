"""통합 개발 지시문 단계 6 — 화면·DB·API·아키텍처 설계.

발주처 문맥 블록(⑤, D-21), 상세 설계 스키마 → 엔티티 저장(화면·테이블·API·모듈·규칙, 코드 발급·갱신·dropped, 추적 연결),
추적·누락·모순·DB 호환성 검사(②·③), 코드 생성 입력 패키지(D-22), 설계서 Excel/MD, 파이프라인 단계, 추적표 열 채움,
QA-TECH-09 구현 방안 구체성(⑥)·QA-CON-15."""

from __future__ import annotations

import io
import json
from datetime import date
from types import SimpleNamespace

from openpyxl import load_workbook

from proposal_agent.agents import FakeAgentRunner
from proposal_agent.agents.schemas import DetailDesignReport
from proposal_agent.design.consistency import design_consistency_issues
from proposal_agent.design.context import client_context, client_context_block
from proposal_agent.generators.excel.design_docs import (
    SHEET_APIS,
    SHEET_COLUMNS,
    SHEET_SCREEN_FIELDS,
    SHEET_TRACE,
    build_design_markdown,
    build_design_workbook,
)
from proposal_agent.generators.excel.style import HEADER_ROW
from proposal_agent.generators.excel.trace_matrix import SHEET_MAIN, build_trace_matrix
from proposal_agent.models import (
    AcceptanceStatus,
    ApiSpec,
    ArchitectureComponent,
    CodeModule,
    DbColumn,
    DbTable,
    DeliverableKind,
    Difficulty,
    Project,
    ProjectMode,
    ProjectSettings,
    Requirement,
    RequirementCategory,
    RequirementType,
    Screen,
)
from proposal_agent.orchestrator import StepRunner
from proposal_agent.services import AppContext, Services
from proposal_agent.services.pipeline_service import PROPOSAL_STEPS
from proposal_agent.validation.qa_checks.tech import _approach_concreteness_issues


def _project() -> Project:
    return Project(
        name="설계",
        mode=ProjectMode.PROPOSAL,
        output_dir="Project_s6",
        workspace_dir="ws_s6",
        start_date=date(2026, 11, 2),
        end_date=date(2027, 4, 30),
        created_by="t",
        client_name="테스트 발주처",
    )


def _req(p: Project, rid: str, name: str, desc: str = "", **kw: object) -> Requirement:
    data: dict[str, object] = {
        "id": rid,
        "project_id": p.id,
        "name": name,
        "description": desc or name,
        "type": RequirementType.FUNCTIONAL,
        "category": RequirementCategory.BUSINESS,
        "rfp_requirement_id": f"S{rid}",
        "created_by": "t",
        "acceptance_status": AcceptanceStatus.ACCEPTED,
        "implementation_approach": "- 접수 화면에서 입력, 실패 시 오류 표시",
        "difficulty": Difficulty("medium"),
    }
    data.update(kw)
    return Requirement(**data)  # type: ignore[arg-type]


REPORT = {
    "screens": [
        {
            "name": "민원 접수 화면",
            "screen_type": "form",
            "purpose": "민원을 접수한다",
            "user_roles": ["민원인"],
            "layout": {
                "description": "상단 입력, 하단 버튼",
                "regions": [{"name": "본문", "components": ["접수 폼"]}],
            },
            "fields": [
                {
                    "name": "title",
                    "label": "제목",
                    "field_type": "text",
                    "required": True,
                    "validation_rules": ["길이 1~100자"],
                    "error_message": "제목을 입력하십시오",
                    "requirement_ids": ["FR-001"],
                }
            ],
            "actions": [
                {
                    "label": "접수",
                    "description": "저장",
                    "expected_result": "상태 접수",
                    "api_ref": "민원 접수 API",
                    "requirement_ids": ["FR-001"],
                }
            ],
            "states": [{"state": "ERROR", "message": "접수 실패"}],
            "functions": [
                {
                    "name": "접수",
                    "trigger": "버튼",
                    "inputs": ["title"],
                    "process": "검증 → 저장 → 3초 이내 응답",
                    "outputs": ["접수번호"],
                    "requirement_ids": ["FR-001"],
                }
            ],
            "requirement_ids": ["FR-001"],
        }
    ],
    "tables": [
        {
            "name": "TB_CASE",
            "logical_name": "민원",
            "description": "민원 접수",
            "columns": [
                {
                    "name": "CASE_ID",
                    "data_type": "NUMBER",
                    "length": "12",
                    "nullable": False,
                    "is_pk": True,
                },
                {"name": "TITLE", "data_type": "VARCHAR2", "length": "100", "nullable": False},
                {
                    "name": "APPLICANT_NM",
                    "data_type": "VARCHAR2",
                    "length": "50",
                    "nullable": False,
                    "is_personal_data": True,
                },
                {
                    "name": "USER_ID",
                    "data_type": "VARCHAR2",
                    "length": "20",
                    "nullable": True,
                    "fk_ref": "TB_USER.USER_ID",
                },
            ],
            "retention_policy": "5년",
            "db_engine": "Oracle 11.2",
            "requirement_ids": ["FR-001", "SFR-002"],
        }
    ],
    "apis": [
        {
            "name": "민원 접수 API",
            "method": "POST",
            "path": "/api/cases",
            "description": "접수",
            "request_fields": ["title:string(1~100)"],
            "response": "{caseId}",
            "auth": "",
            "error_codes": ["C001: 제목 누락"],
            "screen_refs": ["민원 접수 화면"],
            "table_refs": ["TB_CASE"],
            "module_path": "src/case/CaseController.java",
            "requirement_ids": ["FR-001"],
        }
    ],
    "modules": [
        {
            "path": "src/case/CaseController.java",
            "layer": "controller",
            "responsibility": "접수 API",
            "depends_on": ["src/case/CaseService.java"],
            "requirement_ids": ["FR-001"],
            "api_refs": ["민원 접수 API"],
            "screen_refs": ["민원 접수 화면"],
            "table_refs": [],
        },
        {
            "path": "src/case/CaseService.java",
            "layer": "service",
            "responsibility": "접수 규칙",
            "depends_on": [],
            "requirement_ids": ["FR-001", "FR-002"],
            "table_refs": ["TB_CASE"],
        },
    ],
    "business_rules": [
        {
            "rule_id": "BR-001",
            "description": "제목 없는 민원은 접수하지 않는다",
            "enforced_in": "src/case/CaseService.java",
            "requirement_ids": ["FR-001"],
        }
    ],
    "assumptions": ["VARCHAR2 길이 가정"],
    "notes": "",
}


def test_client_context_block_uses_analysis_and_settings() -> None:
    p = _project()
    analysis = {
        "overview": "민원 전산화",
        "current_issues": ["수기 처리 누락", "이메일 접수"],
        "integration_targets": ["알림톡"],
        "named_solutions": [],
    }
    st = ProjectSettings(
        has_existing_system=True,
        existing_system_note="2015년 구축 ASP 시스템",
        integration_systems=["세움터"],
    )
    ctx = client_context(analysis, st)
    assert ctx["integration_targets"] == ["알림톡", "세움터"] and ctx["existing_system"]["exists"]
    block = client_context_block(p, analysis, st)
    assert "<client_context" in block and "수기 처리 누락" in block and "2015년 구축 ASP" in block
    assert "명시된 솔루션·기술: (없음)" in block  # 없는 항목은 추정하지 않도록 명시
    assert "(없음)" in client_context_block(p, None, None)


def test_schema_and_consistency_rules() -> None:
    rep = DetailDesignReport.model_validate(REPORT)
    assert rep.modules[0].name == "CaseController.java" and rep.apis[0].error_codes
    p = _project()
    reqs = [
        _req(p, "FR-001", "민원 접수"),
        _req(p, "FR-002", "배정"),
        _req(p, "FR-003", "통계", is_mandatory=True),
    ]
    screens = [
        Screen(
            project_id=p.id,
            code="SCR-001",
            name="민원 접수 화면",
            related_requirement_ids=["FR-001"],
            states=[{"state": "ERROR", "message": "x"}],
        )
    ]
    tables = [
        DbTable(
            project_id=p.id,
            code="TBL-001",
            name="TB_CASE",
            db_engine="Oracle 11.2",
            related_requirement_ids=["FR-001"],
            columns=[
                DbColumn(name="APPLICANT_NM", is_personal_data=True),
                DbColumn(name="USER_ID", fk_ref="TB_USER.USER_ID"),
            ],
        ),
    ]
    apis = [
        ApiSpec(
            project_id=p.id,
            code="API-001",
            name="민원 접수 API",
            auth="",
            error_codes=[],
            related_requirement_ids=["FR-001", "FR-099"],
            related_screen_ids=["SCR-001"],
            related_table_ids=["TBL-001", "TBL-404"],
        )
    ]
    modules = [
        CodeModule(
            project_id=p.id,
            code="MOD-001",
            name="A",
            path="a.java",
            depends_on=["b.java"],
            related_requirement_ids=["FR-002"],
        ),
        CodeModule(project_id=p.id, code="MOD-002", name="B", path="b.java", depends_on=["a.java"]),
    ]
    fails, revs = design_consistency_issues(
        reqs,
        screens,
        tables,
        apis,
        modules,
        component_techs=["PostgreSQL 15"],
        research_stack=[],
        rfp_specified_tech=[],
    )
    joined = "\n".join(fails)
    assert "FK 참조 테이블 TB_USER 없음" in joined and "FR-099" in joined and "TBL-404" in joined
    assert "순환" in joined and "FR-003 통계: 화면·API·테이블 연결 없음" in joined
    assert any("FR-002" in f for f in fails)  # 모듈 연결만 있는 필수 기능 요구도 실패
    rj = "\n".join(revs)
    assert "개인정보 테이블을 다루는데 인증·권한이 비어 있음" in rj and "오류 코드 없음" in rj
    assert "Oracle 11.2" in rj and "postgres" in rj  # DB 종류가 아키텍처와 다름 → 호환성 검토


def test_persist_entities_traces_codegen_package_and_docs(
    services: Services, proposal_project: Project
) -> None:
    p = proposal_project
    repos = services.ctx.repos
    for rid, name in (("FR-001", "민원 접수"), ("FR-002", "배정")):
        repos.requirements.insert(_req(p, rid, name))
    repos.components.insert(
        ArchitectureComponent(
            project_id=p.id, id="DB01", layer="data", name="DB", technology="Oracle 11.2"
        )
    )
    dd = services.detail_design
    counts = dd.persist(p, REPORT, actor="t")
    assert counts == {"screens": 1, "tables": 1, "apis": 1, "modules": 2, "rules": 1, "traces": 1}
    s = dd.screens(p)[0]
    assert (
        s.code == "SCR-001"
        and s.source == "design"
        and s.fields[0]["validation_rules"] == ["길이 1~100자"]
    )
    t = dd.tables(p)[0]
    assert (
        t.code == "TBL-001"
        and t.columns[2].is_personal_data
        and t.related_requirement_ids == ["FR-001", "FR-002"]
    )  # SFR-002 → FR-002
    a = dd.apis(p)[0]
    assert (
        a.related_screen_ids == ["SCR-001"]
        and a.related_table_ids == ["TBL-001"]
        and a.request_schema["fields"]
    )
    mods = dd.modules(p)
    assert [m.code for m in mods] == ["MOD-001", "MOD-002"] and mods[0].depends_on == [
        "src/case/CaseService.java"
    ]
    assert dd.rules(p)[0].name == "BR-001"
    traces = repos.traces.find(p.id, requirement_id="FR-001")
    assert any(tr.target_id == s.id for tr in traces)
    # 재저장: 같은 이름은 갱신(버전 ↑), 빠진 것은 dropped
    rep2 = json.loads(json.dumps(REPORT))
    rep2["screens"][0]["purpose"] = "민원을 접수하고 접수번호를 발급한다"
    rep2["apis"] = []
    dd.persist(p, rep2, actor="t")
    assert dd.screens(p)[0].purpose.endswith("발급한다") and dd.screens(p)[0].version == 2
    assert dd.apis(p) == [] and repos.api_specs.find(p.id)[0].status == "dropped"
    fails, revs = dd.consistency(p)
    assert any("FK 참조 테이블 TB_USER 없음" in f for f in fails)
    # 코드 생성 패키지 (D-22)
    pkg = dd.codegen_package(p)
    assert (
        pkg["ui_design"]["screens"][0]["screen_id"] == "SCR-001"
        and pkg["architecture"]["modules"][0]["path"] == "src/case/CaseController.java"
    )
    assert (
        pkg["requirements"]["requirements"][0]["requirement_id"] == "FR-001"
        and pkg["architecture"]["test_plan"] == []
    )
    rel = dd.export_codegen_package(p, actor="t")
    assert rel.endswith(".json") and services.ctx.paths.resolve_output(rel).exists()
    # 설계서 Excel·MD
    args = (
        repos.requirements.find(p.id, order_by="id"),
        dd.screens(p),
        dd.tables(p),
        dd.apis(p),
        dd.modules(p),
        dd.rules(p),
    )
    wb = load_workbook(io.BytesIO(build_design_workbook(*args)))
    assert (
        SHEET_SCREEN_FIELDS in wb.sheetnames
        and SHEET_COLUMNS in wb.sheetnames
        and SHEET_APIS in wb.sheetnames
    )
    rows = list(wb[SHEET_TRACE].iter_rows(min_row=2, values_only=True))  # 설계서는 머리글 1행
    assert rows[0][0] == "FR-001" and rows[0][3] == "SCR-001" and rows[1][3] == "미정"
    md = build_design_markdown(p.name, *args, issues=dd.consistency(p))
    assert "## 2. DB 설계" in md and "TB_CASE" in md and "## 6. 추적·모순 검사" in md
    # 추적표의 설계 열이 채워진다 (단계 4 '미정' 해소)
    ctx = services.generation.build_context(p)
    ws = load_workbook(
        io.BytesIO(build_trace_matrix(ctx, repos.traces.find(p.id), dd.screens(p), [], []))
    )[SHEET_MAIN]
    head = [c.value for c in ws[HEADER_ROW]]
    row = next(
        r
        for r in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True)
        if r[head.index("요구사항 ID")] == "FR-001"
    )
    assert (
        row[head.index("DB 테이블 ID")] == "TBL-001"
        and row[head.index("코드 모듈 ID")] == "MOD-001, MOD-002"
    )
    paths = services.generation.generate_design_docs(p, actor="t")
    assert (
        len(paths) == 2
        and any(x.endswith(".xlsx") for x in paths)
        and any(x.endswith(".md") for x in paths)
    )
    kinds = {d.kind for d in services.ctx.repos.deliverables.find(p.id)}
    assert DeliverableKind.DESIGN_DOC_XLSX in kinds and DeliverableKind.DESIGN_DOC_MD in kinds


def test_pipeline_step_and_runner_with_fake_agent(
    ctx: AppContext, proposal_project: Project
) -> None:
    keys = [s.key for s in PROPOSAL_STEPS]
    assert keys.index("detail_design") == keys.index("impl_design") + 1
    runner = FakeAgentRunner({"detail_designer": REPORT}, cost_per_call=0.0)
    s = Services.from_context(ctx, runner=runner)
    p = proposal_project
    s.ctx.repos.requirements.insert(_req(p, "FR-001", "민원 접수"))
    r = StepRunner(s)
    try:
        outcome = r.run_detail_design(p, actor="t")
    except Exception as exc:  # noqa: BLE001 — 게이트 1 미승인 환경이면 메시지로 확인
        assert "gate1" in str(exc).lower() or "게이트" in str(exc)
        return
    assert "저장: 화면 1" in outcome.summary and s.detail_design.has_design(p)
    assert "detail_design" in StepRunner.DISPATCHABLE


def test_qa_tech_09_concreteness() -> None:
    p = _project()
    r_ok = _req(
        p,
        "FR-001",
        "민원 접수",
        "민원을 3초 이내에 접수한다",
        implementation_approach="- 민원 접수 화면에서 입력, 3초 이내 응답, 실패 시 오류 표시·재시도",
    )
    r_bad = _req(
        p,
        "FR-002",
        "배정",
        "담당자를 24시간 이내 배정한다",
        implementation_approach="- 자동 배정한다",
    )
    screens = [
        Screen(
            project_id=p.id,
            code="SCR-001",
            name="민원 접수 화면",
            related_requirement_ids=["FR-001"],
            source="design",
        )
    ]
    inp = SimpleNamespace(
        requirements=[r_ok, r_bad],
        screens=screens,
        db_tables=[],
        api_specs=[],
        ctx=SimpleNamespace(designs=[]),
    )
    fails, revs, ids = _approach_concreteness_issues(inp)  # type: ignore[arg-type]
    assert fails == ["FR-002: 구현 방안·설계에 실제 화면·테이블·API 이름이 없음"]
    assert any("24시간" in x for x in revs) and any(
        "FR-002: 예외·실패 처리 서술 없음" in x for x in revs
    )
    assert not any("FR-001" in x for x in revs) and ids == ["FR-002"]
    empty = SimpleNamespace(
        requirements=[r_ok], screens=[], db_tables=[], api_specs=[], ctx=SimpleNamespace(designs=[])
    )
    f2, r2, _ = _approach_concreteness_issues(empty)  # type: ignore[arg-type]
    assert f2 == [] and r2[0].startswith("설계 객체(화면·테이블·API) 없음")
