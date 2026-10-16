"""Phase 12: 제안서 템플릿 입력·프로파일, RFP 지정 목차 트리, 장 계획, 템플릿 기반 제출용 PPT, 템플릿 준수 QA, 제안 문안 에이전트 연동."""

from __future__ import annotations

import json
import re
from collections.abc import Iterator
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.util import Emu

from fixtures import builders as bld
from proposal_agent.agents import AgentRequest, FakeAgentRunner
from proposal_agent.config import ProjectPaths, Settings, ValidationError
from proposal_agent.generators.pptx.outline import OutlineSection, plan_chapters
from proposal_agent.generators.pptx.template import TemplateDeckBuilder, profile_template
from proposal_agent.ingest.rfp_profile import profile_blocks
from proposal_agent.models import (
    BlockLocation,
    BlockType,
    DeliverableKind,
    DocumentKind,
    ExtractedBlock,
    FileFormat,
    Project,
    ProjectMode,
    QaResult,
)
from proposal_agent.orchestrator import StepRunner
from proposal_agent.services import AppContext, Services
from proposal_agent.validation import pptx_texts
from unit.test_requirements_phase2 import verifier_all_ok
from unit.test_rfp_profile import _extract_only

WRITER_SECTIONS: list[dict[str, Any]] = [
    {
        "slide_group": "understanding",
        "title": "사업의 개요",
        "chapter": "Ⅰ. 제안 개요",
        "section": "1. 사업 이해도",
        "number": "1.1",
        "lead_message": "리드: 민원 사건관리 시스템을 안정적으로 구축한다",
        "bullets": ["사건 접수·배정·종결 관리 신규 구축", "통계 조회 및 접근 통제"],
        "requirement_ids": ["SFR-001", "SFR-002"],
        "footnotes": ["세부 범위는 발주처와 협의"],
        "data_refs": ["SFR-001"],
    },
    {
        "slide_group": "generic",
        "title": "장비 구성 내역",
        "chapter": "Ⅱ. 사업 수행 방안",
        "section": "1. 시스템 구성 방안",
        "number": "1.1",
        "lead_message": "리드: 장비는 2중화 구성",
        "table_headers": ["구분", "주요기능", "비고"],
        "table_rows": [["WAS", "업무 처리", "2대"], ["DB", "데이터 저장", "2대"]],
    },
    {
        "slide_group": "schedule",
        "title": "추진 일정",
        "chapter": "Ⅲ. 사업 관리 방안",
        "number": "2.1",
        "lead_message": "리드: 8개월 일정",
    },
]


def _researcher(request: AgentRequest) -> dict[str, Any]:
    return {
        "items": [
            {
                "topic": "사건 접수·배정 워크플로 구현",
                "related_requirement_ids": ["SFR-001", "SFR-002"],
                "summary": "워크플로 엔진 없이 상태 전이 테이블로 구현 가능",
                "confirmed_facts": ["상태 전이 방식은 공식 가이드에 예시가 있음"],
                "applicability_conditions": [],
                "unverified_items": [],
                "impact_on_implementation": "상태 전이 테이블 + 배정 알림 배치로 구현",
                "impact_on_proposal": "",
                "feasibility": "feasible",
                "risks": [],
                "recommended_stack": ["Spring Boot"],
                "sources": [
                    {
                        "url": "https://docs.example.com/workflow",
                        "title": "Workflow guide",
                        "product": "Spring Boot",
                        "version": "3.2",
                        "trust": "official_doc",
                        "excerpt": "state transition example",
                    }
                ],
            }
        ],
        "comparisons": [],
        "notes": "",
    }


def _writer(request: AgentRequest) -> dict[str, Any]:
    return {
        "profile": "submission",
        "sections": WRITER_SECTIONS,
        "key_messages": ["안정 구축", "보안 준수", "일정 준수"],
        "unverified_claims": [],
    }


@pytest.fixture
def tpl(tmp_path: Path) -> Iterator[tuple[Services, Project, FakeAgentRunner]]:
    ctx = AppContext.create(Settings(), paths=ProjectPaths(root=tmp_path), configure_logging=False)
    runner = FakeAgentRunner(
        {
            "requirement_extractor": _extract_only(set(bld.STRUCTURED_RFP_IDS)),
            "requirement_verifier": verifier_all_ok,
            "proposal_writer": _writer,
            "tech_researcher": _researcher,
        }
    )
    s = Services.from_context(ctx, runner=runner)
    p = s.projects.create(
        name="템플릿 제안",
        mode=ProjectMode.PROPOSAL,
        start_date=date(2026, 11, 2),
        end_date=date(2027, 6, 30),
        client_name="가상 협회",
    )
    s.documents.register_upload(
        p, kind=DocumentKind.RFP, file_name="rfp.docx", data=bld.build_structured_rfp_docx()
    )
    s.documents.register_upload(
        p,
        kind=DocumentKind.PROPOSAL_TEMPLATE,
        file_name="고객양식_제출용.pptx",
        data=bld.build_template_pptx(),
    )
    r = StepRunner(s)
    r.run_extraction(p)
    r.run_requirement_extraction(p)
    r.run_verification(p)
    yield s, p, runner
    ctx.close()


def _slides_from_bytes(data: bytes) -> list[list[str]]:
    import io

    prs = Presentation(io.BytesIO(data))
    return [
        [
            sh.text_frame.text.strip()
            for sh in sl.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()
        ]
        for sl in prs.slides
    ]


def _slides(path: Path) -> list[list[str]]:
    prs = Presentation(str(path))
    return [
        [
            sh.text_frame.text.strip()
            for sh in sl.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()
        ]
        for sl in prs.slides
    ]


# --- 프로파일 ---------------------------------------------------------------------------------
def test_template_profile_reads_layout_headers_lead_fonts_and_chapters() -> None:
    prof = profile_template(bld.build_template_pptx(), file_name="t.pptx")
    assert prof.orientation == "portrait" and prof.purpose == "submission" and prof.slide_count == 7
    assert prof.chapter_layout == "Title Slide" and prof.content_layout == "Blank"
    assert len(prof.headers) == 2
    assert abs(prof.headers[0].top - 3.8) < 0.3 and abs(prof.headers[1].top - 4.5) < 0.3
    assert prof.font_sizes["section"] == 10 and prof.font_sizes["subsection"] == 13
    assert prof.lead is not None and prof.font_sizes["lead"] == 11 and prof.font_sizes["body"] == 9
    assert prof.fonts == {"body": bld.TEMPLATE_FONT, "title": bld.TEMPLATE_FONT}
    assert prof.footnote is not None and prof.footnote.top > 27
    assert prof.table_style["header_fill"] == "1F4E78" and prof.table_patterns == [
        "구분 | 위험요소 | 대처 방안"
    ]
    assert prof.chapter_title_box is not None and prof.chapter_numeral_box is not None
    assert prof.colors["primary"] == "313D83"
    assert [c.numeral for c in prof.chapters] == ["Ⅰ"]
    assert prof.chapters[0].sections == ["1. 사업 이해도", "2. 추진 전략"]
    assert prof.chapters[0].subsections == [
        "1.1 사업의 개요",
        "1.2 사업의 범위",
        "2.1 추진 전략",
        "2.2 핵심 성공 요소",
    ]
    assert prof.body.top > prof.headers[1].bottom and prof.body.bottom < 27.5
    assert any("(계속)" in c for c in prof.conventions) and any(
        "리드 메시지" in c for c in prof.conventions
    )
    assert prof.warnings == []
    # 가로 템플릿은 발표용으로 추정
    assert profile_template(bld.build_template_pptx(portrait=False)).purpose == "presentation"


def test_template_upload_rules_and_extraction_skip(
    tpl: tuple[Services, Project, FakeAgentRunner],
) -> None:
    s, p, _ = tpl
    docs = s.documents.list_documents(p.id)
    tdoc = next(d for d in docs if d.kind == DocumentKind.PROPOSAL_TEMPLATE)
    assert tdoc.is_required_input is False and tdoc.template_purpose is None
    # 추출 실행 기록에 템플릿이 없다 (추출 대상 아님)
    assert tdoc.id not in s.extraction.latest_runs(p)
    hit = s.templates.for_purpose(p, "submission")
    assert hit is not None and hit.document.id == tdoc.id and hit.profile.purpose == "submission"
    assert s.templates.for_purpose(p, "presentation") is None
    with pytest.raises(ValidationError, match="submission 또는 presentation"):
        s.documents.register_upload(
            p,
            kind=DocumentKind.PROPOSAL_TEMPLATE,
            file_name="x.pptx",
            data=bld.build_template_pptx(portrait=False),
            template_purpose="deck",
        )
    with pytest.raises(Exception, match="허용되지 않는 형식"):
        s.documents.register_upload(
            p, kind=DocumentKind.PROPOSAL_TEMPLATE, file_name="x.docx", data=b"PK\x03\x04 no"
        )
    # 용도 지정 업로드
    doc2 = s.documents.register_upload(
        p,
        kind=DocumentKind.PROPOSAL_TEMPLATE,
        file_name="발표.pptx",
        data=bld.build_template_pptx(portrait=False),
        template_purpose="presentation",
    )
    assert doc2.template_purpose == "presentation"
    assert s.templates.for_purpose(p, "presentation") is not None


# --- RFP 지정 목차 트리 · 장 계획 -------------------------------------------------------------------
def _blk(
    order: int, text: str, kind: BlockType, *, cells: list[list[str]] | None = None
) -> ExtractedBlock:
    return ExtractedBlock(
        project_id="prj_x",
        source_document_id="doc_x",
        file_name="rfp.pdf",
        file_format=FileFormat.PDF,
        order=order,
        block_type=kind,
        text=text,
        table_cells=cells,
        location=BlockLocation(page=32),
    )


def test_toc_tree_groups_sections_even_when_chapter_rows_follow() -> None:
    blocks = [
        _blk(1, "Ⅴ. 제안서 작성요령", BlockType.HEADING),
        _blk(2, "4 제안서 목차 및 작성방법", BlockType.HEADING),
        _blk(3, "1. 제안배경", BlockType.LIST_ITEM),
        _blk(4, "○ 제안의 배경과 목적을 기술", BlockType.LIST_ITEM),
        _blk(5, "2. 제안전략 ○ 전략을 요약하여 기술", BlockType.LIST_ITEM),
        _blk(6, "1. 일반현황", BlockType.LIST_ITEM),
        _blk(7, "2. 사업수행 실적", BlockType.LIST_ITEM),
        _blk(8, "- 29 -", BlockType.LIST_ITEM),
        _blk(
            9,
            "작성 목차 | 작성 방법\nⅠ. 제안개요 |",
            BlockType.TABLE,
            cells=[["작성 목차", "작성 방법"], ["Ⅰ. 제안개요", ""]],
        ),
        _blk(10, "Ⅱ. 제안업체 일반 |", BlockType.TABLE, cells=[["Ⅱ. 제안업체 일반", ""]]),
    ]
    prof = profile_blocks(blocks)
    tree = prof.proposal_toc_tree
    assert [c["numeral"] for c in tree] == ["Ⅰ", "Ⅱ"]  # RFP 자체 장 제목(Ⅴ)은 제외
    assert [s["title"] for s in tree[0]["sections"]] == ["제안배경", "제안전략"]
    assert tree[0]["sections"][0]["guide"] == ["제안의 배경과 목적을 기술"]
    assert tree[0]["sections"][1]["guide"] == ["전략을 요약하여 기술"]
    assert [s["title"] for s in tree[1]["sections"]] == ["일반현황", "사업수행 실적"]
    assert all("29" not in g for s in tree[1]["sections"] for g in s["guide"])


def test_plan_chapters_places_outline_and_keeps_rfp_sections(
    tpl: tuple[Services, Project, FakeAgentRunner],
) -> None:
    s, p, _ = tpl
    tree = s.extraction.rfp_proposal_toc_tree(p)
    assert [c["title"] for c in tree] == bld.STRUCTURED_RFP_TOC
    assert [x["title"] for x in tree[0]["sections"]] == bld.STRUCTURED_RFP_TOC_SECTIONS["제안 개요"]
    outline = [
        OutlineSection(key="purpose", title="사업 목적 및 이해", source="default"),
        OutlineSection(key="requirements", title="요구사항 대응", source="default"),
        OutlineSection(key="schedule", title="추진 일정", source="default"),
        OutlineSection(key="organization", title="추진 조직 및 인력", source="default"),
        OutlineSection(key="risk", title="위험 관리", source="default"),
        OutlineSection(
            key="generic_0",
            title="장비 구성",
            source="agent",
            chapter="Ⅱ. 사업 수행 방안",
            number="1.1",
        ),
    ]
    plan = plan_chapters(outline, tree, None)
    assert [c.numeral for c in plan] == ["Ⅰ", "Ⅱ", "Ⅲ", "Ⅳ"]
    by_title = {sec.title: [o.key for o in sec.subsections] for ch in plan for sec in ch.sections}
    assert by_title["사업 이해도"] == ["purpose"]
    assert by_title["요구사항 대응 방안"] == ["requirements"]
    assert by_title["시스템 구성 방안"] == ["generic_0"]  # 에이전트 지정 장·절 번호
    assert by_title["추진 일정"] == ["schedule"]
    assert by_title["조직 및 인력"] == ["organization"]
    assert by_title["교육 훈련"] == ["placeholder"] and by_title["유지보수"] == ["placeholder"]
    # 위험 관리: 맞는 절이 없으면 '관리' 장에 새 절
    risk_ch = next(
        ch for ch in plan if any(o.key == "risk" for sec in ch.sections for o in sec.subsections)
    )
    assert risk_ch.numeral == "Ⅲ"
    # 목차 트리가 없으면 단일 장
    single = plan_chapters(outline[:2], None, None)
    assert len(single) == 1 and [sec.title for sec in single[0].sections] == [
        "사업 목적 및 이해",
        "요구사항 대응",
    ]


# --- 생성 · QA -------------------------------------------------------------------------------------
def test_submission_deck_follows_template_and_rfp_toc_and_passes_template_qa(
    tpl: tuple[Services, Project, FakeAgentRunner],
) -> None:
    s, p, _ = tpl
    r = StepRunner(s)
    r.run_generation(p)
    latest = s.generation.latest_deliverables(p)
    path = s.ctx.paths.resolve_output(latest[DeliverableKind.SUBMISSION_PPTX].file_path)
    prs = Presentation(str(path))
    assert abs(Emu(prs.slide_width).cm - 21.0) < 0.1 and abs(Emu(prs.slide_height).cm - 29.7) < 0.1
    slides = _slides(path)
    numerals = [t for sl in slides for t in sl if re.fullmatch(r"[ⅠⅡⅢⅣ]", t)]
    assert numerals == ["Ⅰ", "Ⅱ", "Ⅲ", "Ⅳ"]
    text = "\n".join(pptx_texts(path))
    for title in bld.STRUCTURED_RFP_TOC:
        assert title in text
    for secs in bld.STRUCTURED_RFP_TOC_SECTIONS.values():
        for sec in secs:
            assert sec in text
    assert "사람 작성 필요" in text and "RFP 작성방법" in text  # 교육 훈련·유지보수 자리표시
    # 본문 슬라이드 머리글 2단 ("N. 절" / "N.M 소절"), 이어지는 슬라이드 "(계속)"
    content = [
        sl
        for sl in slides[1:]
        if sl and not sl[0].startswith("목차") and not any(re.fullmatch(r"[ⅠⅡⅢⅣ]", t) for t in sl)
    ]
    assert content and all(
        re.match(r"\d+\. ", sl[0]) and re.match(r"\d+\.\d+ ", sl[1]) for sl in content
    )
    # 글머리가 한 장 용량을 넘으면 같은 소절 번호로 "(계속)" 슬라이드
    hit = s.templates.for_purpose(p, "submission")
    assert hit is not None
    tb = TemplateDeckBuilder(hit.data, hit.profile)
    tb.begin_chapter("Ⅰ", "제안 개요", ["사업 이해도"])
    tb.begin_section("사업 이해도")
    assert tb.bullet_slides("사업의 개요", [f"글머리 {i}" for i in range(40)], "", lead="리드") >= 3
    cont = _slides_from_bytes(tb.to_bytes())
    assert cont[2][1] == "1.1 사업의 개요 (계속)" and cont[1][1] == "1.1 사업의 개요"
    for req in s.requirements.list_requirements(p):
        assert req.id in text
    # 템플릿 준수 QA
    run = s.qa.run(p, max_fix_rounds=0).qa_run
    assert run is not None
    by_id = {c.check_id: c for c in run.checks}
    assert by_id["QA-FILE-11"].executed and by_id["QA-FILE-11"].result is QaResult.PASS, by_id[
        "QA-FILE-11"
    ].issues
    assert by_id["QA-FILE-02"].result is QaResult.PASS, by_id[
        "QA-FILE-02"
    ].issues  # 템플릿 용지(세로)에 맞음
    assert by_id["QA-FILE-03"].result is not QaResult.FAIL, by_id["QA-FILE-03"].issues
    outline_json = sorted(
        (s.ctx.paths.project_output_dir(p.output_dir) / "07_Submission_Proposal_PPT").glob(
            "목차결정_*.json"
        )
    )
    payload = json.loads(outline_json[-1].read_text(encoding="utf-8"))
    assert payload["templates"] == {"submission": "고객양식_제출용.pptx"}
    assert [c["numeral"] for c in payload["chapters"]] == ["Ⅰ", "Ⅱ", "Ⅲ", "Ⅳ"]
    assert not any("기본 양식(16:9)" in n for n in payload["review_items"])


def test_writer_agent_gets_template_brief_and_deck_uses_lead_table_footnotes(
    tpl: tuple[Services, Project, FakeAgentRunner],
) -> None:
    s, p, runner = tpl
    r = StepRunner(s)
    s.strategy.evaluate(p)  # Phase 18: 목차·전략(Gate P) 승인 후에만 문안 작성
    s.gates.approve_human_review(p, "gateP", approver="pm", comments="목차 확인")
    with pytest.raises(ValidationError, match="기술 조사 기록이 없어"):
        r.run_proposal_writing(p)  # 구현 가능한 내용만 쓰도록 조사(웹 검색) 선행 강제
    r.run_tech_research(p)
    r.run_proposal_writing(p)
    # 조사 결과가 요구사항 구현 방안을 채운다 (근거 표기)
    sfr1 = next(r for r in s.requirements.list_requirements(p) if r.rfp_requirement_id == "SFR-001")
    assert sfr1.implementation_approach and "근거: 조사" in sfr1.implementation_approach
    call = next(c for c in runner.calls if c.role == "proposal_writer")
    for tag in ("<proposal_template>", "<rfp_toc>", "<data>", "<feasibility>"):
        assert tag in call.prompt
    assert "SFR-001" in call.prompt.split("<feasibility>")[1].split("</feasibility>")[0]
    assert (
        "머리글 2단" in call.prompt
        and "Ⅰ. 제안 개요" in call.prompt
        and '"rfp_id": "SFR-001"' in call.prompt
    )
    assert "1. 사업 이해도 — 작성방법: 사업 이해도에 대해 기술" in call.prompt
    r.run_generation(p)
    latest = s.generation.latest_deliverables(p)
    path = s.ctx.paths.resolve_output(latest[DeliverableKind.SUBMISSION_PPTX].file_path)
    text = "\n".join(pptx_texts(path))
    assert "리드: 민원 사건관리 시스템을 안정적으로 구축한다" in text
    assert "대응 요구사항: SFR-001, SFR-002" in text and "※ 세부 범위는 발주처와 협의" in text
    assert "구분 | 주요기능 | 비고" in text and "WAS | 업무 처리 | 2대" in text
    slides = _slides(path)
    lead_slide = next(sl for sl in slides if any("리드: 장비는 2중화 구성" in t for t in sl))
    assert lead_slide[0].startswith("1. 시스템 구성 방안") and lead_slide[1].startswith(
        "1.1 장비 구성 내역"
    )
    run = s.qa.run(p, max_fix_rounds=0).qa_run
    assert run is not None
    assert next(c for c in run.checks if c.check_id == "QA-FILE-11").result is QaResult.PASS


def test_without_template_default_deck_and_qa_not_run(tmp_path: Path) -> None:
    ctx = AppContext.create(Settings(), paths=ProjectPaths(root=tmp_path), configure_logging=False)
    runner = FakeAgentRunner(
        {
            "requirement_extractor": _extract_only(set(bld.STRUCTURED_RFP_IDS)),
            "requirement_verifier": verifier_all_ok,
        }
    )
    s = Services.from_context(ctx, runner=runner)
    p = s.projects.create(
        name="템플릿 없음",
        mode=ProjectMode.PROPOSAL,
        start_date=date(2026, 11, 2),
        end_date=date(2027, 6, 30),
    )
    s.documents.register_upload(
        p, kind=DocumentKind.RFP, file_name="rfp.docx", data=bld.build_structured_rfp_docx()
    )
    r = StepRunner(s)
    r.run_extraction(p)
    r.run_requirement_extraction(p)
    r.run_verification(p)
    result = s.generation.generate_all(p)
    assert any("기본 양식(16:9)" in n for n in result.review_items)
    path = s.ctx.paths.resolve_output(
        s.generation.latest_deliverables(p)[DeliverableKind.SUBMISSION_PPTX].file_path
    )
    prs = Presentation(str(path))
    assert prs.slide_width > prs.slide_height
    text = "\n".join(pptx_texts(path))
    assert "Ⅰ. 제안 개요" in text and "교육 훈련" in text  # 템플릿이 없어도 RFP 지정 목차는 지킨다
    run = s.qa.run(p, max_fix_rounds=0).qa_run
    assert run is not None
    c11 = next(c for c in run.checks if c.check_id == "QA-FILE-11")
    assert c11.applicable is False and "미등록" in c11.not_run_reason
    ctx.close()
