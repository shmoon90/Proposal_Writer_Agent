"""generators·validation: Excel 4종 구조·서식, PPT 2종 검사, 구성도 PNG, 목차 결정, 일관성·추적성 검사, 최종본 분리, Gate 차단."""

from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pptx import Presentation

from fixtures import builders as bld
from proposal_agent.agents import FakeAgentRunner
from proposal_agent.config import ValidationError
from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.excel.requirement_spec import (
    REQUIRED_COLUMNS,
    SHEET_FUNCTIONAL,
    SHEET_NON_FUNCTIONAL,
)
from proposal_agent.generators.pptx.outline import (
    decide_outline,
    default_outline,
    rfp_toc_from_blocks,
)
from proposal_agent.models import (
    AcceptanceStatus,
    DeliverableKind,
    DocumentKind,
    GateStatus,
    IntegrityStatus,
    Project,
    ProjectMode,
    QaStatus,
    StepStatus,
    ValidationLevel,
)
from proposal_agent.orchestrator import GateBlockedError, StepRunner
from proposal_agent.services import AppContext, Services
from proposal_agent.validation import QaOutcome, check_consistency, check_pptx, pptx_texts
from unit.test_requirements_phase2 import extractor_from_text, verifier_all_ok


@pytest.fixture
def services(ctx: AppContext) -> Services:
    return Services.from_context(
        ctx,
        runner=FakeAgentRunner(
            {"requirement_extractor": extractor_from_text, "requirement_verifier": verifier_all_ok}
        ),
    )


@pytest.fixture
def ready(services: Services) -> Project:
    p = services.projects.create(
        name="산출물 테스트",
        mode=ProjectMode.PROPOSAL,
        start_date=date(2026, 11, 2),
        end_date=date(2027, 8, 31),
        client_name="OO기관",
    )
    services.documents.register_upload(
        p, kind=DocumentKind.RFP, file_name="rfp.docx", data=bld.build_docx()
    )
    r = StepRunner(services)
    r.run_extraction(p)
    r.run_requirement_extraction(p)
    r.run_verification(p)
    fr1 = services.requirements.get(p, "FR-001")
    fr1.acceptance_status = AcceptanceStatus.REJECTED
    fr1.notes = "범위 외"
    services.ctx.repos.requirements.update(fr1)
    return p


def _path(services: Services, d) -> Path:  # type: ignore[no-untyped-def]
    return services.ctx.paths.resolve_output(d.file_path)


def test_generation_blocked_until_gate1(services: Services) -> None:
    p = services.projects.create(name="차단", mode=ProjectMode.PROPOSAL)
    with pytest.raises(GateBlockedError):
        StepRunner(services).run_generation(p)
    with pytest.raises(GateBlockedError):
        StepRunner(services).run_qa(p)
    assert services.ctx.repos.step_runs.find(p.id) == []


def test_generate_all_creates_nine_deliverables_in_right_folders(
    services: Services, ready: Project
) -> None:
    outcome = StepRunner(services).run_generation(ready)
    assert outcome.run.status == StepStatus.APPROVED
    latest = services.generation.latest_deliverables(ready)
    assert len(latest) == 9
    for kind, d in latest.items():
        path = _path(services, d)
        assert (
            path.exists()
            and d.sha256
            and d.version == "v1.0"
            and d.integrity_status == IntegrityStatus.PENDING
        )
        folder = path.parent.name
        expected = {
            "xlsx": ("09_Excel_Deliverables", "06_WBS_Effort_Schedule"),
            "pptx": ("07_Submission_Proposal_PPT", "08_Presentation_PPT"),
            "png": ("10_Diagrams_Charts_Images",),
        }[path.suffix[1:]]
        assert folder in expected, (kind, folder)
    # 재생성은 덮어쓰지 않고 v1.1
    StepRunner(services).run_generation(ready)
    assert (
        services.generation.latest_deliverables(ready)[DeliverableKind.SUBMISSION_PPTX].version
        == "v1.1"
    )
    assert (
        services.ctx.paths.project_output_dir(ready.output_dir)
        / "07_Submission_Proposal_PPT"
        / "제안서_제출용_v1.0.pptx"
    ).exists()


def test_requirement_spec_sheets_columns_validation_and_notes(
    services: Services, ready: Project
) -> None:
    StepRunner(services).run_generation(ready)
    d = services.generation.latest_deliverables(ready)[DeliverableKind.REQUIREMENT_SPEC_XLSX]
    wb = load_workbook(io.BytesIO(_path(services, d).read_bytes()))
    assert {SHEET_FUNCTIONAL, SHEET_NON_FUNCTIONAL, "제약·평가·제출조건", "변경이력"} <= set(
        wb.sheetnames
    )
    ws = wb[SHEET_FUNCTIONAL]
    headers = [c.value for c in ws[1]]
    assert all(col in headers for col in REQUIRED_COLUMNS)
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert (
        ws.data_validations.dataValidation
        and "수용" in ws.data_validations.dataValidation[0].formula1
    )
    assert ws.column_dimensions["C"].width >= 40
    ids = {
        ws.cell(row=r, column=1).value
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=1).value and str(ws.cell(row=r, column=1).value).startswith("FR-")
    }
    assert ids == {
        r.id for r in services.requirements.list_requirements(ready) if r.id.startswith("FR-")
    }
    rejected_row = next(
        r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "FR-001"
    )
    assert ws.cell(row=rejected_row, column=5).value == "미수용" and "사유" in str(
        ws.cell(row=rejected_row, column=14).value
    )
    assert wb[SHEET_NON_FUNCTIONAL].max_row >= 2 and all(
        str(wb[SHEET_NON_FUNCTIONAL].cell(row=r, column=1).value).startswith("NFR-")
        for r in range(2, 5)
    )


def test_trace_matrix_marks_pending_and_not_run(services: Services, ready: Project) -> None:
    StepRunner(services).run_generation(ready)
    d = services.generation.latest_deliverables(ready)[DeliverableKind.TRACE_MATRIX_XLSX]
    wb = load_workbook(io.BytesIO(_path(services, d).read_bytes()))
    ws = wb["추적표"]
    headers = [c.value for c in ws[1]]
    for col in (
        "요구사항 ID",
        "요구사항 명",
        "화면 ID",
        "단위테스트 ID",
        "단위테스트 결과",
        "통합테스트 ID",
        "통합테스트 결과",
    ):
        assert col in headers
    row = [c.value for c in ws[2]]
    assert (
        row[2] == "미정" and row[4] == "미실행" and row[6] == "미실행" and row[7] != "미정"
    )  # WBS는 연결됨
    assert "테스트상세" in wb.sheetnames and "미추적" in wb.sheetnames


def test_wbs_gantt_and_staffing_consistency(services: Services, ready: Project) -> None:
    StepRunner(services).run_generation(ready)
    latest = services.generation.latest_deliverables(ready)
    wb = load_workbook(
        io.BytesIO(_path(services, latest[DeliverableKind.WBS_GANTT_XLSX]).read_bytes())
    )
    ws = wb["WBS"]
    headers = [c.value for c in ws[1]]
    for col in (
        "WBS ID",
        "상위 WBS ID",
        "요구사항 ID",
        "Task 명",
        "담당 역할",
        "선행 업무",
        "후행 업무",
        "업무 유형",
        "계획 시작일",
        "계획 종료일",
        "실제 시작일",
        "실제 종료일",
        "기간(영업일)",
        "진행 상태",
        "산정 근거",
        "비고",
    ):
        assert col in headers
    assert "간트" in wb.sheetnames and "의존관계" in wb.sheetnames and "산정근거" in wb.sheetnames
    gantt = wb["간트"]
    assert (
        gantt.cell(row=2, column=1).value == ws.cell(row=2, column=1).value
        and gantt.cell(row=2, column=3).value == ws.cell(row=2, column=10).value
    )
    filled = sum(
        1
        for c in gantt[2]
        if c.fill and c.fill.fgColor and c.fill.fgColor.rgb not in (None, "00000000")
    )
    assert filled >= 1
    sp = load_workbook(
        io.BytesIO(_path(services, latest[DeliverableKind.STAFFING_PLAN_XLSX]).read_bytes())
    )
    plan = sp["투입계획"]
    ph = [c.value for c in plan[1]]
    for col in (
        "조직",
        "역할",
        "기술 등급",
        "필요 기술",
        "투입 인원",
        "담당 업무",
        "투입 시작일",
        "투입 종료일",
        "예상 공수(MM)",
        "산정 근거",
        "단가(원/MM, 입력)",
        "비용 합계(원)",
        "비고",
    ):
        assert col in ph
    cost_col = ph.index("비용 합계(원)") + 1
    assert str(plan.cell(row=2, column=cost_col).value).startswith("=IF(")  # 단가 없으면 빈 값
    summary = sp["공수요약"]
    total_row = next(
        r for r in range(2, summary.max_row + 1) if summary.cell(row=r, column=1).value == "합계"
    )
    wbs_total = sum(
        Decimal(str(w.effort_md))
        for w in services.ctx.repos.wbs.tree(ready.id)
        if w.parent_id is None
    )
    assert Decimal(str(summary.cell(row=total_row, column=3).value)) == wbs_total


def test_pptx_decks_open_and_share_key_figures(services: Services, ready: Project) -> None:
    StepRunner(services).run_generation(ready)
    latest = services.generation.latest_deliverables(ready)
    sub = Presentation(str(_path(services, latest[DeliverableKind.SUBMISSION_PPTX])))
    pres = Presentation(str(_path(services, latest[DeliverableKind.PRESENTATION_PPTX])))
    assert sub.slide_width > sub.slide_height and len(sub.slides) >= 10
    assert len(pres.slides) < len(sub.slides)
    sub_text = "\n".join(pptx_texts(_path(services, latest[DeliverableKind.SUBMISSION_PPTX])))
    pres_text = "\n".join(pptx_texts(_path(services, latest[DeliverableKind.PRESENTATION_PPTX])))
    for r in services.requirements.list_requirements(ready):
        assert r.id in sub_text
    for key in ("MM", "2026-11-02 ~ 2027-08-31", "투입"):
        assert key in sub_text and key in pres_text
    assert "시스템 구성도" in sub_text and "핵심 메시지" in pres_text
    outcome = QaOutcome()
    check_pptx(_path(services, latest[DeliverableKind.SUBMISSION_PPTX]), outcome, "제출용")
    assert not [
        f for f in outcome.findings if f.rule in ("pptx.bounds", "pptx.orientation", "pptx.open")
    ]


def test_outline_prefers_rfp_toc_then_default(services: Services, ready: Project) -> None:
    ctx = services.generation.build_context(ready)
    blocks = services.extraction.current_blocks(ready)
    assert rfp_toc_from_blocks(blocks) == []  # 모의 RFP에는 제안서 목차 없음
    default = default_outline(ctx)
    keys = [s.key for s in default]
    assert (
        keys[:2] == ["purpose", "requirements"]
        and "current" not in keys
        and {
            "system_diagram",
            "sw_diagram",
            "service_diagram",
            "schedule",
            "organization",
            "risk",
        }
        <= set(keys)
    )
    assert "security" in keys  # 암호화 요구사항 존재
    assert "evaluation" not in keys  # RFP 분석 결과·평가 요구사항 없음 → 제외
    mapped = decide_outline(
        ctx,
        [
            "1. 사업 이해",
            "2. 요구사항 대응 방안",
            "3. 기술 아키텍처",
            "4. 추진 일정",
            "5. 우리만의 특별 섹션",
        ],
    )
    assert [s.source for s in mapped][:5] == ["rfp_toc"] * 5
    assert {s.key for s in mapped} >= {
        "purpose",
        "requirements",
        "architecture",
        "schedule",
        "system_diagram",
    }
    assert any(s.key.startswith("generic") for s in mapped)
    ctx.proposal_content = {
        "sections": [{"slide_group": "understanding", "title": "에이전트 섹션"}]
    }
    assert decide_outline(ctx, [])[0].source == "agent"


def test_qa_passes_then_consistency_check_detects_mismatch(
    services: Services, ready: Project
) -> None:
    r = StepRunner(services)
    r.run_generation(ready)
    qa = r.run_qa(ready)
    assert qa.gate and qa.gate.status == GateStatus.NEEDS_HUMAN_REVIEW
    latest = services.generation.latest_deliverables(ready)
    assert all(
        d.integrity_status == IntegrityStatus.PASSED and d.qa_status == QaStatus.PASSED
        for d in latest.values()
    )
    report = (
        services.ctx.paths.project_output_dir(ready.output_dir)
        / "11_QA_Validation_Reports"
        / "QA보고서_v1.0.md"
    )
    assert report.exists() and "열림" in report.read_text(encoding="utf-8")
    # 인위적 불일치: PPT 텍스트에서 기간을 바꾸면 error
    ctx = services.generation.build_context(ready)
    outcome = QaOutcome()
    texts = {
        "submission": pptx_texts(_path(services, latest[DeliverableKind.SUBMISSION_PPTX])),
        "presentation": [
            t.replace("2027-08-31", "2027-12-31")
            for t in pptx_texts(_path(services, latest[DeliverableKind.PRESENTATION_PPTX]))
        ],
    }
    check_consistency(ctx, {}, texts, outcome)
    assert any(
        f.rule == "consistency.ppt_figures" and f.level == ValidationLevel.ERROR
        for f in outcome.findings
    )
    missing_text = {"submission": [re.sub(r"FR-001", "", t) for t in texts["submission"]]}
    outcome2 = QaOutcome()
    check_consistency(ctx, {}, missing_text, outcome2)
    assert any(f.rule == "traceability.requirements" for f in outcome2.findings)


def test_finalize_requires_gate4_approval_and_separates_final(
    services: Services, ready: Project
) -> None:
    r = StepRunner(services)
    r.run_generation(ready)
    with pytest.raises(GateBlockedError):
        r.run_finalize(ready, approver="pm", comments="x")
    r.run_qa(ready)
    with pytest.raises(GateBlockedError):
        r.run_finalize(ready, approver="pm", comments="승인 전")
    services.gates.approve_human_review(ready, "gate4", approver="pm", comments="검토 완료")
    outcome = r.run_finalize(ready, approver="pm", comments="확정")
    final = services.ctx.paths.project_output_dir(ready.output_dir) / "13_Final_Deliverables"
    assert outcome.run.status == StepStatus.APPROVED and len(list(final.iterdir())) == 9
    assert (
        services.ctx.paths.project_output_dir(ready.output_dir)
        / "09_Excel_Deliverables"
        / "요구사항정의서_v1.0.xlsx"
    ).exists()  # 임시본 보존
    with pytest.raises(Exception):  # noqa: B017 - 같은 버전 재확정은 덮어쓰지 않음
        services.generation.finalize(ready, approver="pm", comments="again")


def test_finalize_blocked_when_integrity_failed(services: Services, ready: Project) -> None:
    StepRunner(services).run_generation(ready)
    d = services.generation.latest_deliverables(ready)[DeliverableKind.SUBMISSION_PPTX]
    d.integrity_status = IntegrityStatus.FAILED
    services.ctx.repos.deliverables.update(d)
    with pytest.raises(ValidationError):
        services.generation.finalize(ready, approver="pm", comments="x")


def test_review_items_surface_unverified_and_defaults(services: Services, ready: Project) -> None:
    ctx = services.generation.build_context(ready)
    items = services.generation.review_items(ctx)
    assert any("기본값" in i for i in items) and any("미정" in i for i in items)
    assert isinstance(ctx, GenerationContext) and ctx.period_start == date(2026, 11, 2)
