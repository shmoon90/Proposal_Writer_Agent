"""Phase 16: 산출물 manifest.json(생성 누적·QA 결과 첨부), 그림 정책(허용 태그는 예외 기록, 태그 없는 그림은 실패),
테스트 상태 기록 스크립트."""

from __future__ import annotations

import io
import json
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.util import Cm

from fixtures import builders as bld
from proposal_agent.agents import FakeAgentRunner
from proposal_agent.generators.base import PICTURE_TAGS, picture_name, picture_tag
from proposal_agent.generators.pptx.builder import DeckBuilder
from proposal_agent.models import (
    AcceptanceStatus,
    DeliverableKind,
    DocumentKind,
    Project,
    ProjectMode,
)
from proposal_agent.orchestrator import StepRunner
from proposal_agent.services import AppContext, Services
from proposal_agent.storage.manifest import load_manifest, manifest_path
from proposal_agent.validation.deliverables import QaOutcome, check_pptx
from unit.test_requirements_phase2 import extractor_from_text, verifier_all_ok


def _fake_services(ctx: AppContext) -> Services:
    return Services.from_context(
        ctx,
        runner=FakeAgentRunner(
            {"requirement_extractor": extractor_from_text, "requirement_verifier": verifier_all_ok}
        ),
    )


def _ready(services: Services) -> Project:
    p = services.projects.create(
        name="manifest",
        mode=ProjectMode.PROPOSAL,
        start_date=date(2026, 11, 2),
        end_date=date(2027, 8, 31),
    )
    services.documents.register_upload(
        p, kind=DocumentKind.RFP, file_name="rfp.docx", data=bld.build_docx()
    )
    r = StepRunner(services)
    r.run_extraction(p)
    r.run_requirement_extraction(p)
    r.run_verification(p)
    for req in services.requirements.list_requirements(p):
        req.acceptance_status = AcceptanceStatus.ACCEPTED
        services.ctx.repos.requirements.update(req)
    return p


# 1. manifest: 생성 실행 누적 + QA 결과 첨부 ------------------------------------------------------
def test_manifest_records_generation_runs_and_qa(ctx: AppContext) -> None:
    services = _fake_services(ctx)
    p = _ready(services)
    runner = StepRunner(services)
    runner.run_generation(p)
    path = manifest_path(services.ctx.paths, p)
    assert path.exists() and path.name == "manifest.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    assert data["project_id"] == p.id and len(data["runs"]) == 1
    run = data["runs"][0]
    kinds = {d["kind"] for d in run["deliverables"]}
    assert {
        DeliverableKind.SUBMISSION_PPTX.value,
        DeliverableKind.PRESENTATION_PPTX.value,
        DeliverableKind.REQUIREMENT_SPEC_XLSX.value,
        DeliverableKind.WBS_GANTT_XLSX.value,
    } <= kinds
    cats = {d["category"] for d in run["deliverables"]}
    assert {"submission", "presentation", "excel", "source"} <= cats
    first = run["deliverables"][0]
    assert first["sha256"] and first["bytes"] and first["version"].startswith("v")
    assert run["qa"] is None and run["run_kind"] == "generation" and run["extra_files"]
    assert all(services.ctx.paths.resolve_output(d["path"]).exists() for d in run["deliverables"])
    runner.run_qa(p)
    data = json.loads(path.read_text(encoding="utf-8"))
    qa = data["runs"][-1]["qa"]
    assert (
        qa
        and qa["verdict"] in ("PASS", "FAIL", "NEEDS_HUMAN_REVIEW")
        and qa["counts"]["total"] >= 35
    )
    assert qa["gates"]["gate4"] and qa["report"].endswith(".md") and qa["picture_exceptions"] == []
    # 두 번째 생성은 run을 덧붙이고 과거 run을 지우지 않는다
    runner.run_generation(p)
    data = load_manifest(services.ctx.paths, p)
    assert (
        len(data["runs"]) == 2
        and data["runs"][0]["qa"] is not None
        and data["runs"][1]["qa"] is None
    )


# 2. 그림 정책: 태그 없는 그림 실패, 허용 태그 그림은 예외 기록 ------------------------------------------
def _deck_with_pictures(tmp_path: Path, *, tagged: bool) -> Path:
    from PIL import Image

    img = tmp_path / "x.png"
    Image.new("RGB", (40, 20), "white").save(img)
    prs = Presentation()
    for i in range(3):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        tb = sl.shapes.add_textbox(Cm(1), Cm(1), Cm(5), Cm(1))
        tb.text_frame.text = f"슬라이드 {i}"
        if i == 1:
            pic = sl.shapes.add_picture(str(img), Cm(2), Cm(3))
            if tagged:
                pic.name = picture_name("photo", "청사 전경")
    path = tmp_path / ("tagged.pptx" if tagged else "plain.pptx")
    buf = io.BytesIO()
    prs.save(buf)
    path.write_bytes(buf.getvalue())
    return path


def test_picture_policy_tagged_photo_allowed_untagged_fails(tmp_path: Path) -> None:
    assert picture_tag("photo:청사 전경") == "photo" and picture_tag("Picture 3") is None
    assert picture_tag("ILLUSTRATION:x") == "illustration" and set(PICTURE_TAGS) >= {
        "photo",
        "decoration",
    }
    out = QaOutcome()
    check_pptx(_deck_with_pictures(tmp_path, tagged=False), out, "제출용")
    assert [f.rule for f in out.findings if f.rule == "pptx.picture"] == ["pptx.picture"]
    out2 = QaOutcome()
    check_pptx(_deck_with_pictures(tmp_path, tagged=True), out2, "제출용")
    assert not [f for f in out2.findings if f.rule == "pptx.picture"]
    photos = [f for f in out2.findings if f.rule == "pptx.photo"]
    assert (
        len(photos) == 1
        and photos[0].level.value == "info"
        and photos[0].message.startswith("예외:")
    )
    assert "사진" in photos[0].message and "슬라이드 2" in photos[0].message


def test_image_slide_tag_names_picture_and_qa_passes(tmp_path: Path) -> None:
    from PIL import Image

    img = tmp_path / "site.png"
    Image.new("RGB", (120, 60), "gray").save(img)
    b = DeckBuilder()
    b.title_slide("테스트", "부제", "")
    b.image_slide("현장 사진", img, "footer", caption="청사 전경", tag="photo")
    try:
        b.image_slide("구조도", img, "footer", caption="x", tag="diagram")  # 허용되지 않는 태그
    except ValueError as exc:
        assert "허용되지 않는 그림 태그" in str(exc)
    else:  # pragma: no cover
        raise AssertionError("허용되지 않는 태그를 거부해야 한다")
    path = tmp_path / "deck.pptx"
    path.write_bytes(b.to_bytes())
    prs = Presentation(str(path))
    names = [sh.name for sl in prs.slides for sh in sl.shapes if "PICTURE" in str(sh.shape_type)]
    assert names == ["photo:청사 전경"]
    out = QaOutcome()
    check_pptx(path, out, "제출용")
    assert not [f for f in out.findings if f.rule == "pptx.picture"]
    assert [f for f in out.findings if f.rule == "pptx.photo"]


# 3. 테스트 상태 기록 스크립트 ------------------------------------------------------------------------
def test_record_test_status_parses_summary_only() -> None:
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "record_test_status",
        Path(__file__).resolve().parents[2] / "scripts" / "record_test_status.py",
    )
    mod = importlib.util.module_from_spec(spec)  # type: ignore[arg-type]
    assert spec and spec.loader
    spec.loader.exec_module(mod)  # type: ignore[union-attr]
    assert mod.parse_summary("...\n253 passed, 1 skipped, 2 deselected in 62.7s\n") == {
        "passed": 253,
        "failed": 0,
        "skipped": 1,
        "deselected": 2,
    }
    assert mod.parse_summary("3 failed, 10 passed in 1s") == {
        "passed": 10,
        "failed": 3,
        "skipped": 0,
        "deselected": 0,
    }
    assert mod.parse_summary("no summary here") is None
