"""Phase 20: 네이티브 차트(공수·간트·통계·VisualSpec chart), 화면 목업 컴포넌트·그룹화(해제 가능),
참고 그림(DocumentKind.PICTURE) 태그 삽입, 그룹 안 도형의 QA 넘침·겹침 검사."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from proposal_agent.agents.schemas import ProposalContent
from proposal_agent.generators.base import GenerationContext, parse_picture_name
from proposal_agent.generators.pptx import shapes
from proposal_agent.generators.pptx.charts import (
    add_category_chart,
    add_gantt_chart,
    chart_from_spec,
    parse_number,
)
from proposal_agent.generators.pptx.decks import build_presentation_deck, build_submission_deck
from proposal_agent.generators.pptx.outline import decide_outline
from proposal_agent.generators.pptx.slide_spec import build_slide_spec
from proposal_agent.models import (
    AcceptanceStatus,
    DocumentKind,
    Project,
    ProjectMode,
    Requirement,
    RequirementType,
    ScopeClass,
    StaffingAllocation,
    StaffRole,
    WBSItem,
    WBSLevel,
)
from proposal_agent.validation.deliverables import QaOutcome, check_pptx

BOX = shapes.Box(left=1.0, top=1.0, width=22.0, height=12.0)
STYLE = shapes.ShapeStyle()


def _blank() -> Any:
    prs = Presentation()
    return prs.slides.add_slide(prs.slide_layouts[6])


def _flat(shapes_: Any) -> list[Any]:
    out: list[Any] = []
    for sh in shapes_:
        if "GROUP" in str(sh.shape_type):
            out.extend(_flat(sh.shapes))
        else:
            out.append(sh)
    return out


def test_native_charts_are_chart_objects() -> None:
    slide = _blank()
    f = add_category_chart(
        slide, BOX, ["PM", "개발", "QA"], [("MM", [3.5, 12, 4])], style=STYLE, title="공수"
    )
    assert f.has_chart and f.chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED
    assert [c for c in f.chart.plots[0].categories] == ["PM", "개발", "QA"]
    assert list(f.chart.plots[0].series[0].values) == [3.5, 12.0, 4.0]
    slide2 = _blank()
    g = add_gantt_chart(
        slide2,
        BOX,
        [
            ("1 분석", date(2026, 11, 2), date(2026, 12, 15)),
            ("2 설계", date(2026, 12, 1), date(2027, 1, 31)),
        ],
        date(2026, 11, 2),
        date(2027, 4, 30),
        style=STYLE,
    )
    assert g.chart.chart_type == XL_CHART_TYPE.BAR_STACKED
    offs, durs = g.chart.plots[0].series
    assert list(offs.values) == [0.0, 29.0] and list(durs.values) == [44.0, 62.0]
    assert g.chart.category_axis.reverse_order
    assert add_gantt_chart(slide2, BOX, [], date(2026, 1, 1), date(2026, 2, 1), style=STYLE) is None
    slide3 = _blank()
    c = chart_from_spec(
        slide3,
        BOX,
        {
            "kind": "chart",
            "chart_type": "pie",
            "headers": ["유형", "건수"],
            "rows": [["기능", "8건"], ["비기능", "3"]],
        },
        style=STYLE,
    )
    assert c.chart.chart_type == XL_CHART_TYPE.PIE
    assert parse_number("1,234.5 MM") == 1234.5 and parse_number("없음") is None
    # VisualSpec kind=chart → draw_visual 도 차트 객체
    slide4 = _blank()
    assert (
        shapes.draw_visual(
            slide4,
            BOX,
            {
                "kind": "chart",
                "chart_type": "column",
                "headers": ["월", "계획", "실적"],
                "rows": [["1월", 3, 2]],
            },
            style=STYLE,
        )
        == "chart"
    )
    assert any(sh.has_chart for sh in slide4.shapes)
    assert "chart" in shapes.AGENT_VISUAL_KINDS and len(shapes.AGENT_VISUAL_KINDS) == 14


def test_screen_mockup_components_are_individual_shapes_in_one_group() -> None:
    slide = _blank()
    spec = {
        "kind": "screen",
        "title": "사건 접수",
        "items": [
            {"title": "접수", "kind": "tab"},
            {"title": "조회", "kind": "tab"},
            {"title": "사건 목록", "group": "nav"},
            {"title": "통계", "kind": "menu"},
            {"title": "민원인", "kind": "input", "text": "이름 입력"},
            {"title": "저장", "kind": "button"},
            {"title": "오류", "kind": "message", "text": "필수 항목 누락"},
            {"title": "접수 목록", "kind": "table", "text": "번호|제목|상태"},
            {"title": "월별 통계", "kind": "chart"},
            {"title": "상세", "text": "폼"},
        ],
    }
    assert shapes.draw_visual(slide, BOX, spec, style=STYLE) == "screen"
    top = list(slide.shapes)
    assert len(top) == 1 and "GROUP" in str(top[0].shape_type) and "화면 목업" in top[0].name
    inner = _flat(slide.shapes)
    texts = [sh.text_frame.text for sh in inner if sh.has_text_frame]
    assert len(inner) >= 20 and not any("PICTURE" in str(sh.shape_type) for sh in inner)
    for expect in ("접수", "▸ 사건 목록", "이름 입력", "저장", "ⓘ 필수 항목 누락", "번호", "상태"):
        assert any(expect in t for t in texts), expect
    # 카드·KPI 도 카드마다 그룹
    slide2 = _blank()
    shapes.draw_cards(slide2, BOX, [{"title": "a", "text": "b"}, {"title": "c"}], style=STYLE)
    assert [str(sh.shape_type).split(" ")[0] for sh in slide2.shapes] == ["GROUP", "GROUP"]
    assert (
        ProposalContent.model_validate(
            {
                "profile": "submission",
                "sections": [{"slide_group": "generic", "title": "x", "visual": spec}],
            }
        )
        .sections[0]
        .visual.items[4]
        .kind
        == "input"
    )


def _project() -> Project:
    return Project(
        name="사건관리",
        mode=ProjectMode.PROPOSAL,
        output_dir="Project_x",
        workspace_dir="ws_x",
        start_date=date(2026, 11, 2),
        end_date=date(2027, 4, 30),
        created_by="t",
    )


def _req(
    p: Project, rid: str, name: str, *, scope: ScopeClass = ScopeClass.MANDATORY
) -> Requirement:
    return Requirement(
        id=rid,
        project_id=p.id,
        name=name,
        description=name,
        type=RequirementType.FUNCTIONAL if rid.startswith("FR") else RequirementType.NON_FUNCTIONAL,
        scope_class=scope,
        created_by="t",
        acceptance_status=AcceptanceStatus.ACCEPTED,
        implementation_approach="웹 화면과 API로 구현",
    )


def test_decks_use_native_charts_and_tagged_pictures(tmp_path: Path) -> None:
    p = _project()
    img = tmp_path / "photo__현장 사진.png"
    Image.new("RGB", (160, 90), (100, 120, 140)).save(img)
    role = StaffRole(project_id=p.id, role_name="개발", headcount_planned=2, created_by="t")
    w1 = WBSItem(
        project_id=p.id,
        code="1",
        name="분석",
        level=WBSLevel.PHASE,
        effort_md=Decimal("20"),
        planned_start=date(2026, 11, 2),
        planned_end=date(2026, 12, 15),
        created_by="t",
    )
    ctx = GenerationContext(
        project=p,
        requirements=[
            _req(p, "FR-001", "사건 접수"),
            _req(p, "FR-002", "통계", scope=ScopeClass.OPTIONAL_PROPOSAL),
            _req(p, "NFR-001", "응답 시간"),
        ],
        wbs=[w1],
        roles=[role],
        allocations=[
            StaffingAllocation(
                project_id=p.id,
                wbs_item_id=w1.id,
                staff_role_id=role.id,
                effort_md=Decimal("20"),
                created_by="t",
            )
        ],
        proposal_content={
            "profile": "submission",
            "sections": [
                {
                    "slide_group": "understanding",
                    "title": "사업 이해",
                    "bullets": ["배경"],
                    "picture": "현장 사진",
                },
            ],
        },
        period_start=date(2026, 11, 2),
        period_end=date(2027, 4, 30),
        pictures=[
            {"file_name": img.name, "path": str(img), "tag": "photo", "caption": "현장 사진"}
        ],
    )
    assert parse_picture_name(img.name) == ("photo", "현장 사진")
    assert parse_picture_name("diagram.png") == ("reference", "diagram")
    outline = decide_outline(ctx, [])
    data = build_submission_deck(ctx, outline)
    prs = Presentation(io.BytesIO(data))
    charts = [sh.chart.chart_type for sl in prs.slides for sh in sl.shapes if sh.has_chart]
    assert XL_CHART_TYPE.BAR_STACKED in charts  # 일정(간트)
    assert XL_CHART_TYPE.BAR_CLUSTERED in charts  # 역할별 공수
    assert XL_CHART_TYPE.COLUMN_STACKED in charts  # 요구사항 통계
    stats = next(
        sh.chart
        for sl in prs.slides
        for sh in sl.shapes
        if sh.has_chart and sh.chart.chart_type == XL_CHART_TYPE.COLUMN_STACKED
    )
    assert [c for c in stats.plots[0].categories] == ["기능", "비기능"]
    assert list(stats.plots[0].series[0].values) == [1.0, 1.0] and list(
        stats.plots[0].series[1].values
    ) == [1.0, 0.0]
    pics = [sh.name for sl in prs.slides for sh in sl.shapes if "PICTURE" in str(sh.shape_type)]
    assert len(pics) == 1 and pics[0].startswith("photo:")  # 소절 뒤 참고 그림 1장(부록 중복 없음)
    spec = build_slide_spec(data, ctx, "submission")
    assert spec["summary"]["rendered_layouts"].get("chart", 0) >= 3
    assert spec["summary"]["pictures"] == 1
    # QA: 태그 달린 그림은 예외(통과), 그룹 안 도형도 검사 대상
    out = QaOutcome()
    path = tmp_path / "sub.pptx"
    path.write_bytes(data)
    check_pptx(path, out, "제출용")
    assert not [f for f in out.findings if f.rule == "pptx.picture"]
    assert [f for f in out.findings if f.rule == "pptx.photo"]
    # 발표용: 참조된 그림만, 차트는 공수 막대
    pres = Presentation(io.BytesIO(build_presentation_deck(ctx, outline)))
    assert sum(1 for sl in pres.slides for sh in sl.shapes if "PICTURE" in str(sh.shape_type)) == 1
    assert any(sh.has_chart for sl in pres.slides for sh in sl.shapes)


def test_picture_document_kind_is_accepted_and_not_extracted(ctx: Any) -> None:
    from proposal_agent.services import Services

    services = Services.from_context(ctx)
    p = services.projects.create(name="그림", mode=ProjectMode.PROPOSAL, client_name="x")
    buf = io.BytesIO()
    Image.new("RGB", (64, 32), "white").save(buf, format="PNG")
    optional = [k for k, req in services.documents.kinds_for(p.mode) if not req]
    assert DocumentKind.PICTURE in optional and DocumentKind.PICTURE.label.startswith("참고 그림")
    doc = services.documents.register_upload(
        p, kind=DocumentKind.PICTURE, file_name="illustration__개념도.png", data=buf.getvalue()
    )
    assert doc.kind == DocumentKind.PICTURE
    outcomes = services.extraction.extract_all(p)
    assert all(o.run.source_document_id != doc.id for o in outcomes)  # 그림은 추출 대상이 아니다
    assert [(x["tag"], x["caption"]) for x in services.generation._pictures(p)] == [
        ("illustration", "개념도")
    ]
