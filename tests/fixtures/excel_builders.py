"""Phase 6 테스트용 기준 문서(Excel) 생성기.

표준 구조(Phase 4 생성기와 같은 시트·헤더)와 비표준 구조(다른 헤더명·헤더 행 위치)를 모두 만든다.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any

from openpyxl import Workbook

SPEC_HEADERS = (
    "요구사항 ID",
    "요구사항 명",
    "요구사항 상세 내용",
    "구현방안",
    "수용여부",
    "범위 분류",
    "필수 여부",
    "중요도",
    "검증 방법",
    "원본 출처",
    "검토 상태",
    "검증 상태",
    "연결 산출물",
    "비고",
)
TRACE_HEADERS = (
    "요구사항 ID",
    "요구사항 명",
    "화면 ID",
    "단위테스트 ID",
    "단위테스트 결과",
    "통합테스트 ID",
    "통합테스트 결과",
    "WBS",
    "비고",
)
WBS_HEADERS = (
    "WBS ID",
    "상위 WBS ID",
    "요구사항 ID",
    "Task 명",
    "Task 상세 내용",
    "담당 역할",
    "선행 업무",
    "후행 업무",
    "업무 유형",
    "계획 시작일",
    "계획 종료일",
    "실제 시작일",
    "실제 종료일",
    "기간(영업일)",
    "진행 상태",
    "산정 근거",
    "공수(M/D)",
    "비고",
)
STAFF_HEADERS = (
    "조직",
    "역할",
    "기술 등급",
    "필요 기술",
    "투입 인원",
    "담당 업무",
    "투입 시작일",
    "투입 종료일",
    "투입 기간(영업일)",
    "예상 공수(MM)",
    "예상 공수(M/D)",
    "산정 근거",
    "단가(원/MM, 입력)",
    "비용 합계(원)",
    "비고",
)

# 기준 요구사항 6건 (ID 체계 SFR-/SNR-)
BASE_REQUIREMENTS: list[dict[str, Any]] = [
    {
        "id": "SFR-001",
        "name": "사용자 로그인",
        "desc": "사용자는 아이디와 비밀번호로 로그인한다. 5회 실패 시 계정을 잠근다.",
        "type": "기능",
        "sheet": "기능요구사항",
    },
    {
        "id": "SFR-002",
        "name": "민원 접수",
        "desc": "민원인은 민원 유형을 선택하고 내용과 첨부파일을 등록하여 민원을 접수한다.",
        "type": "기능",
        "sheet": "기능요구사항",
    },
    {
        "id": "SFR-003",
        "name": "민원 처리 현황 조회",
        "desc": "민원인은 접수번호로 처리 현황과 담당 부서를 조회한다.",
        "type": "기능",
        "sheet": "기능요구사항",
    },
    {
        "id": "SFR-004",
        "name": "통계 대시보드",
        "desc": "관리자는 기간별 민원 접수·처리 건수를 차트로 조회한다.",
        "type": "기능",
        "sheet": "기능요구사항",
    },
    {
        "id": "SFR-005",
        "name": "팩스 접수",
        "desc": "민원 담당자는 팩스로 들어온 민원을 수기로 등록한다.",
        "type": "기능",
        "sheet": "기능요구사항",
    },
    {
        "id": "SNR-001",
        "name": "응답 시간",
        "desc": "주요 화면은 3초 이내에 응답한다.",
        "type": "비기능",
        "sheet": "비기능요구사항",
    },
]

BASE_TRACES: list[dict[str, Any]] = [
    {"id": "SFR-001", "screen": "SCR-001", "ut": "UT-001", "it": "IT-001", "wbs": "2.1.1"},
    {"id": "SFR-002", "screen": "SCR-002", "ut": "UT-002", "it": "IT-001", "wbs": "2.1.2"},
    {"id": "SFR-003", "screen": "SCR-003", "ut": "UT-003", "it": "IT-002", "wbs": "2.1.3"},
    {"id": "SFR-004", "screen": "SCR-004", "ut": "UT-004", "it": "IT-002", "wbs": "2.2.1"},
    {"id": "SFR-005", "screen": "SCR-005", "ut": "UT-005", "it": "", "wbs": "2.2.2"},
    {"id": "SNR-001", "screen": "", "ut": "", "it": "IT-003", "wbs": "3.1.1"},
]

# (code, parent, req_ids, name, role, preds, start, end, duration, effort)
BASE_WBS: list[tuple[str, str, str, str, str, str, date | None, date | None, int, float]] = [
    ("1", "", "", "분석·설계", "PL", "", date(2026, 11, 2), date(2026, 11, 27), 20, 40),
    ("1.1", "1", "", "요구사항 분석", "PL", "", date(2026, 11, 2), date(2026, 11, 13), 10, 20),
    (
        "1.1.1",
        "1.1",
        "SFR-001, SFR-002, SFR-003",
        "업무 분석",
        "PL",
        "",
        date(2026, 11, 2),
        date(2026, 11, 13),
        10,
        20,
    ),
    ("1.2", "1", "", "설계", "아키텍트", "", date(2026, 11, 16), date(2026, 11, 27), 10, 20),
    (
        "1.2.1",
        "1.2",
        "SFR-001, SFR-002, SFR-003, SFR-004",
        "화면·DB 설계",
        "아키텍트",
        "1.1.1",
        date(2026, 11, 16),
        date(2026, 11, 27),
        10,
        20,
    ),
    ("2", "", "", "구현", "백엔드", "", date(2026, 11, 30), date(2027, 2, 12), 55, 110),
    ("2.1", "2", "", "민원 기능 개발", "백엔드", "", date(2026, 11, 30), date(2027, 1, 15), 35, 70),
    (
        "2.1.1",
        "2.1",
        "SFR-001",
        "로그인 개발",
        "백엔드",
        "1.2.1",
        date(2026, 11, 30),
        date(2026, 12, 11),
        10,
        20,
    ),
    (
        "2.1.2",
        "2.1",
        "SFR-002",
        "민원 접수 개발",
        "백엔드",
        "2.1.1",
        date(2026, 12, 14),
        date(2026, 12, 31),
        14,
        28,
    ),
    (
        "2.1.3",
        "2.1",
        "SFR-003",
        "처리 현황 조회 개발",
        "백엔드",
        "2.1.2",
        date(2027, 1, 1),
        date(2027, 1, 15),
        11,
        22,
    ),
    (
        "2.2",
        "2",
        "",
        "관리 기능 개발",
        "프런트엔드",
        "",
        date(2027, 1, 18),
        date(2027, 2, 12),
        20,
        40,
    ),
    (
        "2.2.1",
        "2.2",
        "SFR-004",
        "통계 대시보드 개발",
        "프런트엔드",
        "2.1.3",
        date(2027, 1, 18),
        date(2027, 1, 29),
        10,
        20,
    ),
    (
        "2.2.2",
        "2.2",
        "SFR-005",
        "팩스 접수 개발",
        "프런트엔드",
        "2.2.1",
        date(2027, 2, 1),
        date(2027, 2, 12),
        10,
        20,
    ),
    ("3", "", "", "시험·이행", "QA", "", date(2027, 2, 15), date(2027, 3, 12), 20, 40),
    ("3.1", "3", "", "통합 시험", "QA", "", date(2027, 2, 15), date(2027, 3, 12), 20, 40),
    (
        "3.1.1",
        "3.1",
        "SNR-001",
        "성능·통합 시험",
        "QA",
        "2.2.2",
        date(2027, 2, 15),
        date(2027, 3, 12),
        20,
        40,
    ),
]

BASE_STAFF: list[
    tuple[str, str, str, str, int, str, date, date, int, float, float, str, float | None]
] = [
    (
        "수행사",
        "PL",
        "특급",
        "업무 분석, 설계",
        1,
        "분석·설계 리드",
        date(2026, 11, 2),
        date(2026, 11, 27),
        20,
        2.0,
        40,
        "WBS 1 합계",
        None,
    ),
    (
        "수행사",
        "아키텍트",
        "고급",
        "설계",
        1,
        "아키텍처",
        date(2026, 11, 16),
        date(2026, 11, 27),
        10,
        1.0,
        20,
        "WBS 1.2",
        None,
    ),
    (
        "수행사",
        "백엔드",
        "중급",
        "Java, Spring",
        2,
        "서버 개발",
        date(2026, 11, 30),
        date(2027, 1, 15),
        35,
        3.5,
        70,
        "WBS 2.1",
        9000000.0,
    ),
    (
        "수행사",
        "프런트엔드",
        "중급",
        "React",
        1,
        "화면 개발",
        date(2027, 1, 18),
        date(2027, 2, 12),
        20,
        2.0,
        40,
        "WBS 2.2",
        8000000.0,
    ),
    (
        "수행사",
        "QA",
        "중급",
        "테스트",
        1,
        "시험",
        date(2027, 2, 15),
        date(2027, 3, 12),
        20,
        2.0,
        40,
        "WBS 3",
        None,
    ),
]


def _wb_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_requirement_spec(
    requirements: list[dict[str, Any]] | None = None, *, with_history: bool = True
) -> bytes:
    reqs = requirements if requirements is not None else BASE_REQUIREMENTS
    wb = Workbook()
    sheets = {
        "기능요구사항": wb.active,
        "비기능요구사항": wb.create_sheet("비기능요구사항"),
        "제약·평가·제출조건": wb.create_sheet("제약·평가·제출조건"),
    }
    wb.active.title = "기능요구사항"
    for ws in sheets.values():
        ws.append(list(SPEC_HEADERS))
    for r in reqs:
        ws = sheets[r.get("sheet", "기능요구사항")]
        ws.append(
            [
                r["id"],
                r["name"],
                r["desc"],
                r.get("impl", "표준 프레임워크로 구현"),
                r.get("accept", "수용"),
                "RFP 필수",
                "Y",
                r.get("priority", "중"),
                "시연",
                "RFP 3.1",
                "approved",
                "verified",
                "",
                r.get("notes", ""),
            ]
        )
    if with_history:
        ws = wb.create_sheet("변경이력")
        ws.append(["버전", "일자", "변경 내용", "작성자"])
        ws.append(["v1.0", "2026-10-01", "최초 작성", "PM"])
    return _wb_bytes(wb)


def build_trace_matrix(traces: list[dict[str, Any]] | None = None) -> bytes:
    rows = traces if traces is not None else BASE_TRACES
    names = {r["id"]: r["name"] for r in BASE_REQUIREMENTS}
    wb = Workbook()
    ws = wb.active
    ws.title = "추적표"
    ws.append(list(TRACE_HEADERS))
    for t in rows:
        ws.append(
            [
                t["id"],
                names.get(t["id"], t["id"]),
                t.get("screen") or "미정",
                t.get("ut") or "미정",
                "미실행",
                t.get("it") or "미정",
                "미실행",
                t.get("wbs") or "미정",
                "",
            ]
        )
    ws2 = wb.create_sheet("테스트상세")
    ws2.append(["테스트 ID", "종류", "이름", "결과"])
    return _wb_bytes(wb)


def build_wbs(
    items: list[tuple[Any, ...]] | None = None,
    *,
    extra_deps: list[tuple[str, str, str, int]] | None = None,
) -> bytes:
    rows = items if items is not None else BASE_WBS
    wb = Workbook()
    ws = wb.active
    ws.title = "WBS"
    ws.append(list(WBS_HEADERS))
    for code, parent, reqs, name, role, preds, start, end, dur, effort in rows:
        ws.append(
            [
                code,
                parent,
                reqs,
                name,
                f"{name} 상세",
                role,
                preds,
                "",
                "선행업무" if preds else "단독업무",
                start,
                end,
                None,
                None,
                dur,
                "계획",
                "기준 산정",
                effort,
                "",
            ]
        )
    ws2 = wb.create_sheet("간트")
    ws2.append(["WBS ID", "Task 명"])
    ws3 = wb.create_sheet("의존관계")
    ws3.append(["선행 WBS", "후행 WBS", "유형", "지연(일)"])
    for code, _p, _r, _n, _role, preds, *_ in rows:
        for p in [x.strip() for x in preds.split(",") if x.strip()]:
            ws3.append([p, code, "FS", 0])
    for p, s, t, lag in extra_deps or []:
        ws3.append([p, s, t, lag])
    ws4 = wb.create_sheet("산정근거")
    ws4.append(["WBS ID", "Task 명", "산정 근거"])
    return _wb_bytes(wb)


def build_staffing_plan(rows: list[tuple[Any, ...]] | None = None) -> bytes:
    data = rows if rows is not None else BASE_STAFF
    wb = Workbook()
    ws = wb.active
    ws.title = "투입계획"
    ws.append(list(STAFF_HEADERS))
    for org, role, grade, skills, hc, duties, start, end, dur, mm, md, basis, unit in data:
        ws.append(
            [
                org,
                role,
                grade,
                skills,
                hc,
                duties,
                start,
                end,
                dur,
                mm,
                md,
                basis,
                unit,
                f'=IF(M{ws.max_row + 1}="","",M{ws.max_row + 1}*J{ws.max_row + 1})',
                "",
            ]
        )
    return _wb_bytes(wb)


def build_nonstandard_spec() -> bytes:
    """다른 양식: 제목 행 2줄, 헤더 이름 다름('번호', '명칭', '내용', '구분'), 시트명 '요구사항목록'."""
    wb = Workbook()
    ws = wb.active
    ws.title = "요구사항목록"
    ws.append(["OO 시스템 구축 요구사항 목록"])
    ws.append([])
    ws.append(["번호", "구분", "명칭", "내용", "대응방안", "수용", "참고"])
    for r in BASE_REQUIREMENTS:
        ws.append([r["id"], r["type"], r["name"], r["desc"], "표준", "수용", ""])
    return _wb_bytes(wb)


def build_ambiguous_spec() -> bytes:
    """매핑 확신이 낮은 양식: 'ID'만 있고 '명칭'·'내용' 대신 '항목'·'세부' 사용 → 확인 필요."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "항목", "세부 사항", "담당"])
    for r in BASE_REQUIREMENTS:
        ws.append([r["id"], r["name"], r["desc"], "PL"])
    return _wb_bytes(wb)


def build_corrupted_xlsx() -> bytes:
    return b"PK\x03\x04 this is not a real workbook"
