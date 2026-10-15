"""테스트용 **모의 문서** 생성기.

실제 고객 RFP가 아니라 자체 작성한 가상 RFP 본문(``SAMPLE_SECTIONS``)으로 DOCX/TXT/PDF/스캔 PDF를 만든다.
실제 문서로 검증한 항목이 아니므로 보고 시 "모의 문서"로 구분한다.
PDF 생성은 reportlab(텍스트 PDF), Pillow(이미지 전용 스캔 PDF), pypdf(병합·암호화)를 쓴다.
"""

from __future__ import annotations

import io
import random
from dataclasses import dataclass, field
from pathlib import Path

KOREAN_FONT = Path(r"C:\Windows\Fonts\malgun.ttf")


@dataclass
class Section:
    heading: str
    paragraphs: list[str] = field(default_factory=list)
    bullets: list[str] = field(default_factory=list)
    table: list[list[str]] | None = None


SAMPLE_TITLE = "차세대 민원 통합 플랫폼 구축 사업 제안요청서"

SAMPLE_SECTIONS: list[Section] = [
    Section(
        "제1장 사업 개요",
        [
            "본 사업은 분산된 민원 시스템을 통합하여 대민 서비스 품질을 높이는 것을 목적으로 한다.",
            "사업 기간은 계약일로부터 10개월이며, 사업 예산은 부가세 포함 12억 원이다.",
        ],
    ),
    Section(
        "제2장 기능 요구사항",
        ["다음 기능 요구사항은 모두 필수 구현 대상이다."],
        [
            "FR-A. 시스템은 사용자 인증을 공공기관 통합인증(SSO)으로 처리해야 한다.",
            "FR-B. 민원 접수 화면에서 첨부파일을 최대 5개, 각 20MB까지 업로드할 수 있어야 한다.",
            "FR-C. 처리 상태 변경 시 신청인에게 SMS와 이메일로 자동 통보해야 한다.",
            "FR-D. 관리자는 민원 유형별 처리 통계를 기간별로 조회하고 Excel로 내려받을 수 있어야 한다.",
        ],
    ),
    Section(
        "제3장 비기능 요구사항",
        [],
        [
            "NFR-A. 평균 응답 시간은 2초 이내, 동시 사용자 1,000명을 지원해야 한다.",
            "NFR-B. 서비스 가용성은 연 99.9% 이상이어야 하며 이중화 구성을 갖춘다.",
            "NFR-C. 개인정보는 암호화 저장하고 접근 이력을 3년간 보관해야 한다.",
        ],
        table=[
            ["항목", "기준", "비고"],
            ["응답 시간", "2초 이내", "피크 시간 기준"],
            ["가용성", "99.9%", "월 단위 측정"],
            ["동시 사용자", "1,000명", "부하 시험으로 검증"],
        ],
    ),
    Section(
        "제4장 제약사항 및 연계",
        [
            "CON-A. 시스템은 기관 표준 프레임워크(전자정부 표준프레임워크 4.x) 기반으로 개발해야 한다.",
            "CON-B. 기존 행정정보공동이용 시스템과 연계 인터페이스를 유지해야 한다.",
        ],
    ),
    Section(
        "제5장 평가 기준",
        ["기술 평가 90점, 가격 평가 10점으로 평가한다."],
        table=[
            ["평가 항목", "배점"],
            ["사업 이해도", "20"],
            ["기술 및 기능", "40"],
            ["수행 역량", "30"],
        ],
    ),
    Section(
        "제6장 제출 조건",
        [],
        [
            "SUB-A. 제안서는 2026년 11월 30일 15:00까지 전자 제출해야 한다.",
            "SUB-B. 제안서 본문은 100쪽 이내로 작성한다.",
        ],
    ),
]

SAMPLE_REQUIREMENT_KEYS = [
    "FR-A",
    "FR-B",
    "FR-C",
    "FR-D",
    "NFR-A",
    "NFR-B",
    "NFR-C",
    "CON-A",
    "CON-B",
    "SUB-A",
    "SUB-B",
]


def sample_text() -> str:
    lines: list[str] = [
        SAMPLE_TITLE,
        "",
        "목차",
        "제1장 사업 개요 ........ 1",
        "제2장 기능 요구사항 ........ 2",
        "",
    ]
    for s in SAMPLE_SECTIONS:
        lines.append(s.heading)
        lines.append("")
        for p in s.paragraphs:
            lines.append(p)
            lines.append("")
        for b in s.bullets:
            lines.append(f"- {b}")
        if s.bullets:
            lines.append("")
        if s.table:
            for row in s.table:
                lines.append(" | ".join(row))
            lines.append("")
    return "\n".join(lines)


def build_txt() -> bytes:
    return sample_text().encode("utf-8")


def build_docx(*, sections: list[Section] | None = None, title: str = SAMPLE_TITLE) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph("목차")
    doc.add_paragraph("제1장 사업 개요 ........ 1")
    for s in sections or SAMPLE_SECTIONS:
        doc.add_heading(s.heading, level=1)
        for p in s.paragraphs:
            doc.add_paragraph(p)
        for b in s.bullets:
            doc.add_paragraph(b, style="List Bullet")
        if s.table:
            table = doc.add_table(rows=len(s.table), cols=len(s.table[0]))
            table.style = "Table Grid"
            for r, row in enumerate(s.table):
                for c, val in enumerate(row):
                    table.cell(r, c).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _register_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if KOREAN_FONT.exists():
        if "Malgun" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("Malgun", str(KOREAN_FONT)))
        return "Malgun"
    return "Helvetica"


def build_text_pdf(
    *, sections: list[Section] | None = None, title: str = SAMPLE_TITLE, pad_pages: int = 0
) -> bytes:
    """텍스트 레이어와 괘선 표가 있는 PDF (reportlab platypus). ``pad_pages``만큼 본문 페이지를 늘린다."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    font = _register_font()
    body = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=15)
    head = ParagraphStyle(
        "head", fontName=font, fontSize=14, leading=18, spaceBefore=8, spaceAfter=6
    )
    title_style = ParagraphStyle("title", fontName=font, fontSize=18, leading=22, spaceAfter=12)
    story: list = [Paragraph(title, title_style)]
    for s in sections or SAMPLE_SECTIONS:
        story.append(Paragraph(s.heading, head))
        for p in s.paragraphs:
            story.append(Paragraph(p, body))
        for b in s.bullets:
            story.append(Paragraph(f"- {b}", body))
        if s.table:
            t = Table(s.table)
            t.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("FONTNAME", (0, 0), (-1, -1), font),
                    ]
                )
            )
            story.append(Spacer(1, 6))
            story.append(t)
        story.append(Spacer(1, 8))
    for i in range(pad_pages):
        story.append(PageBreak())
        story.append(
            Paragraph(
                f"부록 {i + 1}. 추가 설명 페이지입니다. 본문 요구사항과 무관한 참고 내용을 담는다.",
                body,
            )
        )
    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50
    ).build(story)
    return buf.getvalue()


def render_text_image(
    lines: list[str], *, width: int = 1240, height: int = 1754, font_size: int = 30
) -> object:
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font = (
        ImageFont.truetype(str(KOREAN_FONT), font_size)
        if KOREAN_FONT.exists()
        else ImageFont.load_default()
    )
    y = 80
    for line in lines:
        draw.text((80, y), line, font=font, fill="black")
        y += int(font_size * 1.8)
    return img


SCAN_LINES = [
    "제2장 기능 요구사항",
    "FR-A. 시스템은 사용자 인증을 SSO로 처리해야 한다.",
    "FR-C. 처리 상태 변경 시 SMS와 이메일로 통보해야 한다.",
]


def build_scanned_pdf(lines: list[str] | None = None, *, pages: int = 1) -> bytes:
    """텍스트 레이어 없이 이미지만 있는 PDF (Pillow가 이미지 전용 PDF를 만든다)."""
    img = render_text_image(lines or SCAN_LINES)
    buf = io.BytesIO()
    extra = [img.copy() for _ in range(pages - 1)]
    img.save(buf, format="PDF", resolution=150.0, save_all=bool(extra), append_images=extra)
    return buf.getvalue()


def build_mixed_pdf() -> bytes:
    """1쪽 텍스트 + 1쪽 스캔."""
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for src in (build_text_pdf(sections=SAMPLE_SECTIONS[:1]), build_scanned_pdf()):
        for page in PdfReader(io.BytesIO(src)).pages:
            writer.add_page(page)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def build_encrypted_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(
        clone_from=PdfReader(io.BytesIO(build_text_pdf(sections=SAMPLE_SECTIONS[:1])))
    )
    writer.encrypt(user_password="secret", owner_password="owner", algorithm="AES-256")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def build_corrupted_bytes(seed: int = 7, size: int = 4096) -> bytes:
    rng = random.Random(seed)
    return b"%PDF-1.7\n" + bytes(rng.getrandbits(8) for _ in range(size))


def build_corrupted_docx() -> bytes:
    return b"PK\x03\x04" + b"\x00" * 64 + b"not-a-zip"
