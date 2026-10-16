"""통합 개발 지시문 단계 8 — 전체 흐름 통합 테스트.

제안(RFP → 요구 → 상세 설계 → 산출물 생성 → QA) → 생성된 정의서·추적표·WBS Excel 을 구축 프로젝트의 **기준선** 으로 적재 →
수정 정의서 Excel 변경 요청 → 분석·승인·반영 → 설계 객체 저장·테스트 계획·JUnit 결과(실행기 스텁) 반입 → 통합 QA(구축) → Revision 복원.
회귀: 같은 입력으로 QA 를 두 번 돌려 결과 건수가 같은지 확인한다."""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from fixtures import builders as bld
from proposal_agent.agents import FakeAgentRunner
from proposal_agent.development.test_runner import (
    DockerSandbox,
    DockerStatus,
    ProcessOutcome,
    TestRunner,
)
from proposal_agent.models import (
    ChangeStatus,
    DeliverableKind,
    DocumentKind,
    ProjectMode,
    TestCaseStatus,
)
from proposal_agent.models.change import CandidateClass, InputKind
from proposal_agent.orchestrator import StepRunner
from proposal_agent.services import AppContext, Services
from unit.test_requirements_phase2 import extractor_from_text, verifier_all_ok

DESIGN = {
    "screens": [
        {
            "name": "민원 접수 화면",
            "screen_type": "form",
            "purpose": "접수",
            "fields": [
                {
                    "name": "title",
                    "label": "제목",
                    "field_type": "text",
                    "required": True,
                    "validation_rules": ["1~100자"],
                    "error_message": "제목",
                    "requirement_ids": [],
                }
            ],
            "states": [{"state": "ERROR", "message": "실패"}],
            "functions": [
                {
                    "name": "접수",
                    "trigger": "버튼",
                    "inputs": ["title"],
                    "process": "검증 → 저장",
                    "outputs": ["번호"],
                    "requirement_ids": [],
                }
            ],
            "requirement_ids": [],
        }
    ],
    "tables": [
        {
            "name": "TB_CASE",
            "logical_name": "민원",
            "columns": [
                {"name": "CASE_ID", "data_type": "NUMBER", "is_pk": True, "nullable": False}
            ],
            "db_engine": "PostgreSQL",
            "requirement_ids": [],
        }
    ],
    "apis": [
        {
            "name": "민원 접수 API",
            "method": "POST",
            "path": "/api/cases",
            "error_codes": ["C001: 제목 누락"],
            "auth": "세션",
            "screen_refs": ["민원 접수 화면"],
            "table_refs": ["TB_CASE"],
            "requirement_ids": [],
        }
    ],
    "modules": [
        {
            "path": "src/main/java/com/app/cases/CaseService.java",
            "layer": "service",
            "responsibility": "접수 규칙",
            "requirement_ids": [],
        }
    ],
    "business_rules": [],
}
SVC = "src/main/java/com/app/cases/CaseService.java"
TEST = "src/test/java/com/app/cases/CaseServiceTest.java"


def _stub_runner() -> TestRunner:
    def process(cmd, *, cwd, env, timeout, stdout_path, stderr_path, on_timeout=None):  # type: ignore[no-untyped-def]
        sf = cwd / "out" / "surefire"
        sf.mkdir(parents=True, exist_ok=True)
        (sf / "TEST-com.app.cases.CaseServiceTest.xml").write_text(
            '<?xml version="1.0"?><testsuite name="com.app.cases.CaseServiceTest" tests="1"><testcase classname="com.app.cases.CaseServiceTest" name="accepts"/></testsuite>',
            encoding="utf-8",
        )
        stdout_path.write_text("Tests run: 1, Failures: 0, Errors: 0, Skipped: 0", encoding="utf-8")
        stderr_path.write_text("", encoding="utf-8")
        return ProcessOutcome(exit_code=0)

    class Ready(DockerSandbox):
        def status(self) -> DockerStatus:
            return DockerStatus(True, True, True, "stub")

    return TestRunner(process_runner=process, docker=Ready("stub:1", executable="docker"))


def _design_with_ids(request):  # type: ignore[no-untyped-def]
    import json
    import re

    ids = list(dict.fromkeys(re.findall(r'"id": "(FR-\d{3})"', request.prompt)))
    first = ids
    d = json.loads(json.dumps(DESIGN))
    d["screens"][0]["requirement_ids"] = first
    d["screens"][0]["fields"][0]["requirement_ids"] = first
    d["screens"][0]["functions"][0]["requirement_ids"] = first
    d["tables"][0]["requirement_ids"] = first
    d["apis"][0]["requirement_ids"] = first
    d["modules"][0]["requirement_ids"] = first
    return d


def test_full_flow_proposal_to_build_to_tests_and_restore(ctx: AppContext) -> None:
    runner = FakeAgentRunner(
        {
            "requirement_extractor": extractor_from_text,
            "requirement_verifier": verifier_all_ok,
            "detail_designer": _design_with_ids,
        },
        cost_per_call=0.0,
    )
    s = Services.from_context(ctx, runner=runner)
    r = StepRunner(s)
    # 1) 제안: RFP → 요구 → Gate 1 → 상세 설계 → 산출물 생성 → QA
    p = s.projects.create(
        name="통합 흐름 제안",
        mode=ProjectMode.PROPOSAL,
        start_date=date(2026, 11, 2),
        end_date=date(2027, 8, 31),
        client_name="OO기관",
    )
    s.documents.register_upload(
        p, kind=DocumentKind.RFP, file_name="rfp.docx", data=bld.build_docx()
    )
    r.run_extraction(p)
    r.run_requirement_extraction(p)
    r.run_verification(p)
    reqs = s.ctx.repos.requirements.find(p.id, order_by="id")
    assert reqs
    out = r.run_detail_design(p, actor="t")
    assert "저장: 화면 1" in out.summary and s.detail_design.has_design(p)
    r.run_generation(p)
    qa1 = r.run_qa(p)
    run1 = s.qa.latest_run(p)
    assert run1 is not None and run1.counts["fail"] == 0, qa1.summary
    latest = s.generation.latest_deliverables(p)
    for kind in (
        DeliverableKind.REQUIREMENT_SPEC_XLSX,
        DeliverableKind.TRACE_MATRIX_XLSX,
        DeliverableKind.WBS_GANTT_XLSX,
        DeliverableKind.DESIGN_DOC_XLSX,
    ):
        assert kind in latest, kind
    # 회귀: 같은 상태에서 QA 재실행 → 같은 건수
    r.run_qa(p)
    run2 = s.qa.latest_run(p)
    assert run2 is not None and {k: run2.counts[k] for k in ("executed", "pass", "fail")} == {
        k: run1.counts[k] for k in ("executed", "pass", "fail")
    }
    # 2) 구축: 제안 산출물 Excel 을 기준선으로 적재
    b = s.projects.create(name="통합 흐름 구축", mode=ProjectMode.BUILD)
    for kind, dk in (
        (DocumentKind.REQUIREMENT_SPEC, DeliverableKind.REQUIREMENT_SPEC_XLSX),
        (DocumentKind.TRACE_MATRIX, DeliverableKind.TRACE_MATRIX_XLSX),
        (DocumentKind.WBS, DeliverableKind.WBS_GANTT_XLSX),
    ):
        data = s.ctx.paths.resolve_output(latest[dk].file_path).read_bytes()
        s.documents.register_upload(
            b, kind=kind, file_name=Path(latest[dk].file_path).name, data=data
        )
    profiles = s.baseline.profiles(b)
    assert all(not pr.blocking or pr.confirmed for pr in profiles.values()), {
        k.value: pr.blocking for k, pr in profiles.items()
    }
    result = s.baseline.load(b, actor="t")
    b = s.ctx.repos.projects.get(b.id)
    base_reqs = s.changes.active_requirements(b)
    assert (
        len(base_reqs) == len([x for x in reqs if x.change_state != "deleted"])
        and b.baseline_revision_id
    )
    first_id = next(x.id for x in base_reqs if x.type.value == "functional")  # 기능 시트의 행
    # 3) 변경: 수정 정의서(설명 수치 변경) → 분석 → 수락 → 승인·반영
    spec_bytes = s.ctx.paths.resolve_output(
        latest[DeliverableKind.REQUIREMENT_SPEC_XLSX].file_path
    ).read_bytes()
    wb = load_workbook(io.BytesIO(spec_bytes))
    ws = wb["기능요구사항"]
    hdr = [c.value for c in ws[2]]
    cid, cdesc = (
        hdr.index("요구사항 ID"),
        hdr.index("요구사항 상세 내용"),
    )
    for row in ws.iter_rows(min_row=3):
        if row[cid].value == first_id:
            row[cdesc].value = str(row[cdesc].value) + " 처리 기한은 3일 이내로 한다."
            break
    buf = io.BytesIO()
    wb.save(buf)
    cr = s.changes.create_request(
        b,
        input_kind=InputKind.DOCUMENT,
        request_text="",
        requester="고객 PM",
        reason="정의서 v1.1",
        attachments=[("수정정의서.xlsx", buf.getvalue())],
        actor="t",
    )
    cr = s.changes.analyze(b, cr, use_agent=False)
    mod = [c for c in cr.candidates if c.classification == CandidateClass.MODIFY]
    assert mod and mod[0].requirement_id == first_id and mod[0].change_kind == "semantic"
    for c in cr.candidates:
        if c.classification != CandidateClass.NO_CHANGE and c.user_decision is None:
            s.changes.decide_candidate(
                b, cr, c.id, decision="accept" if c.id == mod[0].id else "reject"
            )
    cr = s.changes.get(b, cr.id)
    assert cr.status == ChangeStatus.PENDING_APPROVAL
    res = s.changes.approve(b, cr, approver="PM", comments="")
    assert (
        res.request.status == ChangeStatus.APPLIED
        and "3일" in s.ctx.repos.requirements.get(first_id, b.id).description
    )
    # 4) 구축 프로젝트 설계·테스트: 상세 설계 저장 → 테스트 계획 → JUnit 결과(스텁) → 추적 사슬
    design = json_with_id([x.id for x in base_reqs if x.type.value == "functional"])
    s.detail_design.persist(b, design, actor="t")
    counts = s.testing.plan_from_design(b, actor="t")
    assert counts["integration"] >= 1 and counts["ui"] == 1
    ws_ = s.development.workspace(b)
    ws_.write(
        SVC,
        "package com.app.cases;\n\npublic class CaseService {\n    public boolean accept(String title) {\n        return title != null && !title.isBlank();\n    }\n}\n",
    )
    ws_.write(
        TEST,
        f'package com.app.cases;\n\nimport static org.junit.jupiter.api.Assertions.assertTrue;\nimport org.junit.jupiter.api.DisplayName;\nimport org.junit.jupiter.api.Test;\n\nclass CaseServiceTest {{\n    @Test\n    @DisplayName("{first_id} TC-001: 제목이 있으면 접수된다")\n    void accepts() {{\n        assertTrue(new CaseService().accept("x"));\n    }}\n}}\n',
    )
    s.development._runner = _stub_runner()
    tr = s.development.run_tests(b, actor="t")
    assert tr.is_success
    tcs = s.testing.test_cases(b)
    assert any(
        t.status == TestCaseStatus.PASSED and t.related_requirement_ids == [first_id] for t in tcs
    )
    qa_b = s.qa.run(b, actor="t").qa_run
    assert qa_b is not None and qa_b.counts["fail"] == 0, [
        c.issues for c in qa_b.checks if c.final_result.value == "FAIL"
    ]
    con16 = next(c for c in qa_b.checks if c.check_id == "QA-CON-16")
    assert con16.executed
    # 5) 복원: 기준선 Revision 으로 되돌리면 변경 반영 전 설명으로 돌아가고, 복원 전 상태는 새 Revision 으로 남는다
    baseline_rev = s.ctx.repos.revisions.get(b.baseline_revision_id, b.id)
    rev = s.baseline.restore_revision(b, baseline_rev.id, actor="t", reason="변경 취소")
    restored = s.ctx.repos.requirements.get(first_id, b.id)
    assert "3일" not in restored.description and rev.label == "복원"
    revs = s.ctx.repos.revisions.find(b.id, order_by="created_at")
    assert [x.label for x in revs[-2:]] == ["복원 전 상태", "복원"]
    assert result.requirements >= 1


def json_with_id(rids: list[str]) -> dict:  # type: ignore[type-arg]
    import json

    d = json.loads(json.dumps(DESIGN))
    for key in ("screens", "tables", "apis", "modules"):
        for item in d[key]:
            item["requirement_ids"] = list(rids)
    d["screens"][0]["fields"][0]["requirement_ids"] = list(rids)
    d["screens"][0]["functions"][0]["requirement_ids"] = list(rids)
    return d
