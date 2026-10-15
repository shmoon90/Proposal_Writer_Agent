"""인프로세스 SDK MCP 서버 ``pwa``: 프로젝트 데이터 조회(읽기 전용)와 산출물 파일 검증.

도구 이름은 ``mcp__pwa__<tool>``로 노출된다. 핸들러는 순수 함수(``tool_*``)로 두고 ``@tool``로 감싸므로
테스트에서 SDK 없이 직접 호출할 수 있다. 쓰기 도구는 제공하지 않는다.
"""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from proposal_agent.config.errors import NotFoundError, PathSecurityError, PWAError
from proposal_agent.config.paths import ProjectPaths
from proposal_agent.models import Project, RequirementStatus
from proposal_agent.storage.repositories import Repositories

MAX_LIST = 200
MAX_TEXT = 4000


@dataclass
class ToolContext:
    repos: Repositories
    paths: ProjectPaths
    project: Project


def _json(payload: Any) -> dict[str, Any]:
    return {
        "content": [{"type": "text", "text": json.dumps(payload, ensure_ascii=False, default=str)}]
    }


def _error(message: str, resolution: str = "") -> dict[str, Any]:
    return {
        "content": [
            {
                "type": "text",
                "text": json.dumps(
                    {"error": message, "resolution": resolution}, ensure_ascii=False
                ),
            }
        ],
        "is_error": True,
    }


# --- 순수 핸들러 ----------------------------------------------------------------------------


def tool_get_project(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    p = ctx.project
    return {
        "id": p.id,
        "name": p.name,
        "mode": str(p.mode),
        "status": str(p.status),
        "client_name": p.client_name,
        "start_date": str(p.start_date) if p.start_date else None,
        "end_date": str(p.end_date) if p.end_date else None,
        "duration_days": p.duration_days,
        "settings": {k: v for k, v in p.settings.items() if "key" not in k.lower()},
    }


def tool_list_requirements(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    filters: dict[str, Any] = {}
    if args.get("type"):
        filters["type"] = str(args["type"])
    if args.get("status"):
        filters["status"] = str(args["status"])
    reqs = ctx.repos.requirements.find(ctx.project.id, order_by="id", **filters)
    include_deleted = bool(args.get("include_deleted", False))
    rows = [
        {
            "id": r.id,
            "name": r.name,
            "type": str(r.type),
            "scope_class": str(r.scope_class),
            "is_mandatory": r.is_mandatory,
            "status": str(r.status),
            "verification_status": str(r.verification_status),
            "priority": str(r.priority),
        }
        for r in reqs
        if include_deleted or r.status != RequirementStatus.DELETED
    ][:MAX_LIST]
    return {"count": len(rows), "requirements": rows}


def tool_get_requirement(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    req_id = str(args.get("id", "")).strip()
    try:
        r = ctx.repos.requirements.get(req_id, ctx.project.id)
    except NotFoundError:
        return {"error": f"요구사항 없음: {req_id}"}
    data = r.model_dump(mode="json")
    data.pop("history", None)
    return data


def tool_search_source_text(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    query = str(args.get("query", "")).strip()
    if len(query) < 2:
        return {"error": "query는 2자 이상이어야 합니다."}
    limit = min(int(args.get("limit", 10) or 10), 50)
    runs = ctx.repos.extraction_runs.latest_by_document(ctx.project.id)
    run_ids = [r.id for r in runs.values() if r.ok]
    blocks = ctx.repos.blocks.search(ctx.project.id, query, run_ids=run_ids or None, limit=limit)
    return {
        "count": len(blocks),
        "blocks": [
            {
                "id": b.id,
                "file": b.file_name,
                "loc": b.location.label(),
                "type": str(b.block_type),
                "method": str(b.extraction_method),
                "confidence": b.confidence,
                "text": b.text[:600],
            }
            for b in blocks
        ],
    }


def tool_get_source_block(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    block_id = str(args.get("block_id", "")).strip()
    try:
        b = ctx.repos.blocks.get(block_id)
    except NotFoundError:
        return {"error": f"블록 없음: {block_id}"}
    if b.project_id != ctx.project.id:
        return {"error": "다른 프로젝트의 블록입니다."}
    return {
        "id": b.id,
        "file": b.file_name,
        "loc": b.location.label(),
        "type": str(b.block_type),
        "method": str(b.extraction_method),
        "confidence": b.confidence,
        "needs_review": b.needs_review,
        "text": b.text[:MAX_TEXT],
        "table_cells": b.table_cells,
    }


def tool_list_research(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    req_id = str(args.get("requirement_id", "") or "").strip()
    items = ctx.repos.research.find(ctx.project.id, order_by="created_at")
    rows = []
    for t in items:
        if req_id and req_id not in t.related_requirement_ids:
            continue
        rows.append(
            {
                "id": t.id,
                "topic": t.topic,
                "feasibility": str(t.feasibility),
                "summary": t.summary[:400],
                "related_requirement_ids": t.related_requirement_ids,
                "investigated_at": str(t.investigated_at),
                "sources": [
                    {
                        "url": s.url,
                        "title": s.title,
                        "trust": str(s.trust),
                        "product": s.product,
                        "version": s.version,
                    }
                    for s in t.sources
                ],
                "unverified_items": t.unverified_items,
            }
        )
    return {"count": len(rows), "research": rows[:MAX_LIST]}


def tool_list_findings(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    req_id = str(args.get("requirement_id", "") or "").strip()
    rows = ctx.repos.validations.find(ctx.project.id, resolved=False)
    out = [
        {
            "id": v.id,
            "target": v.target_id,
            "rule": v.rule_id,
            "level": str(v.level),
            "message": v.message,
            "requirement_ids": v.detail.get("requirement_ids", []),
        }
        for v in rows
        if not req_id or v.target_id == req_id or req_id in v.detail.get("requirement_ids", [])
    ]
    return {"count": len(out), "findings": out[:MAX_LIST]}


def tool_list_gates(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    out = {}
    for key in ("gate0", "gate1"):
        ev = ctx.repos.gate_evaluations.latest(ctx.project.id, key)
        out[key] = (
            None
            if ev is None
            else {"status": str(ev.status), "reasons": ev.reasons, "approval_id": ev.approval_id}
        )
    return out


def tool_get_prior_output(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    """이전 역할 실행 결과(JSON 산출물)를 돌려준다. 같은 작업을 반복하지 않도록 재사용한다."""
    role = str(args.get("role", "")).strip()
    rows = ctx.repos.agent_outputs.find(
        ctx.project.id, role=role, order_by="created_at", descending=True, limit=1
    )
    if not rows:
        return {"role": role, "found": False}
    rec = rows[0]
    try:
        path = ctx.paths.resolve_output(rec.file_path)
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError, PathSecurityError) as exc:
        return {"role": role, "found": True, "error": f"파일을 읽을 수 없음: {exc}"}
    text = json.dumps(payload, ensure_ascii=False)
    return {
        "role": role,
        "found": True,
        "version": rec.version,
        "created_at": str(rec.created_at),
        "output": payload if len(text) <= 20000 else {"truncated": True, "preview": text[:20000]},
    }


_OPENERS: dict[str, str] = {
    ".xlsx": "xlsx",
    ".pptx": "pptx",
    ".docx": "docx",
    ".pdf": "pdf",
    ".png": "image",
    ".json": "json",
    ".md": "text",
}


def verify_file(path: Path) -> dict[str, Any]:
    """파일을 실제로 열어 구조를 확인한다 (xlsx/pptx/docx/pdf/png/json/md)."""
    info: dict[str, Any] = {"path": str(path), "exists": path.exists()}
    if not path.exists():
        info["ok"] = False
        info["error"] = "파일이 없습니다."
        return info
    data = path.read_bytes()
    info.update({"size_bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()})
    kind = _OPENERS.get(path.suffix.lower())
    try:
        if kind == "xlsx":
            import io

            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True)
            info["sheets"] = {ws.title: ws.max_row for ws in wb.worksheets}
        elif kind == "pptx":
            import io

            from pptx import Presentation

            prs = Presentation(io.BytesIO(data))
            info["slides"] = len(prs.slides)
        elif kind == "docx":
            import io

            from docx import Document

            doc = Document(io.BytesIO(data))
            info["paragraphs"] = len(doc.paragraphs)
            info["tables"] = len(doc.tables)
        elif kind == "pdf":
            import io

            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            info["pages"] = len(reader.pages)
            info["encrypted"] = bool(reader.is_encrypted)
        elif kind == "image":
            import io

            from PIL import Image

            with Image.open(io.BytesIO(data)) as im:
                im.verify()
                info["format"] = im.format
        elif kind == "json":
            json.loads(data.decode("utf-8"))
            info["json"] = True
        elif kind == "text":
            info["chars"] = len(data.decode("utf-8"))
        else:
            info["note"] = "구조 검증을 지원하지 않는 형식 (존재·해시만 확인)"
        info["ok"] = True
    except Exception as exc:  # noqa: BLE001 - 손상 파일 등
        info["ok"] = False
        info["error"] = f"{type(exc).__name__}: {str(exc)[:200]}"
    return info


def tool_verify_deliverable_file(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    rel = str(args.get("relative_path", "")).strip()
    if not rel:
        return {"error": "relative_path가 필요합니다 (프로젝트 산출물 폴더 기준)."}
    try:
        path = ctx.paths.resolve_in_project(ctx.project.output_dir, rel)
    except PWAError as exc:
        return {"error": exc.message, "resolution": exc.resolution}
    info = verify_file(path)
    info["path"] = rel
    return info


# --- SDK 서버 조립 ---------------------------------------------------------------------------

TOOL_SPECS: tuple[tuple[str, str, dict[str, Any], Any], ...] = (
    ("get_project", "프로젝트 기본 정보(이름·모드·기간·상태)를 돌려준다.", {}, tool_get_project),
    (
        "list_requirements",
        "요구사항 목록(요약). type/status로 필터, include_deleted로 삭제 항목 포함.",
        {"type": str, "status": str, "include_deleted": bool},
        tool_list_requirements,
    ),
    ("get_requirement", "요구사항 1건 전체(출처 포함).", {"id": str}, tool_get_requirement),
    (
        "search_source_text",
        "RFP 원문 블록을 부분 문자열로 검색한다.",
        {"query": str, "limit": int},
        tool_search_source_text,
    ),
    ("get_source_block", "원문 블록 1건 전문.", {"block_id": str}, tool_get_source_block),
    (
        "list_research",
        "기술 조사 결과 목록. requirement_id로 필터.",
        {"requirement_id": str},
        tool_list_research,
    ),
    (
        "list_findings",
        "미해결 검증 지적 목록. requirement_id로 필터.",
        {"requirement_id": str},
        tool_list_findings,
    ),
    ("list_gates", "Gate 0/1 최신 상태.", {}, tool_list_gates),
    (
        "get_prior_output",
        "이전 역할(role)의 최신 실행 결과 JSON.",
        {"role": str},
        tool_get_prior_output,
    ),
    (
        "verify_deliverable_file",
        "프로젝트 산출물 파일을 열어 존재·해시·구조(시트/슬라이드/페이지)를 확인한다.",
        {"relative_path": str},
        tool_verify_deliverable_file,
    ),
)


def build_project_server(ctx: ToolContext, *, name: str = "pwa") -> Any:
    """``McpSdkServerConfig``를 만든다. 핸들러는 예외를 삼켜 is_error 응답으로 바꾼다."""
    from claude_agent_sdk import create_sdk_mcp_server, tool

    sdk_tools = []
    for tool_name, description, schema, fn in TOOL_SPECS:

        def make(fn_: Any) -> Any:
            async def handler(args: dict[str, Any]) -> dict[str, Any]:
                try:
                    payload = fn_(ctx, args or {})
                except PWAError as exc:
                    return _error(exc.message, exc.resolution)
                except Exception as exc:  # noqa: BLE001
                    return _error(f"{type(exc).__name__}: {str(exc)[:200]}")
                if isinstance(payload, dict) and "error" in payload and len(payload) <= 2:
                    return _error(str(payload["error"]), str(payload.get("resolution", "")))
                return _json(payload)

            return handler

        sdk_tools.append(tool(tool_name, description, schema)(make(fn)))
    return create_sdk_mcp_server(name=name, version="0.3.0", tools=sdk_tools)


TOOL_NAMES: tuple[str, ...] = tuple(f"mcp__pwa__{spec[0]}" for spec in TOOL_SPECS)
