"""산출물 품질 검사 (코드).

- PPT: 열림, 슬라이드 크기·방향(16:9 가로), 도형이 슬라이드 밖으로 나가거나 텍스트 도형끼리 크게 겹치는지, 텍스트 넘침 추정
- Excel: 열림, 필수 시트·컬럼, 수식 형식, 날짜 순서, ID가 공통 데이터와 일치, WBS 상위 참조
- 일관성: WBS 총 공수 == 인력 계획 총 공수, 제출용·발표용 PPT 핵심 수치 일치, PPT 기술명이 구성 요소와 일치
- 추적성: 필수 요구사항이 제출용 PPT 요구사항 대응 표에 모두 있는지
결과는 ValidationResult(target=deliverable)와 QA 보고서(md)로 남긴다.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.excel.requirement_spec import (
    REQUIRED_COLUMNS,
    SHEET_FUNCTIONAL,
    SHEET_HISTORY,
    SHEET_NON_FUNCTIONAL,
    SHEET_OTHER,
)
from proposal_agent.generators.excel.staffing_plan import SHEET_MONTHLY, SHEET_PLAN, SHEET_SUMMARY
from proposal_agent.generators.excel.trace_matrix import SHEET_MAIN, SHEET_TESTS, SHEET_UNTRACED
from proposal_agent.generators.excel.wbs_gantt import (
    SHEET_BASIS,
    SHEET_DEPS,
    SHEET_GANTT,
    SHEET_WBS,
)
from proposal_agent.generators.pptx.decks import key_figures
from proposal_agent.models import DeliverableKind, ValidationLevel

REQUIRED_SHEETS: dict[DeliverableKind, dict[str, tuple[str, ...]]] = {
    DeliverableKind.REQUIREMENT_SPEC_XLSX: {
        SHEET_FUNCTIONAL: REQUIRED_COLUMNS,
        SHEET_NON_FUNCTIONAL: REQUIRED_COLUMNS,
        SHEET_OTHER: REQUIRED_COLUMNS,
        SHEET_HISTORY: ("요구사항 ID",),
    },
    DeliverableKind.TRACE_MATRIX_XLSX: {
        SHEET_MAIN: (
            "요구사항 ID",
            "요구사항 명",
            "화면 ID",
            "단위테스트 ID",
            "단위테스트 결과",
            "통합테스트 ID",
            "통합테스트 결과",
        ),
        SHEET_TESTS: ("테스트 ID",),
        SHEET_UNTRACED: ("요구사항 ID",),
    },
    DeliverableKind.WBS_GANTT_XLSX: {
        SHEET_WBS: (
            "WBS ID",
            "상위 WBS ID",
            "요구사항 ID",
            "Task 명",
            "계획 시작일",
            "계획 종료일",
            "기간(영업일)",
            "진행 상태",
            "산정 근거",
        ),
        SHEET_GANTT: ("WBS ID",),
        SHEET_DEPS: ("선행 WBS", "후행 WBS"),
        SHEET_BASIS: ("WBS ID", "산정 근거"),
    },
    DeliverableKind.STAFFING_PLAN_XLSX: {
        SHEET_PLAN: (
            "역할",
            "기술 등급",
            "투입 인원",
            "담당 업무",
            "투입 시작일",
            "투입 종료일",
            "예상 공수(MM)",
            "산정 근거",
            "단가(원/MM, 입력)",
            "비용 합계(원)",
        ),
        SHEET_SUMMARY: ("역할", "공수(M/D)"),
        SHEET_MONTHLY: ("역할(MM)",),
    },
}


@dataclass
class Finding:
    level: ValidationLevel
    rule: str
    message: str
    location: str = ""


@dataclass
class QaOutcome:
    findings: list[Finding] = field(default_factory=list)
    opened: dict[str, bool] = field(default_factory=dict)
    figures: dict[str, dict[str, str]] = field(default_factory=dict)

    @property
    def errors(self) -> list[Finding]:
        return [f for f in self.findings if f.level == ValidationLevel.ERROR]

    @property
    def warnings(self) -> list[Finding]:
        return [f for f in self.findings if f.level == ValidationLevel.WARNING]

    def add(self, level: ValidationLevel, rule: str, message: str, location: str = "") -> None:
        self.findings.append(Finding(level, rule, message, location))


# --- PPT -------------------------------------------------------------------------------------


def pptx_texts(path: Path) -> list[str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[str] = []
    for slide in prs.slides:
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        out.append("\n".join(parts))
    return out


def _has_picture(shapes: Any, depth: int = 0) -> bool:
    """슬라이드(그룹 안 포함)에 그림 도형이 있는가. 마스터·레이아웃의 그림은 보지 않는다."""
    for sh in shapes:
        st = str(getattr(sh, "shape_type", "") or "")
        if "PICTURE" in st:
            return True
        if "GROUP" in st and depth < 4 and _has_picture(sh.shapes, depth + 1):
            return True
    return False


def _first_font_pt(shape: Any) -> float | None:
    """텍스트 상자의 첫 run(없으면 문단) 글꼴 크기(pt)."""
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.text.strip() and r.font.size:
                return float(r.font.size.pt)
        if p.font.size:
            return float(p.font.size.pt)
    return None


def check_pptx(
    path: Path,
    outcome: QaOutcome,
    label: str,
    *,
    expected_size_cm: tuple[float, float] | None = None,
) -> list[str]:
    """PPT 열기·규격·경계·넘침·겹침 검사. ``expected_size_cm``(템플릿 용지)가 있으면 16:9 대신 그 크기와 대조한다."""
    from pptx import Presentation
    from pptx.util import Emu

    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        outcome.opened[label] = False
        outcome.add(
            ValidationLevel.ERROR,
            "pptx.open",
            f"{label}: 파일을 열 수 없음 ({type(exc).__name__}: {exc})",
            path.name,
        )
        return []
    outcome.opened[label] = True
    w, h = int(prs.slide_width or 0), int(prs.slide_height or 0)
    if expected_size_cm is not None:
        ew, eh = expected_size_cm
        aw, ah = Emu(w).cm, Emu(h).cm
        if abs(aw - ew) > 0.2 or abs(ah - eh) > 0.2:
            outcome.add(
                ValidationLevel.ERROR,
                "pptx.size",
                f"{label}: 슬라이드 {aw:.1f}×{ah:.1f} cm — 템플릿 {ew:.1f}×{eh:.1f} cm 와 다름",
                path.name,
            )
    else:
        if w <= h:
            outcome.add(
                ValidationLevel.ERROR,
                "pptx.orientation",
                f"{label}: 슬라이드가 가로 방향이 아님",
                path.name,
            )
        ratio = w / h
        if abs(ratio - 16 / 9) > 0.05:
            outcome.add(
                ValidationLevel.WARNING,
                "pptx.size",
                f"{label}: 슬라이드 비율 {ratio:.2f} (16:9 아님)",
                path.name,
            )
    if len(prs.slides) < 3:
        outcome.add(
            ValidationLevel.ERROR,
            "pptx.slides",
            f"{label}: 슬라이드가 {len(prs.slides)}장뿐",
            path.name,
        )
    texts: list[str] = []
    pictures: list[int] = []
    for idx, slide in enumerate(prs.slides, start=1):
        boxes: list[tuple[int, int, int, int, bool]] = []
        slide_text: list[str] = []
        if _has_picture(slide.shapes):
            pictures.append(idx)
        for shape in slide.shapes:
            if shape.left is None or shape.top is None:
                continue
            left, top, width, height = (
                int(shape.left),
                int(shape.top),
                int(shape.width or 0),
                int(shape.height or 0),
            )
            if left < 0 or top < 0 or left + width > w + 1000 or top + height > h + 1000:
                outcome.add(
                    ValidationLevel.ERROR,
                    "pptx.bounds",
                    f"{label} {idx}장: 도형 '{shape.name}'이 슬라이드 밖으로 나감",
                    f"slide {idx}",
                )
            has_text = shape.has_text_frame and bool(shape.text_frame.text.strip())
            if has_text:
                slide_text.append(shape.text_frame.text)
                text = shape.text_frame.text
                width_in = max(width / 914400, 0.1)
                height_in = max(height / 914400, 0.1)
                # 글꼴 크기 기반 추정: 줄 높이 ≈ 1.3em, 한 줄 글자 수 ≈ 폭 / 0.75em (한글·영문 혼용 평균)
                size_pt = _first_font_pt(shape) or 18.0
                line_h_in = size_pt * 1.3 / 72
                cpl = max(1, int(width_in * 72 / (size_pt * 0.75)))
                lines = sum(max(1, -(-len(line) // cpl)) for line in text.splitlines())
                if lines * line_h_in > height_in * 1.25 and height_in < 5:
                    outcome.add(
                        ValidationLevel.WARNING,
                        "pptx.overflow",
                        f"{label} {idx}장: 텍스트가 도형을 넘칠 수 있음 ({lines}줄 추정)",
                        f"slide {idx}",
                    )
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    slide_text.append(" | ".join(c.text for c in row.cells))
            boxes.append((left, top, width, height, has_text))
        for i, a in enumerate(boxes):
            for b in boxes[i + 1 :]:
                if not (a[4] and b[4]):
                    continue
                ix = max(0, min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0]))
                iy = max(0, min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]))
                inter = ix * iy
                smaller = max(1, min(a[2] * a[3], b[2] * b[3]))
                if inter / smaller > 0.3:
                    outcome.add(
                        ValidationLevel.WARNING,
                        "pptx.overlap",
                        f"{label} {idx}장: 텍스트 도형이 겹침",
                        f"slide {idx}",
                    )
                    break
        texts.append("\n".join(slide_text))
    if pictures:
        outcome.add(
            ValidationLevel.ERROR,
            "pptx.picture",
            f"{label}: 그림(비트맵) 도형이 있어 편집할 수 없음 — {len(pictures)}장 (예: {', '.join(str(i) for i in pictures[:6])})",
            path.name,
        )
    return texts


# --- Excel ------------------------------------------------------------------------------------


def check_xlsx(
    path: Path, kind: DeliverableKind, outcome: QaOutcome, ctx: GenerationContext
) -> dict[str, list[list[Any]]]:
    from openpyxl import load_workbook

    label = kind.value
    try:
        wb = load_workbook(io.BytesIO(path.read_bytes()))
    except Exception as exc:  # noqa: BLE001
        outcome.opened[label] = False
        outcome.add(
            ValidationLevel.ERROR,
            "xlsx.open",
            f"{label}: 파일을 열 수 없음 ({type(exc).__name__}: {exc})",
            path.name,
        )
        return {}
    outcome.opened[label] = True
    sheets: dict[str, list[list[Any]]] = {}
    for sheet, cols in REQUIRED_SHEETS.get(kind, {}).items():
        if sheet not in wb.sheetnames:
            outcome.add(
                ValidationLevel.ERROR, "xlsx.sheet", f"{label}: 시트 '{sheet}' 없음", path.name
            )
            continue
        ws = wb[sheet]
        headers = [str(c.value) if c.value is not None else "" for c in ws[1]]
        missing = [c for c in cols if c not in headers]
        if missing:
            outcome.add(
                ValidationLevel.ERROR,
                "xlsx.columns",
                f"{label}/{sheet}: 필수 컬럼 누락 {missing}",
                sheet,
            )
        rows: list[list[Any]] = []
        for row in ws.iter_rows(min_row=2):
            values = [c.value for c in row]
            if not any(v is not None for v in values):
                break
            rows.append(values)
        sheets[sheet] = [headers, *rows]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    if "#REF" in c.value or c.value.count("(") != c.value.count(")"):
                        outcome.add(
                            ValidationLevel.ERROR,
                            "xlsx.formula",
                            f"{label}/{sheet}: 수식 오류 {c.coordinate} {c.value}",
                            c.coordinate,
                        )
    # 종류별 데이터 검사
    req_ids = {r.id for r in ctx.requirements}
    if kind == DeliverableKind.REQUIREMENT_SPEC_XLSX:
        found: set[str] = set()
        for sheet in (SHEET_FUNCTIONAL, SHEET_NON_FUNCTIONAL, SHEET_OTHER):
            data = sheets.get(sheet)
            if not data:
                continue
            idx = data[0].index("요구사항 ID") if "요구사항 ID" in data[0] else None
            if idx is None:
                continue
            found |= {str(r[idx]) for r in data[1:] if r[idx]}
        if found != req_ids:
            outcome.add(
                ValidationLevel.ERROR,
                "xlsx.ids",
                f"{label}: 요구사항 ID 집합이 공통 데이터와 다름 (누락 {sorted(req_ids - found)[:5]}, 초과 {sorted(found - req_ids)[:5]})",
                path.name,
            )
    if kind == DeliverableKind.WBS_GANTT_XLSX and SHEET_WBS in sheets:
        data = sheets[SHEET_WBS]
        h = data[0]
        ci = {name: h.index(name) for name in h if name}
        codes = {str(r[ci["WBS ID"]]) for r in data[1:]}
        for r in data[1:]:
            parent = r[ci["상위 WBS ID"]]
            if parent and str(parent) not in codes:
                outcome.add(
                    ValidationLevel.ERROR,
                    "xlsx.wbs_parent",
                    f"{label}: {r[ci['WBS ID']]}의 상위 {parent} 없음",
                    SHEET_WBS,
                )
            s, e = r[ci["계획 시작일"]], r[ci["계획 종료일"]]
            if (
                s is not None
                and e is not None
                and hasattr(s, "toordinal")
                and hasattr(e, "toordinal")
                and e < s
            ):
                outcome.add(
                    ValidationLevel.ERROR,
                    "xlsx.dates",
                    f"{label}: {r[ci['WBS ID']]} 종료일이 시작일보다 빠름",
                    SHEET_WBS,
                )
            if not r[ci["산정 근거"]]:
                outcome.add(
                    ValidationLevel.ERROR,
                    "xlsx.basis",
                    f"{label}: {r[ci['WBS ID']]} 산정 근거 없음",
                    SHEET_WBS,
                )
            for rid in str(r[ci["요구사항 ID"]] or "").split(","):
                rid = rid.strip()
                if rid and rid not in req_ids:
                    outcome.add(
                        ValidationLevel.WARNING,
                        "xlsx.req_ref",
                        f"{label}: {r[ci['WBS ID']]}가 없는 요구사항 {rid} 참조",
                        SHEET_WBS,
                    )
    if kind == DeliverableKind.TRACE_MATRIX_XLSX and SHEET_MAIN in sheets:
        data = sheets[SHEET_MAIN]
        idx = data[0].index("요구사항 ID")
        found = {str(r[idx]) for r in data[1:]}
        if found != req_ids:
            outcome.add(
                ValidationLevel.ERROR,
                "xlsx.ids",
                f"{label}: 추적표 요구사항 집합 불일치",
                path.name,
            )
        res_idx = [data[0].index(c) for c in ("단위테스트 결과", "통합테스트 결과")]
        for r in data[1:]:
            for i in res_idx:
                if r[i] == "성공":
                    outcome.add(
                        ValidationLevel.ERROR,
                        "xlsx.test_result",
                        f"{label}: 수행하지 않은 테스트가 '성공'으로 표시됨 ({r[idx]})",
                        SHEET_MAIN,
                    )
    return sheets


# --- 일관성·추적성 -----------------------------------------------------------------------------


def _find_total_md(sheets: dict[str, list[list[Any]]], sheet: str, col: str) -> Decimal | None:
    data = sheets.get(sheet)
    if not data or col not in data[0]:
        return None
    i = data[0].index(col)
    name_i = 0
    for r in data[1:]:
        if str(r[name_i]) == "합계" and r[i] is not None:
            return Decimal(str(r[i]))
    return None


def check_consistency(
    ctx: GenerationContext,
    excel: dict[DeliverableKind, dict[str, list[list[Any]]]],
    ppt_texts: dict[str, list[str]],
    outcome: QaOutcome,
) -> None:
    wbs_total = ctx.total_effort_md
    staffing_total = _find_total_md(
        excel.get(DeliverableKind.STAFFING_PLAN_XLSX, {}), SHEET_SUMMARY, "공수(M/D)"
    )
    if staffing_total is not None and abs(staffing_total - wbs_total) > Decimal("0.05"):
        outcome.add(
            ValidationLevel.ERROR,
            "consistency.effort",
            f"WBS 총 공수 {wbs_total} M/D ≠ 인력 계획 합계 {staffing_total} M/D",
        )
    alloc_total = sum((a.effort_md for a in ctx.allocations), Decimal("0"))
    if ctx.allocations and abs(alloc_total - wbs_total) > Decimal("0.05"):
        outcome.add(
            ValidationLevel.ERROR,
            "consistency.allocation",
            f"배정 공수 합계 {alloc_total} ≠ WBS 총 공수 {wbs_total}",
        )
    figs = key_figures(ctx)
    outcome.figures["expected"] = figs
    joined = {label: "\n".join(t) for label, t in ppt_texts.items()}
    for label, text in joined.items():
        for key in ("period", "total_mm", "headcount"):
            value = figs[key]
            if value and value not in text and not (key == "period" and "기간 미정" in value):
                outcome.add(
                    ValidationLevel.ERROR,
                    "consistency.ppt_figures",
                    f"{label} PPT에 {key} 값 '{value}'이 없음 — Excel·다른 PPT와 불일치 가능",
                )
    techs = {c.technology for c in ctx.components if c.technology and "미정" not in c.technology}
    for label, text in joined.items():
        for tech in techs:
            if tech not in text:
                outcome.add(
                    ValidationLevel.WARNING,
                    "consistency.technology",
                    f"{label} PPT에 구성 요소 기술 '{tech}' 언급 없음",
                )
    sub = joined.get("submission", "")
    missing = [r.id for r in ctx.mandatory_requirements() if not re.search(re.escape(r.id), sub)]
    if missing:
        outcome.add(
            ValidationLevel.ERROR,
            "traceability.requirements",
            f"제출용 PPT 요구사항 대응에 필수 요구사항 누락: {', '.join(missing[:10])}",
        )
    pres = joined.get("presentation", "")
    for marker in ("검토 필요", "미정"):
        n = pres.count(marker)
        if n > 5:
            outcome.add(
                ValidationLevel.WARNING,
                "quality.placeholders",
                f"발표용 PPT에 '{marker}' 표기 {n}회 — 확정 후 재생성 권장",
            )


def render_report(
    project_name: str,
    outcome: QaOutcome,
    deliverables: list[tuple[str, str]],
    review_items: list[str],
) -> str:
    lines = [
        f"# QA 보고서 — {project_name}",
        "",
        f"- 오류: {len(outcome.errors)}건 / 경고: {len(outcome.warnings)}건",
        "",
    ]
    lines.append("## 파일 열기 검사")
    for label, ok in outcome.opened.items():
        lines.append(f"- {label}: {'열림' if ok else '실패'}")
    lines += ["", "## 산출물"]
    for kind, path in deliverables:
        lines.append(f"- {kind}: `{path}`")
    lines += ["", "## 검사 결과"]
    for f in outcome.findings:
        lines.append(
            f"- [{f.level.value}] {f.rule}: {f.message}"
            + (f" ({f.location})" if f.location else "")
        )
    if not outcome.findings:
        lines.append("- 지적 사항 없음")
    lines += ["", "## 사람 검토 필요"]
    lines += [f"- {r}" for r in review_items] or ["- 없음"]
    if outcome.figures.get("expected"):
        lines += ["", "## 핵심 수치 (공통 데이터)"]
        lines += [f"- {k}: {v}" for k, v in outcome.figures["expected"].items()]
    return "\n".join(lines) + "\n"


# --- 제안서 템플릿 준수 (Phase 12) ----------------------------------------------------------------


def check_template_compliance(
    path: Path, profile: Any, toc_tree: list[dict[str, Any]] | None, label: str
) -> tuple[list[str], list[str]]:
    """생성 PPT가 템플릿 프로파일과 맞는지 (실패 목록, 검토 목록).

    - 용지 크기 (±0.2 cm)                                → 실패
    - 본문 슬라이드에 템플릿 머리글 위치(±0.6 cm)의 텍스트 상자 2단이 있는가 → 실패 (빠진 슬라이드 수)
    - 장 표지의 로마숫자 순서가 RFP 지정 목차와 같은가     → 실패
    - 글꼴이 템플릿 글꼴(제목·본문)인 run 비율 80% 미만      → 검토
    """
    from pptx import Presentation
    from pptx.util import Emu

    fails: list[str] = []
    reviews: list[str] = []
    prs = Presentation(str(path))
    w, h = Emu(int(prs.slide_width or 0)).cm, Emu(int(prs.slide_height or 0)).cm
    if abs(w - profile.slide_width_cm) > 0.2 or abs(h - profile.slide_height_cm) > 0.2:
        fails.append(
            f"{label}: 용지 {w:.1f}×{h:.1f} cm ≠ 템플릿 {profile.slide_width_cm}×{profile.slide_height_cm} cm"
        )
    headers = list(profile.headers)
    numerals: list[str] = []
    missing_header: list[int] = []
    font_total = font_ok = 0
    fonts = {profile.fonts.get("title"), profile.fonts.get("body")} - {None}
    roman = set("ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ")
    for idx, slide in enumerate(prs.slides, start=1):
        texts = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        is_chapter = slide.slide_layout.name == profile.chapter_layout or any(
            (r.font.size.pt if r.font.size else 0) >= 24
            for s in texts
            for p in s.text_frame.paragraphs
            for r in p.runs
        )
        for s in texts:
            t = s.text_frame.text.strip()
            if is_chapter and t and all(ch in roman for ch in t):
                numerals.append(t)
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    name = r.font.name or p.font.name  # run에 없으면 문단 기본 글꼴
                    if r.text.strip() and name:
                        font_total += 1
                        if name in fonts:
                            font_ok += 1
        if is_chapter or idx <= 2 or not headers:
            continue
        found = 0
        for hb in headers:
            if any(
                abs(Emu(int(s.top or 0)).cm - hb.top) <= 0.6
                and abs(Emu(int(s.left or 0)).cm - hb.left) <= 1.0
                for s in texts
            ):
                found += 1
        if found < len(headers):
            missing_header.append(idx)
    if missing_header:
        fails.append(
            f"{label}: 머리글 {len(headers)}단이 없는 본문 슬라이드 {len(missing_header)}장 "
            f"(예: {', '.join(str(i) for i in missing_header[:8])})"
        )
    if toc_tree:
        want = [str(c.get("numeral", "")) for c in toc_tree if c.get("numeral")]
        if want and numerals != want:
            fails.append(f"{label}: 장 순서 {numerals} ≠ RFP 지정 목차 {want}")
    if font_total and fonts and font_ok / font_total < 0.8:
        reviews.append(
            f"{label}: 템플릿 글꼴({', '.join(sorted(fonts))}) 사용 비율 {font_ok / font_total:.0%} — 육안 확인"
        )
    return fails, reviews
