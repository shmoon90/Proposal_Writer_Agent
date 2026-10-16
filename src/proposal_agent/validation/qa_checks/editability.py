"""QA-EDIT: 편집 가능성 검사 (independent_qa 분할)."""

from __future__ import annotations

from typing import Any

from proposal_agent.models import (
    DeliverableKind,
    ProjectMode,
    QaCategory,
    QaCheck,
)
from proposal_agent.validation.qa_checks.common import (
    FIX_REGENERATE,
    CheckArtifacts,
    QaInputs,
    _cap,
    _check,
    _deck_paths,
    _finish,
    _kind_files,
    _not_run,
)


def check_editability(inp: QaInputs, art: CheckArtifacts) -> list[QaCheck]:
    """QA-EDIT-01~10 편집 가능성 체크리스트 (Phase 21): 생성 PPT의 객체 구조를 사람이 PowerPoint에서 고칠 수 있는가."""
    cat = QaCategory.EDITABILITY
    gate = "gate4"
    specs = [
        (
            "QA-EDIT-01",
            "텍스트 편집",
            "모든 글이 텍스트 상자·표 셀·도형 텍스트에 있고(글자가 든 그림 없음) 빈 텍스트 상자가 과다하지 않은가",
            True,
        ),
        (
            "QA-EDIT-02",
            "표 객체",
            "표는 네이티브 표 객체이며 표를 그림으로 넣은 슬라이드가 없는가 (제출용은 표 1개 이상)",
            True,
        ),
        (
            "QA-EDIT-03",
            "도형·연결선",
            "도식이 자동 도형·연결선·텍스트 상자로 그려져 개별 선택·이동·수정이 가능한가 (자유형 과다 없음)",
            False,
        ),
        (
            "QA-EDIT-04",
            "차트 객체",
            "차트는 네이티브 차트(데이터 시트 편집 가능)이며 차트 이미지가 없는가",
            True,
        ),
        (
            "QA-EDIT-05",
            "그룹 해제 가능",
            "그룹 도형이 자식 2개 이상의 평범한 그룹(중첩 3단 이하)이라 그룹 해제로 개별 편집이 되는가",
            False,
        ),
        (
            "QA-EDIT-06",
            "그림 정책",
            "그림은 허용 태그(photo/illustration/decoration/reference)가 붙은 사진·일러스트뿐인가 (정보성 자료 이미지 없음)",
            True,
        ),
        (
            "QA-EDIT-07",
            "글꼴",
            "본문 글꼴이 템플릿·기본 글꼴 집합(맑은 고딕·템플릿 글꼴) 안에 있어 글꼴 치환 없이 편집되는가",
            False,
        ),
        (
            "QA-EDIT-08",
            "발표자 노트",
            "발표용 본문 슬라이드에 발표자 노트가 절반 이상 있는가",
            False,
        ),
        (
            "QA-EDIT-09",
            "레이아웃·마스터",
            "모든 슬라이드가 레이아웃을 쓰고 슬라이드 밖·음수 좌표 도형이 없는가",
            False,
        ),
        (
            "QA-EDIT-10",
            "재열기 안정성",
            "python-pptx로 열고 저장 왕복 후 다시 열어도 슬라이드 수·도형 수가 같은가",
            True,
        ),
    ]
    checks = [
        _check(
            cid,
            cat,
            target,
            criteria,
            gate=gate,
            mandatory=mand,
            fixable=FIX_REGENERATE if mand else None,
            fix_suggestion="산출물 재생성 (도식·표·차트는 코드가 객체로 그림)",
        )
        for cid, target, criteria, mand in specs
    ]
    if inp.mode != ProjectMode.PROPOSAL:
        for c in checks:
            _not_run(c, "구축 단계는 PPT를 생성하지 않음")
        return checks
    decks = _deck_paths(inp)
    if not decks:
        for c in checks:
            _not_run(c, "생성된 PPT 없음", applicable=True)
        return checks
    from pptx import Presentation

    from proposal_agent.generators.base import PICTURE_TAGS
    from proposal_agent.validation.deliverables import _flatten

    fails: dict[int, list[str]] = {i: [] for i in range(10)}
    revs: dict[int, list[str]] = {i: [] for i in range(10)}
    files = _kind_files(inp, DeliverableKind.SUBMISSION_PPTX, DeliverableKind.PRESENTATION_PPTX)
    known_fonts = {"맑은 고딕", "Malgun Gothic", "Arial", "Calibri"}
    for prof in inp.template_profiles.values():
        for attr in ("title_font", "body_font"):
            v = getattr(prof, attr, None)
            if v:
                known_fonts.add(str(v))
    for label, path in decks:
        try:
            prs = Presentation(str(path))
        except Exception as exc:  # noqa: BLE001
            fails[9].append(f"{label}: 열기 실패 {exc}")
            continue
        n_slides = len(prs.slides)
        empty_text = 0
        tables = 0
        charts = 0
        freeform = 0
        auto = 0
        groups_bad: list[str] = []
        pictures_bad: list[str] = []
        pictures_info: list[str] = []
        fonts: set[str] = set()
        notes_body = 0
        body_slides = 0
        no_layout = 0
        out_of_bounds = 0
        sw, sh_ = int(prs.slide_width or 0), int(prs.slide_height or 0)
        shape_total = 0
        for idx, slide in enumerate(prs.slides, start=1):
            if slide.slide_layout is None:
                no_layout += 1
            flat = _flatten(slide.shapes)
            shape_total += len(flat)
            slide_text = 0
            for sh in flat:
                st = str(sh.shape_type or "")
                if sh.has_text_frame:
                    t = sh.text_frame.text.strip()
                    if t:
                        slide_text += len(t)
                        for para in sh.text_frame.paragraphs:
                            for run in para.runs:
                                if run.font.name:
                                    fonts.add(str(run.font.name))
                    elif "TEXT_BOX" in st:
                        empty_text += 1
                if getattr(sh, "has_table", False) and sh.has_table:
                    tables += 1
                elif getattr(sh, "has_chart", False) and sh.has_chart:
                    charts += 1
                elif "PICTURE" in st:
                    name = str(sh.name or "")
                    tag = name.split(":", 1)[0] if ":" in name else ""
                    if tag in PICTURE_TAGS:
                        pictures_info.append(f"{label} {idx}장 {name}")
                    else:
                        pictures_bad.append(f"{label} {idx}장 '{name or '(이름 없음)'}'")
                elif "FREEFORM" in st:
                    freeform += 1
                elif "AUTO_SHAPE" in st or "LINE" in st or "CONNECTOR" in st:
                    auto += 1
                if sh.left is not None and sh.top is not None:
                    if (
                        int(sh.left) < 0
                        or int(sh.top) < 0
                        or int(sh.left) + int(sh.width or 0) > sw + 1000
                        or int(sh.top) + int(sh.height or 0) > sh_ + 1000
                    ):
                        out_of_bounds += 1
            for sh in slide.shapes:  # 그룹 구조 (최상위만)
                if "GROUP" in str(sh.shape_type):
                    kids = list(sh.shapes)
                    depth = _group_depth(sh)
                    if len(kids) < 2 or depth > 3:
                        groups_bad.append(
                            f"{label} {idx}장 '{sh.name}' 자식 {len(kids)}·깊이 {depth}"
                        )
            if idx > 1 and slide_text >= 60:
                body_slides += 1
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    notes_body += 1
        # 01 텍스트
        if empty_text > max(3, n_slides):
            revs[0].append(f"{label}: 빈 텍스트 상자 {empty_text}개")
        # 02 표
        if label == "제출용" and tables == 0:
            fails[1].append(f"{label}: 네이티브 표가 없음")
        # 03 도형
        if freeform > auto:
            revs[2].append(f"{label}: 자유형 {freeform} > 자동 도형 {auto}")
        # 04 차트
        if any("차트" in x or "chart" in x.lower() for x in pictures_bad):
            fails[3].append(
                f"{label}: 차트 그림 {[x for x in pictures_bad if '차트' in x or 'chart' in x.lower()][:2]}"
            )
        if charts == 0 and label == "제출용":
            revs[3].append(f"{label}: 네이티브 차트 없음 (공수·통계·간트 데이터가 비어 있음)")
        # 05 그룹
        if groups_bad:
            revs[4].extend(groups_bad[:5])
        # 06 그림
        if pictures_bad:
            fails[5].append(
                f"{label}: 태그 없는 그림 {len(pictures_bad)} — " + "; ".join(pictures_bad[:3])
            )
        # 07 글꼴
        unknown = sorted(f for f in fonts if f not in known_fonts)
        if unknown:
            revs[6].append(f"{label}: 기본·템플릿 외 글꼴 {', '.join(unknown[:5])}")
        # 08 노트
        if label == "발표용" and body_slides and notes_body < body_slides * 0.5:
            revs[7].append(f"{label}: 발표자 노트 {notes_body}/{body_slides} 본문 슬라이드")
        # 09 레이아웃
        if no_layout:
            revs[8].append(f"{label}: 레이아웃 없는 슬라이드 {no_layout}")
        if out_of_bounds:
            fails[8].append(f"{label}: 슬라이드 밖 도형 {out_of_bounds}")
        # 10 재열기
        try:
            import io as _io

            buf = _io.BytesIO()
            prs.save(buf)
            again = Presentation(_io.BytesIO(buf.getvalue()))
            again_total = sum(len(_flatten(sl.shapes)) for sl in again.slides)
            if len(again.slides) != n_slides or again_total != shape_total:
                fails[9].append(
                    f"{label}: 저장 왕복 후 슬라이드 {len(again.slides)}/{n_slides}, 도형 {again_total}/{shape_total}"
                )
        except Exception as exc:  # noqa: BLE001
            fails[9].append(f"{label}: 저장 왕복 실패 {type(exc).__name__}: {exc}")
        checks[5].issues.extend(f"예외: {x}" for x in pictures_info[:6])
    for i, c in enumerate(checks):
        extra = list(c.issues)
        _finish(c, fail=_cap(fails[i], 8), review=_cap(revs[i], 8), files=files)
        if i == 5:
            c.issues.extend(extra)
    return checks


def _group_depth(shape: Any, depth: int = 1) -> int:
    best = depth
    for kid in shape.shapes:
        if "GROUP" in str(kid.shape_type):
            best = max(best, _group_depth(kid, depth + 1))
    return best
