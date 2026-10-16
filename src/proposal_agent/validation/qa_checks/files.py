"""QA-FILE: 파일 품질·페이지 제한·렌더링 검사 (independent_qa 분할)."""

from __future__ import annotations

import re
from pathlib import Path

from proposal_agent.config.errors import PathSecurityError
from proposal_agent.config.paths import ensure_within
from proposal_agent.generators.excel.style import HEADER_ROW
from proposal_agent.generators.excel.wbs_gantt import (
    SHEET_REPORT,
    SHEET_WBS,
    gantt_first_col,
    read_week_header,
    span_weeks,
)
from proposal_agent.models import (
    DeliverableKind,
    ProjectMode,
    QaCategory,
    QaCheck,
    ValidationLevel,
)
from proposal_agent.requirements.id_policy import normalize_text
from proposal_agent.validation.deliverables import (
    check_consistency,
    check_pptx,
    check_template_compliance,
    check_xlsx,
)
from proposal_agent.validation.qa_checks.common import (
    FIX_REGENERATE,
    CheckArtifacts,
    QaInputs,
    _as_date,
    _cap,
    _check,
    _deck_paths,
    _finish,
    _is_filled,
    _kind_files,
    _not_run,
)


# --- 파일 품질 ---------------------------------------------------------------------------------
def _template_size(inp: QaInputs, purpose: str) -> tuple[float, float] | None:
    prof = inp.template_profiles.get(purpose)
    if prof is None:
        return None
    return (float(prof.slide_width_cm), float(prof.slide_height_cm))


def _template_check(
    inp: QaInputs, art: CheckArtifacts, path: Path, purpose: str, label: str
) -> None:
    prof = inp.template_profiles.get(purpose)
    if prof is None:
        return
    try:
        fails, reviews = check_template_compliance(
            path, prof, inp.rfp_toc_tree if purpose == "submission" else None, label
        )
    except Exception as exc:  # noqa: BLE001
        fails, reviews = [f"{label}: 템플릿 준수 검사 실패 ({type(exc).__name__}: {exc})"], []
    art.template_fails.extend(fails)
    art.template_reviews.extend(reviews)


def run_file_checks(inp: QaInputs) -> tuple[list[QaCheck], CheckArtifacts]:
    cat = QaCategory.FILE_QUALITY
    gate = "gate4"
    art = CheckArtifacts()
    checks = [
        _check(
            "QA-FILE-01",
            cat,
            "산출물 파일 열기",
            "PPT·Excel·PNG 파일이 정상적으로 열리는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="산출물 재생성",
        ),
        _check(
            "QA-FILE-02",
            cat,
            "슬라이드 방향·규격",
            "가로 16:9, 슬라이드 3장 이상",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="PPT 재생성",
        ),
        _check(
            "QA-FILE-03",
            cat,
            "텍스트·도형 잘림",
            "도형이 슬라이드 안에 있고 텍스트가 넘치지 않는가",
            gate=gate,
            fix_suggestion="해당 슬라이드 문안을 줄이거나 PowerPoint에서 육안 확인",
        ),
        _check(
            "QA-FILE-04",
            cat,
            "표·차트 겹침",
            "텍스트 도형끼리 크게 겹치지 않는가",
            gate=gate,
            mandatory=False,
            fix_suggestion="레이아웃 조정 후 재생성, PowerPoint 육안 확인",
        ),
        _check(
            "QA-FILE-05",
            cat,
            "Excel 시트·컬럼",
            "필수 시트와 컬럼이 모두 있는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="Excel 재생성",
        ),
        _check(
            "QA-FILE-06",
            cat,
            "날짜·수식",
            "종료일 ≥ 시작일, 수식 형식 오류 없음, 산정 근거 있음",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="계획 데이터 점검 후 Excel 재생성",
        ),
        _check(
            "QA-FILE-07",
            cat,
            "요구사항 ID",
            "Excel의 요구사항 ID가 기준 데이터와 같고 형식이 올바른가, 미실행 테스트가 '성공'으로 표시되지 않았는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="Excel 재생성",
        ),
        _check(
            "QA-FILE-08",
            cat,
            "간트 차트 ↔ 일정 데이터",
            "간트 시트의 시작·종료와 채움 구간이 WBS 시트 날짜와 같은가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="WBS·간트 재생성",
        ),
        _check(
            "QA-FILE-09",
            cat,
            "파일 경로",
            "산출물 경로가 프로젝트 Output 폴더 안에 있고 실제로 존재하는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="누락 파일을 재생성 (경로 조작은 거부)",
        ),
        _check(
            "QA-FILE-10",
            cat,
            "산출물 누락",
            "모드별 필수 산출물이 모두 생성되었는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="산출물 재생성",
        ),
        _check(
            "QA-FILE-12",
            cat,
            "편집 가능성",
            "생성 PPT에 그림(비트맵) 도형이 없고 모든 요소가 도형·표·텍스트로 편집 가능한가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="산출물 재생성 (도식은 코드가 도형으로 그림)",
        ),
        _check(
            "QA-FILE-13",
            cat,
            "페이지 제한 (RFP)",
            "RFP가 지정한 제안서 페이지 제한(쪽 수) 안에 제출용 PPT 슬라이드 수(표지·목차 제외)가 드는가. 90% 이상이면 검토",
            gate=gate,
            fix_suggestion="소절 글머리·표 행을 줄여 재생성하거나 부록으로 이동",
        ),
        _check(
            "QA-FILE-14",
            cat,
            "실제 렌더링·재열기",
            "PowerPoint(COM)로 제출용·발표용 PPT를 실제로 열어 슬라이드마다 PNG로 내보내고, 슬라이드 수가 python-pptx와 같으며 저장 왕복 후 재열기가 되는가. PowerPoint가 없거나 설정이 꺼져 있으면 미실행",
            gate=gate,
            mandatory=False,
            fix_suggestion="PNG(11_QA_Validation_Reports/render)를 보고 깨진 슬라이드를 확인한 뒤 재생성",
        ),
        _check(
            "QA-FILE-11",
            cat,
            "제안서 템플릿 준수",
            "등록된 제안서 템플릿과 용지 크기·머리글 2단 위치·장 순서(RFP 지정 목차)·글꼴이 맞는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="템플릿을 다시 확인하고 산출물 재생성",
        ),
    ]
    out = art.outcome
    project_dir = inp.paths.project_output_dir(inp.project.output_dir)
    path_fail: list[str] = []
    for kind, d in inp.deliverables.items():
        try:
            path = inp.paths.resolve_output(d.file_path)
            ensure_within(project_dir, path)
        except PathSecurityError as exc:
            path_fail.append(f"{kind.value}: 경로 거부 {d.file_path} ({exc})")
            continue
        if not path.exists():
            path_fail.append(f"{kind.value}: 파일 없음 {d.file_path}")
            out.opened[kind.value] = False
            continue
        if kind.value.endswith("xlsx"):
            art.excel[kind] = check_xlsx(path, kind, out, inp.ctx)
        elif kind == DeliverableKind.SUBMISSION_PPTX:
            art.ppt_texts["submission"] = check_pptx(
                path, out, "제출용", expected_size_cm=_template_size(inp, "submission")
            )
            _template_check(inp, art, path, "submission", "제출용")
        elif kind == DeliverableKind.PRESENTATION_PPTX:
            art.ppt_texts["presentation"] = check_pptx(
                path, out, "발표용", expected_size_cm=_template_size(inp, "presentation")
            )
            _template_check(inp, art, path, "presentation", "발표용")
        elif kind.value.endswith("png"):
            from proposal_agent.agents.tools import verify_file

            info = verify_file(path)
            out.opened[kind.value] = bool(info.get("ok"))
            if not info.get("ok"):
                out.add(
                    ValidationLevel.ERROR,
                    "png.open",
                    f"{kind.value}: {info.get('error')}",
                    d.file_path,
                )
    missing = [k.value for k in inp.required_kinds if k not in inp.deliverables]
    for k in missing:
        out.add(ValidationLevel.ERROR, "deliverable.missing", f"산출물 없음: {k}")
    if art.excel or art.ppt_texts:
        check_consistency(inp.ctx, art.excel, art.ppt_texts, out)
    gantt_fail = _check_gantt(inp, art)

    by_rule: dict[str, list[tuple[ValidationLevel, str]]] = {}
    for x in out.findings:
        by_rule.setdefault(x.rule, []).append((x.level, x.message))

    def split(rules: tuple[str, ...]) -> tuple[list[str], list[str]]:
        fails: list[str] = []
        revs: list[str] = []
        for r in rules:
            for lvl, msg in by_rule.get(r, []):
                (fails if lvl == ValidationLevel.ERROR else revs).append(msg)
        return fails, revs

    all_files = [d.file_path for d in inp.deliverables.values()]
    f, r = split(("pptx.picture",))
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[10], "구축 단계는 PPT를 생성하지 않음")
    else:
        _finish(
            checks[10],
            fail=f,
            review=r,
            files=_kind_files(
                inp, DeliverableKind.SUBMISSION_PPTX, DeliverableKind.PRESENTATION_PPTX
            ),
        )
        # 허용 그림(사진·일러스트·장식·참고) 예외는 판정을 바꾸지 않고 보고서·manifest에 기록한다
        checks[10].issues.extend(x.message for x in out.findings if x.rule == "pptx.photo")
    # 13: 페이지 제한 (Phase 21)
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[11], "구축 단계는 PPT를 생성하지 않음")
    elif not inp.page_limit:
        _not_run(checks[11], "RFP에 페이지 제한 문구 없음")
    else:
        f13, r13 = _page_limit_issues(inp, art)
        _finish(
            checks[11],
            fail=f13,
            review=r13,
            files=_kind_files(inp, DeliverableKind.SUBMISSION_PPTX),
        )
    # 14: 실제 렌더링·재열기 (Phase 21)
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[12], "구축 단계는 PPT를 생성하지 않음")
    elif not inp.render_check:
        _not_run(checks[12], "렌더링 검증이 꺼져 있음 (PWA_QA_RENDER_CHECK=1 또는 화면에서 켜기)")
    else:
        f14, r14, files14 = _render_issues(inp, art)
        if f14 is None:
            _not_run(checks[12], r14[0] if r14 else "렌더링 엔진 없음", applicable=True)
        else:
            _finish(checks[12], fail=f14, review=r14, files=files14)
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[13], "구축 단계는 PPT를 생성하지 않음")
    elif not inp.template_profiles:
        _not_run(checks[13], "제안서 템플릿(PPT) 미등록 — 코드 기본 양식으로 생성됨")
    else:
        _finish(
            checks[13],
            fail=list(art.template_fails),
            review=list(art.template_reviews),
            files=_kind_files(
                inp, DeliverableKind.SUBMISSION_PPTX, DeliverableKind.PRESENTATION_PPTX
            ),
        )
    f, r = split(("pptx.open", "xlsx.open", "png.open"))
    f += [m for m in path_fail if "파일 없음" in m]
    _finish(checks[0], fail=f, review=r, files=all_files)
    f, r = split(("pptx.orientation", "pptx.size", "pptx.slides"))
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[1], "구축 단계는 PPT를 생성하지 않음")
    else:
        _finish(
            checks[1],
            fail=f,
            review=r,
            files=_kind_files(
                inp, DeliverableKind.SUBMISSION_PPTX, DeliverableKind.PRESENTATION_PPTX
            ),
        )
    f, r = split(("pptx.bounds", "pptx.overflow"))
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[2], "구축 단계는 PPT를 생성하지 않음")
    else:
        _finish(
            checks[2],
            fail=f,
            review=r,
            files=_kind_files(
                inp, DeliverableKind.SUBMISSION_PPTX, DeliverableKind.PRESENTATION_PPTX
            ),
        )
    f, r = split(("pptx.overlap",))
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[3], "구축 단계는 PPT를 생성하지 않음")
    else:
        _finish(checks[3], fail=f, review=r)
    f, r = split(("xlsx.sheet", "xlsx.columns"))
    _finish(
        checks[4],
        fail=f,
        review=r,
        files=[d.file_path for k, d in inp.deliverables.items() if k.value.endswith("xlsx")],
    )
    f, r = split(("xlsx.dates", "xlsx.formula", "xlsx.basis"))
    _finish(
        checks[5],
        fail=f,
        review=r,
        files=_kind_files(inp, DeliverableKind.WBS_GANTT_XLSX, DeliverableKind.STAFFING_PLAN_XLSX),
    )
    f, r = split(("xlsx.ids", "xlsx.req_ref", "xlsx.wbs_parent", "xlsx.test_result"))
    bad_ids = _invalid_ids(inp)
    if bad_ids:
        f.append("형식이 맞지 않는 요구사항 ID: " + ", ".join(bad_ids[:8]))
    _finish(checks[6], fail=f, review=r, req_ids=bad_ids)
    if DeliverableKind.WBS_GANTT_XLSX not in inp.deliverables:
        _not_run(checks[7], "WBS·간트 파일 없음", applicable=True)
    else:
        _finish(checks[7], fail=gantt_fail, files=_kind_files(inp, DeliverableKind.WBS_GANTT_XLSX))
    _finish(checks[8], fail=path_fail, files=all_files)
    _finish(checks[9], fail=[f"산출물 없음: {k}" for k in missing])
    return checks, art


def _content_slide_count(path: Path) -> tuple[int, int]:
    """(전체 슬라이드, 본문 슬라이드). 표지·목차·장 구분(제목이 목차/CONTENTS/장 표지이거나 본문 텍스트가 거의 없는) 제외."""
    from pptx import Presentation

    prs = Presentation(str(path))
    total = len(prs.slides)
    body = 0
    for idx, slide in enumerate(prs.slides, start=1):
        texts = [sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame]
        joined = " ".join(t for t in texts if t)
        has_obj = any(
            (getattr(sh, "has_table", False) and sh.has_table)
            or (getattr(sh, "has_chart", False) and sh.has_chart)
            or "GROUP" in str(sh.shape_type)
            or "PICTURE" in str(sh.shape_type)
            for sh in slide.shapes
        )
        first = (texts[0] if texts else "")[:20]
        if idx == 1 or first in ("목차", "CONTENTS") or (len(joined) < 60 and not has_obj):
            continue
        body += 1
    return total, body


def _page_limit_issues(inp: QaInputs, art: CheckArtifacts) -> tuple[list[str], list[str]]:
    limit = int(inp.page_limit or 0)
    fails: list[str] = []
    revs: list[str] = []
    for label, path in _deck_paths(inp):
        if label != "제출용":
            continue
        total, body = _content_slide_count(path)
        msg = f"제출용 본문 슬라이드 {body}장 (전체 {total}장) / 제한 {limit}쪽"
        if body > limit:
            fails.append(msg + " — 초과")
        elif body >= limit * 0.9:
            revs.append(msg + " — 여유 10% 미만")
    return fails, revs


def _render_issues(
    inp: QaInputs, art: CheckArtifacts
) -> tuple[list[str] | None, list[str], list[str]]:
    """QA-FILE-14. (None, [사유], []) 이면 엔진 없음 → 미실행."""
    from proposal_agent.validation.render import com_available, render_pptx

    ok, why = com_available()
    if not ok:
        return None, [why], []
    fails: list[str] = []
    revs: list[str] = []
    files: list[str] = []
    base = inp.render_dir or (
        inp.paths.project_output_dir(inp.project.output_dir) / "11_QA_Validation_Reports" / "render"
    )
    for label, path in _deck_paths(inp):
        res = render_pptx(path, base / path.stem)
        art.render_results[label] = res
        if not res.available:
            return None, res.errors or ["렌더링 엔진 없음"], []
        if res.ok:
            try:
                files.append(inp.paths.relative_output(base / path.stem))
            except PathSecurityError:
                files.append(str(base / path.stem))
        else:
            fails.append(f"{label}: {res.summary} — " + "; ".join(res.errors[:3]))
    return fails, revs, files


def _invalid_ids(inp: QaInputs) -> list[str]:
    """형식(접두사-번호)이 아니거나, 대소문자·공백만 다른 중복 ID."""
    ids = [r.id for r in inp.all_requirements]
    bad = [i for i in ids if not re.fullmatch(r"[A-Za-z]{1,6}-?\d{2,6}", i)]
    seen: dict[str, str] = {}
    for i in ids:
        key = normalize_text(i)
        if key in seen and seen[key] != i:
            bad.append(f"{i} (≈ {seen[key]})")
        seen.setdefault(key, i)
    return bad


def _check_gantt(inp: QaInputs, art: CheckArtifacts) -> list[str]:
    d = inp.deliverables.get(DeliverableKind.WBS_GANTT_XLSX)
    if d is None:
        return []
    from openpyxl import load_workbook

    try:
        path = inp.paths.resolve_output(d.file_path)
        wb = load_workbook(path)
    except Exception as exc:  # noqa: BLE001
        return [f"간트 파일 열기 실패: {exc}"]
    if SHEET_WBS not in wb.sheetnames or SHEET_REPORT not in wb.sheetnames:
        return ["WBS 또는 약식WBS 시트 없음"]
    ws_w, ws_r = wb[SHEET_WBS], wb[SHEET_REPORT]
    wh = [c.value for c in ws_w[HEADER_ROW]]  # Phase 15: 1행 묶음 머리글·월, 2행 컬럼·주차
    if "WBS ID" not in wh or "계획 시작일" not in wh:
        return ["WBS 시트 헤더 없음"]
    ci, si, ei = wh.index("WBS ID"), wh.index("계획 시작일"), wh.index("계획 종료일")
    first_col = gantt_first_col(ws_w)
    if first_col is None:
        return ["WBS 시트에 간트 열 없음"]
    weeks = read_week_header(ws_w, first_col)
    # 기준 데이터(WBS 항목)의 날짜와 시트 날짜·채움 구간을 대조한다
    items = {w.code: w for w in inp.ctx.wbs}
    problems: list[str] = []
    seen: set[str] = set()
    for row in ws_w.iter_rows(min_row=HEADER_ROW + 1):
        code = row[ci].value
        if code is None:
            break
        code = str(code)
        seen.add(code)
        w = items.get(code)
        if w is None:
            problems.append(f"WBS 시트 행 {code}가 기준 데이터에 없음")
            continue
        ws_s, ws_e = _as_date(row[si].value), _as_date(row[ei].value)
        if (ws_s, ws_e) != (w.planned_start, w.planned_end):
            problems.append(f"{code}: 시트 {ws_s}~{ws_e} ≠ 기준 {w.planned_start}~{w.planned_end}")
            continue
        if ws_s and ws_e and weeks:
            cells = row[first_col - 1 : first_col - 1 + len(weeks)]
            filled = [
                i for i, cell in enumerate(cells) if _is_filled(cell) or cell.value is not None
            ]
            expected = span_weeks(weeks, ws_s, ws_e)
            if filled != expected:
                problems.append(
                    f"{code}: 간트 채움 주 {len(filled)}개 ≠ 일정 기준 {len(expected)}개"
                )
    missing = sorted(set(items) - seen)
    if missing:
        problems.append(f"WBS 시트에 없는 항목: {', '.join(missing[:6])}")
    # 약식WBS: 시작·종료일이 기준 데이터의 어떤 항목과 일치해야 한다
    rh = [c.value for c in ws_r[HEADER_ROW]]
    if "항목" in rh and "시작일" in rh:
        ni, rs, re_ = rh.index("항목"), rh.index("시작일"), rh.index("종료일")
        by_name = {w.name: w for w in inp.ctx.wbs}
        for row in ws_r.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
            name = row[ni]
            if name is None:
                break
            w = by_name.get(str(name))
            if w and (_as_date(row[rs]), _as_date(row[re_])) != (w.planned_start, w.planned_end):
                problems.append(f"약식WBS '{str(name)[:20]}' 날짜가 기준 데이터와 다름")
    else:
        problems.append("약식WBS 시트 헤더 없음")
    return _cap(problems)
