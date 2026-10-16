"""실제 렌더링·재열기 검증 (Phase 21).

python-pptx의 객체 구조 검사만으로는 "PowerPoint에서 실제로 열리고 그려지는가"를 알 수 없다. 이 모듈은 이 PC에
설치된 PowerPoint를 COM(pywin32)으로 띄워 덱을 **읽기 전용·창 없이** 열고 슬라이드마다 PNG로 내보낸 뒤, 파일을
python-pptx로 다시 열어 슬라이드 수가 같은지 확인한다. PowerPoint·pywin32가 없으면 ``available=False``로 돌려주고
QA는 그 항목을 **미실행**으로 기록한다(통과로 세지 않는다). 원본 파일은 수정하지 않는다.
"""

from __future__ import annotations

import io
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

PNG_WIDTH = 1280


@dataclass
class RenderResult:
    path: str
    available: bool = False
    ok: bool = False
    slides_pptx: int = 0
    slides_rendered: int = 0
    png_paths: list[str] = field(default_factory=list)
    reopened: bool = False
    errors: list[str] = field(default_factory=list)
    duration_s: float = 0.0
    engine: str = ""

    @property
    def summary(self) -> str:
        if not self.available:
            return "렌더링 엔진 없음 (PowerPoint COM 미사용)"
        state = "성공" if self.ok else "실패"
        return (
            f"{state}: 슬라이드 {self.slides_pptx}장 중 {self.slides_rendered}장 PNG, "
            f"재열기 {'성공' if self.reopened else '실패'}, {self.duration_s:.1f}s"
            + (f", 오류 {len(self.errors)}" if self.errors else "")
        )


def com_available() -> tuple[bool, str]:
    """PowerPoint COM 서버와 pywin32가 있는지 (PowerPoint를 띄우지 않고 레지스트리만 본다)."""
    try:
        import win32com.client  # type: ignore[import-untyped]  # noqa: F401
    except ImportError:
        return False, "pywin32 미설치"
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"PowerPoint.Application\CLSID"):
            return True, "PowerPoint COM 등록됨"
    except OSError:
        return False, "PowerPoint.Application COM 미등록"


def render_pptx(path: Path, out_dir: Path, *, max_slides: int | None = None) -> RenderResult:
    """덱을 PowerPoint로 열어 슬라이드별 PNG(out_dir/slide_NN.png)로 내보내고 재열기를 확인한다."""
    started = time.time()
    res = RenderResult(path=str(path))
    ok, why = com_available()
    if not ok:
        res.errors.append(why)
        return res
    res.available = True
    res.engine = "PowerPoint COM"
    try:
        from pptx import Presentation

        res.slides_pptx = len(Presentation(str(path)).slides)
    except Exception as exc:  # noqa: BLE001
        res.errors.append(f"python-pptx 열기 실패: {type(exc).__name__}: {exc}")
        res.duration_s = round(time.time() - started, 1)
        return res
    out_dir.mkdir(parents=True, exist_ok=True)
    app: Any = None
    pres: Any = None
    try:
        import pythoncom  # type: ignore[import-untyped]
        import win32com.client

        pythoncom.CoInitialize()
        app = win32com.client.DispatchEx("PowerPoint.Application")
        pres = app.Presentations.Open(
            str(path.resolve()), True, False, False
        )  # ReadOnly, Untitled, WithWindow
        count = int(pres.Slides.Count)
        if count != res.slides_pptx:
            res.errors.append(
                f"슬라이드 수 불일치: PowerPoint {count} ≠ python-pptx {res.slides_pptx}"
            )
        limit = count if max_slides is None else min(count, max_slides)
        width = PNG_WIDTH
        height = int(
            round(width * float(pres.PageSetup.SlideHeight) / float(pres.PageSetup.SlideWidth))
        )
        for i in range(1, limit + 1):
            png = out_dir / f"slide_{i:02d}.png"
            try:
                pres.Slides(i).Export(str(png.resolve()), "PNG", width, height)
            except Exception as exc:  # noqa: BLE001
                res.errors.append(f"슬라이드 {i} PNG 내보내기 실패: {exc}")
                continue
            if png.exists() and png.stat().st_size > 0:
                res.png_paths.append(str(png))
            else:
                res.errors.append(f"슬라이드 {i} PNG 파일이 비어 있음")
        res.slides_rendered = len(res.png_paths)
    except Exception as exc:  # noqa: BLE001
        res.errors.append(f"PowerPoint COM 오류: {type(exc).__name__}: {exc}")
    finally:
        try:
            if pres is not None:
                pres.Close()
        except Exception:  # noqa: BLE001
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:  # noqa: BLE001
            pass
        try:
            import pythoncom

            pythoncom.CoUninitialize()
        except Exception:  # noqa: BLE001
            pass
    # 재열기: 원본이 그대로 열리고 저장 왕복(메모리) 후에도 슬라이드 수가 같은가
    try:
        from pptx import Presentation

        prs = Presentation(str(path))
        buf = io.BytesIO()
        prs.save(buf)
        again = Presentation(io.BytesIO(buf.getvalue()))
        res.reopened = len(again.slides) == res.slides_pptx
        if not res.reopened:
            res.errors.append("재열기 후 슬라이드 수가 다름")
    except Exception as exc:  # noqa: BLE001
        res.errors.append(f"재열기 실패: {type(exc).__name__}: {exc}")
    for png_file in res.png_paths[:3]:  # PNG 가 실제 이미지인지 (앞 3장)
        try:
            from PIL import Image

            with Image.open(png_file) as im:
                im.verify()
        except Exception as exc:  # noqa: BLE001
            res.errors.append(f"PNG 검증 실패 {Path(png_file).name}: {exc}")
    expected = res.slides_pptx if max_slides is None else min(res.slides_pptx, max_slides)
    res.ok = res.reopened and res.slides_rendered == expected and expected > 0 and not res.errors
    res.duration_s = round(time.time() - started, 1)
    return res


__all__ = ["PNG_WIDTH", "RenderResult", "com_available", "render_pptx"]
