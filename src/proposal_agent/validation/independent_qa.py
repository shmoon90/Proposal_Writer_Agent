"""통합 QA · 독립 검증 엔진 (Phase 7).

산출물을 만든 에이전트의 자기 평가와 무관하게, 코드가 원본 RFP 블록·요구사항 데이터·기술 조사 근거·계획 데이터·
생성 파일을 직접 대조한다. 결과는 검증 항목(``QaCheck``)으로 남기고 게이트 2·3·4·5를 코드로 판정한다.

- 분류: RFP 적합성(QA-RFP-*), 기술 타당성(QA-TECH-*), 문서 간 일관성(QA-CON-*), 파일 품질(QA-FILE-*), 에이전트 검토(QA-AI-*)
- 자동 수정: 기준 데이터를 바꾸지 않는 '산출물 재생성'만 허용한다. 수정 후 관련 항목을 재검증하고, 설정 횟수 안에 해결되지
  않으면 통과시키지 않고 NEEDS_HUMAN_REVIEW로 둔다. 원본 요구사항과 충돌하거나 기술 위험을 높이는 수정은 제안만 한다.
- 실행하지 않은 항목은 ``executed=False``로 두고 통과로 세지 않는다.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from proposal_agent.config.errors import PathSecurityError
from proposal_agent.config.paths import ProjectPaths, ensure_within
from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.excel.style import HEADER_ROW
from proposal_agent.generators.excel.wbs_gantt import (
    SHEET_REPORT,
    SHEET_WBS,
    gantt_first_col,
    read_week_header,
    span_weeks,
)
from proposal_agent.models import (
    AcceptanceStatus,
    ArchitectureComponent,
    ChangeRequest,
    ChangeStatus,
    Deliverable,
    DeliverableKind,
    DependencyType,
    ExtractedBlock,
    FeasibilityVerdict,
    Project,
    ProjectMode,
    QaCategory,
    QaCheck,
    QaResult,
    Requirement,
    RequirementCategory,
    RequirementType,
    Revision,
    ScopeClass,
    SourceDocument,
    TechnologyResearch,
    ValidationLevel,
)
from proposal_agent.models.change import CandidateClass
from proposal_agent.planning.calendar import BusinessCalendar
from proposal_agent.requirements.id_policy import normalize_text
from proposal_agent.validation.deliverables import (
    QaOutcome,
    check_consistency,
    check_pptx,
    check_template_compliance,
    check_xlsx,
)

MANDATORY_PATTERNS = (
    "하여야 한다",
    "해야 한다",
    "하여야 함",
    "해야 함",
    "반드시",
    "필수",
    "하여야 하며",
    "해야 하며",
    "제공하여야",
    "지원하여야",
)
SUBMISSION_WORDS = (
    "제출 서류",
    "제출서류",
    "제출 방법",
    "제출방법",
    "제출 기한",
    "제출기한",
    "제안서 제출",
    "제출 장소",
    "제출처",
)
EVALUATION_WORDS = (
    "평가 기준",
    "평가기준",
    "배점",
    "평가 항목",
    "평가항목",
    "평가 방법",
    "심사 기준",
)
VERSION_RE = re.compile(r"\b(\d+(?:\.\d+){1,2})\b")
# Phase 19: QA-CON-10 추상 표현(근거 없이 쓰이는 과장 표현) · 단위 붙은 수치
ABSTRACT_WORDS = (
    "최적의",
    "최고의",
    "최상의",
    "완벽한",
    "완벽하게",
    "혁신적",
    "획기적",
    "세계적",
    "압도적",
    "탁월한",
    "유일한",
    "최첨단",
    "업계 최고",
    "검증된 역량",
    "풍부한 경험",
    "최고 수준",
)
NUMBER_CLAIM_RE = re.compile(
    r"(?<![A-Za-z0-9.-])(\d[\d,]*(?:\.\d+)?)[ \t]*(%|퍼센트|배|초|ms|밀리초|분|시간|건|명|년|개월|주|일|억|만|천|TPS|GB|TB|MB|회|대|식|자리|페이지|장|종|배 이상)"
)  # 줄바꿈을 건너 'FR-013' + '대표…' 를 '013대'로 읽지 않도록 ID·줄바꿈 제외
# Phase 19: QA-CON-11 약어 — 대문자 2~8자(숫자·'/' 허용), ID(FR-001)·로마 숫자·일반 약어 제외
ACRONYM_RE = re.compile(r"(?<![A-Za-z0-9/-])[A-Z][A-Z0-9]{1,7}(?:/[A-Z]{1,3})?(?![A-Za-z0-9/-])")
ROMAN_RE = re.compile(r"^[IVX]+$")
COMMON_ACRONYMS = frozenset(
    {
        "RFP",
        "PPT",
        "WBS",
        "PM",
        "PL",
        "QA",
        "IT",
        "SI",
        "ID",
        "PC",
        "OS",
        "MM",
        "MD",
        "M/D",
        "M/M",
        "S/W",
        "H/W",
        "SW",
        "HW",
        "TOC",
        "NO",
        "OK",
        "AS",
        "TO",
        "IS",
        "BE",
        "CR",
        "TR",
        "KPI",
        "RACI",
        "URL",
        "WEB",
        "CONTENTS",
    }
)
ID_IN_TEXT_RE = re.compile(r"\b[A-Z]{2,5}-\d{2,5}\b")

PROPOSAL_KINDS: tuple[DeliverableKind, ...] = (
    DeliverableKind.REQUIREMENT_SPEC_XLSX,
    DeliverableKind.TRACE_MATRIX_XLSX,
    DeliverableKind.WBS_GANTT_XLSX,
    DeliverableKind.STAFFING_PLAN_XLSX,
    DeliverableKind.SUBMISSION_PPTX,
    DeliverableKind.PRESENTATION_PPTX,
    DeliverableKind.SYSTEM_DIAGRAM_PNG,
    DeliverableKind.SW_DIAGRAM_PNG,
    DeliverableKind.SERVICE_DIAGRAM_PNG,
)
BUILD_KINDS: tuple[DeliverableKind, ...] = (
    DeliverableKind.REQUIREMENT_SPEC_XLSX,
    DeliverableKind.TRACE_MATRIX_XLSX,
    DeliverableKind.WBS_GANTT_XLSX,
)
FIX_REGENERATE = "regenerate"


@dataclass
class QaInputs:
    """검증에 필요한 모든 입력. 서비스가 모아서 넘기고, 검사 함수는 저장소를 직접 만지지 않는다."""

    project: Project
    ctx: GenerationContext
    paths: ProjectPaths
    calendar: BusinessCalendar
    blocks: list[ExtractedBlock] = field(default_factory=list)
    documents: list[SourceDocument] = field(default_factory=list)
    all_requirements: list[Requirement] = field(default_factory=list)
    research: list[TechnologyResearch] = field(default_factory=list)
    deliverables: dict[DeliverableKind, Deliverable] = field(default_factory=dict)
    change_requests: list[ChangeRequest] = field(default_factory=list)
    revisions: list[Revision] = field(default_factory=list)
    rfp_analysis: dict[str, Any] | None = None
    id_prefixes: dict[str, str] = field(default_factory=dict)
    rfp_profile: Any | None = None  # ingest.rfp_profile.RfpProfile
    template_profiles: dict[str, Any] = field(default_factory=dict)  # purpose → TemplateProfile
    rfp_toc_tree: list[dict[str, Any]] = field(default_factory=list)
    gate_p_status: str = (
        ""  # Gate P(목차·전략 승인) 상태: PASS | NEEDS_HUMAN_REVIEW | FAIL | ''(미평가) (Phase 18)
    )
    strategy_brief: dict[str, Any] | None = None
    # Phase 21
    render_check: bool = False  # PowerPoint COM 실제 렌더링·재열기(QA-FILE-14) 수행 여부
    render_dir: Path | None = None  # PNG 출력 폴더 (11_QA_Validation_Reports/render)
    page_limit: int | None = None  # RFP 페이지 제한(쪽). 없으면 QA-FILE-13 미실행

    @property
    def mode(self) -> ProjectMode:
        return self.project.mode

    @property
    def requirements(self) -> list[Requirement]:
        return self.ctx.requirements

    @property
    def components(self) -> list[ArchitectureComponent]:
        return self.ctx.components

    @property
    def required_kinds(self) -> tuple[DeliverableKind, ...]:
        if self.mode == ProjectMode.BUILD:
            kinds = list(BUILD_KINDS)
            if self.deliverables.get(DeliverableKind.STAFFING_PLAN_XLSX) or self.ctx.allocations:
                kinds.append(DeliverableKind.STAFFING_PLAN_XLSX)
            return tuple(kinds)
        return PROPOSAL_KINDS


@dataclass
class CheckArtifacts:
    template_fails: list[str] = field(default_factory=list)
    template_reviews: list[str] = field(default_factory=list)
    outcome: QaOutcome = field(default_factory=QaOutcome)
    excel: dict[DeliverableKind, dict[str, list[list[Any]]]] = field(default_factory=dict)
    ppt_texts: dict[str, list[str]] = field(default_factory=dict)
    render_results: dict[str, Any] = field(default_factory=dict)  # 라벨 → RenderResult (Phase 21)


def _check(
    check_id: str,
    category: QaCategory,
    target: str,
    criteria: str,
    *,
    gate: str,
    mandatory: bool = True,
    fixable: str | None = None,
    fix_suggestion: str = "",
) -> QaCheck:
    return QaCheck(
        check_id=check_id,
        category=category,
        target=target,
        criteria=criteria,
        gate=gate,
        mandatory=mandatory,
        fixable=fixable,
        fix_suggestion=fix_suggestion,
    )


def _finish(
    check: QaCheck,
    *,
    fail: list[str] | None = None,
    review: list[str] | None = None,
    req_ids: list[str] | None = None,
    files: list[str] | None = None,
) -> QaCheck:
    fail = fail or []
    review = review or []
    check.issues = [*fail, *review]
    check.requirement_ids = sorted(set(req_ids or []))
    check.files = sorted(set(files or []))
    if fail:
        check.result = QaResult.FAIL
    elif review:
        check.result = QaResult.NEEDS_HUMAN_REVIEW
    else:
        check.result = QaResult.PASS
        check.fixable = None
    return check


def _not_run(check: QaCheck, reason: str, *, applicable: bool = False) -> QaCheck:
    """실행하지 않은 항목. ``applicable=False``(모드상 해당 없음)는 게이트를 막지 않고, True면 검토 필요로 남는다."""
    check.executed = False
    check.applicable = applicable
    check.not_run_reason = reason
    check.result = QaResult.NEEDS_HUMAN_REVIEW
    check.fixable = None
    return check


def _cap(items: list[str], n: int = 10) -> list[str]:
    return items[:n] + ([f"… 외 {len(items) - n}건"] if len(items) > n else [])


# --- RFP 적합성 --------------------------------------------------------------------------------


def _block_text_index(blocks: list[ExtractedBlock]) -> dict[str, str]:
    return {b.id: normalize_text(b.text) for b in blocks}


def check_rfp_fit(inp: QaInputs, ppt_sub_text: str) -> list[QaCheck]:
    gate = "gate5"
    cat = QaCategory.RFP_FIT
    checks = [
        _check(
            "QA-RFP-01",
            cat,
            "RFP 원문 ↔ 요구사항 목록",
            "의무 표현('하여야 한다' 등)이 있는 원문 블록이 모두 요구사항 출처로 연결되어 있는가",
            gate=gate,
            fix_suggestion="누락 블록을 검토해 요구사항을 추가(요구사항 화면 수동 추가)하거나 '해당 없음' 사유를 기록",
        ),
        _check(
            "QA-RFP-02",
            cat,
            "요구사항 범위 분류",
            "RFP 필수가 아닌 항목(부가 제안·범위 외)이 필수로 표시되지 않았는가, 필수 요구사항에 원문 출처가 있는가",
            gate=gate,
            fix_suggestion="범위 분류·필수 여부를 수정하고 출처를 보강",
        ),
        _check(
            "QA-RFP-03",
            cat,
            "요구사항 인용문 ↔ 원문",
            "요구사항의 원문 인용(rfp_quote)이 실제 원문 블록에 존재하는가",
            gate=gate,
            fix_suggestion="인용을 원문과 일치하도록 수정하거나 요구사항을 재추출",
        ),
        _check(
            "QA-RFP-04",
            cat,
            "제출 조건",
            "RFP가 제출 조건을 요구하면 제출 조건 요구사항이 있고 제출용 제안서에 반영되었는가",
            gate=gate,
            mandatory=False,
            fix_suggestion="제출 조건 요구사항(SUB-)을 추가하고 제안서 제출 안내 섹션에 반영",
        ),
        _check(
            "QA-RFP-05",
            cat,
            "평가 기준 대응",
            "RFP 평가 기준이 요구사항(EVA-)과 제출용 제안서에 대응되는가",
            gate=gate,
            mandatory=False,
            fix_suggestion="평가 항목별 대응 섹션을 제안서에 추가",
        ),
        _check(
            "QA-RFP-06",
            cat,
            "원본 출처 참조",
            "요구사항 출처(document_id·block_id)가 실제 등록 문서·추출 블록을 가리키는가",
            gate=gate,
            fix_suggestion="추출을 다시 실행하고 요구사항을 재검증",
        ),
        _check(
            "QA-RFP-07",
            cat,
            "RFP 요구사항 ID 커버리지",
            "RFP가 부여한 요구사항 ID(예: SFR-001)마다 추출 요구사항이 하나 이상 연결되어 있는가",
            gate=gate,
            fix_suggestion="누락 ID의 절을 확인해 요구사항을 추가 추출·수동 추가",
        ),
        _check(
            "QA-RFP-08",
            cat,
            "목차·전략 승인(Gate P)",
            "전체 제안서 작성 전에 목차 초안·주요 전략·차별화·핵심 기술·위험·제출/발표 구성 차이를 담은 브리프를 사람이 승인했는가 (요청서 §4)",
            gate="gate5",
            fix_suggestion="산출물 화면에서 '전략·목차 브리프 작성(에이전트)' 후 검토·승인하고 재생성",
        ),
    ]
    if not inp.blocks or inp.mode != ProjectMode.PROPOSAL:
        return [
            _not_run(c, "RFP 원문 블록이 없음 (구축 단계 프로젝트 또는 추출 전)") for c in checks
        ]
    norm_index = _block_text_index(inp.blocks)
    reqs = inp.requirements
    doc_ids = {d.id for d in inp.documents}
    block_ids = set(norm_index)
    covered: set[str] = set()
    quotes = [normalize_text(r.rfp_quote)[:40] for r in reqs if r.rfp_quote]
    for r in reqs:
        for s in r.source_refs:
            if s.block_id:
                covered.add(s.block_id)
    # 01: 의무 표현 블록 커버리지
    missing_blocks: list[str] = []
    for b in inp.blocks:
        if not any(p in b.text for p in MANDATORY_PATTERNS) or len(b.text) < 15:
            continue
        if b.id in covered:
            continue
        nb = norm_index[b.id]
        if any(q and q in nb for q in quotes):
            continue
        missing_blocks.append(f"{b.location.page or '-'}p/{b.id[:12]}: {b.text[:60]}")
    _finish(
        checks[0],
        review=[
            f"출처로 연결되지 않은 의무 표현 블록 {len(missing_blocks)}건",
            *_cap(missing_blocks, 8),
        ]
        if missing_blocks
        else [],
        files=[d.file_name for d in inp.documents if d.kind.value == "rfp"],
    )
    # 02: 과잉 해석
    over = [
        f"{r.id}: {r.scope_class.value}인데 필수로 표시"
        for r in reqs
        if r.is_mandatory
        and r.scope_class in (ScopeClass.OPTIONAL_PROPOSAL, ScopeClass.OUT_OF_SCOPE)
    ]
    nosrc = [
        f"{r.id}: 필수 요구사항에 원문 출처 없음"
        for r in reqs
        if r.is_mandatory
        and not r.source_refs
        and r.scope_class == ScopeClass.MANDATORY
        and not r.added_by
    ]
    _finish(checks[1], fail=over, review=nosrc, req_ids=[x.split(":")[0] for x in over + nosrc])
    # 03: 인용 ↔ 원문
    mismatch: list[str] = []
    for r in reqs:
        if not r.rfp_quote:
            continue
        q = normalize_text(r.rfp_quote)[:40]
        if not q:
            continue
        targets = [
            norm_index[s.block_id] for s in r.source_refs if s.block_id and s.block_id in norm_index
        ]
        pool = targets or list(norm_index.values())
        if not any(q in t for t in pool):
            mismatch.append(f"{r.id}: 인용 '{r.rfp_quote[:40]}…'이 원문 블록에 없음")
    _finish(checks[2], fail=mismatch, req_ids=[m.split(":")[0] for m in mismatch])
    # 04: 제출 조건
    rfp_text = "\n".join(b.text for b in inp.blocks)
    sub_reqs = [r for r in reqs if r.type == RequirementType.SUBMISSION]
    rev04: list[str] = []
    if any(w in rfp_text for w in SUBMISSION_WORDS) and not sub_reqs:
        rev04.append("RFP에 제출 조건 표현이 있으나 제출 조건 요구사항(SUB-)이 없음")
    for r in sub_reqs:
        if (
            ppt_sub_text
            and r.id not in ppt_sub_text
            and normalize_text(r.name)[:10] not in normalize_text(ppt_sub_text)
        ):
            rev04.append(f"{r.id} '{r.name}'이 제출용 제안서에 언급되지 않음")
    _finish(checks[3], review=rev04, req_ids=[r.id for r in sub_reqs])
    # 05: 평가 기준
    eva_reqs = [r for r in reqs if r.type == RequirementType.EVALUATION]
    rev05: list[str] = []
    if any(w in rfp_text for w in EVALUATION_WORDS) and not eva_reqs:
        rev05.append("RFP에 평가 기준 표현이 있으나 평가 기준 요구사항(EVA-)이 없음")
    for r in eva_reqs:
        if ppt_sub_text and r.id not in ppt_sub_text:
            rev05.append(f"{r.id} '{r.name}' 평가 기준에 대응하는 제안서 내용 없음")
    _finish(checks[4], review=rev05, req_ids=[r.id for r in eva_reqs])
    # 06: 출처 참조 유효성
    dangling: list[str] = []
    for r in reqs:
        for s in r.source_refs:
            if s.document_id not in doc_ids:
                dangling.append(f"{r.id}: 문서 {s.document_id[:12]} 없음")
            elif s.block_id and s.block_id not in block_ids:
                dangling.append(f"{r.id}: 블록 {s.block_id[:12]} 없음 (재추출됨?)")
    _finish(checks[5], fail=dangling, req_ids=[d.split(":")[0] for d in dangling])
    # 07: RFP ID 커버리지 (프로파일이 ID 체계를 찾은 경우만)
    prof = inp.rfp_profile
    if prof is None or not getattr(prof, "requirement_index", None):
        _not_run(checks[6], "RFP 자체 요구사항 ID 체계를 찾지 못함 (문장 단위 RFP)")
    else:
        covered_ids = {r.rfp_requirement_id for r in reqs if r.rfp_requirement_id}
        covered_blocks = {s.block_id for r in reqs for s in r.source_refs if s.block_id}
        uncovered = [
            f"{e.rfp_id} {e.name}".strip()
            for e in prof.requirement_index
            if e.rfp_id not in covered_ids and not (set(e.block_ids) & covered_blocks)
        ]
        _finish(
            checks[6],
            fail=[f"연결된 요구사항이 없는 RFP ID {len(uncovered)}건: " + ", ".join(uncovered[:10])]
            if uncovered
            else [],
        )
    # 08: 목차·전략 승인 (Gate P)
    if inp.gate_p_status == "PASS":
        brief = inp.strategy_brief or {}
        _finish(
            checks[7],
            review=[]
            if brief
            else [
                "브리프 없이 승인됨 — 코드 기본 목차·데이터 기반 문안 (목차·전략을 사람이 직접 확인했는지 기록 필요)"
            ],
        )
    elif inp.gate_p_status == "FAIL":
        _finish(checks[7], fail=["목차·전략 브리프가 반려됨 — 사유를 반영해 다시 작성·승인 필요"])
    else:
        _finish(
            checks[7],
            review=[
                "목차·전략(Gate P)이 아직 승인되지 않음 — 현재 산출물은 시안이며 승인 후 재생성 필요"
                + (" (브리프 작성됨, 검토 대기)" if inp.strategy_brief else " (브리프 미작성)")
            ],
        )
    return checks


# --- 기술 타당성 -------------------------------------------------------------------------------


def check_tech_feasibility(inp: QaInputs, ppt_texts: dict[str, list[str]]) -> list[QaCheck]:
    gate = "gate2"
    cat = QaCategory.TECH_FEASIBILITY
    checks = [
        _check(
            "QA-TECH-01",
            cat,
            "기술 조사 구현 가능성 ↔ 수용 요구사항",
            "구현 불가(infeasible)로 조사된 요구사항이 수용으로 표시되지 않았는가, 미확인(unknown)인 필수 요구사항이 없는가",
            gate=gate,
            fix_suggestion="해당 요구사항의 수용 여부를 '조건부 수용/미수용'으로 바꾸거나 추가 조사를 수행",
        ),
        _check(
            "QA-TECH-02",
            cat,
            "조사 미확인 사항",
            "확인되지 않은 솔루션 기능·주장이 남아 있지 않은가",
            gate=gate,
            mandatory=False,
            fix_suggestion="미확인 항목을 공식 문서로 확인하거나 제안서에서 '확인 필요'로 표기",
        ),
        _check(
            "QA-TECH-03",
            cat,
            "버전 호환성",
            "같은 제품이 서로 다른 버전으로 조사·명시되지 않았는가",
            gate=gate,
            mandatory=False,
            fix_suggestion="적용 버전을 하나로 확정하고 구성 요소·조사 기록을 맞춤",
        ),
        _check(
            "QA-TECH-04",
            cat,
            "기술 간 연계",
            "구성 요소 연결이 존재하는 구성 요소를 가리키고, RFP 연계 대상마다 연계 구성 요소가 있는가",
            gate=gate,
            fix_suggestion="아키텍처 구성 요소·연결을 보완",
        ),
        _check(
            "QA-TECH-05",
            cat,
            "보안·성능·운영 위험",
            "조사에서 식별된 위험과 보안·성능 요구사항에 대응 구성 요소가 있는가",
            gate=gate,
            mandatory=False,
            fix_suggestion="위험 대응 방안을 제안서·아키텍처에 반영",
        ),
        _check(
            "QA-TECH-07",
            cat,
            "기능 요구사항 구현 근거",
            "수용된 필수 기능 요구사항마다 구현 방안이 있고, 구현 가능(feasible/with_risk)으로 조사된 근거 기록이 연결되어 있는가 (조사 선행 원칙)",
            gate=gate,
            fix_suggestion="기술 조사를 실행해 해당 요구사항을 조사 주제에 연결하고 구현 방안을 확정",
        ),
        _check(
            "QA-TECH-06",
            cat,
            "기술 주장 근거",
            "구성 요소·제안서가 언급한 기술이 출처 있는 조사 기록으로 뒷받침되는가",
            gate=gate,
            mandatory=False,
            fix_suggestion="기술 조사를 실행하거나 근거 없는 기술명을 '검토 필요'로 바꿈",
        ),
    ]
    reqs = {r.id: r for r in inp.requirements}
    research = inp.research
    # 01
    fail01: list[str] = []
    rev01: list[str] = []
    for tr in research:
        for rid in tr.related_requirement_ids:
            r = reqs.get(rid)
            if r is None:
                continue
            if tr.feasibility == FeasibilityVerdict.INFEASIBLE and r.acceptance_status in (
                AcceptanceStatus.ACCEPTED,
                AcceptanceStatus.PROPOSED,
            ):
                fail01.append(f"{rid}: 조사 '{tr.topic}' 구현 불가인데 수용/제안 상태")
            elif tr.feasibility == FeasibilityVerdict.UNKNOWN and r.is_mandatory:
                rev01.append(f"{rid}: 조사 '{tr.topic}' 구현 가능성 미확인")
            elif tr.feasibility == FeasibilityVerdict.FEASIBLE_WITH_RISK:
                rev01.append(f"{rid}: '{tr.topic}' 위험 동반 구현 가능 — 위험 수용 여부 확인")
    if not research and inp.mode == ProjectMode.PROPOSAL:
        rev01.append("기술 조사 기록이 없어 구현 가능성을 확인하지 못함 (tech_researcher 미실행)")
    _finish(
        checks[0],
        fail=fail01,
        review=rev01,
        req_ids=[x.split(":")[0] for x in fail01 + rev01 if ":" in x],
    )
    # 02
    unverified = [f"'{tr.topic}': {u}" for tr in research for u in tr.unverified_items]
    _finish(checks[1], review=_cap(unverified) if unverified else [])
    # 03
    versions: dict[str, set[str]] = {}
    for tr in research:
        for s in tr.sources:
            if s.product and s.version:
                versions.setdefault(s.product.casefold(), set()).add(s.version)
    for c in inp.components:
        m = VERSION_RE.search(c.technology)
        if m:
            name = VERSION_RE.sub("", c.technology).strip(" /v").casefold()
            if name:
                versions.setdefault(name, set()).add(m.group(1))
    conflicts = [f"{p}: {', '.join(sorted(v))}" for p, v in versions.items() if len(v) > 1]
    _finish(checks[2], review=conflicts)
    # 04
    comp_ids = {c.id for c in inp.components}
    fail04 = [
        f"{c.name}: 존재하지 않는 구성 요소 {conn.to_component_id} 연결"
        for c in inp.components
        for conn in c.connections
        if conn.to_component_id not in comp_ids
    ]
    rev04: list[str] = []
    targets = (inp.rfp_analysis or {}).get("integration_targets") or []
    integ = [c for c in inp.components if c.layer == "integration"]
    if targets and not integ:
        rev04.append(
            f"RFP 연계 대상 {len(targets)}건({', '.join(map(str, targets[:3]))}…)에 대응하는 연계(integration) 구성 요소 없음"
        )
    _finish(checks[3], fail=fail04, review=rev04)
    # 05
    risks = [f"'{tr.topic}': {rk}" for tr in research for rk in tr.risks]
    sec_reqs = [
        r
        for r in inp.requirements
        if r.category
        in (
            RequirementCategory.SECURITY,
            RequirementCategory.PERFORMANCE,
            RequirementCategory.AVAILABILITY,
        )
    ]
    layers = {c.layer for c in inp.components}
    rev05 = _cap(risks)
    if sec_reqs and "security" not in layers and "infra" not in layers:
        rev05.append(
            f"보안·성능·가용성 요구사항 {len(sec_reqs)}건이 있으나 보안/인프라 구성 요소가 없음"
        )
    _finish(checks[4], review=rev05, req_ids=[r.id for r in sec_reqs] if rev05 else [])
    # 07: 기능 요구사항 구현 근거 (조사 선행)
    fail07: list[str] = []
    rev07: list[str] = []
    ok_verdicts = (FeasibilityVerdict.FEASIBLE, FeasibilityVerdict.FEASIBLE_WITH_RISK)
    for r in inp.requirements:
        if r.type != RequirementType.FUNCTIONAL or not r.is_mandatory:
            continue
        if r.acceptance_status not in (
            AcceptanceStatus.ACCEPTED,
            AcceptanceStatus.PARTIALLY_ACCEPTED,
        ):
            continue
        linked = [
            t
            for t in research
            if r.id in t.related_requirement_ids
            or (r.rfp_requirement_id and r.rfp_requirement_id in t.related_requirement_ids)
        ]
        if not r.implementation_approach:
            rev07.append(f"{r.id}: 구현 방안 미작성")
        elif r.approach_source == "rule" or "[자동 초안]" in r.implementation_approach:
            rev07.append(
                f"{r.id}: 구현 방안 미확정(규칙 초안 [자동 초안] — 에이전트 작성 또는 사람 확정 필요)"
            )
        if not linked:
            rev07.append(f"{r.id}: 근거 조사 없음 (웹 검색 기반 기술 조사 필요)")
        elif not any(t.feasibility in ok_verdicts for t in linked):
            rev07.append(
                f"{r.id}: 연결된 조사가 구현 가능으로 확인되지 않음 ({', '.join(t.feasibility.value for t in linked[:3])})"
            )
    _finish(checks[5], fail=fail07, review=_cap(rev07), req_ids=[x.split(":")[0] for x in rev07])
    # 06
    evidence = " ".join(
        [
            *(f for tr in research for f in tr.confirmed_facts),
            *(s for tr in research for s in tr.recommended_stack),
            *(src.product for tr in research for src in tr.sources),
            *(tr.summary for tr in research),
        ]
    ).casefold()
    techs = {
        c.technology
        for c in inp.components
        if c.technology and "미정" not in c.technology and "검토" not in c.technology
    }
    unbacked = [
        t
        for t in techs
        if not any(
            tok and tok.casefold() in evidence for tok in re.split(r"[\s/,+]+", t) if len(tok) >= 3
        )
    ]
    rev06 = [f"근거 조사 없는 기술: {t}" for t in sorted(unbacked)]
    if not techs and inp.mode == ProjectMode.PROPOSAL:
        rev06.append("구성 요소 기술이 모두 미정 — 아키텍처 확정 필요")
    _finish(checks[6], review=rev06)
    return checks


# --- 문서 간 일관성 ----------------------------------------------------------------------------


def check_consistency_rules(inp: QaInputs, art: CheckArtifacts) -> list[QaCheck]:
    cat = QaCategory.CONSISTENCY
    checks = [
        _check(
            "QA-CON-01",
            cat,
            "PPT ↔ 요구사항 정의서",
            "제출용 PPT가 모든 필수 요구사항 ID를 다루고 없는 ID를 언급하지 않으며, PPT 핵심 수치가 Excel과 같은가",
            gate="gate5",
            fixable=FIX_REGENERATE,
            fix_suggestion="공통 데이터로 산출물을 재생성",
        ),
        _check(
            "QA-CON-02",
            cat,
            "요구사항 정의서 ↔ 추적표",
            "두 파일의 요구사항 ID 집합이 같고 기준 데이터와 일치하는가",
            gate="gate4",
            fixable=FIX_REGENERATE,
            fix_suggestion="공통 데이터로 Excel을 재생성",
        ),
        _check(
            "QA-CON-03",
            cat,
            "요구사항 ↔ WBS",
            "수용된 필수 요구사항마다 WBS Task가 있고, WBS가 없는 요구사항을 참조하지 않는가",
            gate="gate3",
            fix_suggestion="계획(WBS)을 다시 세워 누락 요구사항에 Task를 배정",
        ),
        _check(
            "QA-CON-04",
            cat,
            "WBS ↔ 일정",
            "선후행(FS) 관계·상위 항목 범위·기간(영업일)이 날짜와 맞는가",
            gate="gate3",
            fix_suggestion="일정을 재계산(계획 재생성)하거나 의존 관계를 수정",
        ),
        _check(
            "QA-CON-05",
            cat,
            "WBS ↔ 인력 투입계획",
            "WBS 총 공수 = 인력 계획 합계 = 배정 합계인가",
            gate="gate3",
            fix_suggestion="계획을 재생성해 배정을 다시 계산",
        ),
        _check(
            "QA-CON-06",
            cat,
            "기술 조사 ↔ 아키텍처",
            "구성 요소 기술이 조사 권고 스택과 모순되지 않는가",
            gate="gate2",
            mandatory=False,
            fix_suggestion="아키텍처 기술을 조사 권고와 맞추거나 선택 사유를 기록",
        ),
        _check(
            "QA-CON-07",
            cat,
            "변경 이력 ↔ 현재 기준 데이터",
            "반영된 변경 요청의 요구사항 상태·Revision·변경 후 파일이 현재 데이터와 맞는가",
            gate="gate5",
            fix_suggestion="변경 요청을 재반영하거나 기준 데이터를 점검",
        ),
        _check(
            "QA-CON-08",
            cat,
            "발표용 PPT 미확정 표기",
            "'검토 필요/미정' 표기가 과다하지 않은가",
            gate="gate5",
            mandatory=False,
            fix_suggestion="기술·수용 여부를 확정한 뒤 재생성",
        ),
        _check(
            "QA-CON-09",
            cat,
            "요구사항 정의서 구현 방안·난이도",
            "수용된 구현 대상 요구사항마다 구현 방안(글머리 요약)·구현 난이도(상/중/하)·검사 기준이 기준 데이터와 정의서 파일에 있는가. 규칙 초안([자동 초안])은 사람 확인 대상(조사 검토 게이트)",
            gate="gate2",
            fix_suggestion="'구현 방안 작성(에이전트)'을 실행하거나 요구사항 화면에서 구현 방안·난이도를 확정한 뒤 재생성",
        ),
        _check(
            "QA-CON-10",
            cat,
            "제안 문안 추상 표현·근거 없는 수치",
            "제안 문안(소절 제목·결론형 제목·리드·글머리·표)에 근거 없는 과장 표현(최적의·완벽한·세계적 등)이 없고, 단위가 붙은 수치가 기준 데이터(요구사항·조사·RFP 본문·WBS·인력)나 소절 근거(evidence)에 있는 값인가. 제안 문안이 없으면 제출용 PPT 본문의 추상 표현만 검사",
            gate="gate5",
            mandatory=False,
            fix_suggestion="과장 표현은 근거 있는 사실로 바꾸고, 수치는 출처(요구사항·조사·RFP)를 evidence에 적거나 unverified_claims로 분리",
        ),
        _check(
            "QA-CON-12",
            cat,
            "정의서 구현방안 ↔ 제안서 반영",
            "요구사항 정의서의 구현방안(요구사항별 줄)이 제출용 PPT에 실제로 실려 있는가. 수용된 필수 기능 요구의 구현방안이 한 줄도 없으면 실패, 일부만 있으면 검토",
            gate="gate5",
            fixable=FIX_REGENERATE,
            fix_suggestion="산출물 재생성 (코드가 요구사항 묶음별 구현 방안 소절에 정의서 전문을 싣는다)",
        ),
        _check(
            "QA-CON-11",
            cat,
            "약어·전문 용어 설명",
            "제안 문안(없으면 제출용 PPT 본문)에 쓴 영문 약어가 용어집(glossary)에 있거나 첫 등장 시 괄호로 풀어쓰여 있는가 (RFP·WBS·PM 등 일반 약어와 요구사항 ID 제외)",
            gate="gate5",
            mandatory=False,
            fix_suggestion="glossary에 용어·풀이·비전문가용 설명을 추가하거나 본문 첫 등장 시 '약어(풀이)'로 표기",
        ),
        _check(
            "QA-CON-13",
            cat,
            "구현 설계 깊이 ↔ 제안서 반영",
            "요구사항 묶음마다 구현 설계(구현 구조 2개 이상·처리 절차 3단계 이상·적용 기술과 선정 이유·데이터/인터페이스)가 있고, 그 구성 요소·절차·적용 기술이 제출용 PPT 본문에 실려 있는가. 설계가 없으면 검토(구현 설계 작성 필요), 최소 요건 미달·미반영은 실패, 출처 미확인 기술 선택과 설계에서 빠진 요구사항은 검토",
            gate="gate5",
            fixable=FIX_REGENERATE,
            fix_suggestion="'구현 설계 작성(에이전트, 웹 검색)'을 실행해 묶음마다 구조·절차·적용 기술·근거를 채우고 산출물을 재생성",
        ),
    ]
    f = art.outcome.findings
    by_rule: dict[str, list[tuple[ValidationLevel, str]]] = {}
    for x in f:
        by_rule.setdefault(x.rule, []).append((x.level, x.message))

    def split(rules: tuple[str, ...]) -> tuple[list[str], list[str]]:
        fails: list[str] = []
        revs: list[str] = []
        for r in rules:
            for lvl, msg in by_rule.get(r, []):
                (fails if lvl == ValidationLevel.ERROR else revs).append(msg)
        return fails, revs

    req_ids = {r.id for r in inp.requirements}
    # 01
    f01, r01 = split(("traceability.requirements", "consistency.ppt_figures"))
    sub_text = "\n".join(art.ppt_texts.get("submission", []))
    rfp_ids = {r.rfp_requirement_id for r in inp.requirements if r.rfp_requirement_id}
    unknown = sorted(
        {
            m
            for m in ID_IN_TEXT_RE.findall(sub_text)
            if _looks_like_req_id(m, inp.id_prefixes) and m not in req_ids and m not in rfp_ids
        }
    )
    if unknown:
        f01.append(f"제출용 PPT가 기준에 없는 요구사항 ID를 언급: {', '.join(unknown[:8])}")
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[0], "구축 단계는 PPT를 생성하지 않음")
    else:
        _finish(
            checks[0],
            fail=f01,
            review=r01,
            files=_kind_files(
                inp, DeliverableKind.SUBMISSION_PPTX, DeliverableKind.PRESENTATION_PPTX
            ),
        )
    # 02
    f02, r02 = split(("xlsx.ids",))
    spec = art.excel.get(DeliverableKind.REQUIREMENT_SPEC_XLSX, {})
    trace = art.excel.get(DeliverableKind.TRACE_MATRIX_XLSX, {})
    spec_ids = _ids_from_sheets(spec, ("기능요구사항", "비기능요구사항", "제약·평가·제출조건"))
    trace_ids = _ids_from_sheets(trace, ("추적표",))
    if spec and trace and spec_ids != trace_ids:
        f02.append(
            f"정의서 ID {len(spec_ids)}건 ≠ 추적표 ID {len(trace_ids)}건 (차이 {sorted(spec_ids ^ trace_ids)[:6]})"
        )
    _finish(
        checks[1],
        fail=f02,
        review=r02,
        files=_kind_files(
            inp, DeliverableKind.REQUIREMENT_SPEC_XLSX, DeliverableKind.TRACE_MATRIX_XLSX
        ),
    )
    # 03
    linked: set[str] = set()
    leaf_codes = []
    parents = {w.parent_id for w in inp.ctx.wbs if w.parent_id}
    for w in inp.ctx.wbs:
        if w.id not in parents and w.status.value != "보류":
            leaf_codes.append(w.code)
            linked.update(w.related_requirement_ids)
    need = [
        r
        for r in inp.requirements
        if r.is_mandatory
        and r.acceptance_status != AcceptanceStatus.REJECTED
        and r.change_state != "deleted"
        and r.type in (RequirementType.FUNCTIONAL, RequirementType.NON_FUNCTIONAL)
    ]
    missing = [r.id for r in need if r.id not in linked]
    f03 = [f"WBS Task가 없는 수용 요구사항: {', '.join(missing[:10])}"] if missing else []
    dangling = sorted(
        {
            rid
            for w in inp.ctx.wbs
            for rid in w.related_requirement_ids
            if rid not in {x.id for x in inp.all_requirements}
        }
    )
    if dangling:
        f03.append(f"WBS가 없는 요구사항을 참조: {', '.join(dangling[:6])}")
    if not inp.ctx.wbs:
        f03.append("WBS가 없음")
    _finish(checks[2], fail=f03, req_ids=missing + dangling)
    # 04
    f04: list[str] = []
    by_id = {w.id: w for w in inp.ctx.wbs}
    for d in inp.ctx.dependencies:
        p, s = by_id.get(d.predecessor_id), by_id.get(d.successor_id)
        if not p or not s or not (p.planned_end and s.planned_start):
            continue
        if (
            d.type == DependencyType.FS
            and s.status.value != "보류"
            and p.status.value != "보류"
            and s.planned_start <= p.planned_end
            and d.lag_days >= 0
        ):
            f04.append(
                f"{s.code} 시작({s.planned_start})이 선행 {p.code} 종료({p.planned_end}) 이후가 아님"
            )
    for w in inp.ctx.wbs:
        par = by_id.get(w.parent_id) if w.parent_id else None
        if (
            par
            and w.planned_start
            and w.planned_end
            and par.planned_start
            and par.planned_end
            and (w.planned_start < par.planned_start or w.planned_end > par.planned_end)
            and w.status.value != "보류"
        ):
            f04.append(f"{w.code} 기간이 상위 {par.code} 범위를 벗어남")
        if w.id not in parents and w.planned_start and w.planned_end and w.duration_days:
            actual = inp.calendar.business_days_between(w.planned_start, w.planned_end)
            if abs(actual - w.duration_days) > 1:
                f04.append(f"{w.code} 기간 {w.duration_days}영업일 ≠ 날짜 차이 {actual}영업일")
    _finish(checks[3], fail=_cap(f04), files=_kind_files(inp, DeliverableKind.WBS_GANTT_XLSX))
    # 05
    f05, r05 = split(("consistency.effort", "consistency.allocation"))
    _finish(
        checks[4],
        fail=f05,
        review=r05,
        files=_kind_files(inp, DeliverableKind.WBS_GANTT_XLSX, DeliverableKind.STAFFING_PLAN_XLSX),
    )
    # 06
    stack = " ".join(s for tr in inp.research for s in tr.recommended_stack).casefold()
    r06 = []
    if stack:
        for c in inp.components:
            if (
                c.technology
                and "미정" not in c.technology
                and not any(
                    tok.casefold() in stack
                    for tok in re.split(r"[\s/,+]+", c.technology)
                    if len(tok) >= 3
                )
            ):
                r06.append(f"{c.name} 기술 '{c.technology}'이 조사 권고 스택에 없음")
    _finish(checks[5], review=r06)
    # 07
    f07: list[str] = []
    all_reqs = {r.id: r for r in inp.all_requirements}
    applied = [cr for cr in inp.change_requests if cr.status == ChangeStatus.APPLIED]
    for cr in applied:
        for rid in cr.applied_requirement_ids:
            r = all_reqs.get(rid)
            if r is None:
                f07.append(f"{cr.id}: 반영 요구사항 {rid}가 기준 데이터에 없음")
                continue
            cands = [
                cand
                for cand in cr.accepted_candidates
                if (cand.requirement_id or cand.user_override.get("requirement_id")) == rid
            ]
            for cand in cands:
                if cand.classification == CandidateClass.ADD and r.change_state != "added":
                    f07.append(f"{cr.id}: {rid} 추가 반영인데 상태가 '{r.change_state}'")
                if (
                    cand.classification == CandidateClass.DELETE_REQUEST
                    and r.change_state != "deleted"
                ):
                    f07.append(f"{cr.id}: {rid} 삭제 요청 반영인데 상태가 '{r.change_state}'")
                if cand.classification == CandidateClass.MODIFY and (
                    r.change_state not in ("modified", "deleted") or not r.previous_versions
                ):
                    f07.append(f"{cr.id}: {rid} 수정 반영인데 이전 버전 기록이 없음")
        for fp in cr.after_files:
            try:
                if not inp.paths.resolve_output(fp).exists():
                    f07.append(f"{cr.id}: 변경 후 파일 없음 {fp}")
            except PathSecurityError:
                f07.append(f"{cr.id}: 변경 후 파일 경로가 Output 밖 {fp}")
        if cr.revision_id and not any(rv.id == cr.revision_id for rv in inp.revisions):
            f07.append(f"{cr.id}: Revision {cr.revision_id} 없음")
    for rv in inp.revisions:
        if rv.snapshot_path and not inp.paths.resolve_output(rv.snapshot_path).exists():
            f07.append(f"Revision {rv.number} 스냅샷 파일 없음")
    if not applied and inp.mode == ProjectMode.BUILD:
        _finish(checks[6], review=["반영된 변경 요청이 없음 (검증할 이력 없음)"])
    else:
        _finish(
            checks[6],
            fail=_cap(f07),
            req_ids=[
                x.split(" ")[1]
                for x in f07
                if x.startswith("CR-") and len(x.split(" ")) > 1 and x.split(" ")[1] in all_reqs
            ],
        )
    # 08
    f08, r08 = split(("quality.placeholders",))
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[7], "구축 단계는 PPT를 생성하지 않음")
    else:
        _finish(checks[7], review=r08 + f08)
    # 09: 정의서 구현 방안·난이도·검사 기준 (Phase 14) — 수용된 구현 대상 요구사항마다 채워져 있는가
    f09: list[str] = []
    r09: list[str] = []
    spec_rows = _spec_rows(spec)
    for r in inp.requirements:
        if (
            r.type.value in ("evaluation", "submission")
            or r.acceptance_status == AcceptanceStatus.REJECTED
            or r.change_state == "deleted"
        ):
            continue
        if not r.implementation_approach:
            f09.append(f"{r.id}: 구현 방안 없음")
        elif r.approach_source == "rule" or "[자동 초안]" in r.implementation_approach:
            r09.append(f"{r.id}: 규칙 초안([자동 초안]) — 사람 확인 필요")
        if r.difficulty is None:
            r09.append(f"{r.id}: 구현 난이도 미기재 (변경 영향 산정 시 '중' 가중치 적용)")
        if not r.acceptance_criteria:
            r09.append(f"{r.id}: 검사 기준 없음")
        row = spec_rows.get(r.id)
        if row is not None:
            if not row.get("구현방안") or row.get("구현방안") == "검토 필요":
                f09.append(f"{r.id}: 정의서 파일의 구현방안 칸이 비어 있음")
            if row.get("구현 난이도") not in ("상", "중", "하"):
                r09.append(f"{r.id}: 정의서 파일의 구현 난이도가 상/중/하가 아님")
    _finish(
        checks[8],
        fail=_cap(f09),
        review=_cap(r09),
        req_ids=[x.split(":")[0] for x in f09 + r09],
        files=_kind_files(inp, DeliverableKind.REQUIREMENT_SPEC_XLSX),
    )
    # 10·11: 제안 문안 추상 표현·근거 없는 수치, 약어 미설명 (Phase 19)
    if inp.mode != ProjectMode.PROPOSAL:
        for i in (9, 10, 11, 12):
            _not_run(checks[i], "구축 단계는 PPT를 생성하지 않음")
        return checks
    content = inp.ctx.proposal_content or {}
    sections = [x for x in content.get("sections") or [] if isinstance(x, dict)]
    ppt_files = _kind_files(inp, DeliverableKind.SUBMISSION_PPTX)
    _finish(
        checks[9],
        review=_cap(_abstract_and_number_issues(inp, sections, sub_text), 12),
        files=ppt_files,
    )
    f12, r12, ids12 = _approach_coverage_issues(inp, sub_text)
    _finish(checks[10], fail=_cap(f12), review=_cap(r12), req_ids=ids12, files=ppt_files)
    _finish(
        checks[11],
        review=_cap(_unexplained_acronyms(content, sections, sub_text), 6),
        files=ppt_files,
    )
    f13, r13, ids13 = _design_depth_issues(inp, sub_text)
    _finish(checks[12], fail=_cap(f13), review=_cap(r13), req_ids=ids13, files=ppt_files)
    return checks


def _design_depth_issues(inp: QaInputs, sub_text: str) -> tuple[list[str], list[str], list[str]]:
    """QA-CON-13: 구현 설계(design_writer)의 깊이와 제출용 PPT 반영 (Phase 23).

    - 설계가 없는데 설계 대상 요구가 있으면 검토(설계 작성 필요).
    - 묶음마다 구조 2·절차 3·적용 기술 1·데이터/인터페이스 1 미만이면 실패; 구성 요소 이름·절차 제목·적용 기술이
      덱 본문에 없으면 실패(재생성). 출처 미확인(verified=false) 기술 선택, 설계에서 빠진 요구사항은 검토.
    """
    from proposal_agent.generators.pptx.req_groups import design_texts
    from proposal_agent.requirements.design import (
        design_min_issues,
        design_targets,
        needs_verification,
        resolve_design_ids,
    )

    fails: list[str] = []
    revs: list[str] = []
    ids: list[str] = []
    designs = [d for d in (inp.ctx.designs or []) if isinstance(d, dict)]
    targets = design_targets(list(inp.requirements))
    if not designs:
        if targets:
            revs.append(
                f"구현 설계 없음 — 설계 대상 요구 {len(targets)}건: '구현 설계 작성(에이전트, 웹 검색)' 실행 필요 "
                "(묶음 소절은 정의서 구현방안 표와 기능 블록 도식만 싣는다)"
            )
        return fails, revs, ids
    deck = normalize_text(sub_text)
    covered: set[str] = set()
    for d in designs:
        title = str(d.get("title", "") or "묶음")[:30]
        rids = resolve_design_ids(d, targets)
        covered.update(rids)
        issues = design_min_issues(d)
        if issues:
            fails.append(f"{title}: {', '.join(issues)}")
            ids.extend(rids)
        if deck:
            missing = [t for t in design_texts(d) if normalize_text(t[:25]) not in deck]
            if missing:
                fails.append(
                    f"{title}: 설계 문구 {len(missing)}건이 제출용 PPT에 없음 ({'; '.join(x[:20] for x in missing[:3])})"
                )
                ids.extend(x for x in rids if x not in ids)
        unverified = [
            str(t.get("choice", ""))[:24]
            for t in d.get("tech_choices") or []
            if isinstance(t, dict) and needs_verification(t)
        ]
        if unverified:
            revs.append(
                f"{title}: 출처 미확인 기술 선택 {len(unverified)}건 ({', '.join(unverified[:3])}) — 공식 문서 확인 필요"
            )
    left = [r.id for r in targets if r.id not in covered]
    if left:
        revs.append(f"설계에 없는 요구사항 {len(left)}건: {', '.join(left[:8])}")
        ids.extend(x for x in left if x not in ids)
    return fails, revs, ids


def _approach_coverage_issues(
    inp: QaInputs, sub_text: str
) -> tuple[list[str], list[str], list[str]]:
    """QA-CON-12: 정의서 구현방안 줄(8자 이상)의 앞 25자가 제출용 PPT 본문에 있는지."""
    deck = normalize_text(sub_text)
    fails: list[str] = []
    revs: list[str] = []
    ids: list[str] = []
    if not deck:
        return fails, revs, ids
    for r in inp.requirements:
        if r.type.value not in ("functional", "non_functional") or r.change_state == "deleted":
            continue
        if (
            r.acceptance_status == AcceptanceStatus.REJECTED
            or not (r.implementation_approach or "").strip()
        ):
            continue
        lines = [
            x.strip().lstrip("-•·▪ ").strip()
            for x in (r.implementation_approach or "").splitlines()
            if len(x.strip()) >= 8
        ]
        if not lines:
            continue
        hit = sum(1 for x in lines if normalize_text(x[:25]) in deck)
        if hit == 0:
            msg = f"{r.id} {r.name[:24]}: 구현방안 {len(lines)}줄 중 제안서에 없음"
            (fails if (r.is_mandatory and r.type.value == "functional") else revs).append(msg)
            ids.append(r.id)
        elif hit < len(lines):
            revs.append(f"{r.id} {r.name[:24]}: 구현방안 {len(lines)}줄 중 {hit}줄만 제안서에 있음")
            ids.append(r.id)
    return fails, revs, ids


def _section_texts(sec: dict[str, Any]) -> list[str]:
    out = [
        str(sec.get("title", "") or ""),
        str(sec.get("key_message", "") or ""),
        str(sec.get("lead_message", "") or ""),
        str(sec.get("body", "") or ""),
        *[str(x) for x in sec.get("bullets", []) or []],
        *[str(c) for r in sec.get("table_rows", []) or [] for c in r],
    ]
    vis = sec.get("visual")
    if isinstance(vis, dict):
        out.append(str(vis.get("title", "") or ""))
        out.extend(f"{x.get('title', '')} {x.get('text', '')}" for x in vis.get("items", []) or [])
        out.extend(str(c) for r in vis.get("rows", []) or [] for c in r)
    td = sec.get("tech_detail")
    if isinstance(td, dict):
        for k in ("components", "interfaces", "data", "operations", "security"):
            out.extend(str(x) for x in td.get(k, []) or [])
        out.append(str(td.get("selection_reason", "") or ""))
        for a in td.get("alternatives", []) or []:
            out.append(f"{a.get('name', '')} {a.get('cost_note', '')} {a.get('basis', '')}")
            out.extend(str(x) for x in (a.get("pros") or []) + (a.get("cons") or []))
    return [x for x in out if x]


def _section_label(sec: dict[str, Any]) -> str:
    return str(sec.get("number") or "").strip() or str(sec.get("title", ""))[:20]


def _base_number_tokens(inp: QaInputs) -> set[str]:
    """기준 데이터 안의 숫자 토큰(쉼표 제거). 제안 문안의 수치가 여기 없으면 '근거 없는 수치'."""
    parts: list[str] = []
    for r in inp.all_requirements or inp.requirements:
        parts.extend(
            (
                r.name,
                r.description,
                r.implementation_approach or "",
                r.acceptance_criteria,
                r.difficulty_basis,
            )
        )
    for t in inp.research:
        parts.extend((t.topic, t.summary, *t.unverified_items, *t.recommended_stack))
        parts.extend(f"{src.title} {src.url}" for src in t.sources)
    parts.extend(b.text for b in inp.blocks)
    parts.extend(w.name for w in inp.ctx.wbs)
    parts.extend(str(getattr(w, "estimate_basis", "") or "") for w in inp.ctx.wbs)
    parts.extend(str(w.effort_md) for w in inp.ctx.wbs)
    parts.append(str(inp.ctx.total_effort_md))
    parts.extend(f"{r.role_name} {r.headcount_planned}" for r in inp.ctx.roles)
    parts.append(str(len(inp.ctx.requirements)))
    parts.append(str(len(inp.ctx.mandatory_requirements())))
    # 범위 분류·유형·수용 상태별 건수 ("부가 제안 2건" 같은 집계 수치의 근거)
    for attr in ("scope_class", "type", "acceptance_status", "category"):
        cnt: dict[str, int] = {}
        for r in inp.ctx.requirements:
            cnt[str(getattr(r, attr, ""))] = cnt.get(str(getattr(r, attr, "")), 0) + 1
        parts.extend(str(v) for v in cnt.values())
    parts.append(str(len(inp.ctx.roles)))
    parts.append(str(sum(r.headcount_planned for r in inp.ctx.roles)))
    if inp.ctx.period_start and inp.ctx.period_end:
        days = (inp.ctx.period_end - inp.ctx.period_start).days
        parts.extend((str(days), str(round(days / 30)), str(round(days / 30.4)), str(days // 7)))
    analysis = inp.ctx.analysis or {}
    parts.append(str(analysis))
    corpus = " ".join(parts)
    return {m.replace(",", "") for m in re.findall(r"\d[\d,]*(?:\.\d+)?", corpus)}


def _abstract_and_number_issues(
    inp: QaInputs, sections: list[dict[str, Any]], ppt_text: str
) -> list[str]:
    issues: list[str] = []
    if not sections:
        found = sorted({w for w in ABSTRACT_WORDS if w in ppt_text})
        if found:
            issues.append("제출용 PPT 본문의 추상 표현 (제안 문안 없음): " + ", ".join(found[:8]))
        return issues
    base = _base_number_tokens(inp)
    for sec in sections:
        text = "\n".join(_section_texts(sec))
        label_ = _section_label(sec)
        words = [w for w in ABSTRACT_WORDS if w in text]
        if words:
            issues.append(f"[{label_}] 추상 표현: {', '.join(words[:4])}")
        evidence_text = " ".join(
            f"{e.get('ref', '')} {e.get('text', '')}"
            for e in sec.get("evidence", []) or []
            if isinstance(e, dict)
        )
        ev_tokens = {m.replace(",", "") for m in re.findall(r"\d[\d,]*(?:\.\d+)?", evidence_text)}
        unsupported: list[str] = []
        for num, unit in NUMBER_CLAIM_RE.findall(text):
            tok = num.replace(",", "")
            if tok in base or tok in ev_tokens or tok in unsupported:
                continue
            unsupported.append(f"{num}{unit}")
        if unsupported:
            issues.append(
                f"[{label_}] 근거 없는 수치: {', '.join(dict.fromkeys(unsupported).keys())}"
                + ("" if sec.get("evidence") else " (evidence 없음)")
            )
    return issues


def _unexplained_acronyms(
    content: dict[str, Any], sections: list[dict[str, Any]], ppt_text: str
) -> list[str]:
    glossary = {
        str(g.get("term", "")).upper() for g in content.get("glossary") or [] if isinstance(g, dict)
    } | {
        str(g.get("full_form", "")).upper()
        for g in content.get("glossary") or []
        if isinstance(g, dict)
    }
    text = "\n".join(t for sec in sections for t in _section_texts(sec)) if sections else ppt_text
    if not text.strip():
        return []
    first_where: dict[str, str] = {}
    if sections:
        for sec in sections:
            for m in ACRONYM_RE.findall("\n".join(_section_texts(sec))):
                first_where.setdefault(m, _section_label(sec))
    unexplained: list[str] = []
    for acr in dict.fromkeys(ACRONYM_RE.findall(text)):
        parts = [p for p in acr.split("/") if p]
        if all(
            p in COMMON_ACRONYMS or ROMAN_RE.match(p) or p.isdigit() or p in glossary for p in parts
        ):
            continue  # 'UI/UX' 처럼 슬래시로 묶인 약어는 각각 설명되어 있으면 통과
        if re.search(rf"{re.escape(acr)}\s*[\(（]|[\(（]\s*{re.escape(acr)}\s*[\)）]", text):
            continue  # 본문에서 '약어(풀이)' 또는 '풀이(약어)'로 설명됨
        unexplained.append(acr + (f"({first_where[acr]})" if acr in first_where else ""))
    if not unexplained:
        return []
    src = "제안 문안" if sections else "제출용 PPT 본문 (제안 문안 없음)"
    return [
        f"{src} 약어 미설명 {len(unexplained)}건: " + ", ".join(unexplained[:15]),
        "용어집(glossary)에 용어·풀이·비전문가용 설명을 추가하거나 첫 등장 시 괄호로 풀어쓰기",
    ]


def _spec_rows(sheets: dict[str, list[list[Any]]]) -> dict[str, dict[str, Any]]:
    """정의서 본문 시트 행을 요구사항 ID → {컬럼명: 값}으로 모은다."""
    out: dict[str, dict[str, Any]] = {}
    for name in ("기능요구사항", "비기능요구사항", "제약·평가·제출조건"):
        data = sheets.get(name)
        if not data or "요구사항 ID" not in data[0]:
            continue
        headers = [str(h) for h in data[0]]
        for row in data[1:]:
            rec = dict(zip(headers, row, strict=False))
            rid = rec.get("요구사항 ID")
            if rid:
                out[str(rid)] = rec
    return out


def _looks_like_req_id(token: str, prefixes: dict[str, str]) -> bool:
    pre = token.split("-")[0]
    return pre in set(prefixes.values()) or pre in {"SFR", "SNR", "FR", "NFR", "CON", "EVA", "SUB"}


def _ids_from_sheets(sheets: dict[str, list[list[Any]]], names: tuple[str, ...]) -> set[str]:
    out: set[str] = set()
    for n in names:
        data = sheets.get(n)
        if not data or "요구사항 ID" not in data[0]:
            continue
        i = data[0].index("요구사항 ID")
        out |= {str(r[i]) for r in data[1:] if r[i]}
    return out


def _kind_files(inp: QaInputs, *kinds: DeliverableKind) -> list[str]:
    return [inp.deliverables[k].file_path for k in kinds if k in inp.deliverables]


# --- 파일 품질 ---------------------------------------------------------------------------------


def _template_size(inp: QaInputs, purpose: str) -> tuple[float, float] | None:
    prof = inp.template_profiles.get(purpose)
    if prof is None:
        return None
    return (float(prof.slide_width_cm), float(prof.slide_height_cm))


def _template_check(
    inp: QaInputs, art: CheckArtifacts, path: Path, purpose: str, label: str
) -> None:
    prof = inp.template_profiles.get(purpose)
    if prof is None:
        return
    try:
        fails, reviews = check_template_compliance(
            path, prof, inp.rfp_toc_tree if purpose == "submission" else None, label
        )
    except Exception as exc:  # noqa: BLE001
        fails, reviews = [f"{label}: 템플릿 준수 검사 실패 ({type(exc).__name__}: {exc})"], []
    art.template_fails.extend(fails)
    art.template_reviews.extend(reviews)


def run_file_checks(inp: QaInputs) -> tuple[list[QaCheck], CheckArtifacts]:
    cat = QaCategory.FILE_QUALITY
    gate = "gate4"
    art = CheckArtifacts()
    checks = [
        _check(
            "QA-FILE-01",
            cat,
            "산출물 파일 열기",
            "PPT·Excel·PNG 파일이 정상적으로 열리는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="산출물 재생성",
        ),
        _check(
            "QA-FILE-02",
            cat,
            "슬라이드 방향·규격",
            "가로 16:9, 슬라이드 3장 이상",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="PPT 재생성",
        ),
        _check(
            "QA-FILE-03",
            cat,
            "텍스트·도형 잘림",
            "도형이 슬라이드 안에 있고 텍스트가 넘치지 않는가",
            gate=gate,
            fix_suggestion="해당 슬라이드 문안을 줄이거나 PowerPoint에서 육안 확인",
        ),
        _check(
            "QA-FILE-04",
            cat,
            "표·차트 겹침",
            "텍스트 도형끼리 크게 겹치지 않는가",
            gate=gate,
            mandatory=False,
            fix_suggestion="레이아웃 조정 후 재생성, PowerPoint 육안 확인",
        ),
        _check(
            "QA-FILE-05",
            cat,
            "Excel 시트·컬럼",
            "필수 시트와 컬럼이 모두 있는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="Excel 재생성",
        ),
        _check(
            "QA-FILE-06",
            cat,
            "날짜·수식",
            "종료일 ≥ 시작일, 수식 형식 오류 없음, 산정 근거 있음",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="계획 데이터 점검 후 Excel 재생성",
        ),
        _check(
            "QA-FILE-07",
            cat,
            "요구사항 ID",
            "Excel의 요구사항 ID가 기준 데이터와 같고 형식이 올바른가, 미실행 테스트가 '성공'으로 표시되지 않았는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="Excel 재생성",
        ),
        _check(
            "QA-FILE-08",
            cat,
            "간트 차트 ↔ 일정 데이터",
            "간트 시트의 시작·종료와 채움 구간이 WBS 시트 날짜와 같은가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="WBS·간트 재생성",
        ),
        _check(
            "QA-FILE-09",
            cat,
            "파일 경로",
            "산출물 경로가 프로젝트 Output 폴더 안에 있고 실제로 존재하는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="누락 파일을 재생성 (경로 조작은 거부)",
        ),
        _check(
            "QA-FILE-10",
            cat,
            "산출물 누락",
            "모드별 필수 산출물이 모두 생성되었는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="산출물 재생성",
        ),
        _check(
            "QA-FILE-12",
            cat,
            "편집 가능성",
            "생성 PPT에 그림(비트맵) 도형이 없고 모든 요소가 도형·표·텍스트로 편집 가능한가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="산출물 재생성 (도식은 코드가 도형으로 그림)",
        ),
        _check(
            "QA-FILE-13",
            cat,
            "페이지 제한 (RFP)",
            "RFP가 지정한 제안서 페이지 제한(쪽 수) 안에 제출용 PPT 슬라이드 수(표지·목차 제외)가 드는가. 90% 이상이면 검토",
            gate=gate,
            fix_suggestion="소절 글머리·표 행을 줄여 재생성하거나 부록으로 이동",
        ),
        _check(
            "QA-FILE-14",
            cat,
            "실제 렌더링·재열기",
            "PowerPoint(COM)로 제출용·발표용 PPT를 실제로 열어 슬라이드마다 PNG로 내보내고, 슬라이드 수가 python-pptx와 같으며 저장 왕복 후 재열기가 되는가. PowerPoint가 없거나 설정이 꺼져 있으면 미실행",
            gate=gate,
            mandatory=False,
            fix_suggestion="PNG(11_QA_Validation_Reports/render)를 보고 깨진 슬라이드를 확인한 뒤 재생성",
        ),
        _check(
            "QA-FILE-11",
            cat,
            "제안서 템플릿 준수",
            "등록된 제안서 템플릿과 용지 크기·머리글 2단 위치·장 순서(RFP 지정 목차)·글꼴이 맞는가",
            gate=gate,
            fixable=FIX_REGENERATE,
            fix_suggestion="템플릿을 다시 확인하고 산출물 재생성",
        ),
    ]
    out = art.outcome
    project_dir = inp.paths.project_output_dir(inp.project.output_dir)
    path_fail: list[str] = []
    for kind, d in inp.deliverables.items():
        try:
            path = inp.paths.resolve_output(d.file_path)
            ensure_within(project_dir, path)
        except PathSecurityError as exc:
            path_fail.append(f"{kind.value}: 경로 거부 {d.file_path} ({exc})")
            continue
        if not path.exists():
            path_fail.append(f"{kind.value}: 파일 없음 {d.file_path}")
            out.opened[kind.value] = False
            continue
        if kind.value.endswith("xlsx"):
            art.excel[kind] = check_xlsx(path, kind, out, inp.ctx)
        elif kind == DeliverableKind.SUBMISSION_PPTX:
            art.ppt_texts["submission"] = check_pptx(
                path, out, "제출용", expected_size_cm=_template_size(inp, "submission")
            )
            _template_check(inp, art, path, "submission", "제출용")
        elif kind == DeliverableKind.PRESENTATION_PPTX:
            art.ppt_texts["presentation"] = check_pptx(
                path, out, "발표용", expected_size_cm=_template_size(inp, "presentation")
            )
            _template_check(inp, art, path, "presentation", "발표용")
        elif kind.value.endswith("png"):
            from proposal_agent.agents.tools import verify_file

            info = verify_file(path)
            out.opened[kind.value] = bool(info.get("ok"))
            if not info.get("ok"):
                out.add(
                    ValidationLevel.ERROR,
                    "png.open",
                    f"{kind.value}: {info.get('error')}",
                    d.file_path,
                )
    missing = [k.value for k in inp.required_kinds if k not in inp.deliverables]
    for k in missing:
        out.add(ValidationLevel.ERROR, "deliverable.missing", f"산출물 없음: {k}")
    if art.excel or art.ppt_texts:
        check_consistency(inp.ctx, art.excel, art.ppt_texts, out)
    gantt_fail = _check_gantt(inp, art)

    by_rule: dict[str, list[tuple[ValidationLevel, str]]] = {}
    for x in out.findings:
        by_rule.setdefault(x.rule, []).append((x.level, x.message))

    def split(rules: tuple[str, ...]) -> tuple[list[str], list[str]]:
        fails: list[str] = []
        revs: list[str] = []
        for r in rules:
            for lvl, msg in by_rule.get(r, []):
                (fails if lvl == ValidationLevel.ERROR else revs).append(msg)
        return fails, revs

    all_files = [d.file_path for d in inp.deliverables.values()]
    f, r = split(("pptx.picture",))
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[10], "구축 단계는 PPT를 생성하지 않음")
    else:
        _finish(
            checks[10],
            fail=f,
            review=r,
            files=_kind_files(
                inp, DeliverableKind.SUBMISSION_PPTX, DeliverableKind.PRESENTATION_PPTX
            ),
        )
        # 허용 그림(사진·일러스트·장식·참고) 예외는 판정을 바꾸지 않고 보고서·manifest에 기록한다
        checks[10].issues.extend(x.message for x in out.findings if x.rule == "pptx.photo")
    # 13: 페이지 제한 (Phase 21)
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[11], "구축 단계는 PPT를 생성하지 않음")
    elif not inp.page_limit:
        _not_run(checks[11], "RFP에 페이지 제한 문구 없음")
    else:
        f13, r13 = _page_limit_issues(inp, art)
        _finish(
            checks[11],
            fail=f13,
            review=r13,
            files=_kind_files(inp, DeliverableKind.SUBMISSION_PPTX),
        )
    # 14: 실제 렌더링·재열기 (Phase 21)
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[12], "구축 단계는 PPT를 생성하지 않음")
    elif not inp.render_check:
        _not_run(checks[12], "렌더링 검증이 꺼져 있음 (PWA_QA_RENDER_CHECK=1 또는 화면에서 켜기)")
    else:
        f14, r14, files14 = _render_issues(inp, art)
        if f14 is None:
            _not_run(checks[12], r14[0] if r14 else "렌더링 엔진 없음", applicable=True)
        else:
            _finish(checks[12], fail=f14, review=r14, files=files14)
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[13], "구축 단계는 PPT를 생성하지 않음")
    elif not inp.template_profiles:
        _not_run(checks[13], "제안서 템플릿(PPT) 미등록 — 코드 기본 양식으로 생성됨")
    else:
        _finish(
            checks[13],
            fail=list(art.template_fails),
            review=list(art.template_reviews),
            files=_kind_files(
                inp, DeliverableKind.SUBMISSION_PPTX, DeliverableKind.PRESENTATION_PPTX
            ),
        )
    f, r = split(("pptx.open", "xlsx.open", "png.open"))
    f += [m for m in path_fail if "파일 없음" in m]
    _finish(checks[0], fail=f, review=r, files=all_files)
    f, r = split(("pptx.orientation", "pptx.size", "pptx.slides"))
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[1], "구축 단계는 PPT를 생성하지 않음")
    else:
        _finish(
            checks[1],
            fail=f,
            review=r,
            files=_kind_files(
                inp, DeliverableKind.SUBMISSION_PPTX, DeliverableKind.PRESENTATION_PPTX
            ),
        )
    f, r = split(("pptx.bounds", "pptx.overflow"))
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[2], "구축 단계는 PPT를 생성하지 않음")
    else:
        _finish(
            checks[2],
            fail=f,
            review=r,
            files=_kind_files(
                inp, DeliverableKind.SUBMISSION_PPTX, DeliverableKind.PRESENTATION_PPTX
            ),
        )
    f, r = split(("pptx.overlap",))
    if inp.mode != ProjectMode.PROPOSAL:
        _not_run(checks[3], "구축 단계는 PPT를 생성하지 않음")
    else:
        _finish(checks[3], fail=f, review=r)
    f, r = split(("xlsx.sheet", "xlsx.columns"))
    _finish(
        checks[4],
        fail=f,
        review=r,
        files=[d.file_path for k, d in inp.deliverables.items() if k.value.endswith("xlsx")],
    )
    f, r = split(("xlsx.dates", "xlsx.formula", "xlsx.basis"))
    _finish(
        checks[5],
        fail=f,
        review=r,
        files=_kind_files(inp, DeliverableKind.WBS_GANTT_XLSX, DeliverableKind.STAFFING_PLAN_XLSX),
    )
    f, r = split(("xlsx.ids", "xlsx.req_ref", "xlsx.wbs_parent", "xlsx.test_result"))
    bad_ids = _invalid_ids(inp)
    if bad_ids:
        f.append("형식이 맞지 않는 요구사항 ID: " + ", ".join(bad_ids[:8]))
    _finish(checks[6], fail=f, review=r, req_ids=bad_ids)
    if DeliverableKind.WBS_GANTT_XLSX not in inp.deliverables:
        _not_run(checks[7], "WBS·간트 파일 없음", applicable=True)
    else:
        _finish(checks[7], fail=gantt_fail, files=_kind_files(inp, DeliverableKind.WBS_GANTT_XLSX))
    _finish(checks[8], fail=path_fail, files=all_files)
    _finish(checks[9], fail=[f"산출물 없음: {k}" for k in missing])
    return checks, art


def _deck_paths(inp: QaInputs) -> list[tuple[str, Path]]:
    """(라벨, 절대 경로) — 제출용·발표용 덱 중 존재하는 것."""
    out: list[tuple[str, Path]] = []
    for kind, label in (
        (DeliverableKind.SUBMISSION_PPTX, "제출용"),
        (DeliverableKind.PRESENTATION_PPTX, "발표용"),
    ):
        d = inp.deliverables.get(kind)
        if d is None:
            continue
        try:
            path = inp.paths.resolve_output(d.file_path)
        except PathSecurityError:
            continue
        if path.exists():
            out.append((label, path))
    return out


def _content_slide_count(path: Path) -> tuple[int, int]:
    """(전체 슬라이드, 본문 슬라이드). 표지·목차·장 구분(제목이 목차/CONTENTS/장 표지이거나 본문 텍스트가 거의 없는) 제외."""
    from pptx import Presentation

    prs = Presentation(str(path))
    total = len(prs.slides)
    body = 0
    for idx, slide in enumerate(prs.slides, start=1):
        texts = [sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame]
        joined = " ".join(t for t in texts if t)
        has_obj = any(
            (getattr(sh, "has_table", False) and sh.has_table)
            or (getattr(sh, "has_chart", False) and sh.has_chart)
            or "GROUP" in str(sh.shape_type)
            or "PICTURE" in str(sh.shape_type)
            for sh in slide.shapes
        )
        first = (texts[0] if texts else "")[:20]
        if idx == 1 or first in ("목차", "CONTENTS") or (len(joined) < 60 and not has_obj):
            continue
        body += 1
    return total, body


def _page_limit_issues(inp: QaInputs, art: CheckArtifacts) -> tuple[list[str], list[str]]:
    limit = int(inp.page_limit or 0)
    fails: list[str] = []
    revs: list[str] = []
    for label, path in _deck_paths(inp):
        if label != "제출용":
            continue
        total, body = _content_slide_count(path)
        msg = f"제출용 본문 슬라이드 {body}장 (전체 {total}장) / 제한 {limit}쪽"
        if body > limit:
            fails.append(msg + " — 초과")
        elif body >= limit * 0.9:
            revs.append(msg + " — 여유 10% 미만")
    return fails, revs


def _render_issues(
    inp: QaInputs, art: CheckArtifacts
) -> tuple[list[str] | None, list[str], list[str]]:
    """QA-FILE-14. (None, [사유], []) 이면 엔진 없음 → 미실행."""
    from proposal_agent.validation.render import com_available, render_pptx

    ok, why = com_available()
    if not ok:
        return None, [why], []
    fails: list[str] = []
    revs: list[str] = []
    files: list[str] = []
    base = inp.render_dir or (
        inp.paths.project_output_dir(inp.project.output_dir) / "11_QA_Validation_Reports" / "render"
    )
    for label, path in _deck_paths(inp):
        res = render_pptx(path, base / path.stem)
        art.render_results[label] = res
        if not res.available:
            return None, res.errors or ["렌더링 엔진 없음"], []
        if res.ok:
            try:
                files.append(inp.paths.relative_output(base / path.stem))
            except PathSecurityError:
                files.append(str(base / path.stem))
        else:
            fails.append(f"{label}: {res.summary} — " + "; ".join(res.errors[:3]))
    return fails, revs, files


def check_editability(inp: QaInputs, art: CheckArtifacts) -> list[QaCheck]:
    """QA-EDIT-01~10 편집 가능성 체크리스트 (Phase 21): 생성 PPT의 객체 구조를 사람이 PowerPoint에서 고칠 수 있는가."""
    cat = QaCategory.EDITABILITY
    gate = "gate4"
    specs = [
        (
            "QA-EDIT-01",
            "텍스트 편집",
            "모든 글이 텍스트 상자·표 셀·도형 텍스트에 있고(글자가 든 그림 없음) 빈 텍스트 상자가 과다하지 않은가",
            True,
        ),
        (
            "QA-EDIT-02",
            "표 객체",
            "표는 네이티브 표 객체이며 표를 그림으로 넣은 슬라이드가 없는가 (제출용은 표 1개 이상)",
            True,
        ),
        (
            "QA-EDIT-03",
            "도형·연결선",
            "도식이 자동 도형·연결선·텍스트 상자로 그려져 개별 선택·이동·수정이 가능한가 (자유형 과다 없음)",
            False,
        ),
        (
            "QA-EDIT-04",
            "차트 객체",
            "차트는 네이티브 차트(데이터 시트 편집 가능)이며 차트 이미지가 없는가",
            True,
        ),
        (
            "QA-EDIT-05",
            "그룹 해제 가능",
            "그룹 도형이 자식 2개 이상의 평범한 그룹(중첩 3단 이하)이라 그룹 해제로 개별 편집이 되는가",
            False,
        ),
        (
            "QA-EDIT-06",
            "그림 정책",
            "그림은 허용 태그(photo/illustration/decoration/reference)가 붙은 사진·일러스트뿐인가 (정보성 자료 이미지 없음)",
            True,
        ),
        (
            "QA-EDIT-07",
            "글꼴",
            "본문 글꼴이 템플릿·기본 글꼴 집합(맑은 고딕·템플릿 글꼴) 안에 있어 글꼴 치환 없이 편집되는가",
            False,
        ),
        (
            "QA-EDIT-08",
            "발표자 노트",
            "발표용 본문 슬라이드에 발표자 노트가 절반 이상 있는가",
            False,
        ),
        (
            "QA-EDIT-09",
            "레이아웃·마스터",
            "모든 슬라이드가 레이아웃을 쓰고 슬라이드 밖·음수 좌표 도형이 없는가",
            False,
        ),
        (
            "QA-EDIT-10",
            "재열기 안정성",
            "python-pptx로 열고 저장 왕복 후 다시 열어도 슬라이드 수·도형 수가 같은가",
            True,
        ),
    ]
    checks = [
        _check(
            cid,
            cat,
            target,
            criteria,
            gate=gate,
            mandatory=mand,
            fixable=FIX_REGENERATE if mand else None,
            fix_suggestion="산출물 재생성 (도식·표·차트는 코드가 객체로 그림)",
        )
        for cid, target, criteria, mand in specs
    ]
    if inp.mode != ProjectMode.PROPOSAL:
        for c in checks:
            _not_run(c, "구축 단계는 PPT를 생성하지 않음")
        return checks
    decks = _deck_paths(inp)
    if not decks:
        for c in checks:
            _not_run(c, "생성된 PPT 없음", applicable=True)
        return checks
    from pptx import Presentation

    from proposal_agent.generators.base import PICTURE_TAGS
    from proposal_agent.validation.deliverables import _flatten

    fails: dict[int, list[str]] = {i: [] for i in range(10)}
    revs: dict[int, list[str]] = {i: [] for i in range(10)}
    files = _kind_files(inp, DeliverableKind.SUBMISSION_PPTX, DeliverableKind.PRESENTATION_PPTX)
    known_fonts = {"맑은 고딕", "Malgun Gothic", "Arial", "Calibri"}
    for prof in inp.template_profiles.values():
        for attr in ("title_font", "body_font"):
            v = getattr(prof, attr, None)
            if v:
                known_fonts.add(str(v))
    for label, path in decks:
        try:
            prs = Presentation(str(path))
        except Exception as exc:  # noqa: BLE001
            fails[9].append(f"{label}: 열기 실패 {exc}")
            continue
        n_slides = len(prs.slides)
        empty_text = 0
        tables = 0
        charts = 0
        freeform = 0
        auto = 0
        groups_bad: list[str] = []
        pictures_bad: list[str] = []
        pictures_info: list[str] = []
        fonts: set[str] = set()
        notes_body = 0
        body_slides = 0
        no_layout = 0
        out_of_bounds = 0
        sw, sh_ = int(prs.slide_width or 0), int(prs.slide_height or 0)
        shape_total = 0
        for idx, slide in enumerate(prs.slides, start=1):
            if slide.slide_layout is None:
                no_layout += 1
            flat = _flatten(slide.shapes)
            shape_total += len(flat)
            slide_text = 0
            for sh in flat:
                st = str(sh.shape_type or "")
                if sh.has_text_frame:
                    t = sh.text_frame.text.strip()
                    if t:
                        slide_text += len(t)
                        for para in sh.text_frame.paragraphs:
                            for run in para.runs:
                                if run.font.name:
                                    fonts.add(str(run.font.name))
                    elif "TEXT_BOX" in st:
                        empty_text += 1
                if getattr(sh, "has_table", False) and sh.has_table:
                    tables += 1
                elif getattr(sh, "has_chart", False) and sh.has_chart:
                    charts += 1
                elif "PICTURE" in st:
                    name = str(sh.name or "")
                    tag = name.split(":", 1)[0] if ":" in name else ""
                    if tag in PICTURE_TAGS:
                        pictures_info.append(f"{label} {idx}장 {name}")
                    else:
                        pictures_bad.append(f"{label} {idx}장 '{name or '(이름 없음)'}'")
                elif "FREEFORM" in st:
                    freeform += 1
                elif "AUTO_SHAPE" in st or "LINE" in st or "CONNECTOR" in st:
                    auto += 1
                if sh.left is not None and sh.top is not None:
                    if (
                        int(sh.left) < 0
                        or int(sh.top) < 0
                        or int(sh.left) + int(sh.width or 0) > sw + 1000
                        or int(sh.top) + int(sh.height or 0) > sh_ + 1000
                    ):
                        out_of_bounds += 1
            for sh in slide.shapes:  # 그룹 구조 (최상위만)
                if "GROUP" in str(sh.shape_type):
                    kids = list(sh.shapes)
                    depth = _group_depth(sh)
                    if len(kids) < 2 or depth > 3:
                        groups_bad.append(
                            f"{label} {idx}장 '{sh.name}' 자식 {len(kids)}·깊이 {depth}"
                        )
            if idx > 1 and slide_text >= 60:
                body_slides += 1
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    notes_body += 1
        # 01 텍스트
        if empty_text > max(3, n_slides):
            revs[0].append(f"{label}: 빈 텍스트 상자 {empty_text}개")
        # 02 표
        if label == "제출용" and tables == 0:
            fails[1].append(f"{label}: 네이티브 표가 없음")
        # 03 도형
        if freeform > auto:
            revs[2].append(f"{label}: 자유형 {freeform} > 자동 도형 {auto}")
        # 04 차트
        if any("차트" in x or "chart" in x.lower() for x in pictures_bad):
            fails[3].append(
                f"{label}: 차트 그림 {[x for x in pictures_bad if '차트' in x or 'chart' in x.lower()][:2]}"
            )
        if charts == 0 and label == "제출용":
            revs[3].append(f"{label}: 네이티브 차트 없음 (공수·통계·간트 데이터가 비어 있음)")
        # 05 그룹
        if groups_bad:
            revs[4].extend(groups_bad[:5])
        # 06 그림
        if pictures_bad:
            fails[5].append(
                f"{label}: 태그 없는 그림 {len(pictures_bad)} — " + "; ".join(pictures_bad[:3])
            )
        # 07 글꼴
        unknown = sorted(f for f in fonts if f not in known_fonts)
        if unknown:
            revs[6].append(f"{label}: 기본·템플릿 외 글꼴 {', '.join(unknown[:5])}")
        # 08 노트
        if label == "발표용" and body_slides and notes_body < body_slides * 0.5:
            revs[7].append(f"{label}: 발표자 노트 {notes_body}/{body_slides} 본문 슬라이드")
        # 09 레이아웃
        if no_layout:
            revs[8].append(f"{label}: 레이아웃 없는 슬라이드 {no_layout}")
        if out_of_bounds:
            fails[8].append(f"{label}: 슬라이드 밖 도형 {out_of_bounds}")
        # 10 재열기
        try:
            import io as _io

            buf = _io.BytesIO()
            prs.save(buf)
            again = Presentation(_io.BytesIO(buf.getvalue()))
            again_total = sum(len(_flatten(sl.shapes)) for sl in again.slides)
            if len(again.slides) != n_slides or again_total != shape_total:
                fails[9].append(
                    f"{label}: 저장 왕복 후 슬라이드 {len(again.slides)}/{n_slides}, 도형 {again_total}/{shape_total}"
                )
        except Exception as exc:  # noqa: BLE001
            fails[9].append(f"{label}: 저장 왕복 실패 {type(exc).__name__}: {exc}")
        checks[5].issues.extend(f"예외: {x}" for x in pictures_info[:6])
    for i, c in enumerate(checks):
        extra = list(c.issues)
        _finish(c, fail=_cap(fails[i], 8), review=_cap(revs[i], 8), files=files)
        if i == 5:
            c.issues.extend(extra)
    return checks


def _group_depth(shape: Any, depth: int = 1) -> int:
    best = depth
    for kid in shape.shapes:
        if "GROUP" in str(kid.shape_type):
            best = max(best, _group_depth(kid, depth + 1))
    return best


def _invalid_ids(inp: QaInputs) -> list[str]:
    """형식(접두사-번호)이 아니거나, 대소문자·공백만 다른 중복 ID."""
    ids = [r.id for r in inp.all_requirements]
    bad = [i for i in ids if not re.fullmatch(r"[A-Za-z]{1,6}-?\d{2,6}", i)]
    seen: dict[str, str] = {}
    for i in ids:
        key = normalize_text(i)
        if key in seen and seen[key] != i:
            bad.append(f"{i} (≈ {seen[key]})")
        seen.setdefault(key, i)
    return bad


def _check_gantt(inp: QaInputs, art: CheckArtifacts) -> list[str]:
    d = inp.deliverables.get(DeliverableKind.WBS_GANTT_XLSX)
    if d is None:
        return []
    from openpyxl import load_workbook

    try:
        path = inp.paths.resolve_output(d.file_path)
        wb = load_workbook(path)
    except Exception as exc:  # noqa: BLE001
        return [f"간트 파일 열기 실패: {exc}"]
    if SHEET_WBS not in wb.sheetnames or SHEET_REPORT not in wb.sheetnames:
        return ["WBS 또는 약식WBS 시트 없음"]
    ws_w, ws_r = wb[SHEET_WBS], wb[SHEET_REPORT]
    wh = [c.value for c in ws_w[HEADER_ROW]]  # Phase 15: 1행 묶음 머리글·월, 2행 컬럼·주차
    if "WBS ID" not in wh or "계획 시작일" not in wh:
        return ["WBS 시트 헤더 없음"]
    ci, si, ei = wh.index("WBS ID"), wh.index("계획 시작일"), wh.index("계획 종료일")
    first_col = gantt_first_col(ws_w)
    if first_col is None:
        return ["WBS 시트에 간트 열 없음"]
    weeks = read_week_header(ws_w, first_col)
    # 기준 데이터(WBS 항목)의 날짜와 시트 날짜·채움 구간을 대조한다
    items = {w.code: w for w in inp.ctx.wbs}
    problems: list[str] = []
    seen: set[str] = set()
    for row in ws_w.iter_rows(min_row=HEADER_ROW + 1):
        code = row[ci].value
        if code is None:
            break
        code = str(code)
        seen.add(code)
        w = items.get(code)
        if w is None:
            problems.append(f"WBS 시트 행 {code}가 기준 데이터에 없음")
            continue
        ws_s, ws_e = _as_date(row[si].value), _as_date(row[ei].value)
        if (ws_s, ws_e) != (w.planned_start, w.planned_end):
            problems.append(f"{code}: 시트 {ws_s}~{ws_e} ≠ 기준 {w.planned_start}~{w.planned_end}")
            continue
        if ws_s and ws_e and weeks:
            cells = row[first_col - 1 : first_col - 1 + len(weeks)]
            filled = [
                i for i, cell in enumerate(cells) if _is_filled(cell) or cell.value is not None
            ]
            expected = span_weeks(weeks, ws_s, ws_e)
            if filled != expected:
                problems.append(
                    f"{code}: 간트 채움 주 {len(filled)}개 ≠ 일정 기준 {len(expected)}개"
                )
    missing = sorted(set(items) - seen)
    if missing:
        problems.append(f"WBS 시트에 없는 항목: {', '.join(missing[:6])}")
    # 약식WBS: 시작·종료일이 기준 데이터의 어떤 항목과 일치해야 한다
    rh = [c.value for c in ws_r[HEADER_ROW]]
    if "항목" in rh and "시작일" in rh:
        ni, rs, re_ = rh.index("항목"), rh.index("시작일"), rh.index("종료일")
        by_name = {w.name: w for w in inp.ctx.wbs}
        for row in ws_r.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
            name = row[ni]
            if name is None:
                break
            w = by_name.get(str(name))
            if w and (_as_date(row[rs]), _as_date(row[re_])) != (w.planned_start, w.planned_end):
                problems.append(f"약식WBS '{str(name)[:20]}' 날짜가 기준 데이터와 다름")
    else:
        problems.append("약식WBS 시트 헤더 없음")
    return _cap(problems)


def _as_date(v: Any) -> date | None:
    if v is None:
        return None
    if hasattr(v, "date") and callable(v.date):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _is_filled(cell: Any) -> bool:
    try:
        fill = cell.fill
        rgb = getattr(fill.fgColor, "rgb", None)
        return bool(
            fill and fill.fill_type == "solid" and rgb and rgb not in ("00000000", "FFFFFFFF")
        )
    except Exception:  # noqa: BLE001
        return False


# --- 실행·게이트 ---------------------------------------------------------------------------------


def run_all_checks(inp: QaInputs) -> tuple[list[QaCheck], CheckArtifacts]:
    file_checks, art = run_file_checks(inp)
    sub_text = "\n".join(art.ppt_texts.get("submission", []))
    checks = [
        *check_rfp_fit(inp, sub_text),
        *check_tech_feasibility(inp, art.ppt_texts),
        *check_consistency_rules(inp, art),
        *file_checks,
        *check_editability(inp, art),
    ]
    return checks, art


GATE_SCOPES: dict[str, Callable[[QaCheck], bool]] = {
    "gate2": lambda c: c.gate == "gate2",
    "gate3": lambda c: c.gate == "gate3",
    "gate4": lambda c: c.gate == "gate4",
    "gate5": lambda c: True,
}


def gate_status(checks: list[QaCheck], gate_key: str) -> tuple[str, list[str]]:
    """게이트별 판정 (코드). FAIL: 필수 항목 실패. NEEDS_HUMAN_REVIEW: 검토 필요·미실행·비필수 실패. PASS: 그 외.
    gate5(최종 독립 QA)는 자동으로 PASS가 되지 않는다 — 사람 승인이 필요하다."""
    scope = [c for c in checks if GATE_SCOPES[gate_key](c) and c.applicable]
    reasons: list[str] = []
    fails = [c for c in scope if c.executed and c.final_result == QaResult.FAIL and c.mandatory]
    soft = [
        c
        for c in scope
        if (c.executed and c.final_result == QaResult.FAIL and not c.mandatory)
        or c.final_result == QaResult.NEEDS_HUMAN_REVIEW
        or not c.executed
    ]
    if fails:
        reasons.append(f"필수 검증 실패 {len(fails)}건: " + ", ".join(c.check_id for c in fails))
        return "FAIL", reasons
    if soft:
        reasons.append(f"검토 필요 {len(soft)}건: " + ", ".join(c.check_id for c in soft))
        return "NEEDS_HUMAN_REVIEW", reasons
    if gate_key == "gate5":
        reasons.append("자동 검증 모두 통과 — 최종 판정은 사람 승인이 필요합니다.")
        return "NEEDS_HUMAN_REVIEW", reasons
    reasons.append(f"검증 {len(scope)}건 통과")
    return "PASS", reasons


__all__ = [
    "BUILD_KINDS",
    "FIX_REGENERATE",
    "PROPOSAL_KINDS",
    "CheckArtifacts",
    "QaInputs",
    "check_consistency_rules",
    "check_rfp_fit",
    "check_tech_feasibility",
    "gate_status",
    "run_all_checks",
    "run_file_checks",
]
