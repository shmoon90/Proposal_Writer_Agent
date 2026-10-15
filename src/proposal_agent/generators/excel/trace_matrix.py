"""요구사항 추적표 Excel (Phase 14 양식).

고객 산출물 관행(검사기준서형 추적표)을 따른다: 요구사항 → 수용 → 설계(화면·구성 요소·담당) → 구현(WBS·단위테스트) →
테스트(통합테스트) → 검사(판정 기준·완료 여부) → 검토 묶음, 묶음별 머리글 색. 비기능 요구사항은 '검사기준(비기능)'
시트에 요구 내용·구현 내용·검사 기준·수행 결과로 다시 정리한다.

제안 단계에서는 화면·테스트 ID가 아직 확정되지 않으므로 실제 존재하는 ID만 연결하고, 없으면 '미정'으로 둔다.
테스트 결과는 미실행/성공/실패/보류이며 수행하지 않은 테스트는 '미실행', 완료 여부는 '미착수'다.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook

from proposal_agent.generators.base import PENDING_ID, GenerationContext, acceptance_label
from proposal_agent.generators.excel.style import (
    HEADER_ROW,
    ROW_FILLS,
    Column,
    add_cover_sheet,
    add_guide_sheet,
    add_note,
    add_revision_sheet,
    write_grouped_table,
    write_table,
)
from proposal_agent.models import (
    IntegrationTest,
    Requirement,
    RequirementTrace,
    RequirementType,
    Screen,
    TestResult,
    UnitTest,
)

SHEET_MAIN = "추적표"
SHEET_NFR = "검사기준(비기능)"
SHEET_TESTS = "테스트상세"
SHEET_UNTRACED = "미추적"
DOC_TITLE = "요구사항 추적표"
RESULT_OPTIONS = tuple(t.value for t in TestResult)
COMPLETION_OPTIONS = ("미착수", "진행중", "완료", "대상제외")
REVIEW_OPTIONS = ("유효", "재검토 필요", "무효")
GROUP_TITLES = {
    "content": "요구사항",
    "accept": "수용",
    "design": "설계",
    "impl": "구현",
    "test": "테스트",
    "check": "검사 기준",
    "note": "검토",
}
COLUMNS: tuple[Column, ...] = (
    Column("No", 5, "center", "content"),
    Column("대분류", 11, "center", "content"),
    Column("중분류", 11, "center", "content"),
    Column("요구사항 ID", 16, "center", "content"),
    Column("요구사항 명", 28, "left", "content"),
    Column("유형", 9, "center", "content"),
    Column("수용여부", 11, "center", "accept"),
    Column("생성/변경일자", 12, "center", "accept"),
    Column("화면 ID", 14, "center", "design"),
    Column("화면명", 22, "left", "design"),
    Column("구성 요소 ID", 14, "center", "design"),
    Column("구성 요소명", 22, "left", "design"),
    Column("담당 역할", 12, "center", "design"),
    Column("WBS", 14, "center", "impl"),
    Column("Task 명", 24, "left", "impl"),
    Column("단위테스트 ID", 16, "center", "impl"),
    Column("단위테스트 결과", 10, "center", "impl", RESULT_OPTIONS),
    Column("통합테스트 ID", 16, "center", "test"),
    Column("통합테스트 명", 24, "left", "test"),
    Column("통합테스트 결과", 10, "center", "test", RESULT_OPTIONS),
    Column("판정 기준", 36, "left", "check"),
    Column("완료여부", 9, "center", "check", COMPLETION_OPTIONS),
    Column("검토 상태", 11, "center", "note", REVIEW_OPTIONS),
    Column("변경 사유", 30, "left", "note"),
    Column("비고", 30, "left", "note"),
)
HEADERS: tuple[str, ...] = tuple(c.header for c in COLUMNS)
NFR_COLUMNS: tuple[Column, ...] = (
    Column("No", 5, "center", "content"),
    Column("대분류", 11, "center", "content"),
    Column("중분류", 11, "center", "content"),
    Column("요구사항 ID", 16, "center", "content"),
    Column("요구사항 명", 24, "left", "content"),
    Column("요구사항 내용", 50, "left", "content"),
    Column("구현 내용", 50, "left", "impl"),
    Column("검사 기준", 30, "left", "check"),
    Column("수행 결과", 10, "center", "check", RESULT_OPTIONS),
    Column("비고", 24, "left", "note"),
)
GUIDE_ROWS: tuple[tuple[str, str, str, str], ...] = (
    (
        "요구사항 ID / 명",
        "요구사항 정의서의 ID와 이름을 그대로 적는다 (정의서에 없는 ID 금지).",
        "FR-001",
        "필수",
    ),
    (
        "화면 ID / 화면명",
        "요구사항이 구현되는 화면(설계서)의 ID. 설계 전이면 '미정'.",
        "SCR-001",
        "필수",
    ),
    ("구성 요소 ID / 명", "아키텍처 구성 요소(서비스·모듈·연계)의 ID.", "ARC-…", "선택"),
    ("담당 역할", "해당 요구사항을 구현하는 WBS Task의 담당 역할.", "응용 개발", "선택"),
    (
        "WBS / Task 명",
        "요구사항을 구현하는 WBS 항목. WBS에 없으면 '미추적' 시트에 표시.",
        "2.1.3",
        "필수",
    ),
    (
        "단위/통합테스트 ID·결과",
        "테스트 케이스 ID와 결과(미실행/성공/실패/보류). 수행 전에는 반드시 '미실행'.",
        "UT-001 / 미실행",
        "필수",
    ),
    ("판정 기준", "정의서 '검사 기준'. 완료를 판정할 수 있는 조건.", "응답 3초 이내", "필수"),
    ("완료여부", "미착수 / 진행중 / 완료 / 대상제외(미수용·삭제).", "미착수", "필수"),
    (
        "검토 상태 / 변경 사유",
        "요구사항이 변경·삭제되면 링크는 유지하고 '재검토 필요'와 사유를 적는다.",
        "재검토 필요 — CR-0001 수정",
        "필수",
    ),
)


def _completion(r: Requirement) -> str:
    if r.change_state == "deleted" or r.acceptance_status.value == "rejected":
        return "대상제외"
    return "미착수"


def _changed_on(r: Requirement) -> str:
    for h in reversed(r.history):
        if h.action.startswith(("change.", "baseline.load", "baseline.reload")):
            return h.at[:10]
    return r.created_at.date().isoformat() if getattr(r, "created_at", None) else ""


def build_trace_matrix(
    ctx: GenerationContext,
    traces: list[RequirementTrace],
    screens: list[Screen],
    unit_tests: list[UnitTest],
    integration_tests: list[IntegrationTest],
) -> bytes:
    screen_names = {s.code or s.id: s.name for s in screens}
    ut_ids = {t.code or t.id for t in unit_tests}
    it_names = {t.code or t.id: t.name for t in integration_tests}
    comp_names = {c.id: (c.name, c.layer) for c in ctx.components}
    by_req: dict[str, dict[str, list[str]]] = {}
    review: dict[str, list[str]] = {}
    for t in traces:
        bucket = by_req.setdefault(t.requirement_id, {})
        bucket.setdefault(t.target_type.value, []).append(t.target_id)
        if t.review_status != "유효":
            review.setdefault(t.requirement_id, []).append(
                f"{t.target_id}: {t.review_status} ({t.review_reason})"
            )
    wbs_by_req: dict[str, list[tuple[str, str, str]]] = {}
    for w in ctx.wbs:
        for rid in w.related_requirement_ids:
            wbs_by_req.setdefault(rid, []).append((w.code, w.name, w.role))

    rows: list[list[Any]] = []
    fills: list[str | None] = []
    strikes: list[bool] = []
    untraced: list[list[Any]] = []
    nfr_rows: list[list[Any]] = []
    band = False
    prev_group: str | None = None
    for no, r in enumerate(ctx.requirements, start=1):
        links = by_req.get(r.id, {})
        scr = [s for s in links.get("screen", []) if s in screen_names]
        uts = [u for u in links.get("unit_test", []) if u in ut_ids]
        its = [i for i in links.get("integration_test", []) if i in it_names]
        comps = [c for c in links.get("component", []) if c in comp_names]
        wbs_items = wbs_by_req.get(r.id, [])
        note = ""
        if not (scr or uts or its):
            note = "제안 단계: 화면·테스트 ID 미확정"
        if (
            not wbs_items
            and r.acceptance_status.value != "rejected"
            and r.change_state != "deleted"
        ):
            untraced.append([r.id, r.name, "WBS 미연결"])
        roles = sorted({role for _, _, role in wbs_items if role})
        rows.append(
            [
                no,
                r.business_area or r.type.label,
                r.category.label,
                r.id,
                r.name,
                r.type.label,
                acceptance_label(r),
                _changed_on(r),
                ", ".join(scr) or PENDING_ID,
                "\n".join(screen_names[s] for s in scr),
                ", ".join(comps),
                "\n".join(f"{comp_names[c][1]}/{comp_names[c][0]}" for c in comps),
                ", ".join(roles),
                ", ".join(code for code, _, _ in wbs_items) or PENDING_ID,
                "\n".join(name for _, name, _ in wbs_items),
                ", ".join(uts) or PENDING_ID,
                TestResult.NOT_RUN.value,
                ", ".join(its) or PENDING_ID,
                "\n".join(it_names[i] for i in its),
                TestResult.NOT_RUN.value,
                r.acceptance_criteria,
                _completion(r),
                "재검토 필요" if review.get(r.id) else "유효",
                "; ".join(review.get(r.id, [])),
                note,
            ]
        )
        group = r.rfp_requirement_id or r.id
        if group != prev_group:
            band = not band
            prev_group = group
        fill = ROW_FILLS.get(r.change_state)
        fills.append(fill if fill else (ROW_FILLS["band"] if band else None))
        strikes.append(r.change_state == "deleted")
        if r.type == RequirementType.NON_FUNCTIONAL:
            nfr_rows.append(
                [
                    len(nfr_rows) + 1,
                    r.business_area or r.type.label,
                    r.category.label,
                    r.id,
                    r.name,
                    r.description,
                    r.implementation_approach or "검토 필요",
                    r.acceptance_criteria,
                    TestResult.NOT_RUN.value,
                    "; ".join(review.get(r.id, [])),
                ]
            )

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_MAIN
    last = write_grouped_table(
        ws, COLUMNS, GROUP_TITLES, rows, freeze_col=4, row_fills=fills, strike_rows=strikes
    )
    add_note(
        ws,
        last + 2,
        "테스트 결과는 실제 수행 전까지 '미실행', 완료여부는 '미착수'. '미정'은 제안 단계에서 아직 확정되지 않은 ID. "
        "행 색: 연녹=추가, 연노랑=변경, 연주황=삭제 요청, 회색+취소선=삭제, 연회색 띠=같은 RFP 번호 묶음.",
    )
    ws_n = wb.create_sheet(SHEET_NFR)
    write_grouped_table(ws_n, NFR_COLUMNS, GROUP_TITLES, nfr_rows, freeze_col=4)

    ws2 = wb.create_sheet(SHEET_TESTS)
    test_rows = [
        [
            t.code or t.id,
            "단위",
            t.name,
            t.target_component,
            t.status if t.status in RESULT_OPTIONS else TestResult.NOT_RUN.value,
            "",
        ]
        for t in unit_tests
    ]
    test_rows += [
        [
            t.code or t.id,
            "통합",
            t.name,
            t.scenario,
            t.status if t.status in RESULT_OPTIONS else TestResult.NOT_RUN.value,
            "",
        ]
        for t in integration_tests
    ]
    write_table(
        ws2,
        ("테스트 ID", "구분", "테스트 명", "대상/시나리오", "결과", "비고"),
        test_rows,
        widths=(14, 8, 30, 40, 10, 30),
        validations={5: RESULT_OPTIONS},
    )

    ws3 = wb.create_sheet(SHEET_UNTRACED)
    write_table(ws3, ("요구사항 ID", "요구사항 명", "사유"), untraced, widths=(12, 30, 30))
    add_guide_sheet(wb, f"{DOC_TITLE} 작성 가이드", GUIDE_ROWS)
    add_revision_sheet(
        wb,
        DOC_TITLE,
        [["1.0", str(ctx.generated_on), "최초 작성 (공통 기준 데이터에서 생성)", "", ""]],
        index=0,
    )
    add_cover_sheet(
        wb,
        doc_title=DOC_TITLE,
        project_name=ctx.project.name,
        client=ctx.project.client_name or "",
        version=ctx.version,
        generated_on=ctx.generated_on,
    )
    wb.active = wb.sheetnames.index(SHEET_MAIN)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = [
    "COLUMNS",
    "HEADERS",
    "HEADER_ROW",
    "SHEET_MAIN",
    "SHEET_NFR",
    "SHEET_TESTS",
    "SHEET_UNTRACED",
    "build_trace_matrix",
]
