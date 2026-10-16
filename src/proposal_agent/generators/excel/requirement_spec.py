"""요구사항 정의서 Excel (Phase 14 양식).

고객 산출물 관행(표지·개정이력·요구사항목록·본문·작성가이드)을 따르되 고객을 특정하는 값은 없다.
본문 시트(기능/비기능/제약·평가·제출)는 **묶음 머리글(1행) + 컬럼 머리글(2행) + 데이터(3행~)** 구조이며
컬럼은 식별 → 요구사항 내용 → 구현(구현방안·난이도·검증·검사 기준) → 수용·우선순위 → 변경 정보 → 협의·비고 순이다.
구현방안은 요구사항마다 짧게(글머리 2~5줄) 들어가며 작성 주체(에이전트/규칙 초안/사람)를 '구현 근거' 컬럼에 남긴다.
변경 상태(추가/변경/삭제)는 행 색으로, 같은 RFP 요구사항의 세부 항목은 교대 띠로 구분한다 — 구축 단계 변경관리의
기준 문서로 쓰이므로 변경 일자·근거·요청자 컬럼을 둔다.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook

from proposal_agent.generators.base import (
    ACCEPTANCE_OPTIONS,
    GenerationContext,
    acceptance_label,
    requirement_note,
)
from proposal_agent.generators.excel.style import (
    HEADER_ROW,
    ROW_FILLS,
    Column,
    add_cover_sheet,
    add_guide_sheet,
    add_note,
    add_revision_sheet,
    write_grouped_table,
    write_table,
)
from proposal_agent.models import Requirement, RequirementCategory, RequirementType
from proposal_agent.requirements.evidence import (
    EVIDENCE_LABELS,
    default_response_strategy,
    evidence_label,
)

SHEET_COVER = "표지"
SHEET_REVISION = "개정이력"
SHEET_LIST = "요구사항목록"
SHEET_FUNCTIONAL = "기능요구사항"
SHEET_NON_FUNCTIONAL = "비기능요구사항"
SHEET_OTHER = "제약·평가·제출조건"
SHEET_HISTORY = "변경이력"
SHEET_CHANGE_LOG = "변경관리대장"
SHEET_GUIDE = "작성가이드"
DOC_TITLE = "요구사항 정의서"
CHANGE_LOG_HEADERS = (
    "CR ID",
    "변경 유형",
    "요청자",
    "요청일",
    "사유",
    "기존 요구사항 ID",
    "변경 전",
    "변경 후",
    "분류 근거",
    "변경 관점(문맥)",
    "구현 난이도",
    "승인 상태",
    "승인자",
    "승인일",
    "관련 산출물",
    "일정 영향",
    "공수 영향",
    "비용 영향",
    "결과",
)
CHANGE_LOG_WIDTHS = (10, 14, 12, 12, 30, 14, 40, 40, 50, 24, 10, 12, 12, 12, 40, 16, 30, 14, 40)
# 호환용: 검사·재적재 코드가 요구하는 필수 컬럼명 (바꾸지 않는다)
REQUIRED_COLUMNS: tuple[str, ...] = (
    "요구사항 ID",
    "요구사항 명",
    "요구사항 상세 내용",
    "구현방안",
    "수용여부",
    "비고",
)
DIFFICULTY_OPTIONS = ("상", "중", "하")
PRIORITY_OPTIONS = ("상", "중", "하")
VERIFICATION_OPTIONS = ("시연", "시험", "검사", "분석")
CHANGE_STATE_LABELS = {
    "none": "신규",
    "added": "추가",
    "modified": "변경",
    "delete_requested": "삭제 요청",
    "deleted": "삭제",
}
CHANGE_STATE_OPTIONS = tuple(CHANGE_STATE_LABELS.values())
CATEGORY_OPTIONS = tuple(c.label for c in RequirementCategory)
GROUP_TITLES = {
    "id": "식별",
    "content": "요구사항 내용",
    "impl": "구현 방안",
    "accept": "수용·우선순위",
    "change": "변경 정보 (구축 단계 기준)",
    "note": "협의·비고",
}
COLUMNS: tuple[Column, ...] = (
    Column("No", 5, "center", "id"),
    Column("요구사항 ID", 16, "center", "id"),
    Column("RFP 고유번호", 12, "center", "id"),
    Column("대분류", 12, "center", "id"),
    Column("중분류", 12, "center", "id", CATEGORY_OPTIONS),
    Column("요구사항 명", 26, "left", "id"),
    Column("요구사항 상세 내용", 52, "left", "content"),
    Column("유형", 9, "center", "content"),
    Column("원본 출처", 18, "left", "content"),
    Column("RFP 원문 근거", 36, "left", "content"),
    Column("근거 구분", 10, "center", "content", tuple(EVIDENCE_LABELS.values())),
    Column("구현방안", 52, "left", "impl"),
    Column("구현 난이도", 8, "center", "impl", DIFFICULTY_OPTIONS),
    Column("난이도 근거", 24, "left", "impl"),
    Column("구현 근거(조사·출처)", 24, "left", "impl"),
    Column("기술검토(14항목)", 30, "left", "impl"),
    Column("검증 방법", 9, "center", "impl", VERIFICATION_OPTIONS),
    Column("검사 기준", 36, "left", "impl"),
    Column("수용여부", 11, "center", "accept", ACCEPTANCE_OPTIONS),
    Column("대응 전략", 30, "left", "accept"),
    Column("중요도", 8, "center", "accept", PRIORITY_OPTIONS),
    Column("필수 여부", 8, "center", "accept", ("필수", "선택")),
    Column("범위 분류", 10, "center", "accept"),
    Column("검토 상태", 10, "center", "accept"),
    Column("검증 상태", 10, "center", "accept"),
    Column("요구사항 변경여부", 10, "center", "change", CHANGE_STATE_OPTIONS),
    Column("변경 내용", 30, "left", "change"),
    Column("변경 일자", 12, "center", "change"),
    Column("변경 근거", 24, "left", "change"),
    Column("추가 요청자", 11, "center", "change"),
    Column("추가일", 11, "center", "change"),
    Column("검토 및 협의 내용", 36, "left", "note"),
    Column("연결 산출물", 16, "left", "note"),
    Column("비고", 36, "left", "note"),
)
HEADERS: tuple[str, ...] = tuple(c.header for c in COLUMNS)
ACCEPT_COL = HEADERS.index("수용여부") + 1
GUIDE_ROWS: tuple[tuple[str, str, str, str], ...] = (
    (
        "요구사항 ID",
        "시스템이 발급한 식별자. 재발급·재사용하지 않는다 (추적표·WBS가 참조).",
        "FR-001",
        "필수",
    ),
    (
        "RFP 고유번호",
        "RFP(과업지시서)가 붙인 요구사항 번호. 한 번호의 세부 항목은 여러 행이 된다.",
        "SFR-001",
        "선택",
    ),
    (
        "대분류 / 중분류",
        "업무 구분(기능·비기능·제약…)과 분류(UI/UX·업무 기능·인터페이스·성능·보안…).",
        "기능 / 인터페이스",
        "필수",
    ),
    (
        "요구사항 명 / 상세 내용",
        "RFP 원문 의미를 보존한 이름과 설명. 수치·기한·범위는 원문 그대로.",
        "",
        "필수",
    ),
    ("원본 출처 / RFP 원문 근거", "원본 문서의 위치(쪽·절)와 근거 문장 인용.", "p.12 3.1", "필수"),
    (
        "구현방안",
        "요구사항을 실제로 구현할 방식의 요약(글머리 2~5줄). 제안서보다 짧게. 기술 조사·아키텍처 근거가 있는 내용만 구체화.",
        "- 관리자 화면에서 코드·메뉴 CRUD …",
        "필수",
    ),
    (
        "구현 난이도 / 난이도 근거",
        "상·중·하. 변경 요청의 공수·일정 영향 산정 가중치. 근거를 함께 적는다.",
        "상 — 외부 연계 2건",
        "필수",
    ),
    (
        "구현 근거(조사·출처)",
        "구현방안의 작성 주체(에이전트/규칙 초안/사람)와 근거 조사 ID·출처 URL. '[자동 초안]'은 사람 확인 대상.",
        "agent · TR-0001",
        "필수",
    ),
    (
        "검증 방법 / 검사 기준",
        "시연·시험·검사·분석 중 하나와, 완료를 판정할 수 있는 조건.",
        "시험 — 응답 3초 이내",
        "필수",
    ),
    (
        "수용여부",
        "수용 / 조건부 수용 / 미수용 / 검토 필요. 미수용·조건부는 비고에 사유.",
        "수용",
        "필수",
    ),
    (
        "요구사항 변경여부",
        "신규(최초 기준선) / 추가 / 변경 / 삭제 요청 / 삭제. 변경 내용·일자·근거를 함께 기록한다.",
        "변경",
        "필수",
    ),
    (
        "검토 및 협의 내용",
        "고객 협의·검토 의견 기록(날짜·결정 사항).",
        "@협의(2026.10.15) …",
        "선택",
    ),
    ("비고", "근거·제약·미수용 사유·검증 상태 등.", "", "선택"),
)


def _source(r: Requirement) -> str:
    return (
        "; ".join(ref.location_label or ref.document_id for ref in r.source_refs) or "(출처 없음)"
    )


def _last_change(r: Requirement) -> tuple[str, str]:
    """(변경 일자, 변경 근거) — 변경관리·편집 이력의 마지막 항목."""
    for h in reversed(r.history):
        if h.action.startswith(("change.", "edited", "baseline.reload")):
            return h.at[:10], h.note or h.action
    return "", ""


def _analysis_summary(r: Requirement) -> str:
    """요구사항별 기술 검토 14항목 요약 (단계 3·4): 채운 항목 수와 핵심 3항목 첫 줄."""
    a = r.analysis or {}
    filled = [(k, v) for k, v in a.items() if v]
    if not filled:
        return ""
    labels = {
        "purpose": "목적",
        "business_rules": "규칙",
        "data_transactions": "데이터",
        "api_integration": "연계",
        "performance_scalability": "성능",
        "failure_recovery": "장애",
        "test_acceptance": "테스트",
    }
    parts = [f"{len(filled)}/14 항목"]
    for k, v in filled:
        if k in labels and len(parts) < 4:
            first = v[0] if isinstance(v, list) and v else str(v)
            parts.append(f"{labels[k]}: {str(first)[:40]}")
    return " | ".join(parts)


def _approach_basis(r: Requirement) -> str:
    src = {
        "agent": "에이전트",
        "research": "조사 기반 초안",
        "rule": "규칙 초안(확인 필요)",
        "user": "사람 작성",
        "baseline": "기준 문서",
    }.get(r.approach_source, r.approach_source or "")
    parts = [p for p in (src, ", ".join(r.approach_basis[:4])) if p]
    return " · ".join(parts)


def _row(no: int, r: Requirement) -> list[Any]:
    changed_on, change_basis = _last_change(r)
    return [
        no,
        r.id,
        r.rfp_requirement_id or "",
        r.business_area or r.type.label,
        r.category.label,
        r.name,
        r.description,
        r.type.label,
        _source(r),
        (r.rfp_quote or "")[:400],
        evidence_label(r),
        r.implementation_approach
        or (
            f"해당 없음 ({r.type.label})"
            if r.type in (RequirementType.EVALUATION, RequirementType.SUBMISSION)
            else "검토 필요"
        ),
        r.difficulty.label if r.difficulty else "",
        r.difficulty_basis,
        _approach_basis(r),
        _analysis_summary(r),
        r.verification_method or "",
        r.acceptance_criteria,
        acceptance_label(r),
        default_response_strategy(r),
        r.priority.label,
        "필수" if r.is_mandatory else "선택",
        r.scope_class.label,
        r.review_status.value,
        r.verification_status.value,
        CHANGE_STATE_LABELS.get(r.change_state, r.change_state),
        r.change_summary,
        changed_on if r.change_state != "none" else "",
        change_basis if r.change_state != "none" else "",
        r.added_by or "",
        r.added_on,
        r.review_notes,
        ", ".join(r.linked_deliverables),
        requirement_note(r),
    ]


def _row_fills(reqs: list[Requirement]) -> tuple[list[str | None], list[bool]]:
    fills: list[str | None] = []
    strikes: list[bool] = []
    band = False
    prev_group: str | None = None
    for r in reqs:
        group = r.rfp_requirement_id or r.id
        if group != prev_group:
            band = not band
            prev_group = group
        fill = ROW_FILLS.get(r.change_state)
        if fill is None and band:
            fill = ROW_FILLS["band"]
        fills.append(fill)
        strikes.append(r.change_state == "deleted")
    return fills, strikes


def _revision_rows(ctx: GenerationContext) -> list[list[Any]]:
    rows: list[list[Any]] = [
        ["1.0", str(ctx.generated_on), "최초 작성 (공통 기준 데이터에서 생성)", "", ""]
    ]
    if ctx.version not in ("v1.0", "1.0"):
        rows.append([ctx.version.lstrip("v"), str(ctx.generated_on), "변경 반영 재생성", "", ""])
    return rows


def _write_list_sheet(wb: Workbook, reqs: list[Requirement]) -> None:
    """요구사항목록: 구분(종류)별 번호·이름 요약 (사람이 훑어보는 목차)."""
    ws = wb.create_sheet(SHEET_LIST)
    rows: list[list[Any]] = []
    for t in RequirementType:
        group = [r for r in reqs if r.type == t]
        for i, r in enumerate(group):
            rows.append(
                [
                    f"{t.label} ({len(group)}건)" if i == 0 else "",
                    r.id,
                    r.name,
                    r.rfp_requirement_id or "",
                ]
            )
    write_table(
        ws, ("구분", "요구사항 ID", "요구사항 명", "RFP 고유번호"), rows, widths=(22, 16, 50, 14)
    )


def build_requirement_spec(ctx: GenerationContext) -> bytes:
    wb = Workbook()
    groups: list[tuple[str, list[Requirement]]] = [
        (SHEET_FUNCTIONAL, [r for r in ctx.requirements if r.type == RequirementType.FUNCTIONAL]),
        (
            SHEET_NON_FUNCTIONAL,
            [r for r in ctx.requirements if r.type == RequirementType.NON_FUNCTIONAL],
        ),
        (
            SHEET_OTHER,
            [
                r
                for r in ctx.requirements
                if r.type not in (RequirementType.FUNCTIONAL, RequirementType.NON_FUNCTIONAL)
            ],
        ),
    ]
    first = True
    for title, reqs in groups:
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        first = False
        rows = [_row(i, r) for i, r in enumerate(reqs, start=1)]
        fills, strikes = _row_fills(reqs)
        last = write_grouped_table(
            ws,
            COLUMNS,
            GROUP_TITLES,
            rows,
            freeze_col=2,
            row_fills=fills,
            strike_rows=strikes,
        )
        add_note(
            ws,
            last + 2,
            f"{ctx.project.name} · 생성 {ctx.generated_on} · {ctx.version} · 수용여부: 수용/조건부 수용/미수용/검토 필요 · "
            "행 색: 연녹=추가, 연노랑=변경, 연주황=삭제 요청, 회색+취소선=삭제(기록 보존), 연회색 띠=같은 RFP 번호 묶음 · "
            "'[자동 초안]' 구현방안은 사람 확인 필요",
        )
    _write_list_sheet(wb, ctx.requirements)
    ws = wb.create_sheet(SHEET_HISTORY)
    hist_rows = []
    for r in ctx.requirements:
        for h in r.history:
            hist_rows.append(
                [r.id, h.at, h.actor, h.action, h.note, str(h.changes) if h.changes else ""]
            )
    write_table(
        ws,
        ("요구사항 ID", "일시", "작업자", "작업", "메모", "변경 내용"),
        hist_rows,
        widths=(12, 22, 22, 18, 40, 50),
    )
    if ctx.change_log:
        ws = wb.create_sheet(SHEET_CHANGE_LOG)
        write_table(ws, CHANGE_LOG_HEADERS, ctx.change_log, widths=CHANGE_LOG_WIDTHS)
    add_guide_sheet(wb, f"{DOC_TITLE} 작성 가이드", GUIDE_ROWS)
    add_revision_sheet(wb, DOC_TITLE, _revision_rows(ctx), index=0)
    add_cover_sheet(
        wb,
        doc_title=DOC_TITLE,
        project_name=ctx.project.name,
        client=ctx.project.client_name or "",
        version=ctx.version,
        generated_on=ctx.generated_on,
    )
    # 시트 순서: 표지 → 개정이력 → 요구사항목록 → 본문 3장 → 변경이력 → (변경관리대장) → 작성가이드
    order = [
        SHEET_COVER,
        SHEET_REVISION,
        SHEET_LIST,
        SHEET_FUNCTIONAL,
        SHEET_NON_FUNCTIONAL,
        SHEET_OTHER,
        SHEET_HISTORY,
    ]
    if SHEET_CHANGE_LOG in wb.sheetnames:
        order.append(SHEET_CHANGE_LOG)
    order.append(SHEET_GUIDE)
    wb._sheets = [wb[name] for name in order]  # noqa: SLF001
    wb.active = wb.sheetnames.index(SHEET_FUNCTIONAL)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = [
    "CHANGE_LOG_HEADERS",
    "CHANGE_STATE_LABELS",
    "COLUMNS",
    "HEADERS",
    "HEADER_ROW",
    "REQUIRED_COLUMNS",
    "SHEET_CHANGE_LOG",
    "SHEET_FUNCTIONAL",
    "SHEET_HISTORY",
    "SHEET_LIST",
    "SHEET_NON_FUNCTIONAL",
    "SHEET_OTHER",
    "build_requirement_spec",
]
