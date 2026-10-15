"""openpyxl 공통 서식: 헤더, 필터, 고정 창, 열 너비, 데이터 유효성, 줄바꿈."""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="맑은 고딕", size=10)
BODY_FONT = Font(name="맑은 고딕", size=10)
NOTE_FONT = Font(name="맑은 고딕", size=9, italic=True, color="7F7F7F")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_table(
    ws: Worksheet,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    *,
    widths: Sequence[int] | None = None,
    freeze: str = "A2",
    start_row: int = 1,
    validations: dict[int, Sequence[str]] | None = None,
    number_formats: dict[int, str] | None = None,
) -> int:
    """헤더+행을 쓰고 필터·고정 창·너비·유효성 검사를 적용한다. 마지막 행 번호를 돌려준다."""
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.fill, cell.font, cell.alignment, cell.border = HEADER_FILL, HEADER_FONT, CENTER, BORDER
    r = start_row
    for r_off, row in enumerate(rows, start=1):
        r = start_row + r_off
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.alignment, cell.border = BODY_FONT, WRAP, BORDER
            if number_formats and c in number_formats:
                cell.number_format = number_formats[c]
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A{start_row}:{last_col}{max(r, start_row + 1)}"
    ws.freeze_panes = freeze
    for c in range(1, len(headers) + 1):
        width = widths[c - 1] if widths and c - 1 < len(widths) else 16
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[start_row].height = 30
    if validations:
        for col, options in validations.items():
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(options) + '"',
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="허용되지 않는 값",
                error="목록에서 선택하십시오.",
            )
            ws.add_data_validation(dv)
            dv.add(
                f"{get_column_letter(col)}{start_row + 1}:{get_column_letter(col)}{max(r, start_row + 1) + 200}"
            )
    return r


def add_note(ws: Worksheet, row: int, text: str) -> None:
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = NOTE_FONT
    cell.alignment = Alignment(wrap_text=False)


# --- 묶음 머리글 표 (요구사항 정의서·추적표, Phase 14) -------------------------------------------
# 고객 산출물 관행: 1행 묶음 머리글(식별/내용/설계/구현/테스트/검사…, 묶음별 색) + 2행 컬럼 머리글 + 3행부터 데이터.
# 색은 샘플에서 쓰인 계열(연파랑·연주황·황갈·보라·녹색·회색)을 따르되 어떤 고객도 특정하지 않는다.
GROUP_FILLS: dict[str, str] = {
    "id": "D9D9D9",  # 식별 (회색)
    "content": "DBEEF3",  # 요구사항 내용 (연파랑)
    "accept": "FDEADB",  # 수용·우선순위 (연주황)
    "design": "DDD9C3",  # 설계 (황갈)
    "impl": "C4BD97",  # 구현 (진한 황갈)
    "test": "CCC1DA",  # 테스트 (보라)
    "check": "D8E4BC",  # 검사 기준 (연녹)
    "change": "E2EFDA",  # 변경 정보 (연녹)
    "note": "D9D9D9",  # 협의·비고 (회색)
}
ROW_FILLS: dict[str, str] = {
    "added": "E2EFDA",  # 추가
    "modified": "FFF2CC",  # 변경
    "delete_requested": "FCE4D6",  # 삭제 요청
    "deleted": "F2F2F2",  # 삭제 (기록 보존)
    "band": "F7F7F7",  # RFP 요구사항 그룹 교대 띠
}
GROUP_HEADER_FONT = Font(bold=True, name="맑은 고딕", size=10, color="000000")
COL_HEADER_FONT = Font(bold=True, name="맑은 고딕", size=10, color="000000")
STRIKE_FONT = Font(name="맑은 고딕", size=10, strike=True, color="7F7F7F")
TOP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")
TOP_CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")
GROUP_ROW = 1
HEADER_ROW = 2
FIRST_DATA_ROW = 3


@dataclass(frozen=True)
class Column:
    header: str
    width: float = 14
    align: str = "left"  # left | center
    group: str = "id"  # GROUP_FILLS 키
    options: tuple[str, ...] | None = None  # 데이터 유효성(목록)
    number_format: str | None = None


def write_grouped_table(
    ws: Worksheet,
    columns: Sequence[Column],
    group_titles: dict[str, str],
    rows: Sequence[Sequence[Any]],
    *,
    freeze_col: int = 1,
    row_fills: Sequence[str | None] | None = None,
    strike_rows: Sequence[bool] | None = None,
    min_row_height: float = 18.0,
) -> int:
    """묶음 머리글(1행)+컬럼 머리글(2행)+데이터(3행~)를 쓴다. 마지막 행 번호를 돌려준다.

    - ``group_titles``: 묶음 키 → 표기(예: "식별", "구현"). 인접한 같은 묶음 컬럼을 병합한다.
    - ``row_fills``: 행별 채움색(변경 상태·그룹 띠). ``strike_rows``: 취소선(삭제 행).
    - 머리글 행 번호는 ``HEADER_ROW``(2)로 고정되어 검사·재적재 코드가 그 행을 찾는다.
    """
    # 1행: 묶음 머리글 (병합)
    c = 1
    while c <= len(columns):
        g = columns[c - 1].group
        end = c
        while end < len(columns) and columns[end].group == g:
            end += 1
        cell = ws.cell(row=GROUP_ROW, column=c, value=group_titles.get(g, g))
        fill = PatternFill("solid", fgColor=GROUP_FILLS.get(g, "D9D9D9"))
        for cc in range(c, end + 1):
            hc = ws.cell(row=GROUP_ROW, column=cc)
            hc.fill, hc.font, hc.alignment, hc.border = fill, GROUP_HEADER_FONT, CENTER, BORDER
        if end > c:
            ws.merge_cells(start_row=GROUP_ROW, start_column=c, end_row=GROUP_ROW, end_column=end)
        cell.value = group_titles.get(g, g)
        c = end + 1
    # 2행: 컬럼 머리글
    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=HEADER_ROW, column=i, value=col.header)
        cell.fill = PatternFill("solid", fgColor=GROUP_FILLS.get(col.group, "D9D9D9"))
        cell.font, cell.alignment, cell.border = COL_HEADER_FONT, CENTER, BORDER
        ws.column_dimensions[get_column_letter(i)].width = col.width
    ws.row_dimensions[GROUP_ROW].height = 20
    ws.row_dimensions[HEADER_ROW].height = 32
    # 데이터
    r = HEADER_ROW
    for idx, row in enumerate(rows):
        r = FIRST_DATA_ROW + idx
        fill_hex = row_fills[idx] if row_fills and idx < len(row_fills) else None
        strike = bool(strike_rows and idx < len(strike_rows) and strike_rows[idx])
        max_lines = 1
        for i, col in enumerate(columns, start=1):
            v = row[i - 1] if i - 1 < len(row) else None
            cell = ws.cell(row=r, column=i, value=v)
            cell.font = STRIKE_FONT if strike else BODY_FONT
            cell.alignment = TOP_CENTER if col.align == "center" else TOP_LEFT
            cell.border = BORDER
            if fill_hex:
                cell.fill = PatternFill("solid", fgColor=fill_hex)
            if col.number_format:
                cell.number_format = col.number_format
            if isinstance(v, str) and v:
                est = sum(
                    max(1, len(ln) // max(8, int(col.width * 0.9)) + 1) for ln in v.splitlines()
                )
                max_lines = max(max_lines, est)
        ws.row_dimensions[r].height = max(min_row_height, min(409.0, 14.5 * max_lines + 4))
    last_col = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{max(r, FIRST_DATA_ROW)}"
    ws.freeze_panes = f"{get_column_letter(freeze_col + 1)}{FIRST_DATA_ROW}"
    for i, col in enumerate(columns, start=1):
        if col.options:
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(col.options) + '"',
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="허용되지 않는 값",
                error="목록에서 선택하십시오.",
            )
            ws.add_data_validation(dv)
            letter = get_column_letter(i)
            dv.add(f"{letter}{FIRST_DATA_ROW}:{letter}{max(r, FIRST_DATA_ROW) + 200}")
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{GROUP_ROW}:{HEADER_ROW}"
    return r


def find_header_row(ws: Worksheet, must_have: Sequence[str], *, limit: int = 6) -> int:
    """상위 ``limit``행 중 ``must_have`` 컬럼명을 가장 많이 포함한 행 (없으면 1)."""
    best, hits_best = 1, -1
    for r in range(1, min(ws.max_row, limit) + 1):
        values = {str(c.value).strip() for c in ws[r] if c.value is not None}
        hits = sum(1 for h in must_have if h in values)
        if hits > hits_best:
            best, hits_best = r, hits
    return best


def add_cover_sheet(
    wb: Any,
    *,
    doc_title: str,
    project_name: str,
    client: str,
    version: str,
    generated_on: Any,
    author: str = "",
) -> Worksheet:
    """표지 시트: 사업명·문서명·버전·작성일·작성/검토/승인란 (사람이 채운다)."""
    ws = wb.create_sheet("표지", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 18
    for col in "CDEFG":
        ws.column_dimensions[col].width = 16
    title = ws.cell(row=6, column=2, value=project_name)
    title.font = Font(name="맑은 고딕", size=20, bold=True)
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("B6:G8")
    sub = ws.cell(row=10, column=2, value=doc_title)
    sub.font = Font(name="맑은 고딕", size=16, bold=True, color="1F4E78")
    sub.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("B10:G11")
    meta = [
        ("발주 기관", client or "(미기재)"),
        ("문서 버전", version),
        ("작성일", str(generated_on)),
        ("작성", author or ""),
        ("검토", ""),
        ("승인", ""),
    ]
    for i, (k, v) in enumerate(meta):
        r = 16 + i
        kc = ws.cell(row=r, column=3, value=k)
        kc.font = Font(name="맑은 고딕", size=11, bold=True)
        kc.fill = PatternFill("solid", fgColor="D9D9D9")
        kc.alignment, kc.border = CENTER, BORDER
        vc = ws.cell(row=r, column=4, value=v)
        vc.font, vc.border = BODY_FONT, BORDER
        vc.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 22
    note = ws.cell(
        row=24,
        column=2,
        value="※ 작성·검토·승인란은 사람이 기입합니다. 본 문서는 공통 기준 데이터에서 생성되었습니다.",
    )
    note.font = NOTE_FONT
    return ws


def add_revision_sheet(
    wb: Any, doc_title: str, rows: Sequence[Sequence[Any]], *, index: int | None = None
) -> Worksheet:
    """개정이력 시트: 문서명 + (버전·날짜·내용·작성자·승인자) 표."""
    ws = wb.create_sheet("개정이력", index) if index is not None else wb.create_sheet("개정이력")
    head = ws.cell(row=1, column=1, value="문서 개정 이력")
    head.font = Font(name="맑은 고딕", size=12, bold=True, color="FFFFFF")
    head.fill = PatternFill("solid", fgColor="808080")
    head.alignment = CENTER
    ws.merge_cells("A1:E1")
    ws.cell(row=2, column=1, value="문서명").font = COL_HEADER_FONT
    ws.cell(row=2, column=1).fill = PatternFill("solid", fgColor="D9D9D9")
    ws.cell(row=2, column=1).alignment = CENTER
    ws.merge_cells("A2:B2")
    ws.cell(row=2, column=3, value=doc_title).font = BODY_FONT
    ws.merge_cells("C2:E2")
    for i, h in enumerate(("버전", "날짜", "내용", "작성자", "승인자"), start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font, c.alignment, c.border = COL_HEADER_FONT, CENTER, BORDER
        c.fill = PatternFill("solid", fgColor="D9D9D9")
    for r_off, row in enumerate(rows, start=4):
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r_off, column=i, value=v)
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = TOP_LEFT if i == 3 else TOP_CENTER
        ws.row_dimensions[r_off].height = 22
    for i in range(4 + len(rows), 4 + len(rows) + 6):  # 사람이 추가할 빈 줄
        for c_i in range(1, 6):
            ws.cell(row=i, column=c_i).border = BORDER
    for col, w in zip("ABCDE", (10, 14, 60, 12, 12), strict=True):
        ws.column_dimensions[col].width = w
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 22
    return ws


def add_guide_sheet(wb: Any, title: str, rows: Sequence[tuple[str, str, str, str]]) -> Worksheet:
    """작성가이드 시트: (항목명, 설명, 예시, 필수) — 사람이 편집할 때 참고."""
    ws = wb.create_sheet("작성가이드")
    head = ws.cell(row=1, column=1, value=title)
    head.font = Font(name="맑은 고딕", size=12, bold=True)
    ws.merge_cells("A1:D1")
    write_table(
        ws,
        ("항목명", "항목 설명", "예시", "필수"),
        rows,
        widths=(22, 70, 36, 8),
        start_row=3,
        freeze="A4",
    )
    return ws
