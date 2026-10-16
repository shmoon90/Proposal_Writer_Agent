"""인력 투입계획 Excel (Phase 15 양식).

- **투입계획**(구성원별): 조직 · 역할 · 구성원 · 기술 등급 · 담당 업무 · 투입 시작/종료월 · 월별 공수(MM) · 계(MM, SUM 수식) · 비고,
  맨 아래 **합계 행**(월별 SUM·총계) — 고객 월별 투입계획 관행(구분·등급·월·계·비고). 구성원 이름이 없으면 '역할N' 자리표시.
- **역할요약**: 역할 단위(조직·역할·등급·기술·인원·담당·기간·MM·M/D·월별·근거·단가 입력·비용 수식·비고).
- **공수요약**·**월별투입**: 역할별 합계(QA가 WBS 총 공수와 대조).
단가가 없으면 비용을 계산하지 않고 단가 입력란을 비워 둔다. 공수는 WBS 배분과 동일한 Decimal 값이다.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.excel.style import (
    BORDER,
    CENTER,
    HEADER_ROW,
    Column,
    add_note,
    write_grouped_table,
    write_table,
)
from proposal_agent.models import StaffRole
from proposal_agent.planning.calendar import BusinessCalendar
from proposal_agent.planning.effort import md_to_mm
from proposal_agent.planning.staffing import cost_for

SHEET_PLAN = "투입계획"
SHEET_ROLES = "역할요약"
SHEET_SUMMARY = "공수요약"
SHEET_MONTHLY = "월별투입"
GRADE_OPTIONS = ("특급", "고급", "중급", "초급", "미정")
GROUP_TITLES = {
    "id": "수행 조직·구성원",
    "content": "담당",
    "accept": "투입 기간",
    "impl": "월별 공수(MM)",
    "check": "합계",
    "note": "비고",
}
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
TOTAL_FONT = Font(name="맑은 고딕", size=10, bold=True)
Q2 = Decimal("0.01")


@dataclass
class RoleAgg:
    role: StaffRole
    md: Decimal = Decimal("0")
    monthly: dict[str, Decimal] = field(default_factory=dict)
    start: date | None = None
    end: date | None = None
    tasks: list[str] = field(default_factory=list)


def _aggregate(ctx: GenerationContext) -> tuple[dict[str, RoleAgg], list[str]]:
    roles = ctx.roles_by_id
    by_role: dict[str, RoleAgg] = {}
    months: set[str] = set()
    for a in ctx.allocations:
        role = roles.get(a.staff_role_id)
        if role is None:
            continue
        agg = by_role.setdefault(role.id, RoleAgg(role=role, start=a.start, end=a.end))
        agg.md += a.effort_md
        for m, v in a.monthly_md.items():
            agg.monthly[m] = agg.monthly.get(m, Decimal("0")) + Decimal(v)
            months.add(m)
        if a.start and (agg.start is None or a.start < agg.start):
            agg.start = a.start
        if a.end and (agg.end is None or a.end > agg.end):
            agg.end = a.end
        w = ctx.wbs_by_id.get(a.wbs_item_id)
        if w:
            agg.tasks.append(w.code)
    for role in ctx.roles:  # 배분이 없는 역할(조직 구성에만 있음)도 행으로 남긴다
        by_role.setdefault(role.id, RoleAgg(role=role))
    if not months and ctx.period_start and ctx.period_end:
        d = ctx.period_start.replace(day=1)
        while d <= ctx.period_end:
            months.add(f"{d.year:04d}-{d.month:02d}")
            d = (d.replace(day=28) + __import__("datetime").timedelta(days=4)).replace(day=1)
    return by_role, sorted(months)


def split_members(monthly_mm: dict[str, Decimal], n: int) -> list[dict[str, Decimal]]:
    """역할 월별 MM을 구성원 n명에게 균등 분배 (소수 둘째 자리, 끝 사람이 나머지를 받아 합계 보존)."""
    n = max(1, n)
    out: list[dict[str, Decimal]] = [{} for _ in range(n)]
    for m, total in monthly_mm.items():
        share = (total / n).quantize(Q2, rounding=ROUND_HALF_UP)
        for i in range(n - 1):
            out[i][m] = share
        out[n - 1][m] = (total - share * (n - 1)).quantize(Q2, rounding=ROUND_HALF_UP)
    return out


def _member_rows(
    by_role: dict[str, RoleAgg], month_list: list[str]
) -> tuple[list[list[Any]], list[str]]:
    rows: list[list[Any]] = []
    warnings: list[str] = []
    for agg in by_role.values():
        role = agg.role
        names = role.member_names()
        monthly_mm = {m: md_to_mm(v) for m, v in agg.monthly.items()}
        parts = split_members(monthly_mm, len(names))
        for name, part in zip(names, parts, strict=True):
            active = [m for m in month_list if part.get(m, Decimal("0")) > 0]
            peak = max((part.get(m, Decimal("0")) for m in month_list), default=Decimal("0"))
            note = "" if role.members else "구성원 이름 미정 — 확정 후 입력"
            if peak > Decimal("1.0"):
                note = (note + "; " if note else "") + f"월 최대 {peak} MM > 1.0 — 인원 추가 검토"
                warnings.append(f"{role.role_name}/{name}: 월 최대 {peak} MM")
            rows.append(
                [
                    role.organization or "수행사",
                    role.role_name,
                    name,
                    role.grade or "미정",
                    role.duties,
                    active[0] if active else "",
                    active[-1] if active else "",
                    *[float(part.get(m, Decimal("0"))) for m in month_list],
                    None,  # 계(MM): 수식
                    note or role.notes,
                ]
            )
    return rows, warnings


def _write_member_sheet(ws: Any, by_role: dict[str, RoleAgg], month_list: list[str]) -> int:
    columns = [
        Column("조직", 10, "center", "id"),
        Column("역할", 12, "center", "id"),
        Column("구성원", 12, "center", "id"),
        Column("기술 등급", 9, "center", "id", GRADE_OPTIONS),
        Column("담당 업무", 26, "left", "content"),
        Column("투입 시작월", 10, "center", "accept"),
        Column("투입 종료월", 10, "center", "accept"),
        *[Column(m, 8, "center", "impl", None, "0.00") for m in month_list],
        Column("계(MM)", 9, "center", "check", None, "0.00"),
        Column("비고", 30, "left", "note"),
    ]
    rows, _ = _member_rows(by_role, month_list)
    last = write_grouped_table(ws, columns, GROUP_TITLES, rows, freeze_col=3)
    first_month_col = 8
    last_month_col = 7 + len(month_list)
    total_col = last_month_col + 1
    fm, lm = get_column_letter(first_month_col), get_column_letter(last_month_col)
    data_first, data_last = 3, max(last, 3)
    for r in range(data_first, last + 1):  # 사람별 합계 (SUM 수식)
        cell = ws.cell(row=r, column=total_col, value=f"=SUM({fm}{r}:{lm}{r})")
        cell.number_format = "0.00"
    tr = last + 1  # 월별 합계 행
    ws.cell(row=tr, column=1, value="합계")
    for c in range(1, total_col + 2):
        cell = ws.cell(row=tr, column=c)
        cell.fill, cell.font, cell.border, cell.alignment = TOTAL_FILL, TOTAL_FONT, BORDER, CENTER
    for c in range(first_month_col, total_col + 1):
        letter = get_column_letter(c)
        cell = ws.cell(row=tr, column=c, value=f"=SUM({letter}{data_first}:{letter}{data_last})")
        cell.number_format = "0.00"
    ws.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=7)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(total_col + 1)}{data_last}"
    add_note(
        ws,
        tr + 2,
        "월별 공수(MM)는 WBS 역할 배분을 구성원 수로 균등 분배한 값(1 MM = 20 M/D). '계'와 '합계' 행은 수식. "
        "구성원 이름·등급은 확정 후 입력하며, 월 최대 1.0 MM를 넘는 행은 인원 추가를 검토한다.",
    )
    return tr


def build_staffing_plan(ctx: GenerationContext) -> bytes:
    by_role, month_list = _aggregate(ctx)
    calendar = BusinessCalendar()
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_PLAN
    _write_member_sheet(ws, by_role, month_list)

    # 역할요약 (역할 단위, 단가 입력·비용 수식)
    ws_r = wb.create_sheet(SHEET_ROLES)
    headers = [
        "조직",
        "역할",
        "기술 등급",
        "필요 기술",
        "투입 인원",
        "담당 업무",
        "투입 시작일",
        "투입 종료일",
        "투입 기간(영업일)",
        "예상 공수(MM)",
        "예상 공수(M/D)",
        *month_list,
        "산정 근거",
        "단가(원/MM, 입력)",
        "비용 합계(원)",
        "비고",
    ]
    rows: list[list[object]] = []
    total_md = Decimal("0")
    for agg in by_role.values():
        role = agg.role
        total_md += agg.md
        mm = md_to_mm(agg.md)
        duration = (
            calendar.business_days_between(agg.start, agg.end) if agg.start and agg.end else None
        )
        cost = cost_for(mm, role.unit_cost_per_mm)
        rows.append(
            [
                role.organization or "수행사",
                role.role_name,
                role.grade or "미정",
                ", ".join(role.skills),
                role.headcount_planned,
                role.duties,
                agg.start,
                agg.end,
                duration,
                float(mm),
                float(agg.md),
                *[float(md_to_mm(agg.monthly.get(m, Decimal("0")))) for m in month_list],
                f"WBS {', '.join(sorted(set(agg.tasks)))[:80]} 역할 배분 합계"
                if agg.tasks
                else "WBS 배분 없음 (조직 구성)",
                float(role.unit_cost_per_mm) if role.unit_cost_per_mm is not None else None,
                float(cost) if cost is not None else None,
                role.notes
                or ("단가 입력 시 비용 자동 계산" if role.unit_cost_per_mm is None else ""),
            ]
        )
    widths = [10, 12, 10, 20, 8, 28, 12, 12, 10, 10, 10, *[8] * len(month_list), 40, 14, 14, 24]
    fmt: dict[int, str] = {10: "0.00", 11: "0.0"}
    for i in range(len(month_list)):
        fmt[12 + i] = "0.00"
    unit_col = len(headers) - 2
    cost_col = len(headers) - 1
    fmt[unit_col] = "#,##0"
    fmt[cost_col] = "#,##0"
    last = write_table(
        ws_r,
        headers,
        rows,
        widths=widths,
        validations={3: GRADE_OPTIONS},
        number_formats=fmt,
    )
    unit_letter = get_column_letter(unit_col)
    for r in range(2, last + 1):  # 단가 × MM, 단가가 비어 있으면 빈 값
        ws_r.cell(
            row=r, column=cost_col, value=f'=IF({unit_letter}{r}="","",{unit_letter}{r}*J{r})'
        )
    add_note(
        ws_r,
        last + 2,
        f"총 공수 {float(total_md):.1f} M/D = {float(md_to_mm(total_md)):.2f} MM (1 MM = 20 M/D 가정). 단가 미입력 시 비용은 계산하지 않음.",
    )

    ws2 = wb.create_sheet(SHEET_SUMMARY)
    sum_rows: list[list[object]] = [
        [agg.role.role_name, agg.role.headcount_planned, float(agg.md), float(md_to_mm(agg.md))]
        for agg in by_role.values()
    ]
    sum_rows.append(
        [
            "합계",
            sum(agg.role.headcount_planned for agg in by_role.values()),
            float(total_md),
            float(md_to_mm(total_md)),
        ]
    )
    write_table(
        ws2,
        ("역할", "투입 인원", "공수(M/D)", "공수(MM)"),
        sum_rows,
        widths=(14, 10, 12, 12),
        number_formats={3: "0.0", 4: "0.00"},
    )
    add_note(
        ws2, len(sum_rows) + 3, "합계는 WBS 산정근거 시트의 총 공수와 같아야 한다 (QA에서 검증)."
    )

    ws3 = wb.create_sheet(SHEET_MONTHLY)
    m_rows: list[list[object]] = [
        [
            agg.role.role_name,
            *[float(md_to_mm(agg.monthly.get(m, Decimal("0")))) for m in month_list],
        ]
        for agg in by_role.values()
    ]
    m_rows.append(
        [
            "합계",
            *[
                float(
                    md_to_mm(
                        sum(
                            (agg.monthly.get(m, Decimal("0")) for agg in by_role.values()),
                            Decimal("0"),
                        )
                    )
                )
                for m in month_list
            ],
        ]
    )
    write_table(
        ws3,
        ("역할(MM)", *month_list),
        m_rows,
        widths=(14, *[9] * len(month_list)),
        number_formats={i + 2: "0.00" for i in range(len(month_list))},
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = [
    "SHEET_MONTHLY",
    "SHEET_PLAN",
    "SHEET_ROLES",
    "SHEET_SUMMARY",
    "build_staffing_plan",
    "split_members",
]
