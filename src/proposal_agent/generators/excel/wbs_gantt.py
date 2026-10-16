"""WBS Excel (Phase 15 양식): 상세 WBS(개발용) + 약식 WBS(보고용) + 의존관계 + 산정근거, 한 파일.

고객 WBS 관행을 따른다.
- **WBS(상세)**: 식별(WBS ID·상위·구분·단계) · 개발 항목(Task 명·상세·수량·요구사항·산출물) · 담당(개발 조직·담당 역할·개발자) ·
  일정(계획/실제 시작·종료일·기간·진척율·진행 상태) · 공수(M/D·산정 근거) · 관계(선행·후행·업무 유형) · 비고, 그 오른쪽에
  **주 단위 간트**(1행 월, 2행 월 내 주차, 셀 채움 = 계획 구간, 진척율만큼 진한 색, 마일스톤은 ◆ 표기).
- **약식WBS(보고용)**: 구분(단계) · 단계(활동) · 항목(작업) · 시작일 · 종료일 + 같은 주 단위 간트. 첫 행은 사업관리 마일스톤
  (착수·보고·오픈·종료 라벨), 구분·단계 셀은 세로 병합.
- 날짜·기간·공수 합계는 planning/ 코드가 계산한 값 그대로다 (시트 안에서 다시 계산하지 않음).
"""

from __future__ import annotations

import io
from datetime import date, timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from proposal_agent.generators.base import GenerationContext
from proposal_agent.generators.excel.style import (
    BORDER,
    CENTER,
    HEADER_ROW,
    Column,
    add_note,
    write_grouped_table,
    write_table,
)
from proposal_agent.models import TaskStatus, WBSItem, WorkType

SHEET_WBS = "WBS"
SHEET_REPORT = "약식WBS"
SHEET_DEPS = "의존관계"
SHEET_BASIS = "산정근거"
SHEET_GANTT = SHEET_WBS  # 호환: 간트는 별도 시트가 아니라 WBS·약식WBS 시트 오른쪽에 붙는다
GROUP_TITLES = {
    "id": "식별",
    "content": "개발 항목",
    "design": "담당",
    "accept": "일정",
    "impl": "공수",
    "test": "관계",
    "note": "비고",
}
STATUS_OPTIONS = tuple(t.value for t in TaskStatus)
COLUMNS: tuple[Column, ...] = (
    Column("WBS ID", 9, "center", "id"),
    Column("상위 WBS ID", 9, "center", "id"),
    Column("구분", 12, "center", "id"),
    Column("단계", 8, "center", "id"),
    Column("Task 명", 30, "left", "content"),
    Column("Task 상세 내용", 34, "left", "content"),
    Column("수량", 11, "center", "content"),
    Column("요구사항 ID", 18, "left", "content"),
    Column("산출물", 16, "left", "content"),
    Column("개발 조직", 11, "center", "design"),
    Column("담당 역할", 13, "center", "design"),
    Column("개발자", 11, "center", "design"),
    Column("계획 시작일", 11, "center", "accept", None, "yyyy-mm-dd"),
    Column("계획 종료일", 11, "center", "accept", None, "yyyy-mm-dd"),
    Column("기간(영업일)", 8, "center", "accept"),
    Column("실제 시작일", 11, "center", "accept", None, "yyyy-mm-dd"),
    Column("실제 종료일", 11, "center", "accept", None, "yyyy-mm-dd"),
    Column("진척율(%)", 8, "center", "accept", None, "0"),
    Column("진행 상태", 9, "center", "accept", STATUS_OPTIONS),
    Column("공수(M/D)", 9, "center", "impl", None, "0.0"),
    Column("산정 근거", 34, "left", "impl"),
    Column("선행 업무", 11, "center", "test"),
    Column("후행 업무", 11, "center", "test"),
    Column("업무 유형", 9, "center", "test", tuple(t.value for t in WorkType)),
    Column("비고", 22, "left", "note"),
)
HEADERS: tuple[str, ...] = tuple(c.header for c in COLUMNS)
REPORT_COLUMNS: tuple[Column, ...] = (
    Column("구분", 12, "center", "id"),
    Column("단계", 12, "center", "id"),
    Column("항목", 30, "left", "content"),
    Column("시작일", 11, "center", "accept", None, "yyyy-mm-dd"),
    Column("종료일", 11, "center", "accept", None, "yyyy-mm-dd"),
)
REPORT_GROUP_TITLES = {"id": "구분", "content": "개발 항목", "accept": "일정"}
LEVEL_LABELS = {"phase": "단계", "activity": "활동", "task": "작업"}
# 간트 색: 단계(진한 파랑)·활동(파랑)·작업(연파랑)·진척(녹색)·마일스톤(주황)·월 머리글(회색)
PHASE_FILL = PatternFill("solid", fgColor="2E75B6")
ACTIVITY_FILL = PatternFill("solid", fgColor="5B9BD5")
TASK_FILL = PatternFill("solid", fgColor="BDD7EE")
DONE_FILL = PatternFill("solid", fgColor="70AD47")
MILESTONE_FILL = PatternFill("solid", fgColor="F4B183")
MONTH_FILL = PatternFill("solid", fgColor="D9D9D9")
WEEK_FILL = PatternFill("solid", fgColor="F2F2F2")
GANTT_WIDTH = 3.4
GANTT_FONT = Font(name="맑은 고딕", size=8)
MILESTONE_FONT = Font(name="맑은 고딕", size=8, bold=True)


# --- 주 단위 간트 공통 ---------------------------------------------------------------------------
def week_starts(start: date, end: date) -> list[date]:
    """start가 속한 주의 월요일부터 end까지 7일 간격 월요일 목록."""
    first = start - timedelta(days=start.weekday())
    out: list[date] = []
    cursor = first
    while cursor <= end:
        out.append(cursor)
        cursor += timedelta(days=7)
    return out


def week_of_month(monday: date) -> int:
    """월요일 기준 월 내 주차 (그 달의 n번째 월요일)."""
    return (monday.day - 1) // 7 + 1


def nth_monday(year: int, month: int, n: int) -> date:
    d = date(year, month, 1)
    d += timedelta(days=(7 - d.weekday()) % 7)  # 첫 월요일
    return d + timedelta(days=7 * (n - 1))


def write_week_header(ws: Worksheet, weeks: list[date], first_col: int) -> None:
    """1행: 월(YYYY-MM, 같은 달 병합) · 2행: 월 내 주차."""
    col = first_col
    i = 0
    while i < len(weeks):
        label = f"{weeks[i].year:04d}-{weeks[i].month:02d}"
        j = i
        while j + 1 < len(weeks) and (weeks[j + 1].year, weeks[j + 1].month) == (
            weeks[i].year,
            weeks[i].month,
        ):
            j += 1
        cell = ws.cell(row=1, column=col + i, value=label)
        for c in range(col + i, col + j + 1):
            hc = ws.cell(row=1, column=c)
            hc.fill, hc.font, hc.alignment, hc.border = MONTH_FILL, MILESTONE_FONT, CENTER, BORDER
        if j > i:
            ws.merge_cells(start_row=1, start_column=col + i, end_row=1, end_column=col + j)
        cell.value = label
        i = j + 1
    for k, wk in enumerate(weeks):
        hc = ws.cell(row=HEADER_ROW, column=col + k, value=week_of_month(wk))
        hc.fill, hc.font, hc.alignment, hc.border = WEEK_FILL, GANTT_FONT, CENTER, BORDER
        ws.column_dimensions[get_column_letter(col + k)].width = GANTT_WIDTH


def read_week_header(ws: Worksheet, first_col: int) -> list[date]:
    """write_week_header로 쓴 머리글(월 라벨 병합 + 주차)을 월요일 날짜 목록으로 복원한다 (QA 대조용)."""
    weeks: list[date] = []
    month: tuple[int, int] | None = None
    col = first_col
    while True:
        label = ws.cell(row=1, column=col).value
        n = ws.cell(row=HEADER_ROW, column=col).value
        if label is not None:
            try:
                y, m = str(label).split("-")
                month = (int(y), int(m))
            except ValueError:
                break
        if month is None or not isinstance(n, int | float):
            break
        weeks.append(nth_monday(month[0], month[1], int(n)))
        col += 1
    return weeks


def gantt_first_col(ws: Worksheet) -> int | None:
    """컬럼 머리글 행에서 '비고'(상세) 또는 '종료일'(약식) 다음 열 = 간트 첫 열."""
    headers = [c.value for c in ws[HEADER_ROW]]
    for key in ("비고", "종료일"):
        if key in headers:
            return headers.index(key) + 2
    return None


def span_weeks(weeks: list[date], start: date | None, end: date | None) -> list[int]:
    if not start or not end:
        return []
    return [i for i, wk in enumerate(weeks) if wk <= end and wk + timedelta(days=6) >= start]


def fill_gantt_row(
    ws: Worksheet,
    row: int,
    first_col: int,
    weeks: list[date],
    w: WBSItem,
    *,
    is_parent: bool,
    label: str | None = None,
) -> None:
    span = span_weeks(weeks, w.planned_start, w.planned_end)
    if w.milestone:
        for i in span:
            cell = ws.cell(row=row, column=first_col + i, value=label or f"◆{w.name[:6]}")
            cell.fill, cell.font = MILESTONE_FILL, MILESTONE_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)
        return
    base = (
        PHASE_FILL
        if (is_parent and w.parent_id is None)
        else ACTIVITY_FILL
        if is_parent
        else TASK_FILL
    )
    done_until = -1
    pct = 100 if w.status == TaskStatus.DONE else w.progress_pct
    if pct > 0 and span:
        done_until = int(round(len(span) * pct / 100)) - 1
    for k, i in enumerate(span):
        ws.cell(row=row, column=first_col + i).fill = DONE_FILL if k <= done_until else base


def _progress(w: WBSItem) -> int:
    return 100 if w.status == TaskStatus.DONE else w.progress_pct


# --- 시트 작성 --------------------------------------------------------------------------------------
def _detail_rows(ctx: GenerationContext) -> tuple[list[list[Any]], list[WBSItem]]:
    by_id = ctx.wbs_by_id
    preds: dict[str, list[str]] = {}
    succs: dict[str, list[str]] = {}
    for d in ctx.dependencies:
        p, s = by_id.get(d.predecessor_id), by_id.get(d.successor_id)
        if p and s:
            preds.setdefault(s.id, []).append(p.code)
            succs.setdefault(p.id, []).append(s.code)
    rows: list[list[Any]] = []
    for w in ctx.wbs:
        parent = by_id.get(w.parent_id) if w.parent_id else None
        rows.append(
            [
                w.code,
                parent.code if parent else "",
                w.category or (parent.category if parent and parent.category else ""),
                "마일스톤" if w.milestone else LEVEL_LABELS.get(w.level.value, w.level.value),
                w.name,
                w.description,
                w.quantity,
                ", ".join(w.related_requirement_ids),
                ", ".join(w.deliverable_names),
                w.organization,
                w.role,
                w.assignee or ("" if w.milestone else "(미정)"),
                w.planned_start,
                w.planned_end,
                w.duration_days,
                w.actual_start,
                w.actual_end,
                _progress(w),
                w.status.value,
                float(w.effort_md),
                w.estimate_basis,
                ", ".join(preds.get(w.id, [])),
                ", ".join(succs.get(w.id, [])),
                w.work_type.value,
                ("자동 생성 — 검토 필요; " if w.auto_generated else "") + w.notes,
            ]
        )
    return rows, list(ctx.wbs)


def _write_detail(ws: Worksheet, ctx: GenerationContext, weeks: list[date]) -> int:
    rows, items = _detail_rows(ctx)
    parents = {w.parent_id for w in ctx.wbs if w.parent_id}
    fills: list[str | None] = []
    for w in items:
        if w.milestone:
            fills.append("FBE5D6")
        elif w.parent_id is None:
            fills.append("DDEBF7")
        elif w.id in parents:
            fills.append("F2F7FC")
        else:
            fills.append(None)
    last = write_grouped_table(ws, COLUMNS, GROUP_TITLES, rows, freeze_col=5, row_fills=fills)
    first_col = len(COLUMNS) + 1
    if weeks:
        write_week_header(ws, weeks, first_col)
        for i, w in enumerate(items):
            fill_gantt_row(ws, 3 + i, first_col, weeks, w, is_parent=w.id in parents)
    ws.freeze_panes = "F3"
    add_note(
        ws,
        last + 2,
        "날짜는 계산값(영업일). 업무 유형은 표시용이며 선후행은 '의존관계' 시트가 기준. 간트: 진한 파랑=단계, 파랑=활동, "
        "연파랑=작업, 녹색=진척(진척율·완료), 주황 ◆=마일스톤. '개발자'·'실제 시작/종료일'·'진척율'은 구축 단계에서 사람이 갱신.",
    )
    return last


def _write_report(ws: Worksheet, ctx: GenerationContext, weeks: list[date]) -> int:
    """약식WBS(보고용): 마일스톤 행 + 단계/활동/작업 행 (구분·단계 세로 병합)."""
    by_id = ctx.wbs_by_id
    children: dict[str | None, list[WBSItem]] = {}
    for w in ctx.wbs:
        children.setdefault(w.parent_id, []).append(w)
    phases = [w for w in ctx.wbs if w.parent_id is None and not w.milestone]
    milestones = [w for w in ctx.wbs if w.milestone and w.parent_id is not None]
    rows: list[list[Any]] = []
    gantt: list[tuple[WBSItem, bool, str | None]] = []  # (item, is_parent, label)
    merges: list[tuple[int, int, int]] = []  # (col, first_row_idx, last_row_idx)
    if milestones:
        rows.append(
            [
                "사업관리",
                "마일스톤",
                " → ".join(m.name for m in milestones),
                min(m.planned_start for m in milestones if m.planned_start),
                max(m.planned_end for m in milestones if m.planned_end),
            ]
        )
        gantt.append((milestones[0], False, "__milestones__"))
    for ph in phases:
        block_start = len(rows)
        acts = children.get(ph.id, [])
        if not acts:
            rows.append([ph.name, "", ph.name, ph.planned_start, ph.planned_end])
            gantt.append((ph, True, None))
        for act in acts:
            tasks = children.get(act.id, [])
            act_start = len(rows)
            if not tasks:
                rows.append([ph.name, act.name, act.name, act.planned_start, act.planned_end])
                gantt.append((act, False, None))
            for t in tasks:
                rows.append([ph.name, act.name, t.name, t.planned_start, t.planned_end])
                gantt.append((t, False, None))
            if len(rows) - act_start > 1:
                merges.append((2, act_start, len(rows) - 1))
        if len(rows) - block_start > 1:
            merges.append((1, block_start, len(rows) - 1))
    last = write_grouped_table(ws, REPORT_COLUMNS, REPORT_GROUP_TITLES, rows, freeze_col=5)
    for col, a, b in merges:
        ws.merge_cells(start_row=3 + a, start_column=col, end_row=3 + b, end_column=col)
        ws.cell(row=3 + a, column=col).alignment = CENTER
    first_col = len(REPORT_COLUMNS) + 1
    if weeks:
        write_week_header(ws, weeks, first_col)
        for i, (w, is_parent, label) in enumerate(gantt):
            if label == "__milestones__":
                for m in milestones:
                    for k in span_weeks(weeks, m.planned_start, m.planned_end):
                        cell = ws.cell(row=3 + i, column=first_col + k)
                        cell.value = (
                            f"{cell.value}/{m.name}" if cell.value else m.name
                        )  # 같은 주의 마일스톤은 함께 표기
                        cell.fill, cell.font = MILESTONE_FILL, MILESTONE_FONT
                        cell.alignment = Alignment(
                            horizontal="left", vertical="center", shrink_to_fit=True
                        )
                continue
            fill_gantt_row(ws, 3 + i, first_col, weeks, w, is_parent=is_parent)
    ws.freeze_panes = "F3"
    add_note(
        ws, last + 2, "보고용 약식 WBS — 상세 WBS와 같은 일정 데이터. 구분·단계 셀은 세로 병합."
    )
    _ = by_id
    return last


def build_wbs_gantt(ctx: GenerationContext) -> bytes:
    by_id = ctx.wbs_by_id
    starts = [w.planned_start for w in ctx.wbs if w.planned_start]
    ends = [w.planned_end for w in ctx.wbs if w.planned_end]
    weeks = week_starts(min(starts), max(ends)) if starts and ends else []
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_WBS
    _write_detail(ws, ctx, weeks)
    ws_r = wb.create_sheet(SHEET_REPORT)
    _write_report(ws_r, ctx, weeks)

    ws3 = wb.create_sheet(SHEET_DEPS)
    dep_rows = [
        [by_id[d.predecessor_id].code, by_id[d.successor_id].code, d.type.value, d.lag_days]
        for d in ctx.dependencies
        if d.predecessor_id in by_id and d.successor_id in by_id
    ]
    write_table(ws3, ("선행 WBS", "후행 WBS", "유형", "지연(일)"), dep_rows, widths=(12, 12, 8, 10))

    ws4 = wb.create_sheet(SHEET_BASIS)
    basis_rows = [
        [
            w.code,
            w.name,
            float(w.effort_md),
            w.estimate_basis,
            "자동 생성" if w.auto_generated else "에이전트/사용자",
        ]
        for w in ctx.wbs
        if not w.milestone
    ]
    write_table(
        ws4,
        ("WBS ID", "Task 명", "공수(M/D)", "산정 근거", "출처"),
        basis_rows,
        widths=(10, 32, 10, 60, 14),
        number_formats={3: "0.0"},
    )
    add_note(
        ws4,
        len(basis_rows) + 3,
        f"총 공수 {float(ctx.total_effort_md):.1f} M/D · 1 M/M = 20 M/D 가정",
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = [
    "COLUMNS",
    "HEADERS",
    "SHEET_BASIS",
    "SHEET_DEPS",
    "SHEET_GANTT",
    "SHEET_REPORT",
    "SHEET_WBS",
    "build_wbs_gantt",
    "gantt_first_col",
    "read_week_header",
    "span_weeks",
    "week_starts",
]
