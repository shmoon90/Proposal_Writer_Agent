"""PPTX 리더 (python-pptx) — 단계 2 (지시문 §0.1 입력 자료: PowerPoint).

슬라이드 순서대로 제목 자리표시자 → 제목 블록, 텍스트 상자(그룹 안 포함) → 문단 블록, 표 → 표 블록으로 읽는다.
``page`` 는 슬라이드 번호다. 그림만 있는 슬라이드는 글자 수 0 으로 기록하며 OCR 은 하지 않는다(그림은 이미지
문서로 따로 올리면 OCR 한다).
"""

from __future__ import annotations

import io
from typing import Any

from proposal_agent.config.logging import get_logger
from proposal_agent.ingest.base import BlockBuilder, ReaderResult
from proposal_agent.models import (
    BlockLocation,
    ExtractionMethod,
    ExtractionStatus,
    FileFormat,
    PageStat,
)

log = get_logger("ingest.pptx")
_GROUP_SHAPE = 6  # MSO_SHAPE_TYPE.GROUP


def _iter_shapes(shapes: Any) -> Any:
    for sh in shapes:
        if getattr(sh, "shape_type", None) == _GROUP_SHAPE:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _is_title(shape: Any) -> bool:
    try:
        if not shape.is_placeholder:
            return False
        return "TITLE" in str(shape.placeholder_format.type)
    except Exception:
        return False


def read_pptx(data: bytes, *, project_id: str, document_id: str, file_name: str) -> ReaderResult:
    result = ReaderResult(
        file_format=FileFormat.PPTX, status=ExtractionStatus.SUCCESS, unit_label="slide"
    )
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        result.status = ExtractionStatus.FAILED
        result.error = {
            "message": "python-pptx가 설치되지 않았습니다.",
            "cause": str(exc),
            "resolution": "pip install python-pptx",
        }
        return result
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        result.status = ExtractionStatus.FAILED
        result.error = {
            "message": "PPTX 파일을 열 수 없습니다.",
            "cause": f"{type(exc).__name__}: {exc}",
            "resolution": "PowerPoint에서 다시 저장(.pptx)하거나 PDF로 변환해 업로드하십시오.",
        }
        return result
    builder = BlockBuilder(project_id, document_id, file_name, FileFormat.PPTX)
    total_chars = 0
    table_index = 0
    for idx, slide in enumerate(prs.slides, start=1):
        chars = 0
        for shape in _iter_shapes(slide.shapes):
            if getattr(shape, "has_table", False) and shape.has_table:
                cells = [[" ".join(c.text.split()) for c in r.cells] for r in shape.table.rows]
                cells = [r for r in cells if any(r)]
                if cells:
                    chars += sum(len(c) for r in cells for c in r)
                    builder.add_table(
                        cells, location=BlockLocation(page=idx, table_index=table_index)
                    )
                    table_index += 1
                continue
            if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
                continue
            text = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
            if not text:
                continue
            chars += len(text)
            lines = [t for t in text.split("\n") if t.strip()]
            if _is_title(shape) and lines:
                builder.add_text(
                    lines[0][:80],
                    location=BlockLocation(
                        page=idx, paragraph_index=builder.next_paragraph_index()
                    ),
                    style_hint="title",
                )
                lines = lines[1:]
            for para in lines:
                builder.add_text(
                    para,
                    location=BlockLocation(
                        page=idx, paragraph_index=builder.next_paragraph_index()
                    ),
                )
        total_chars += chars
        result.page_stats.append(
            PageStat(page=idx, char_count=chars, method=ExtractionMethod.TEXT, confidence=1.0)
        )
    result.blocks = builder.blocks
    result.page_count = len(prs.slides)
    if total_chars == 0:
        result.warnings.append(
            "텍스트가 없는 PPTX 입니다 (그림만 있는 슬라이드는 OCR 하지 않음 — 이미지 문서로 올리면 OCR)."
        )
    return result


__all__ = ["read_pptx"]
