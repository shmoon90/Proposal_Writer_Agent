"""구축 단계 기준 문서(Excel) 파서: 시트·컬럼 자동 매핑 제안, 데이터 품질 검사, 결정적 파싱.

표준 구조(이 시스템이 만든 파일)는 자동 매핑되고, 다른 양식은 동의어·유사도로 매핑을 제안한다.
확신이 낮은 매핑(``needs_confirmation``)은 사용자가 확정해야 적재한다.
"""

from __future__ import annotations

import difflib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from proposal_agent.config.errors import IngestError
from proposal_agent.models import DocumentKind

# 대상 필드 → 동의어 (정규화 비교). 첫 항목이 표준 헤더.
FIELD_SYNONYMS: dict[DocumentKind, dict[str, tuple[str, ...]]] = {
    DocumentKind.REQUIREMENT_SPEC: {
        "requirement_id": (
            "요구사항 ID",
            "요구사항ID",
            "REQ ID",
            "ID",
            "식별자",
            "요구사항번호",
            "번호",
        ),
        "name": ("요구사항 명", "요구사항명", "명칭", "제목", "요구사항"),
        "description": (
            "요구사항 상세 내용",
            "상세내용",
            "상세 내용",
            "내용",
            "설명",
            "요구사항 설명",
            "상세",
        ),
        "implementation": ("구현방안", "구현 방안", "대응방안", "대응 방안"),
        "acceptance": ("수용여부", "수용 여부", "수용"),
        "notes": ("비고", "참고", "메모"),
        "type": ("종류", "구분", "유형", "요구사항 유형"),
        "scope": ("범위 분류", "범위"),
        "mandatory": ("필수 여부", "필수"),
        "priority": ("중요도", "우선순위"),
        "added_by": ("추가 요청자", "요청자"),
        "added_on": ("추가일", "추가 일자", "등록일"),
    },
    DocumentKind.TRACE_MATRIX: {
        "requirement_id": ("요구사항 ID", "요구사항ID", "REQ ID", "ID"),
        "name": ("요구사항 명", "요구사항명", "명칭"),
        "screen_id": ("화면 ID", "화면ID", "화면"),
        "unit_test_id": ("단위테스트 ID", "단위테스트ID", "단위 테스트", "UT ID"),
        "unit_test_result": ("단위테스트 결과", "단위 결과"),
        "integration_test_id": ("통합테스트 ID", "통합테스트ID", "통합 테스트", "IT ID"),
        "integration_test_result": ("통합테스트 결과", "통합 결과"),
        "wbs": ("WBS", "WBS ID", "작업"),
        "review_status": ("검토 상태", "검토상태"),
        "review_reason": ("변경 사유", "사유"),
    },
    DocumentKind.WBS: {
        "wbs_id": ("WBS ID", "WBS 코드", "WBS", "ID", "코드"),
        "parent_id": ("상위 WBS ID", "상위 WBS", "상위", "상위 코드"),
        "requirement_id": ("요구사항 ID", "요구사항ID", "요구사항"),
        "name": ("Task 명", "Task명", "작업명", "작업 명", "활동", "업무명", "과업"),
        "description": ("Task 상세 내용", "상세 내용", "내용", "설명"),
        "role": ("담당 역할", "담당", "역할", "담당자"),
        "predecessors": ("선행 업무", "선행", "선행 작업"),
        "successors": ("후행 업무", "후행", "후행 작업"),
        "work_type": ("업무 유형", "유형"),
        "planned_start": ("계획 시작일", "시작일", "착수일", "계획시작"),
        "planned_end": ("계획 종료일", "종료일", "완료일", "계획종료"),
        "actual_start": ("실제 시작일", "실제시작"),
        "actual_end": ("실제 종료일", "실제종료"),
        "duration": ("기간(영업일)", "기간", "일수"),
        "status": ("진행 상태", "상태", "진행률"),
        "basis": ("산정 근거", "산정근거", "근거"),
        "effort_md": ("공수(M/D)", "공수", "MD", "M/D"),
        "notes": ("비고", "참고"),
    },
    DocumentKind.STAFFING_PLAN: {
        "organization": ("조직", "소속"),
        "role": ("역할", "직무", "담당 역할"),
        "grade": ("기술 등급", "등급"),
        "skills": ("필요 기술", "기술", "스킬"),
        "headcount": ("투입 인원", "인원", "인력"),
        "duties": ("담당 업무", "업무"),
        "start": ("투입 시작일", "시작일", "투입시작"),
        "end": ("투입 종료일", "종료일", "투입종료"),
        "effort_mm": ("예상 공수(MM)", "공수(MM)", "MM", "공수"),
        "effort_md": ("예상 공수(M/D)", "공수(M/D)"),
        "basis": ("산정 근거", "근거"),
        "unit_cost": ("단가(원/MM, 입력)", "단가", "단가(원/MM)"),
        "cost": ("비용 합계(원)", "비용", "금액"),
        "notes": ("비고",),
    },
}
# WBS 파일의 "의존관계" 시트는 별도 컬럼 집합을 쓴다.
DEP_SHEET = "의존관계"
DEP_SYNONYMS: dict[str, tuple[str, ...]] = {
    "pred": ("선행 WBS", "선행", "선행 작업", "predecessor"),
    "succ": ("후행 WBS", "후행", "후행 작업", "successor"),
    "dep_type": ("유형", "의존 유형", "type"),
    "lag": ("지연(일)", "지연", "lag"),
}
DEP_REQUIRED = ("pred", "succ")


def synonyms_for(kind: DocumentKind, sheet: str) -> dict[str, tuple[str, ...]]:
    if kind == DocumentKind.WBS and sheet == DEP_SHEET:
        return DEP_SYNONYMS
    return FIELD_SYNONYMS[kind]


def required_for(kind: DocumentKind, sheet: str) -> tuple[str, ...]:
    if kind == DocumentKind.WBS and sheet == DEP_SHEET:
        return DEP_REQUIRED
    return REQUIRED_FIELDS[kind]


REQUIRED_FIELDS: dict[DocumentKind, tuple[str, ...]] = {
    DocumentKind.REQUIREMENT_SPEC: ("requirement_id", "name", "description"),
    DocumentKind.TRACE_MATRIX: ("requirement_id",),
    DocumentKind.WBS: ("wbs_id", "name"),
    DocumentKind.STAFFING_PLAN: ("role",),
}
# 표준 시트명 (이 시스템이 생성). 시트 → 종류 힌트
STANDARD_SHEETS: dict[str, DocumentKind] = {
    "기능요구사항": DocumentKind.REQUIREMENT_SPEC,
    "비기능요구사항": DocumentKind.REQUIREMENT_SPEC,
    "제약·평가·제출조건": DocumentKind.REQUIREMENT_SPEC,
    "추적표": DocumentKind.TRACE_MATRIX,
    "WBS": DocumentKind.WBS,
    "의존관계": DocumentKind.WBS,
    "투입계획": DocumentKind.STAFFING_PLAN,
}
SKIP_SHEETS = {"변경이력", "간트", "산정근거", "공수요약", "월별투입", "미추적", "테스트상세"}
_NORM = re.compile(r"[\s\W_]+")


def normalize_header(value: Any) -> str:
    return _NORM.sub("", str(value or "")).casefold()


@dataclass
class ColumnMapping:
    field: str
    header: str | None
    column_index: int | None
    confidence: float
    needs_confirmation: bool
    reason: str = ""


@dataclass
class SheetMapping:
    sheet: str
    kind: DocumentKind
    header_row: int
    columns: list[ColumnMapping] = field(default_factory=list)
    sample_rows: list[list[Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def needs_confirmation(self) -> bool:
        return any(c.needs_confirmation for c in self.columns)

    def index_of(self, field_name: str) -> int | None:
        for c in self.columns:
            if c.field == field_name:
                return c.column_index
        return None

    def to_dict(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet,
            "kind": self.kind.value,
            "header_row": self.header_row,
            "columns": {
                c.field: c.column_index for c in self.columns if c.column_index is not None
            },
        }


@dataclass
class WorkbookProfile:
    kind: DocumentKind
    sheets: list[SheetMapping]
    quality: list[str] = field(default_factory=list)
    formulas: int = 0

    @property
    def needs_confirmation(self) -> bool:
        return any(s.needs_confirmation for s in self.sheets)


def _find_header_row(ws: Any, kind: DocumentKind) -> int:
    """상위 10행 중 동의어가 가장 많이 맞는 행."""
    vocab = {normalize_header(s) for syns in synonyms_for(kind, ws.title).values() for s in syns}
    best_row, best_hits = 1, -1
    for r in range(1, min(ws.max_row, 10) + 1):
        hits = sum(1 for c in ws[r] if normalize_header(c.value) in vocab)
        if hits > best_hits:
            best_row, best_hits = r, hits
    return best_row


def propose_mapping(ws: Any, kind: DocumentKind) -> SheetMapping:
    header_row = _find_header_row(ws, kind)
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[header_row]]
    norm = [normalize_header(h) for h in headers]
    mapping = SheetMapping(sheet=ws.title, kind=kind, header_row=header_row)
    used: set[int] = set()
    required_fields = required_for(kind, ws.title)
    for field_name, synonyms in synonyms_for(kind, ws.title).items():
        chosen: ColumnMapping | None = None
        for rank, syn in enumerate(synonyms):
            n = normalize_header(syn)
            for i, h in enumerate(norm):
                if h and h == n and i not in used:
                    chosen = ColumnMapping(
                        field_name,
                        headers[i],
                        i,
                        1.0 if rank == 0 else 0.9,
                        False,
                        "표준 헤더" if rank == 0 else f"동의어 '{syn}'",
                    )
                    break
            if chosen:
                break
        if chosen is None:
            best_i, best_ratio = None, 0.0
            for i, h in enumerate(norm):
                if not h or i in used:
                    continue
                ratio = max(
                    difflib.SequenceMatcher(None, h, normalize_header(s)).ratio() for s in synonyms
                )
                if ratio > best_ratio:
                    best_i, best_ratio = i, ratio
            if best_i is not None and best_ratio >= 0.72:
                chosen = ColumnMapping(
                    field_name,
                    headers[best_i],
                    best_i,
                    round(best_ratio * 0.8, 2),
                    True,
                    f"유사 헤더 '{headers[best_i]}' (유사도 {best_ratio:.2f}) — 확인 필요",
                )
        if chosen is None:
            required = field_name in required_fields
            chosen = ColumnMapping(
                field_name,
                None,
                None,
                0.0,
                required,
                "헤더를 찾지 못함" + (" (필수)" if required else ""),
            )
        if chosen.column_index is not None:
            used.add(chosen.column_index)
        mapping.columns.append(chosen)
    for row in ws.iter_rows(min_row=header_row + 1, max_row=min(ws.max_row, header_row + 5)):
        values = [c.value for c in row]
        if any(v is not None for v in values):
            mapping.sample_rows.append(values)
    return mapping


def detect_kind(ws: Any, declared: DocumentKind | None) -> DocumentKind | None:
    if ws.title in STANDARD_SHEETS:
        return STANDARD_SHEETS[ws.title]
    if ws.title in SKIP_SHEETS:
        return None
    best, best_hits = None, 0
    hits_by_kind: dict[DocumentKind, int] = {}
    for kind in (
        DocumentKind.REQUIREMENT_SPEC,
        DocumentKind.TRACE_MATRIX,
        DocumentKind.WBS,
        DocumentKind.STAFFING_PLAN,
    ):
        row = _find_header_row(ws, kind)
        vocab = {normalize_header(s) for syns in FIELD_SYNONYMS[kind].values() for s in syns}
        hits = sum(1 for c in ws[row] if normalize_header(c.value) in vocab)
        hits_by_kind[kind] = hits
        if hits > best_hits:
            best, best_hits = kind, hits
    if declared is not None and (best_hits < 2 or hits_by_kind.get(declared, 0) >= best_hits - 1):
        return declared  # 사용자가 지정한 종류를 우선 (근소한 차이는 지정 종류로)
    return best if best_hits >= 2 else None


def profile_workbook(data: bytes, declared_kind: DocumentKind) -> WorkbookProfile:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=False)
    except Exception as exc:  # noqa: BLE001
        raise IngestError(
            "Excel 파일을 열 수 없습니다.",
            cause=f"{type(exc).__name__}: {exc}",
            resolution=".xlsx 형식인지, 암호화되지 않았는지 확인하십시오.",
        ) from exc
    profile = WorkbookProfile(kind=declared_kind, sheets=[])
    for ws in wb.worksheets:
        kind = detect_kind(ws, declared_kind)
        if kind is None or kind != declared_kind:
            if ws.title not in SKIP_SHEETS and kind is not None:
                profile.quality.append(
                    f"시트 '{ws.title}'는 {kind.value} 구조로 보여 {declared_kind.value} 적재에서 제외합니다."
                )
            continue
        mapping = propose_mapping(ws, kind)
        _check_quality(ws, mapping, profile)
        profile.sheets.append(mapping)
    if not profile.sheets:
        profile.quality.append(
            f"{declared_kind.value} 구조로 인식되는 시트가 없습니다. 헤더 행과 컬럼명을 확인하십시오."
        )
    return profile


def _check_quality(ws: Any, mapping: SheetMapping, profile: WorkbookProfile) -> None:
    id_field = (
        "requirement_id"
        if mapping.kind in (DocumentKind.REQUIREMENT_SPEC, DocumentKind.TRACE_MATRIX)
        else "wbs_id"
        if mapping.kind == DocumentKind.WBS
        else "role"
    )
    idx = mapping.index_of(id_field) if mapping.sheet != DEP_SHEET else None
    seen: dict[str, int] = {}
    empties = 0
    date_fields = [
        f
        for f in ("planned_start", "planned_end", "start", "end", "added_on")
        if mapping.index_of(f) is not None
    ]
    for r, row in enumerate(
        ws.iter_rows(min_row=mapping.header_row + 1), start=mapping.header_row + 1
    ):
        values = [c.value for c in row]
        if not any(v is not None for v in values):
            break
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                profile.formulas += 1
        if idx is not None:
            key = str(values[idx]).strip() if idx < len(values) and values[idx] is not None else ""
            if not key:
                empties += 1
            elif key in seen:
                mapping.warnings.append(f"{mapping.sheet} {r}행: ID '{key}' 중복 ({seen[key]}행)")
            else:
                seen[key] = r
        for f in date_fields:
            ci = mapping.index_of(f)
            if (
                ci is not None
                and ci < len(values)
                and values[ci] not in (None, "")
                and parse_date(values[ci]) is None
            ):
                mapping.warnings.append(
                    f"{mapping.sheet} {r}행: '{f}' 날짜 형식 해석 불가 ({values[ci]!r})"
                )
    if empties:
        mapping.warnings.append(f"{mapping.sheet}: 식별자 없는 행 {empties}개 (적재 제외)")
    for c in mapping.columns:
        if c.column_index is None and c.field in required_for(mapping.kind, mapping.sheet):
            mapping.warnings.append(f"{mapping.sheet}: 필수 컬럼 '{c.field}' 매핑 없음")


def parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S", "%y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


@dataclass
class ParsedRow:
    sheet: str
    row: int
    values: dict[str, Any]


def parse_rows(
    data: bytes, profile: WorkbookProfile, *, overrides: dict[str, dict[str, int]] | None = None
) -> list[ParsedRow]:
    """확정된 매핑(필요 시 사용자 override: {sheet: {field: col_index}})으로 행을 읽는다. 수식은 계산값(data_only)을 쓴다."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    rows: list[ParsedRow] = []
    for mapping in profile.sheets:
        ws = wb[mapping.sheet]
        cols = {c.field: c.column_index for c in mapping.columns if c.column_index is not None}
        if overrides and mapping.sheet in overrides:
            cols.update({k: v for k, v in overrides[mapping.sheet].items() if v is not None})
        unresolved = [
            c.field
            for c in mapping.columns
            if c.needs_confirmation
            and c.field not in (overrides or {}).get(mapping.sheet, {})
            and c.field in required_for(mapping.kind, mapping.sheet)
        ]
        if unresolved:
            raise IngestError(
                f"시트 '{mapping.sheet}'의 필수 컬럼 매핑이 확정되지 않았습니다: {', '.join(unresolved)}",
                resolution="컬럼 매핑을 확인·확정한 뒤 적재하십시오.",
            )
        for r, row in enumerate(
            ws.iter_rows(min_row=mapping.header_row + 1), start=mapping.header_row + 1
        ):
            values = [c.value for c in row]
            if not any(v is not None for v in values):
                break
            record = {f: (values[i] if i < len(values) else None) for f, i in cols.items()}
            record["_sheet"] = mapping.sheet
            rows.append(ParsedRow(sheet=mapping.sheet, row=r, values=record))
    return rows
